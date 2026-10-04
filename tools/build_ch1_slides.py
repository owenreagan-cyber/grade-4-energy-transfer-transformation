"""Build Chapter 1 "Energy Causes Change" — GOLD-STANDARD slide deck (PPTX, 16:9).

Phase 4. Consumes the FROZEN Unit 1 Visual System v1.0 tokens and mirrors the
FINALIZED Chapter 1 Guided Notes (immutable BQ/V/N anchors) exactly.

Run: .venv/bin/python tools/build_ch1_slides.py
Then: soffice --headless --convert-to pdf  (see render step)
"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
OUT = os.path.join(ROOT, "production", "slides")

# ---------------------------------------------------------------- FROZEN Ch1 tokens
PRIMARY = RGBColor(0x1F, 0x2A, 0x5A)   # deep navy
ACCENT  = RGBColor(0xA6, 0x7C, 0x16)   # warm gold
NEUTRAL = RGBColor(0x7A, 0x56, 0x20)   # muted amber
SURFACE = RGBColor(0xF5, 0xEE, 0xDC)   # cream
INK     = RGBColor(0x1A, 0x1D, 0x26)
INK_SOFT= RGBColor(0x4A, 0x52, 0x68)
PAPER   = RGBColor(0xFF, 0xFF, 0xFF)
MIST    = RGBColor(0xF6, 0xF7, 0xF9)
RULE    = RGBColor(0xD8, 0xDC, 0xE4)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)

# gold 10% tint (card definition/visual panel) — a light warm fill
GOLD_TINT = RGBColor(0xF3, 0xE8, 0xCF)
NAVY_TINT = RGBColor(0xE3, 0xE7, 0xF2)

FONT = "Avenir Next"

# Type scale (task §4 + frozen tokens)
T_MAJOR  = 48   # chapter title
T_TITLE  = 40   # slide title
T_HEAD   = 32   # notes section heading
T_BODY   = 30   # instructional body
T_BODY_S = 28   # body (denser lines)
T_OPT    = 28   # MCQ options
T_DEF    = 26   # vocab definition on card
T_KICK   = 20   # kicker / subtitle
T_PAGE   = 16   # Reader page refs
T_TAB    = 12   # chapter tab / chrome
T_HUD    = 18   # HUD bubble text

CHAPTER_TAB = "CHAPTER 1 · ENERGY CAUSES CHANGE"

# ---------------------------------------------------------------- presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

SLIDE_NO = [0]  # mutable counter

def new_slide(bg=SURFACE):
    SLIDE_NO[0] += 1
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background(); r.shadow.inherit = False
    return s

# ---------------------------------------------------------------- low-level helpers
def rect(s, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line_color=None, line_w=None, rounded=False):
    sp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else shape,
                            Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        try:
            sp.adjustments[0] = 0.08
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line_color is not None:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(line_w or 1.25)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp

def tb(s, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    return box, tf

def run(p, text, size=T_BODY, color=INK, bold=False, italic=False, underline=False, font=FONT, caps=False):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color; r.font.name = font
    if underline:
        r.font.underline = True
    if caps:
        r.font.caps = True
    return r

def text(s, x, y, w, h, lines, size=T_BODY, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False, gap=10, caps=False):
    """lines = str or list of str (each becomes a paragraph)."""
    if isinstance(lines, str):
        lines = [lines]
    box, tf = tb(s, x, y, w, h, anchor=anchor)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(gap)
        run(p, ln, size=size, color=color, bold=bold, italic=italic, caps=caps)
    return box

def rich(s, x, y, w, h, segs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, gap=12):
    """segs = list of paragraphs; each paragraph = list of (text, fmt_dict)."""
    box, tf = tb(s, x, y, w, h, anchor=anchor)
    for i, para_segs in enumerate(segs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(gap)
        for (txt, fmt) in para_segs:
            run(p, txt, size=fmt.get('size', T_BODY), color=fmt.get('color', INK),
                bold=fmt.get('bold', False), italic=fmt.get('italic', False),
                underline=fmt.get('underline', False))
    return box

def blank_str(n):
    return "_" * n

# ---------------------------------------------------------------- chrome / header
def chapter_tab(s, x=0.55, y=0.40, w=3.55, h=0.44):
    sp = rect(s, x, y, w, h, PRIMARY, rounded=True)
    tf = sp.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.12); tf.margin_right = Inches(0.12)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, CHAPTER_TAB, size=T_TAB, color=WHITE, bold=True)
    return sp

def mode_badge(s, mode, x=12.55, y=0.40, w=0.55, h=0.44, fill=PRIMARY, color=WHITE):
    sp = rect(s, x, y, w, h, fill, rounded=True)
    tf = sp.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, mode, size=T_HUD, color=color, bold=True)
    return sp

HUD_W = {  # 18pt bold Avenir widths + 0.34 padding, right-aligned to 12.5 (clear of mode badge)
    "TURN & TALK": 2.00, "OBSERVE": 1.56, "PREDICT": 1.44, "EVIDENCE": 1.62,
    "WRITE": 1.15, "CHECK": 1.23, "RECITE": 1.26, "TRACE": 1.20, "DRAW": 1.10,
    "EXPLAIN": 1.42, "VOTE": 1.02, "WATCH": 1.24, "DATA": 1.12, "ENGINEER": 1.64,
}
def hud_bubble(s, text_label, x=None, y=0.40, h=0.44):
    w = HUD_W.get(text_label, 1.6)
    if x is None:
        x = 12.5 - w
    sp = rect(s, x, y, w, h, ACCENT, rounded=True)
    tf = sp.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.08)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, text_label, size=T_HUD, color=WHITE, bold=True)
    return sp

def slide_number(s, n=None):
    n = n or SLIDE_NO[0]
    text(s, 12.2, 7.02, 0.9, 0.4, str(n), size=T_PAGE, color=INK_SOFT, align=PP_ALIGN.RIGHT)

def anchor_chip(s, code, x=0.55, y=1.02, w=0.78, h=0.5):
    sp = rect(s, x, y, w, h, ACCENT, rounded=True)
    tf = sp.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, code, size=24, color=WHITE, bold=True)
    return sp

def header(s, title, kicker=None, mode=None, hud=None, anchor=None, title_size=T_TITLE):
    chapter_tab(s)
    if mode:
        mode_badge(s, mode)
    if hud:
        hud_bubble(s, hud)
    tx = 0.55
    if anchor:
        anchor_chip(s, anchor)
        tx = 1.5
    text(s, tx, 1.02, 12.75 - tx, 0.7, title, size=title_size, color=PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if kicker:
        text(s, 0.55, 1.72, 11.7, 0.4, kicker, size=T_KICK, color=INK_SOFT, italic=True)
    slide_number(s)

# ---------------------------------------------------------------- model / diagrams
def model_strip(s, items, x=0.75, y=None, w=11.8, box_h=0.82, gap=0.5, size=26, fills=None):
    """Horizontal arrow chain (boxes + gold arrows)."""
    n = len(items)
    box_w = (w - (n - 1) * gap) / n
    cx = x
    for i, it in enumerate(items):
        f = (fills[i] if fills else (PRIMARY if i % 2 == 0 else NEUTRAL))
        sp = rect(s, cx, y, box_w, box_h, f, rounded=True)
        tf = sp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.08)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        run(p, it, size=size, color=WHITE, bold=True)
        if i < n - 1:
            ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(cx + box_w + 0.02),
                                    Inches(y + box_h/2 - 0.16), Inches(gap - 0.04), Inches(0.32))
            ar.fill.solid(); ar.fill.fore_color.rgb = ACCENT
            ar.line.fill.background(); ar.shadow.inherit = False
        cx += box_w + gap
    return box_h

def before_after(s, x=0.75, y=2.5, w=5.6, h=2.6, left="BEFORE", right="AFTER", left_hint="", right_hint=""):
    """Two-panel BEFORE / AFTER comparison."""
    gap = 0.35
    bw = (w - gap) / 2
    for idx, (lab, hint) in enumerate([(left, left_hint), (right, right_hint)]):
        cx = x + idx * (bw + gap)
        rect(s, cx, y, bw, h, PAPER, line_color=PRIMARY, line_w=1.25, rounded=True)
        text(s, cx + 0.2, y + 0.15, bw - 0.4, 0.4, lab, size=24, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
        if hint:
            text(s, cx + 0.2, y + 0.62, bw - 0.4, h - 0.8, hint, size=T_BODY_S, color=INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- vocab card
def vocab_card(s, x, y, w, h, term, definition, page, visual=None):
    """One vocabulary TEACH card (frozen component)."""
    rect(s, x, y, w, h, PAPER, line_color=RULE, line_w=1.25, rounded=True)
    # navy top band
    rect(s, x, y, w, 0.16, PRIMARY)
    # TEACH pill (top-right)
    pill = rect(s, x + w - 1.15, y + 0.5, 0.9, 0.4, NAVY_TINT, rounded=True)
    tf = pill.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, "TEACH", size=14, color=PRIMARY, bold=True)
    # label
    text(s, x + 0.35, y + 0.42, w - 1.7, 0.4, "VOCABULARY", size=18, color=NEUTRAL, bold=True, caps=True)
    # huge term
    text(s, x + 0.35, y + 0.78, w - 0.7, 1.05, term, size=54, color=PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    # definition
    text(s, x + 0.35, y + 1.95, w - 0.7, 0.95, definition, size=T_DEF, color=INK)
    # visual panel (gold-tinted)
    if visual:
        rect(s, x + 0.35, y + 3.0, w - 0.7, h - 3.7, GOLD_TINT, line_color=ACCENT, line_w=0.75, rounded=True)
        visual(s, x + 0.35, y + 3.0, w - 0.7, h - 3.7)
    # reader page
    text(s, x + 0.35, y + h - 0.55, w - 0.7, 0.4, page, size=T_PAGE, color=INK_SOFT, bold=True)

def card_visual_chain(panel_words):
    """Return a visual-draw function that renders a small arrow chain inside a panel."""
    def draw(s, x, y, w, h):
        n = len(panel_words)
        cw = (w - 0.5 * (n - 1) - 0.3) / n
        cx = x + 0.15
        cy = y + h / 2 - 0.28
        for i, word in enumerate(panel_words):
            sp = rect(s, cx, cy, cw, 0.56, PRIMARY if i == 0 else NEUTRAL, rounded=True)
            tf = sp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = Inches(0.05); tf.margin_right = Inches(0.05)
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            run(p, word, size=15, color=WHITE, bold=True)
            if i < n - 1:
                ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(cx + cw + 0.02),
                                        Inches(cy + 0.14), Inches(0.46), Inches(0.28))
                ar.fill.solid(); ar.fill.fore_color.rgb = ACCENT
                ar.line.fill.background(); ar.shadow.inherit = False
            cx += cw + 0.5
    return draw

# ---------------------------------------------------------------- notes mirror
def notes_heading(s, code, title_text, reader_ref):
    anchor_chip(s, code)
    text(s, 1.5, 1.02, 11.0, 0.6, title_text, size=T_HEAD, color=PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.5, 1.62, 8.0, 0.35, reader_ref, size=T_KICK, color=INK_SOFT, italic=True)

def write_line(s, x, y, w, color=INK_SOFT):
    rect(s, x, y, w, 0.035, color)

def set_note(slide, lines):
    slide.notes_slide.notes_text_frame.text = "\n".join(lines)

# ================================================================================
# SLIDE 1 — TITLE
# ================================================================================
s = new_slide(PRIMARY)
# chapter tab (inverted: gold-on-navy)
sp = rect(s, 0.9, 0.75, 4.4, 0.55, ACCENT, rounded=True)
tf = sp.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = Inches(0.15); tf.margin_right = Inches(0.15)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
run(p, "CHAPTER 1 · GRADE 4 · WEEK 1", size=14, color=WHITE, bold=True)
text(s, 0.9, 1.75, 11.5, 1.0, "Energy Causes Change", size=T_MAJOR, color=WHITE, bold=True)
text(s, 0.9, 2.85, 11.5, 0.5, "Discovery / Evidence  ·  Where do we see energy at work?", size=T_KICK, color=RGBColor(0xD8, 0xDC, 0xE4))
# small model teaser
model_strip(s, ["ENERGY", "CHANGE", "EVIDENCE"], x=0.9, y=4.2, w=7.5, box_h=0.7, gap=0.45, size=22)
text(s, 0.9, 5.9, 11.5, 0.5, "Unit 1 · Energy Transfer & Transformation", size=T_KICK, color=RGBColor(0xD8, 0xDC, 0xE4), italic=True)
slide_number(s)
set_note(s, ["MODE Title · orient", "TIME 2 min", "STUDENT RESOURCE Guided Notes (not yet open)",
             "ASK What do you think energy does?", "LISTEN FOR words like move, change, power",
             "NEXT Flip the lights as you open the Big Question."])

# ================================================================================
# SLIDE 2 — BIG QUESTION (intro)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s)
mode_badge(s, "BQ")
hud_bubble(s, "TURN & TALK")
text(s, 0.55, 1.1, 11.7, 0.5, "Big Question", size=T_KICK, color=NEUTRAL, bold=True, caps=True)
# navy BQ banner
rect(s, 0.55, 1.7, 12.23, 1.9, PRIMARY, rounded=True)
text(s, 0.9, 1.95, 11.5, 1.4, "Where can we observe evidence of energy causing change?",
     size=34, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.9, 4.0, 11.5, 0.5, "Keep it in mind all chapter. We will come back to it.", size=T_BODY_S, color=INK_SOFT, align=PP_ALIGN.CENTER)
text(s, 0.55, 4.9, 11.7, 0.6, "Think: What is ONE change you have seen happen?", size=T_BODY, color=INK, align=PP_ALIGN.CENTER, bold=True)
slide_number(s)
set_note(s, ["MODE BQ · introduce", "TIME 3 min", "STUDENT RESOURCE Guided Notes p.1 (BQ box)",
             "ASK What change have you seen happen today?", "LISTEN FOR any observable change (light, sound, movement)",
             "REVEAL Do not answer the BQ yet — we collect evidence first.", "NEXT My first thinking."])

# ================================================================================
# SLIDE 3 — BQ · MY FIRST THINKING (notes mirror)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "WRITE")
notes_heading(s, "BQ", "Big Question — My first thinking", "Guided Notes p. 1")
text(s, 1.5, 2.3, 11.0, 1.0, "Where can we observe evidence of energy causing change?", size=T_BODY, color=PRIMARY, bold=True)
text(s, 1.5, 3.5, 6.0, 0.55, "My first thinking", size=T_BODY, color=INK, bold=True)
write_line(s, 1.5, 4.1, 10.5)
write_line(s, 1.5, 4.8, 10.5)
text(s, 1.5, 5.55, 10.5, 0.5, "Write 1–2 quick ideas. No wrong answer yet.", size=T_BODY_S, color=INK_SOFT, italic=True)
slide_number(s)
set_note(s, ["MODE Notes · BQ anchor", "TIME 2 min", "STUDENT RESOURCE Guided Notes p.1 top",
             "ASK What change have you seen? Write your first idea.", "LISTEN FOR students naming a change (not 'energy')",
             "MISCONCEPTION Accept any change; do not correct yet.", "NEXT Observe the light switch."])

# ================================================================================
# SLIDE 4 — PHENOMENON: dark → light
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "D"); hud_bubble(s, "OBSERVE")
text(s, 0.55, 1.02, 11.7, 0.7, "What changed?", size=T_TITLE, color=PRIMARY, bold=True)
text(s, 0.55, 1.75, 11.7, 0.4, "Reader p. 1 — you walk into a dark room and flip a switch.", size=T_KICK, color=INK_SOFT, italic=True)
before_after(s, x=0.75, y=2.5, w=7.4, h=2.6,
             left="DARK room", right="BRIGHT room",
             left_hint="quiet · hard to see", right_hint="light fills the room")
text(s, 8.7, 2.5, 3.8, 2.6, "Something became different.\nThat is a change.", size=T_BODY, color=INK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
slide_number(s)
set_note(s, ["MODE Demo · observe", "TIME 3 min", "STUDENT RESOURCE Reader p.1",
             "ASK What changed when I flipped the switch?", "LISTEN FOR 'the room got bright' / 'light came on'",
             "REVEAL The room changed from dark to light — a change.", "NEXT Notice & wonder."])

# ================================================================================
# SLIDE 5 — NOTICE & WONDER
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "D"); hud_bubble(s, "TURN & TALK")
text(s, 0.55, 1.02, 11.7, 0.7, "I notice…  I wonder…", size=T_TITLE, color=PRIMARY, bold=True)
rect(s, 0.75, 2.2, 5.6, 3.2, PAPER, line_color=PRIMARY, line_w=1.25, rounded=True)
text(s, 1.05, 2.5, 5.0, 0.5, "I NOTICE", size=T_HEAD, color=ACCENT, bold=True, caps=True)
text(s, 1.05, 3.2, 5.0, 1.8, "The room is different now.\nLight is everywhere.", size=T_BODY_S, color=INK)
rect(s, 7.0, 2.2, 5.6, 3.2, PAPER, line_color=PRIMARY, line_w=1.25, rounded=True)
text(s, 7.3, 2.5, 5.0, 0.5, "I WONDER", size=T_HEAD, color=ACCENT, bold=True, caps=True)
text(s, 7.3, 3.2, 5.0, 1.8, "What made the light\nchange so fast?", size=T_BODY_S, color=INK)
text(s, 0.55, 5.8, 11.7, 0.5, "Share one notice and one wonder with a partner.", size=T_BODY, color=INK, align=PP_ALIGN.CENTER)
slide_number(s)
set_note(s, ["MODE Demo · discuss", "TIME 2 min", "STUDENT RESOURCE none",
             "ASK What did you notice? What do you wonder?", "LISTEN FOR change words and curiosity",
             "NEXT Electric pencil sharpener demo."])

# ================================================================================
# SLIDE 6 — DEMO: observe dull pencil
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "D"); hud_bubble(s, "OBSERVE")
text(s, 0.55, 1.02, 11.7, 0.7, "Electric Pencil Sharpener", size=T_TITLE, color=PRIMARY, bold=True)
text(s, 0.55, 1.75, 11.7, 0.4, "Observe the pencil BEFORE we sharpen it.", size=T_KICK, color=INK_SOFT, italic=True)
rect(s, 2.0, 2.5, 9.3, 3.3, PAPER, line_color=PRIMARY, line_w=1.25, rounded=True)
text(s, 2.3, 2.8, 8.7, 0.5, "BEFORE", size=T_HEAD, color=ACCENT, bold=True, caps=True)
text(s, 2.3, 3.5, 8.7, 1.6, "A dull pencil — tip is flat, and it is quiet.",
     size=T_BODY, color=INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.55, 6.15, 11.7, 0.5, "Look closely. What do you see? What do you hear?", size=T_BODY, color=INK, align=PP_ALIGN.CENTER)
slide_number(s)
set_note(s, ["MODE Demo · observe", "TIME 2 min", "STUDENT RESOURCE none (live demo)",
             "ASK What does the pencil look and sound like right now?", "LISTEN FOR dull, flat, quiet, still",
             "NEXT Predict before sharpening."])

# ================================================================================
# SLIDE 7 — DEMO: predict
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "D"); hud_bubble(s, "PREDICT")
text(s, 0.55, 1.02, 11.7, 0.7, "Predict: what will change?", size=T_TITLE, color=PRIMARY, bold=True)
text(s, 0.55, 1.75, 11.7, 0.4, "Think before we turn it on.", size=T_KICK, color=INK_SOFT, italic=True)
rect(s, 1.5, 2.6, 10.3, 2.6, PAPER, line_color=PRIMARY, line_w=1.25, rounded=True)
text(s, 1.8, 2.9, 9.7, 0.5, "When the sharpener runs, the pencil will…", size=T_BODY, color=INK, bold=True)
text(s, 1.8, 3.6, 9.7, 1.3, "get  ____________________", size=T_BODY, color=INK)
text(s, 0.55, 5.6, 11.7, 0.5, "Make a prediction. Say it to your partner.", size=T_BODY, color=INK, align=PP_ALIGN.CENTER)
slide_number(s)
set_note(s, ["MODE Demo · predict", "TIME 2 min", "STUDENT RESOURCE none",
             "ASK What will change about the pencil? The sharpener?", "LISTEN FOR shorter, sharper, noise",
             "REVEAL Do not confirm yet — run the demo.", "NEXT Observe the change."])

# ================================================================================
# SLIDE 8 — DEMO: what changed (before/after)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "D"); hud_bubble(s, "OBSERVE")
text(s, 0.55, 1.02, 11.7, 0.7, "What changed?", size=T_TITLE, color=PRIMARY, bold=True)
before_after(s, x=0.75, y=2.4, w=7.4, h=3.0,
             left="BEFORE", right="AFTER",
             left_hint="dull tip · quiet", right_hint="shorter · sharper · humming")
rect(s, 8.7, 2.4, 3.8, 3.0, GOLD_TINT, line_color=ACCENT, line_w=0.75, rounded=True)
text(s, 8.95, 2.6, 3.3, 0.5, "CHANGES", size=20, color=NEUTRAL, bold=True, caps=True)
text(s, 8.95, 3.2, 3.3, 2.0, "• the tip is sharper\n• the pencil is shorter\n• it makes a sound", size=T_BODY_S, color=INK)
slide_number(s)
set_note(s, ["MODE Demo · observe", "TIME 3 min", "STUDENT RESOURCE none",
             "ASK Name every change you saw or heard.", "LISTEN FOR shorter, sharper, noise",
             "REVEAL Record all three changes.", "NEXT Turn changes into evidence."])

# ================================================================================
# SLIDE 9 — DEMO: evidence
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "D"); hud_bubble(s, "EVIDENCE")
text(s, 0.55, 1.02, 11.7, 0.7, "How do we know energy was involved?", size=T_TITLE, color=PRIMARY, bold=True)
text(s, 0.55, 1.75, 11.7, 0.4, "Energy can be invisible. We look for the change.", size=T_KICK, color=INK_SOFT, italic=True)
model_strip(s, ["ENERGY", "CHANGE", "EVIDENCE"], x=0.75, y=2.7, w=11.8, box_h=1.0, gap=0.5, size=28)
text(s, 0.55, 4.3, 11.7, 1.4,
     "The pencil got shorter and sharper and made a sound.\nThose changes are evidence that energy was involved.",
     size=T_BODY, color=INK, align=PP_ALIGN.CENTER)
slide_number(s)
set_note(s, ["MODE Demo · evidence", "TIME 3 min", "STUDENT RESOURCE Reader p.2 (Change Is Evidence of Energy)",
             "ASK How do we know energy was here if we can't see it?", "LISTEN FOR 'we saw the change'",
             "MISCONCEPTION We see the change, not the energy itself.", "NEXT Vocabulary: energy + change."])

# ================================================================================
# SLIDE 10 — V1 VOCABULARY CARDS (energy + change)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "V")
text(s, 0.55, 1.02, 11.7, 0.7, "Vocabulary — Energy & Change", size=T_TITLE, color=PRIMARY, bold=True)
text(s, 0.55, 1.72, 11.7, 0.4, "Reader p. 1", size=T_KICK, color=INK_SOFT, italic=True)
cw = 5.95
vocab_card(s, 0.55, 2.25, cw, 4.5, "energy",
           "the ability to cause change",
           "Reader p. 1",
           visual=card_visual_chain(["ENERGY", "CHANGE"]))
vocab_card(s, 0.55 + cw + 0.33, 2.25, cw, 4.5, "change",
           "to become different",
           "Reader p. 1",
           visual=card_visual_chain(["BEFORE", "AFTER"]))
slide_number(s)
set_note(s, ["MODE Vocabulary · TEACH", "TIME 4 min", "STUDENT RESOURCE Guided Notes V1",
             "ASK Read each definition together. Use each word in a sentence.",
             "LISTEN FOR 'energy causes change' and 'to become different'",
             "REVEAL Full definitions shown — these are the two V1 terms.", "NEXT Record V1 on notes."])

# ================================================================================
# SLIDE 11 — V1 NOTES (write)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "WRITE")
notes_heading(s, "V1", "Vocabulary — Energy & Change", "Guided Notes · Reader p. 1")
rich(s, 1.5, 2.6, 11.0, 1.0, [
    [("energy — the ability to cause ", {}), (blank_str(7), {"underline": True, "color": INK}), (" .", {})],
])
rich(s, 1.5, 3.6, 11.0, 1.0, [
    [("change — to become ", {}), (blank_str(9), {"underline": True, "color": INK}), (" .", {})],
])
text(s, 1.5, 4.9, 11.0, 0.5, "Write the missing word in each blank.", size=T_BODY_S, color=INK_SOFT, italic=True)
slide_number(s)
set_note(s, ["MODE Notes · V1 anchor", "TIME 3 min", "STUDENT RESOURCE Guided Notes V1 (Reader p.1)",
             "ASK Fill in each blank.", "LISTEN FOR 'change' and 'different'",
             "REVEAL Reveal after students write.", "NEXT Check answers."])

# ================================================================================
# SLIDE 12 — V1 NOTES (check/reveal)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "CHECK")
notes_heading(s, "V1", "Vocabulary — Energy & Change", "Guided Notes · Reader p. 1")
rich(s, 1.5, 2.6, 11.0, 1.0, [
    [("energy — the ability to cause ", {}), ("change", {"bold": True, "color": ACCENT}), (" .", {})],
])
rich(s, 1.5, 3.6, 11.0, 1.0, [
    [("change — to become ", {}), ("different", {"bold": True, "color": ACCENT}), (" .", {})],
])
text(s, 1.5, 4.9, 11.0, 0.5, "change  ·  different", size=T_BODY_S, color=INK_SOFT, italic=True)
slide_number(s)
set_note(s, ["MODE Notes · V1 reveal", "TIME 1 min", "STUDENT RESOURCE Guided Notes V1",
             "REVEAL change · different", "NEXT Energy Causes Change (N1)."])

# ================================================================================
# SLIDE 13 — N1 CONCEPT: energy causes change
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "R"); hud_bubble(s, "RECITE")
text(s, 0.55, 1.02, 11.7, 0.7, "What is energy?", size=T_TITLE, color=PRIMARY, bold=True)
text(s, 0.55, 1.75, 11.7, 0.4, "Reader pp. 1–2", size=T_KICK, color=INK_SOFT, italic=True)
rect(s, 0.75, 2.5, 11.8, 2.2, PRIMARY, rounded=True)
text(s, 1.2, 2.9, 10.9, 1.4, "Energy is the ability to cause change.", size=38, color=WHITE, bold=True,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.75, 5.0, 11.8, 1.0, "Say it together: “Energy is the ability to cause change.”", size=T_BODY, color=INK, align=PP_ALIGN.CENTER)
slide_number(s)
set_note(s, ["MODE Read · concept", "TIME 3 min", "STUDENT RESOURCE Reader p.1",
             "ASK What is energy?", "LISTEN FOR 'the ability to cause change'",
             "REVEAL Energy is the ability to cause change.", "NEXT Change is evidence."])

# ================================================================================
# SLIDE 14 — N1 MODEL: ENERGY → CHANGE → EVIDENCE
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "R"); hud_bubble(s, "TRACE")
text(s, 0.55, 1.02, 11.7, 0.7, "Change is evidence of energy", size=T_TITLE, color=PRIMARY, bold=True)
text(s, 0.55, 1.75, 11.7, 0.4, "Reader p. 2 — the seashore changes all around us.", size=T_KICK, color=INK_SOFT, italic=True)
model_strip(s, ["ENERGY", "CHANGE", "EVIDENCE"], x=0.75, y=2.7, w=11.8, box_h=1.1, gap=0.5, size=30)
text(s, 0.75, 4.3, 11.8, 1.6,
     "When you see a change — water moving, sand shifting, sound —\nyou are seeing evidence of energy.",
     size=T_BODY, color=INK, align=PP_ALIGN.CENTER)
slide_number(s)
set_note(s, ["MODE Read · model", "TIME 3 min", "STUDENT RESOURCE Reader p.2",
             "ASK Trace the arrow: energy causes a change, the change is evidence.",
             "LISTEN FOR students using 'evidence' not 'proof'", "NEXT Record N1 on notes."])

# ================================================================================
# SLIDE 15 — N1 NOTES (write)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "WRITE")
notes_heading(s, "N1", "Energy Causes Change", "Guided Notes · Reader pp. 1–2")
rich(s, 1.5, 2.35, 11.0, 0.8, [
    [("Energy has the ability to ", {}), (blank_str(6), {"underline": True, "color": INK}), ("  ", {}),
     (blank_str(8), {"underline": True, "color": INK}), (" .", {})],
])
text(s, 1.5, 3.25, 11.0, 0.5, "When you see a change, you are seeing evidence of energy.", size=T_BODY_S, color=INK, italic=True)
model_strip(s, ["ENERGY", "CHANGE", "EVIDENCE"], x=1.5, y=4.0, w=8.0, box_h=0.7, gap=0.4, size=20)
rich(s, 1.5, 5.1, 11.0, 0.8, [
    [("A change is ", {}), (blank_str(9), {"underline": True, "color": INK}), (" that energy is involved.", {})],
])
slide_number(s)
set_note(s, ["MODE Notes · N1 anchor", "TIME 4 min", "STUDENT RESOURCE Guided Notes N1 (Reader pp.1–2)",
             "ASK Fill in the three blanks.", "LISTEN FOR cause · change · evidence",
             "REVEAL Reveal after writing.", "NEXT Check answers."])

# ================================================================================
# SLIDE 16 — N1 NOTES (check/reveal)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "CHECK")
notes_heading(s, "N1", "Energy Causes Change", "Guided Notes · Reader pp. 1–2")
rich(s, 1.5, 2.35, 11.0, 0.8, [
    [("Energy has the ability to ", {}), ("cause", {"bold": True, "color": ACCENT}), ("  ", {}),
     ("change", {"bold": True, "color": ACCENT}), (" .", {})],
])
text(s, 1.5, 3.25, 11.0, 0.5, "When you see a change, you are seeing evidence of energy.", size=T_BODY_S, color=INK, italic=True)
model_strip(s, ["ENERGY", "CHANGE", "EVIDENCE"], x=1.5, y=4.0, w=8.0, box_h=0.7, gap=0.4, size=20)
rich(s, 1.5, 5.1, 11.0, 0.8, [
    [("A change is ", {}), ("evidence", {"bold": True, "color": ACCENT}), (" that energy is involved.", {})],
])
text(s, 1.5, 6.05, 11.0, 0.4, "cause · change · evidence", size=T_BODY_S, color=INK_SOFT, italic=True)
slide_number(s)
set_note(s, ["MODE Notes · N1 reveal", "TIME 1 min", "STUDENT RESOURCE Guided Notes N1",
             "REVEAL cause · change · evidence", "NEXT Forms of energy (N2)."])

# ================================================================================
# SLIDE 17 — N2 CONCEPT: energy has forms
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "R")
text(s, 0.55, 1.02, 11.7, 0.7, "Energy has different forms", size=T_TITLE, color=PRIMARY, bold=True)
text(s, 0.55, 1.75, 11.7, 0.4, "Reader p. 3 — think of a spinning carousel.", size=T_KICK, color=INK_SOFT, italic=True)
forms = ["motion", "sound", "light", "heat", "electrical energy"]
cw = 2.2; gap = 0.3; total = 5*cw + 4*gap
x0 = 0.55 + (12.23 - total)/2
for i, f in enumerate(forms):
    cx = x0 + i*(cw+gap)
    sp = rect(s, cx, 2.7, cw, 1.2, PRIMARY, rounded=True)
    tf = sp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.08)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, f, size=20, color=WHITE, bold=True)
text(s, 0.55, 4.4, 11.7, 1.4,
     "The carousel spins (motion), plays music (sound), shines bright bulbs\n(light), warms those bulbs (heat), and plugs in (electrical energy).",
     size=T_BODY_S, color=INK, align=PP_ALIGN.CENTER)
slide_number(s)
set_note(s, ["MODE Read · forms", "TIME 4 min", "STUDENT RESOURCE Reader p.3",
             "ASK What forms of energy do you see or hear on a carousel?", "LISTEN FOR motion, sound, light, heat, electrical",
             "NEXT Record examples on notes (N2)."])

# ================================================================================
# SLIDE 18 — N2 FORMS: clap check
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "C"); hud_bubble(s, "OBSERVE")
text(s, 0.55, 1.02, 11.7, 0.7, "Clap three times", size=T_TITLE, color=PRIMARY, bold=True)
text(s, 0.55, 1.75, 11.7, 0.4, "What forms of energy do you observe?", size=T_KICK, color=INK_SOFT, italic=True)
rect(s, 1.5, 2.5, 10.3, 3.2, PAPER, line_color=PRIMARY, line_w=1.25, rounded=True)
text(s, 1.8, 2.85, 9.7, 0.5, "When you clap…", size=T_HEAD, color=PRIMARY, bold=True)
text(s, 1.8, 3.6, 9.7, 1.7, "motion  —  your hands move\nsound  —  you hear a clap\nheat  —  your hands may feel warm", size=T_BODY, color=INK)
slide_number(s)
set_note(s, ["MODE Check · forms", "TIME 2 min", "STUDENT RESOURCE Reader p.3",
             "ASK What forms of energy can you observe when you clap?",
             "LISTEN FOR motion, sound, and maybe heat", "NEXT Record N2 table."])

# ================================================================================
# SLIDE 19 — N2 NOTES (table)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "WRITE")
notes_heading(s, "N2", "Energy Occurs in Different Forms", "Guided Notes · Reader p. 3")
text(s, 1.5, 2.3, 11.0, 0.5, "Write or draw one example that shows a change for each form.", size=T_BODY_S, color=INK, italic=True)
# table header
hx, hy, hw, hh = 1.5, 2.95, 3.4, 0.6
ex, ew = hx + hw + 0.15, 11.0 - hw - 0.15
rect(s, hx, hy, hw, hh, PRIMARY, rounded=False)
text(s, hx+0.15, hy+0.12, hw-0.3, 0.4, "FORM", size=20, color=WHITE, bold=True)
rect(s, ex, hy, ew, hh, PRIMARY, rounded=False)
text(s, ex+0.15, hy+0.12, ew-0.3, 0.4, "ONE EXAMPLE THAT SHOWS A CHANGE", size=20, color=WHITE, bold=True)
rows = ["motion", "sound", "light", "heat", "electrical energy"]
ry = hy + hh
for i, rname in enumerate(rows):
    rh = 0.62
    cy = ry + i*rh
    rect(s, hx, cy, hw, rh, PAPER, line_color=RULE, line_w=0.75)
    text(s, hx+0.15, cy+0.13, hw-0.3, 0.4, rname, size=20, color=PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, ex, cy, ew, rh, PAPER, line_color=RULE, line_w=0.75)
slide_number(s)
set_note(s, ["MODE Notes · N2 anchor", "TIME 5 min", "STUDENT RESOURCE Guided Notes N2 (Reader p.3)",
             "ASK Write or draw one example per form.", "LISTEN FOR carousel examples or clap examples",
             "NEXT Stored energy (V2)."])

# ================================================================================
# SLIDE 20 — V2 VOCABULARY CARD (stored energy)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "V")
text(s, 0.55, 1.02, 11.7, 0.7, "Vocabulary — Stored Energy", size=T_TITLE, color=PRIMARY, bold=True)
text(s, 0.55, 1.72, 11.7, 0.4, "Reader p. 4", size=T_KICK, color=INK_SOFT, italic=True)
vocab_card(s, 3.0, 2.25, 7.3, 4.5, "stored energy",
           "energy that has the ability to cause change at a later time",
           "Reader p. 4",
           visual=card_visual_chain(["STILL", "LATER", "CHANGE"]))
slide_number(s)
set_note(s, ["MODE Vocabulary · TEACH", "TIME 3 min", "STUDENT RESOURCE Guided Notes V2 (Reader p.4)",
             "ASK Does a battery have energy if it is sitting still?", "LISTEN FOR 'yes — it can cause change later'",
             "REVEAL stored energy = ability to cause change at a later time.", "NEXT Record V2."])

# ================================================================================
# SLIDE 21 — V2 NOTES (write)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "WRITE")
notes_heading(s, "V2", "Vocabulary — Stored Energy", "Guided Notes · Reader p. 4")
rich(s, 1.5, 2.65, 11.0, 1.15, [
    [("stored energy — energy that has the ability to cause change at a ", {}),
     (blank_str(6), {"underline": True, "color": INK}), (" time.", {})],
])
text(s, 1.5, 4.15, 11.0, 0.5, "Write the missing word in the blank.", size=T_BODY_S, color=INK_SOFT, italic=True)
slide_number(s)
set_note(s, ["MODE Notes · V2 anchor", "TIME 2 min", "STUDENT RESOURCE Guided Notes V2 (Reader p.4)",
             "ASK Fill in the blank.", "LISTEN FOR 'later'", "REVEAL Reveal after writing.", "NEXT Check answer."])

# ================================================================================
# SLIDE 22 — V2 NOTES (check/reveal)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "CHECK")
notes_heading(s, "V2", "Vocabulary — Stored Energy", "Guided Notes · Reader p. 4")
rich(s, 1.5, 2.65, 11.0, 1.15, [
    [("stored energy — energy that has the ability to cause change at a ", {}),
     ("later", {"bold": True, "color": ACCENT}), (" time.", {})],
])
text(s, 1.5, 4.15, 11.0, 0.5, "later", size=T_BODY_S, color=INK_SOFT, italic=True)
slide_number(s)
set_note(s, ["MODE Notes · V2 reveal", "TIME 1 min", "STUDENT RESOURCE Guided Notes V2",
             "REVEAL later", "NEXT Objects can store energy (N3)."])

# ================================================================================
# SLIDE 23 — N3 CONCEPT: still objects can store energy (predict)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "R"); hud_bubble(s, "PREDICT")
text(s, 0.55, 1.02, 11.7, 0.7, "Does a still battery have energy?", size=T_TITLE, color=PRIMARY, bold=True)
text(s, 0.55, 1.75, 11.7, 0.4, "Reader p. 4 — a battery sits still and shows no change… yet.", size=T_KICK, color=INK_SOFT, italic=True)
rect(s, 1.5, 2.6, 10.3, 3.0, PAPER, line_color=PRIMARY, line_w=1.25, rounded=True)
text(s, 1.8, 2.95, 9.7, 0.5, "A battery sitting still shows no change — yet.", size=T_BODY, color=INK, bold=True)
text(s, 1.8, 3.75, 9.7, 1.4,
     "Inside a battery, chemicals store energy.\nWhen it goes in a flashlight, the stored energy is released.",
     size=T_BODY_S, color=INK)
slide_number(s)
set_note(s, ["MODE Read · predict", "TIME 3 min", "STUDENT RESOURCE Reader p.4",
             "ASK Does a still battery have energy?", "LISTEN FOR 'yes, stored energy'",
             "MISCONCEPTION Only moving objects have energy → pair moving + stored.", "NEXT Position."])

# ================================================================================
# SLIDE 24 — N3 CONCEPT: stored energy of position
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "R")
text(s, 0.55, 1.02, 11.7, 0.7, "The stored energy of position", size=T_TITLE, color=PRIMARY, bold=True)
text(s, 0.55, 1.75, 11.7, 0.4, "Reader p. 5 — an object held up can fall and cause a change.", size=T_KICK, color=INK_SOFT, italic=True)
rect(s, 1.5, 2.5, 10.3, 3.0, GOLD_TINT, line_color=ACCENT, line_w=0.75, rounded=True)
text(s, 1.8, 2.85, 9.7, 0.5, "Raised apple → let go → it falls", size=T_HEAD, color=PRIMARY, bold=True)
text(s, 1.8, 3.6, 9.7, 1.4,
     "An apple held above the ground has stored energy\nbecause of its position. Let it go and it will fall.",
     size=T_BODY_S, color=INK)
slide_number(s)
set_note(s, ["MODE Read · position", "TIME 3 min", "STUDENT RESOURCE Reader p.5",
             "ASK Why does a raised apple have stored energy?", "LISTEN FOR 'because of its position'",
             "NEXT Build the stored-energy model."])

# ================================================================================
# SLIDE 25 — N3 MODEL: STORED ENERGY → RELEASE → CHANGE
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "R"); hud_bubble(s, "TRACE")
text(s, 0.55, 1.02, 11.7, 0.7, "Stored energy can cause a change later", size=T_TITLE, color=PRIMARY, bold=True)
model_strip(s, ["STORED ENERGY", "RELEASE", "CHANGE"], x=0.75, y=2.7, w=11.8, box_h=1.1, gap=0.5, size=28)
text(s, 0.75, 4.3, 11.8, 1.6,
     "Stretch a rubber band — it now stores energy.\nLet it go — it moves and makes a sound: the change.",
     size=T_BODY, color=INK, align=PP_ALIGN.CENTER)
slide_number(s)
set_note(s, ["MODE Read · model", "TIME 2 min", "STUDENT RESOURCE Reader p.4",
             "ASK Trace: stored → release → change.", "LISTEN FOR rubber band SNAP example",
             "NEXT Record N3 on notes."])

# ================================================================================
# SLIDE 26 — N3 NOTES (write)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "WRITE")
notes_heading(s, "N3", "Objects Can Store Energy", "Guided Notes · Reader pp. 4–5")
model_strip(s, ["STORED ENERGY", "RELEASE", "CHANGE"], x=1.5, y=2.15, w=8.0, box_h=0.55, gap=0.4, size=18)
# table
hx, hy, hw, hh = 1.5, 2.85, 4.2, 0.5
ex, ew = hx + hw + 0.15, 11.0 - hw - 0.15
rect(s, hx, hy, hw, hh, PRIMARY)
text(s, hx+0.12, hy+0.08, hw-0.24, 0.35, "EXAMPLE", size=18, color=WHITE, bold=True)
rect(s, ex, hy, ew, hh, PRIMARY)
text(s, ex+0.12, hy+0.08, ew-0.24, 0.35, "WHAT CHANGE COULD HAPPEN LATER?", size=18, color=WHITE, bold=True)
rows = ["battery (in a flashlight)", "stretched rubber band", "raised object (apple)"]
ry = hy + hh
for i, rname in enumerate(rows):
    rh = 0.55
    cy = ry + i*rh
    rect(s, hx, cy, hw, rh, PAPER, line_color=RULE, line_w=0.75)
    text(s, hx+0.12, cy+0.13, hw-0.24, 0.4, rname, size=18, color=PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, ex, cy, ew, rh, PAPER, line_color=RULE, line_w=0.75)
rich(s, 1.5, 5.2, 11.0, 0.7, [
    [("Objects also have stored energy because of their ", {"size": T_BODY_S}),
     (blank_str(8), {"underline": True, "color": INK, "size": T_BODY_S}), (" .", {"size": T_BODY_S})],
])
slide_number(s)
set_note(s, ["MODE Notes · N3 anchor", "TIME 5 min", "STUDENT RESOURCE Guided Notes N3 (Reader pp.4–5)",
             "ASK Fill in the blank; note the three examples.", "LISTEN FOR 'position'",
             "REVEAL Reveal after writing.", "NEXT Check answer."])

# ================================================================================
# SLIDE 27 — N3 NOTES (check/reveal)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "CHECK")
notes_heading(s, "N3", "Objects Can Store Energy", "Guided Notes · Reader pp. 4–5")
model_strip(s, ["STORED ENERGY", "RELEASE", "CHANGE"], x=1.5, y=2.15, w=8.0, box_h=0.55, gap=0.4, size=18)
hx, hy, hw, hh = 1.5, 2.85, 4.2, 0.5
ex, ew = hx + hw + 0.15, 11.0 - hw - 0.15
rect(s, hx, hy, hw, hh, PRIMARY)
text(s, hx+0.12, hy+0.08, hw-0.24, 0.35, "EXAMPLE", size=18, color=WHITE, bold=True)
rect(s, ex, hy, ew, hh, PRIMARY)
text(s, ex+0.12, hy+0.08, ew-0.24, 0.35, "WHAT CHANGE COULD HAPPEN LATER?", size=18, color=WHITE, bold=True)
rows = ["battery (in a flashlight)", "stretched rubber band", "raised object (apple)"]
ry = hy + hh
for i, rname in enumerate(rows):
    rh = 0.55
    cy = ry + i*rh
    rect(s, hx, cy, hw, rh, PAPER, line_color=RULE, line_w=0.75)
    text(s, hx+0.12, cy+0.13, hw-0.24, 0.4, rname, size=18, color=PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, ex, cy, ew, rh, PAPER, line_color=RULE, line_w=0.75)
rich(s, 1.5, 5.2, 11.0, 0.7, [
    [("Objects also have stored energy because of their ", {"size": T_BODY_S}),
     ("position", {"bold": True, "color": ACCENT, "size": T_BODY_S}), (" .", {"size": T_BODY_S})],
])
text(s, 1.5, 6.15, 11.0, 0.4, "position", size=T_BODY_S, color=INK_SOFT, italic=True)
slide_number(s)
set_note(s, ["MODE Notes · N3 reveal", "TIME 1 min", "STUDENT RESOURCE Guided Notes N3",
             "REVEAL position", "NEXT Recognizing evidence (N4)."])

# ================================================================================
# SLIDE 28 — N4 CONCEPT: recognizing evidence
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "R"); hud_bubble(s, "EVIDENCE")
text(s, 0.55, 1.02, 11.7, 0.7, "Recognizing evidence", size=T_TITLE, color=PRIMARY, bold=True)
text(s, 0.55, 1.75, 11.7, 0.4, "Reader pp. 1–6 — we see the change, not the energy.", size=T_KICK, color=INK_SOFT, italic=True)
rect(s, 0.75, 2.5, 11.8, 2.6, PRIMARY, rounded=True)
text(s, 1.2, 2.85, 10.9, 1.9,
     "Energy can be invisible.\nWe recognize energy by the change we observe.",
     size=T_BODY, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.75, 5.4, 11.8, 0.7, "Evidence = the change you see, hear, or feel.", size=T_BODY_S, color=INK, align=PP_ALIGN.CENTER, italic=True)
slide_number(s)
set_note(s, ["MODE Read · evidence", "TIME 2 min", "STUDENT RESOURCE Reader pp.1–6",
             "ASK How do we know energy is involved when we can't see it?", "LISTEN FOR 'we see the change'",
             "MISCONCEPTION Evidence, not proof.", "NEXT Draw before/after (N4)."])

# ================================================================================
# SLIDE 29 — N4 NOTES (before/after draw)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "DRAW")
notes_heading(s, "N4", "Recognizing Evidence", "Guided Notes · Reader pp. 1–6")
text(s, 1.5, 2.25, 11.0, 1.0, "Draw a change you observed. Show what it looked like BEFORE and AFTER.", size=T_BODY_S, color=INK, italic=True)
before_after(s, x=1.5, y=3.3, w=8.0, h=1.9, left="BEFORE", right="AFTER", left_hint="draw", right_hint="draw")
rich(s, 1.5, 5.4, 11.0, 0.5, [
    [("What changed?  ", {"bold": True}), (blank_str(16), {"underline": True, "color": INK})],
])
slide_number(s)
set_note(s, ["MODE Notes · N4 anchor", "TIME 4 min", "STUDENT RESOURCE Guided Notes N4 (Reader pp.1–6)",
             "ASK Draw one change before and after; write what changed.",
             "LISTEN FOR clear before/after pairs", "NEXT Check your thinking."])

# ================================================================================
# SLIDE 30 — CHECK YOUR THINKING (write)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "WRITE")
notes_heading(s, "N4", "Check Your Thinking", "Guided Notes · Reader pp. 1–6")
rect(s, 1.5, 2.4, 11.0, 1.7, GOLD_TINT, line_color=ACCENT, line_w=0.75, rounded=True)
text(s, 1.8, 2.55, 10.4, 0.45, "CHECK YOUR THINKING", size=22, color=NEUTRAL, bold=True, caps=True)
text(s, 1.8, 3.05, 10.4, 1.05, "How do we know energy is involved if we cannot see energy itself?", size=T_BODY, color=INK, bold=True)
text(s, 1.5, 4.35, 11.0, 0.5, "Start with: “We know because we see …”", size=T_BODY, color=INK, italic=True)
write_line(s, 1.5, 5.15, 10.5)
write_line(s, 1.5, 5.85, 10.5)
slide_number(s)
set_note(s, ["MODE Notes · Check Your Thinking", "TIME 3 min", "STUDENT RESOURCE Guided Notes N4",
             "ASK Complete the sentence starter.", "LISTEN FOR 'a change, like a ball starting to move'",
             "NEXT Reveal a sample response."])

# ================================================================================
# SLIDE 31 — CHECK YOUR THINKING (sample + classify)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "C"); hud_bubble(s, "CHECK")
text(s, 0.55, 1.02, 11.7, 0.7, "We know because we see a change", size=T_TITLE, color=PRIMARY, bold=True)
rect(s, 1.5, 2.4, 10.3, 2.2, GOLD_TINT, line_color=ACCENT, line_w=0.75, rounded=True)
text(s, 1.8, 2.6, 9.7, 0.5, "SAMPLE", size=18, color=NEUTRAL, bold=True, caps=True)
text(s, 1.8, 3.15, 9.7, 1.3,
     "“We know because we see a change — like a ball starting to move or a light turning on.”",
     size=T_BODY_S, color=INK, italic=True)
text(s, 1.5, 4.9, 10.3, 0.6, "Any change you observe is evidence of energy.", size=T_BODY_S, color=INK, align=PP_ALIGN.CENTER)
slide_number(s)
set_note(s, ["MODE Check · evidence", "TIME 2 min", "STUDENT RESOURCE Guided Notes N4",
             "REVEAL Sample response (accept any change-as-evidence).", "NEXT Apply: classify evidence."])

# ================================================================================
# SLIDE 32 — APPLICATION: which change shows energy?
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "C"); hud_bubble(s, "EVIDENCE")
text(s, 0.55, 1.02, 11.7, 0.7, "Which change shows energy at work?", size=T_TITLE, color=PRIMARY, bold=True)
opts = [
    ("A toy car starts to roll.", True),
    ("A rock keeps sitting still.", False),
    ("A room stays dark.", False),
]
y = 2.5
for i, (txt, is_yes) in enumerate(opts):
    cy = y + i * 1.25
    rect(s, 1.5, cy, 10.3, 1.0, PAPER, line_color=PRIMARY, line_w=1.25, rounded=True)
    text(s, 1.8, cy + 0.22, 9.7, 0.6, txt, size=T_BODY, color=INK, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.55, 6.4, 11.7, 0.5, "Look for a change — that is the evidence of energy.", size=T_BODY_S, color=INK_SOFT, align=PP_ALIGN.CENTER, italic=True)
slide_number(s)
set_note(s, ["MODE Apply · evidence", "TIME 3 min", "STUDENT RESOURCE none",
             "ASK Which one shows a change (energy at work)?", "LISTEN FOR 'the toy car'",
             "REVEAL The toy car starts to roll — a change.", "NEXT Return to the Big Question."])

# ================================================================================
# SLIDE 33 — BQ · NOW I THINK (return)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "WRITE")
notes_heading(s, "BQ", "BQ · Now I Think", "Guided Notes p. 3")
rect(s, 1.5, 2.3, 11.0, 1.75, PRIMARY, rounded=True)
text(s, 1.8, 2.45, 10.4, 0.4, "BIG QUESTION", size=18, color=ACCENT, bold=True, caps=True)
text(s, 1.8, 2.9, 10.4, 1.05, "Where can we observe evidence of energy causing change?", size=T_BODY, color=WHITE, bold=True)
text(s, 1.5, 4.3, 11.0, 0.5, "Now I think…", size=T_BODY, color=INK, bold=True)
write_line(s, 1.5, 5.0, 10.5)
write_line(s, 1.5, 5.7, 10.5)
write_line(s, 1.5, 6.4, 10.5)
slide_number(s)
set_note(s, ["MODE Notes · BQ return", "TIME 3 min", "STUDENT RESOURCE Guided Notes p.3",
             "ASK Answer the Big Question now.", "LISTEN FOR any change named as evidence",
             "NEXT End-of-chapter review."])

# ================================================================================
# END-OF-CHAPTER REVIEW (4–6 MCQ, question state + answer state)
# ================================================================================
REVIEW = [
    {"q": "Energy is best described as —",
     "opts": ["the ability to cause change", "a visible glow", "something only batteries have", "a kind of fuel"],
     "correct": 0, "type": "concept",
     "explain": "Energy is the ability to cause change (Reader p. 1)."},
    {"q": "A pencil is dull, then you sharpen it. What change do you observe?",
     "opts": ["It gets longer", "It gets shorter and sharper", "It changes color", "It disappears"],
     "correct": 1, "type": "visual",
     "explain": "The pencil becomes shorter and sharper — a change."},
    {"q": "Which object has stored energy that can cause a change LATER?",
     "opts": ["A ball already rolling", "A toy already moving", "A stretched rubber band", "A flashlight already on"],
     "correct": 2, "type": "scenario",
     "explain": "The rubber band stores energy; it changes when released."},
    {"q": "A carousel spins and plays music. Which two forms of energy do you observe?",
     "opts": ["motion and silence", "sound and stillness", "only light", "motion and sound"],
     "correct": 3, "type": "comparison",
     "explain": "The spinning is motion; the music is sound (Reader p. 3)."},
    {"q": "How do we know energy is involved if we cannot see it?",
     "opts": ["Energy can always be seen", "We observe a change happen", "We guess without evidence", "Nothing can change"],
     "correct": 1, "type": "concept",
     "explain": "We see a change — like movement — as evidence."},
    {"q": "A battery sits still in a flashlight. Does it have energy?",
     "opts": ["No — still objects have no energy", "Only when it is moving", "Yes — stored energy that can cause change later", "Only when the flashlight is on"],
     "correct": 2, "type": "scenario",
     "explain": "A battery stores energy; it changes when released."},
]

def review_question(slide_q, num, item):
    chapter_tab(slide_q); mode_badge(slide_q, "C")
    text(slide_q, 0.55, 1.02, 11.7, 0.7, f"Review {num} — {item['type']}", size=T_HEAD, color=PRIMARY, bold=True)
    text(slide_q, 0.55, 1.66, 11.7, 1.05, item["q"], size=T_BODY, color=INK, bold=True)
    y = 2.78
    letters = ["A", "B", "C", "D"]
    for i, opt in enumerate(item["opts"]):
        cy = y + i * 0.88
        rect(slide_q, 1.0, cy, 11.3, 0.74, PAPER, line_color=RULE, line_w=1.0, rounded=True)
        text(slide_q, 1.25, cy + 0.13, 0.6, 0.5, letters[i], size=T_OPT, color=ACCENT, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(slide_q, 1.95, cy + 0.11, 10.1, 0.54, opt, size=T_OPT, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    slide_number(slide_q)

def review_answer(slide_a, num, item):
    chapter_tab(slide_a); mode_badge(slide_a, "C"); hud_bubble(slide_a, "CHECK")
    text(slide_a, 0.55, 1.02, 11.7, 0.7, f"Review {num} — answer", size=T_HEAD, color=PRIMARY, bold=True)
    text(slide_a, 0.55, 1.66, 11.7, 1.05, item["q"], size=T_BODY, color=INK, bold=True)
    y = 2.78
    letters = ["A", "B", "C", "D"]
    for i, opt in enumerate(item["opts"]):
        cy = y + i * 0.88
        is_correct = (i == item["correct"])
        fill = GOLD_TINT if is_correct else PAPER
        line_c = ACCENT if is_correct else RULE
        rect(slide_a, 1.0, cy, 11.3, 0.74, fill, line_color=line_c, line_w=1.5 if is_correct else 1.0, rounded=True)
        text(slide_a, 1.25, cy + 0.13, 0.6, 0.5, letters[i], size=T_OPT, color=ACCENT, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(slide_a, 1.95, cy + 0.11, 10.1, 0.54, opt, size=T_OPT, color=INK, bold=is_correct, anchor=MSO_ANCHOR.MIDDLE)
    rect(slide_a, 0.55, 6.35, 12.23, 0.72, MIST, line_color=RULE, line_w=1.0, rounded=True)
    text(slide_a, 0.8, 6.47, 11.7, 0.5, item["explain"], size=T_BODY_S, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    slide_number(slide_a)

for idx, item in enumerate(REVIEW, 1):
    sq = new_slide(SURFACE)
    review_question(sq, idx, item)
    set_note(sq, [f"MODE Check · review Q{idx}", "STUDENT RESOURCE none",
                  "ASK Read aloud; students vote A/B/C/D.", "REVEAL Do not reveal yet.",
                  "NEXT Reveal the answer."])
    sa = new_slide(SURFACE)
    review_answer(sa, idx, item)
    set_note(sa, [f"MODE Check · review Q{idx} answer", "REVEAL " + REVIEW[idx-1]["opts"][item["correct"]],
                  "NEXT Next question (or close)."])

# ================================================================================
# save
# ================================================================================
out = os.path.join(OUT, "01-ch1-energy-causes-change-slides.pptx")
prs.save(out)
print("built slides:", len(prs.slides._sldIdLst))
print("saved:", out)
