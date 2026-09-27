"""Build Chapter 4 Interactive Science Notebook (5 inserts)."""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from ch4_common import *
from ch4_common import _set_run

OUT = os.path.join(os.path.dirname(__file__), '..', 'production', 'student-notebook', 'chapter-04')

def _border(cell, color="666666", dashed=True):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'dashed' if dashed else 'single')
        el.set(qn('w:sz'), '12')
        el.set(qn('w:color'), color)
        borders.append(el)
    tcPr.append(borders)

def insert_card(d, instruction):
    add_par(d, instruction, size=9, bold=True, color=GRAY, space_after=2)
    t = d.add_table(rows=1, cols=1)
    t.alignment = 1
    cell = t.cell(0, 0)
    cell.width = Inches(7.0)
    _border(cell)
    return cell

def cell_add(cell, text="", size=11, bold=False, italic=False, color=INK, align=None,
             before=0, after=4):
    p = cell.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if text:
        _set_run(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return p

def cell_title(cell, text, sub):
    cell_add(cell, text, size=13, bold=True, color=HEADER, after=0)
    cell_add(cell, sub, size=10, italic=True, color=GRAY, after=4)

def cell_section(cell, text):
    cell_add(cell, text, size=12, bold=True, color=ACCENT, before=6, after=2)

def blank(cell, n=2):
    for _ in range(n):
        cell_add(cell, " " * 46, size=12, after=8)

def build():
    d = Document()
    setup_portrait(d)
    add_par(d, "Interactive Science Notebook — Chapter 4 · Energy Transfer", size=16, bold=True, color=HEADER, space_after=2)
    add_par(d, "Trim each insert on the dashed line and glue into your notebook. "
              "Every interactive part has a flat, low-motor alternative.",
            size=10, italic=True, color=GRAY, space_after=10)

    # ---------------- INSERT 1 ----------------
    cell = insert_card(d, "INSERT 1 — What Is Energy Transfer?    ·    trim on the dashed line    ·    glue the BACK")
    cell_title(cell, "What Is Energy Transfer?", "Chapter 4 · Week 4")
    cell_section(cell, "VOCABULARY")
    cell_add(cell, "Energy transfer happens when energy moves from one ______________ "
                   "or ______________ to another.", size=12, after=4)
    cell_section(cell, "MODEL")
    cell_add(cell, "SOURCE  →  ENERGY MOVES  →  TARGET", size=12, bold=True, after=2)
    cell_section(cell, "DRAW")
    cell_add(cell, "Draw ONE example of energy transfer. Label the source and the target.", size=11, after=2)
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(p.add_run("[  draw your example here  ]"), size=10, italic=True, color=GRAY)

    # ---------------- INSERT 2 ----------------
    cell = insert_card(d, "INSERT 2 — Sound + Light Transfer    ·    trim on the dashed line    ·    glue the BACK")
    cell_title(cell, "Sound + Light Transfer", "Chapter 4 · Week 4")
    cell_section(cell, "SOUND")
    cell_add(cell, "source → transfer → observed change", size=11, italic=True, after=2)
    cell_add(cell, "Draw the sound setup and the change you saw:", size=11, after=2)
    p = cell.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(p.add_run("[  draw here  ]"), size=10, italic=True, color=GRAY)
    cell_section(cell, "LIGHT")
    cell_add(cell, "source → transfer → observed change", size=11, italic=True, after=2)
    cell_add(cell, "Draw the light path and the change you saw:", size=11, after=2)
    p = cell.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(p.add_run("[  draw here  ]"), size=10, italic=True, color=GRAY)

    # ---------------- INSERT 3 ----------------
    cell = insert_card(d, "INSERT 3 — Heat + Electricity Transfer    ·    trim on the dashed line    ·    glue the BACK")
    cell_title(cell, "Heat + Electricity Transfer", "Chapter 4 · Week 5")
    cell_section(cell, "HEAT")
    cell_add(cell, "warmer  →  energy transfer  →  cooler", size=11, italic=True, after=2)
    cell_add(cell, "Evidence I observed:", size=11, bold=True, after=2)
    cell_add(cell, "The temperature changed: ____________________", size=12, after=6)
    cell_section(cell, "ELECTRICITY")
    cell_add(cell, "battery  →  circuit  →  device", size=11, italic=True, after=2)
    cell_add(cell, "Evidence I observed:", size=11, bold=True, after=2)
    cell_add(cell, "The device changed: (light / sound / motion) ____________________", size=12, after=6)

    # ---------------- INSERT 4 (interactive) ----------------
    cell = insert_card(d, "INSERT 4 — Transfer vs. Transformation    ·    trim on the dashed line    ·    "
                          "cut out the bottom cards and sort them under the flaps    ·    "
                          "flat option: draw a line from each system to its column")
    cell_title(cell, "Transfer vs. Transformation", "Chapter 4 · Week 5 · lift the flap to check")
    cell_section(cell, "THE THREE IDEAS")
    cell_add(cell, "TRANSFER — energy moves.", size=11, bold=True, after=0)
    cell_add(cell, "TRANSFORMATION — energy changes form.", size=11, bold=True, after=0)
    cell_add(cell, "BOTH — energy moves AND changes form.", size=11, bold=True, after=6)
    cell_section(cell, "SORT THE SYSTEMS")
    cell_add(cell, "Sort each system into TRANSFER, TRANSFORMATION, or BOTH. "
                   "(BOTH is allowed — do not force either/or.)", size=11, after=2)
    t = cell.add_table(rows=2, cols=3)
    t.style = 'Table Grid'
    for i, h in enumerate(["TRANSFER", "TRANSFORMATION", "BOTH"]):
        set_cell_text(t.cell(0, i), h, size=11, bold=True, color=RGBColor(255,255,255))
    for i in range(3):
        set_cell_text(t.cell(1, i), " ", size=11)
    style_table(t)
    cell_add(cell, " ", size=11, after=2)
    cell_add(cell, "Systems to sort:  flashlight · buzzer · motor · warm water warms a spoon · sound source (cup + paper)",
             size=11, italic=True, after=2)

    # ---------------- INSERT 5 ----------------
    cell = insert_card(d, "INSERT 5 — Evidence + CER    ·    trim on the dashed line    ·    glue the BACK")
    cell_title(cell, "Evidence + CER", "Chapter 4 · Week 5")
    cell_section(cell, "FOUR STEPS")
    cell_add(cell, "NOTICE — what did I observe?", size=11, bold=True, after=0)
    cell_add(cell, "EVIDENCE — what changed?", size=11, bold=True, after=0)
    cell_add(cell, "EXPLAIN — why did it change?", size=11, bold=True, after=0)
    cell_add(cell, "CONNECT — how does this show energy transfer?", size=11, bold=True, after=6)
    cell_section(cell, "CER PLANNING BOX")
    cell_add(cell, "Claim: I know ____________________________________", size=12, after=4)
    cell_add(cell, "because ________________________________________", size=12, after=4)
    cell_add(cell, "Evidence (the change I observed): ________________", size=12, after=4)
    cell_add(cell, "Reasoning (how the change shows transfer): ________", size=12, after=4)

    d.save(os.path.join(OUT, "04-ch4-interactive-notebook.docx"))
    print("built notebook")

if __name__ == "__main__":
    build()
