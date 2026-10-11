"""Build the Chapter 1 PREMIUM v1.2 nine-slide design spike (PPTX, 16:9).

Phase 3B. Reuses the canonical Chapter 1 visual tokens (navy/gold/cream, Avenir
Next, type scale) and component architecture, but produces ONLY the nine
approved premium spike targets. It integrates the Chief-approved source art from
work/premium/chapter-01/v1.2/assets/approved/.

This is a REVIEW derivative. It does NOT touch the frozen 45-slide deck, does
NOT modify guided notes, and does NOT propagate to the full deck.

Run: .venv/bin/python tools/build_ch1_premium_spike.py
"""
import os
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
WS = ROOT / "work" / "premium" / "chapter-01" / "v1.2"
APPROVED = WS / "assets" / "approved"
PROD = WS / "assets" / "production"
PROD.mkdir(parents=True, exist_ok=True)
SPIKE = WS / "spike"
SPIKE.mkdir(parents=True, exist_ok=True)
OUT = SPIKE / "Chapter_1_Premium_v1.2_Spike.pptx"

# ---------------------------------------------------------------- FROZEN Ch1 tokens
PRIMARY = RGBColor(0x1F, 0x2A, 0x5A)   # deep navy
ACCENT  = RGBColor(0xA6, 0x7C, 0x16)   # warm gold
NEUTRAL = RGBColor(0x7A, 0x56, 0x20)   # muted amber
SURFACE = RGBColor(0xF5, 0xEE, 0xDC)   # cream
INK     = RGBColor(0x1A, 0x1D, 0x26)
INK_SOFT= RGBColor(0x4A, 0x52, 0x68)
PAPER   = RGBColor(0xFF, 0xFF, 0xFF)
MIST    = RGBColor(0xF6, 0xF7, 0xF9)
RULE    = RGBColor(0xD8, 0xDC, 0xE4)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)

GOLD_TINT = RGBColor(0xF3, 0xE8, 0xCF)
NAVY_TINT = RGBColor(0xE3, 0xE7, 0xF2)

FONT = "Avenir Next"

T_MAJOR  = 48
T_TITLE  = 40
T_HEAD   = 32
T_BODY   = 30
T_BODY_S = 28
T_DEF    = 28
T_KICK   = 24
T_PAGE   = 16
T_TAB    = 12
T_HUD    = 24

CHAPTER_TAB = "CHAPTER 1 · ENERGY CAUSES CHANGE"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
_sz = prs.part._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}sldSz')
if _sz is not None and 'type' in _sz.attrib:
    del _sz.attrib['type']
BLANK = prs.slide_layouts[6]

SLIDE_NO = [0]

def new_slide(bg=SURFACE):
    SLIDE_NO[0] += 1
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background(); r.shadow.inherit = False
    return s

# ---------------------------------------------------------------- helpers
def rect(s, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line_color=None, line_w=None, rounded=False):
    sp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else shape,
                            Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        try:
            sp.adjustments[0] = 0.08
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line_color is not None:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(line_w or 1.25)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp

def tb(s, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    return box, tf

def run(p, text, size=T_BODY, color=INK, bold=False, italic=False, font=FONT, caps=False):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color; r.font.name = font
    if caps:
        r.font.caps = True
    return r

def text(s, x, y, w, h, lines, size=T_BODY, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False, gap=6, caps=False):
    if isinstance(lines, str):
        lines = [lines]
    box, tf = tb(s, x, y, w, h, anchor=anchor)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(gap)
        run(p, ln, size=size, color=color, bold=bold, italic=italic, caps=caps)
    return box

def rich(s, x, y, w, h, segs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, gap=12):
    box, tf = tb(s, x, y, w, h, anchor=anchor)
    for i, para_segs in enumerate(segs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(gap)
        for (txt, fmt) in para_segs:
            run(p, txt, size=fmt.get('size', T_BODY), color=fmt.get('color', INK),
                bold=fmt.get('bold', False), italic=fmt.get('italic', False))
    return box

def blank_str(n):
    return "_" * n

def chapter_tab(s, x=0.55, y=0.40, w=3.55, h=0.44):
    sp = rect(s, x, y, w, h, PRIMARY, rounded=True)
    tf = sp.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.12); tf.margin_right = Inches(0.12)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, CHAPTER_TAB, size=T_TAB, color=WHITE, bold=True)
    return sp

HUD_W = {"OBSERVE": 2.15, "WRITE": 1.6, "CHECK": 1.7, "TRACE": 1.65, "TEACH": 1.75, "V": 1.0}
def hud_bubble(s, label, x=None, y=0.36, h=0.52):
    w = HUD_W.get(label, 2.2)
    if x is None:
        x = 12.4 - w
    sp = rect(s, x, y, w, h, ACCENT, rounded=True)
    tf = sp.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.08)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, label, size=T_HUD, color=WHITE, bold=True)
    return sp

def mode_badge(s, mode, x=12.5, y=0.36, w=0.6, h=0.52, fill=PRIMARY, color=WHITE):
    sp = rect(s, x, y, w, h, fill, rounded=True)
    tf = sp.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, mode, size=T_HUD, color=color, bold=True)
    return sp

def slide_number(s, n=None):
    n = n or SLIDE_NO[0]
    text(s, 12.2, 7.02, 0.9, 0.4, str(n), size=T_PAGE, color=INK_SOFT, align=PP_ALIGN.RIGHT)

def anchor_chip(s, code, x=0.55, y=1.02, w=0.78, h=0.5):
    sp = rect(s, x, y, w, h, ACCENT, rounded=True)
    tf = sp.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, code, size=24, color=WHITE, bold=True)
    return sp

def header(s, title, kicker=None, mode=None, hud=None, anchor=None, title_size=T_TITLE):
    chapter_tab(s)
    if mode:
        mode_badge(s, mode)
    if hud:
        hud_bubble(s, hud)
    tx = 0.55
    if anchor:
        anchor_chip(s, anchor)
        tx = 1.5
    text(s, tx, 1.02, 12.75 - tx, 0.7, title, size=title_size, color=PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if kicker:
        text(s, 0.55, 1.72, 11.7, 0.4, kicker, size=T_KICK, color=INK_SOFT, italic=True)
    slide_number(s)

def set_note(s, lines):
    s.notes_slide.notes_text_frame.text = "\n".join(lines)

# ---------------------------------------------------------------- image pipeline
def fit_crop(src_path, aspect, out_name):
    """Center-crop src to target aspect (w/h) and write a normalized JPEG copy."""
    im = Image.open(src_path).convert("RGB")
    sw, sh = im.size
    src_asp = sw / sh
    if abs(src_asp - aspect) < 0.005:
        out = im
    elif src_asp > aspect:  # too wide -> crop sides
        nw = int(round(sh * aspect))
        x0 = (sw - nw) // 2
        out = im.crop((x0, 0, x0 + nw, sh))
    else:  # too tall -> crop top/bottom
        nh = int(round(sw / aspect))
        y0 = (sh - nh) // 2
        out = im.crop((0, y0, sw, y0 + nh))
    out_path = PROD / out_name
    out.save(out_path, quality=95)
    return out_path

def dim_image(src_path, factor, out_name):
    im = Image.open(src_path).convert("RGB")
    im = Image.eval(im, lambda v: int(v * factor))
    out_path = PROD / out_name
    im.save(out_path, quality=95)
    return out_path

def zoom_crop(src_path, zoom, cy, out_name):
    """Center-crop to a 4:3 sub-region (zoom-in) to protect an evidence cue.

    zoom = enlargement factor (>1). cy = vertical center of the crop window in
    [0,1] (0.5 = centered; <0.5 biases the window upward). Aspect is preserved,
    so the crop still displays undistorted in the on-slide 4:3 image box.
    """
    im = Image.open(src_path).convert("RGB")
    w, h = im.size
    s = 1.0 / zoom  # fraction of original retained
    cw = int(round(w * s))
    ch = int(round(h * s))
    left = int(round((0.5 - s / 2) * w))
    top = int(round((cy - s / 2) * h))
    left = max(0, min(left, w - cw))
    top = max(0, min(top, h - ch))
    out = im.crop((left, top, left + cw, top + ch))
    out_path = PROD / out_name
    out.save(out_path, quality=95)
    return out_path

def picture(s, path, x, y, w, h):
    s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))

# ---------------------------------------------------------------- premium vocab card
def premium_vocab_card(s, x, y, w, h, term, definition, page, img_path):
    """Collectible science card: TOP term / CENTER hero art (~50%) / BOTTOM def / FOOTER."""
    rect(s, x, y, w, h, PAPER, line_color=RULE, line_w=1.5, rounded=True)
    rect(s, x, y, w, 0.14, PRIMARY)  # navy top band
    # KEY WORD badge (top-left, restrained)
    pill = rect(s, x + 0.35, y + 0.3, 1.9, 0.46, ACCENT, rounded=True)
    tf = pill.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, "KEY WORD", size=18, color=WHITE, bold=True)
    # huge term
    text(s, x + 0.35, y + 0.8, w - 0.7, 0.95, term, size=68, color=PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    # hero image (center, ~50% of card area)
    iw = w - 0.7
    ih = 3.75
    aspect = iw / ih
    img = fit_crop(img_path, aspect, f"{Path(img_path).stem}_card.jpg")
    picture(s, img, x + 0.35, y + 1.8, iw, ih)
    # definition (bottom; 2 lines allowed for the longest definition)
    text(s, x + 0.35, y + 5.6, w - 0.7, 0.75, definition, size=26, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    # footer (reader page + chapter identity)
    text(s, x + 0.35, y + h - 0.45, w - 0.7, 0.35, f"{page}   ·   Unit 1 · Energy Transfer & Transformation",
         size=T_PAGE, color=INK_SOFT, bold=True)

# ================================================================================
# SLIDE 1 — CHAPTER OPENER (CH1-HERO-01 full bleed)
# ================================================================================
s = new_slide(PRIMARY)
picture(s, APPROVED / "CH1-HERO-01.jpg", 0, 0, 13.333, 7.5)
# chapter identifier (gold pill, top-left dark region)
sp = rect(s, 0.9, 0.75, 4.4, 0.55, ACCENT, rounded=True)
tf = sp.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = Inches(0.15); tf.margin_right = Inches(0.15)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
run(p, "CHAPTER 1 · GRADE 4 · WEEK 1", size=14, color=WHITE, bold=True)
# title in dark negative-space region (top-left)
text(s, 0.9, 1.7, 7.4, 1.05, "Energy Causes Change", size=T_MAJOR, color=WHITE, bold=True)
text(s, 0.9, 2.85, 7.0, 0.5, "Discovery / Evidence", size=T_KICK, color=RGBColor(0xD8, 0xDC, 0xE4), bold=True)
text(s, 0.9, 3.4, 7.6, 0.45, "Where can we observe evidence of energy causing change?", size=22, color=WHITE, italic=True)
# restrained unresolved model motif (PPTX shapes, not baked into image)
mx, my, mh = 0.9, 4.6, 0.62
mw = 1.9
for i, word in enumerate(["ENERGY", "CHANGE", "?"]):
    cx = mx + i * (mw + 0.55)
    chip = rect(s, cx, my, mw, mh, PRIMARY if i < 2 else ACCENT, rounded=True)
    tf = chip.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.06); tf.margin_right = Inches(0.06)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, word, size=20, color=WHITE, bold=True)
    if i < 2:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(cx + mw + 0.04), Inches(my + 0.17), Inches(0.47), Inches(0.28))
        ar.fill.solid(); ar.fill.fore_color.rgb = ACCENT
        ar.line.fill.background(); ar.shadow.inherit = False
text(s, 0.9, 5.55, 11.5, 0.5, "Unit 1 · Energy Transfer & Transformation", size=T_KICK, color=RGBColor(0xD8, 0xDC, 0xE4), italic=True)
slide_number(s)
set_note(s, ["MODE Title · orient", "SPIKE 1 of 9 — chapter opener (CH1-HERO-01 full bleed, title in dark negative space, unresolved motif)"])

# ================================================================================
# SLIDE 2 — DARK -> BRIGHT PHENOMENON (CH1-DARK-BRIGHT-01)
# ================================================================================
s = new_slide(SURFACE)
header(s, "What changed?", kicker="Reader p. 1 — you walk into a dark room and flip a switch.", hud="OBSERVE")
# matched comparison image (16:9 native, uncropped)
picture(s, APPROVED / "CH1-DARK-BRIGHT-01.jpg", 0.55, 2.35, 7.4, 7.4 * (9 / 16))
# DARK / BRIGHT captions under the image
text(s, 0.55, 6.6, 3.7, 0.4, "DARK room · hard to see", size=16, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)
text(s, 4.25, 6.6, 3.7, 0.4, "BRIGHT room · light fills the room", size=16, color=NEUTRAL, bold=True, align=PP_ALIGN.CENTER)
# explanation panel (right)
rect(s, 8.3, 2.35, 4.48, 4.16, NAVY_TINT, line_color=PRIMARY, line_w=1.0, rounded=True)
text(s, 8.65, 2.75, 3.8, 3.4, "Something became different.\n\nThat is a change.", size=T_BODY, color=PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
slide_number(s)
set_note(s, ["MODE Demo · observe", "SPIKE 2 of 9 — dark->bright matched comparison (CH1-DARK-BRIGHT-01, uncropped)"])

# ================================================================================
# SLIDE 3 — PENCIL BEFORE/AFTER (CH1-PENCIL-EVIDENCE-01)
# ================================================================================
s = new_slide(SURFACE)
header(s, "What changed?", hud="OBSERVE")
# combined before/after evidence image (1.556 native, uncropped)
pencil_img = APPROVED / "CH1-PENCIL-EVIDENCE-01.jpg"
iw = 6.4
ih = iw * (1234 / 1920)
ix, iy = 0.75, 2.3
picture(s, pencil_img, ix, iy, iw, ih)
# BEFORE / AFTER labels + hints (two-line captions under each image half)
cap_y = iy + ih + 0.12
cap_w = 3.5
text(s, ix, cap_y, cap_w, 0.45, "BEFORE", size=22, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
text(s, ix, cap_y + 0.4, cap_w, 0.4, "dull tip · sharpener quiet", size=13, color=INK_SOFT, align=PP_ALIGN.CENTER)
text(s, ix + iw - cap_w, cap_y, cap_w, 0.45, "AFTER", size=22, color=NEUTRAL, bold=True, align=PP_ALIGN.CENTER)
text(s, ix + iw - cap_w, cap_y + 0.4, cap_w, 0.4, "shorter · sharper · sharpener humming", size=13, color=INK_SOFT, align=PP_ALIGN.CENTER)
# CHANGES panel (right)
rect(s, 7.45, 2.3, 5.13, ih + 0.15, GOLD_TINT, line_color=ACCENT, line_w=0.75, rounded=True)
text(s, 7.75, 2.5, 4.5, 0.5, "CHANGES", size=24, color=NEUTRAL, bold=True, caps=True)
text(s, 7.75, 3.15, 4.7, 2.6, "• the tip is sharper\n• the pencil is shorter\n• the sharpener makes a sound", size=24, color=INK)
slide_number(s)
set_note(s, ["MODE Demo · observe", "SPIKE 3 of 9 — pencil before/after evidence (CH1-PENCIL-EVIDENCE-01, both states preserved)"])

# ================================================================================
# SLIDE 4 — ENERGY VOCABULARY CARD (CH1-VOC-ENERGY-01)
# ================================================================================
s = new_slide(SURFACE)
premium_vocab_card(s, 2.167, 0.3, 9.0, 6.9, "energy",
                   "the ability to cause change",
                   "Reader p. 1",
                   APPROVED / "CH1-VOC-ENERGY-01.jpg")
slide_number(s)
set_note(s, ["MODE Vocabulary · TEACH", "SPIKE 4 of 9 — ENERGY card (CH1-VOC-ENERGY-01 hero art)"])

# ================================================================================
# SLIDE 5 — CHANGE VOCABULARY CARD (CH1-VOC-CHANGE-01)
# ================================================================================
s = new_slide(SURFACE)
premium_vocab_card(s, 2.167, 0.3, 9.0, 6.9, "change",
                   "to become different",
                   "Reader p. 1",
                   APPROVED / "CH1-VOC-CHANGE-01.jpg")
slide_number(s)
set_note(s, ["MODE Vocabulary · TEACH", "SPIKE 5 of 9 — CHANGE card (CH1-VOC-CHANGE-01 hero art)"])

# ================================================================================
# SLIDE 6 — ENERGY -> CHANGE -> EVIDENCE MODEL (CH1-MODEL-ECE-01 subdued bg)
# ================================================================================
s = new_slide(PRIMARY)
bg = dim_image(APPROVED / "CH1-MODEL-ECE-01.jpg", 0.42, "CH1-MODEL-ECE-01_subdued.jpg")
picture(s, bg, 0, 0, 13.333, 7.5)
chapter_tab(s)
text(s, 0.55, 1.0, 12.23, 0.7, "Look at the seashore", size=34, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
text(s, 0.55, 1.75, 12.23, 0.5, "Reader p. 2 — the seashore changes all around us.", size=22, color=RGBColor(0xE8, 0xEB, 0xF2), italic=True, align=PP_ALIGN.CENTER)
# three navy nodes + gold arrows
nodes = ["ENERGY", "CHANGE", "EVIDENCE"]
nw, nh, ngap = 3.2, 1.2, 0.9
total = 3 * nw + 2 * ngap
x0 = (13.333 - total) / 2
ny = 3.05
for i, word in enumerate(nodes):
    cx = x0 + i * (nw + ngap)
    node = rect(s, cx, ny, nw, nh, PRIMARY, rounded=True)
    node.line.color.rgb = WHITE
    node.line.width = Pt(1.25)
    tf = node.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.1); tf.margin_right = Inches(0.1)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, word, size=28, color=WHITE, bold=True)
    if i < 2:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(cx + nw + 0.06), Inches(ny + nh/2 - 0.18), Inches(ngap - 0.12), Inches(0.36))
        ar.fill.solid(); ar.fill.fore_color.rgb = ACCENT
        ar.line.fill.background(); ar.shadow.inherit = False
text(s, 0.75, 4.7, 11.83, 0.6, "Waves move. Sand shifts. Water makes sound.", size=26, color=WHITE, align=PP_ALIGN.CENTER)
text(s, 0.75, 5.45, 11.83, 0.6, "A change you observe is evidence that energy was involved.", size=28, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
slide_number(s)
set_note(s, ["MODE Read · model (resolved showcase)", "SPIKE 6 of 9 — ECE model (CH1-MODEL-ECE-01 subdued bg + navy nodes/gold arrows). Reveal timing preserved in canonical architecture (slide 14 '?' -> slide 16 EVIDENCE)."])

# ================================================================================
# SLIDE 7 — FIVE FORMS OF ENERGY (CH1-FORMS-*)
# ================================================================================
s = new_slide(SURFACE)
header(s, "Energy has different forms", kicker="Reader p. 3 — think of a spinning carousel.")
forms = [
    ("motion", "CH1-FORMS-MOTION-01.jpg", None),
    ("sound", "CH1-FORMS-SOUND-01.jpg", None),
    ("light", "CH1-FORMS-LIGHT-01.jpg", None),
    ("heat", "CH1-FORMS-HEAT-01.jpg", (1.28, 0.42)),
    ("electrical energy", "CH1-FORMS-ELECTRICAL-01.jpg", (1.28, 0.5)),
]
cw, gap = 2.3, 0.183
total = 5 * cw + 4 * gap
x0 = 0.55 + (12.23 - total) / 2
card_y, card_h = 2.5, 2.6
for i, (label, fname, zoom) in enumerate(forms):
    cx = x0 + i * (cw + gap)
    rect(s, cx, card_y, cw, card_h, PAPER, line_color=PRIMARY, line_w=1.25, rounded=True)
    # image (native 4:3; heat/electrical get a modest zoom crop to protect cues)
    img = APPROVED / fname
    if zoom:
        img = zoom_crop(img, zoom[0], zoom[1], f"{Path(fname).stem}_zoom.jpg")
    iw = cw - 0.16
    ih = iw * (3 / 4)
    picture(s, img, cx + 0.08, card_y + 0.08, iw, ih)
    # label band
    band = rect(s, cx, card_y + 0.08 + ih + 0.08, cw, card_h - 0.24 - ih, PRIMARY)
    tfb = band.text_frame; tfb.word_wrap = True; tfb.vertical_anchor = MSO_ANCHOR.MIDDLE
    tfb.margin_left = Inches(0.06); tfb.margin_right = Inches(0.06)
    p = tfb.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, label, size=22, color=WHITE, bold=True)
text(s, 0.55, 5.5, 12.23, 1.2, "On the carousel, the horses spin and move up and down.\nThat is motion.", size=T_BODY_S, color=INK, align=PP_ALIGN.CENTER)
slide_number(s)
set_note(s, ["MODE Read · forms", "SPIKE 7 of 9 — five-form family (uniform card system, labels exact, clues preserved)"])

# ================================================================================
# SLIDE 8 — STORED ENERGY VOCABULARY CARD (CH1-VOC-STORED-ENERGY-01)
# ================================================================================
s = new_slide(SURFACE)
premium_vocab_card(s, 2.167, 0.3, 9.0, 6.9, "stored energy",
                   "energy that has the ability to cause change at a later time",
                   "Reader p. 4",
                   APPROVED / "CH1-VOC-STORED-ENERGY-01.jpg")
slide_number(s)
set_note(s, ["MODE Vocabulary · TEACH", "SPIKE 8 of 9 — STORED ENERGY card (CH1-VOC-STORED-ENERGY-01 hero art)"])

# ================================================================================
# SLIDE 9 — WRITE/CHECK RESTRAINT BENCHMARK (V1 write, no image)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "WRITE")
anchor_chip(s, "V1")
text(s, 1.5, 1.02, 11.0, 0.6, "Vocabulary — Energy & Change", size=T_HEAD, color=PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
text(s, 1.5, 1.62, 8.0, 0.35, "Guided Notes · Reader p. 1", size=T_KICK, color=INK_SOFT, italic=True)
rich(s, 1.5, 2.6, 11.0, 1.0, [
    [("energy — the ability to cause ", {}), (blank_str(7), {"underline": True, "color": INK}), (" .", {})],
])
rich(s, 1.5, 3.6, 11.0, 1.0, [
    [("change — to become ", {}), (blank_str(9), {"underline": True, "color": INK}), (" .", {})],
])
text(s, 1.5, 4.9, 11.0, 0.5, "Write the missing word in each blank.", size=T_BODY_S, color=INK_SOFT, italic=True)
slide_number(s)
set_note(s, ["MODE Notes · V1 anchor (restraint benchmark)", "SPIKE 9 of 9 — WRITE/CHECK V1, no premium image, exact notes mirror"])

# ================================================================================
prs.save(OUT)
print("built spike slides:", len(prs.slides._sldIdLst))
print("saved:", OUT)
