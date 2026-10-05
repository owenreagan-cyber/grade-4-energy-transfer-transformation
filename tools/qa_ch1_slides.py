"""Phase 4 — Chapter 1 Slideshow automated QA (deterministic, structural/text).

Headless gates against the FINAL Ch1 slideshow (PPTX + rendered PDF):
  slide dims · font floors · safe-margin/overflow · visual-token compliance ·
  BQ/V/N anchor sync + order · exact note wording · answer reveal (no front-load) ·
  Reader page refs · vocab grouping · review A/B/C/D structure ·
  no stale INB/26-item language · no banned advanced vocabulary ·
  render QA (PDF page count).

Run: .venv/bin/python tools/qa_ch1_slides.py
"""
import os, re, subprocess, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import ch1_pacing as PACE
import qa_render_geometry as GEO
from pptx import Presentation
from pptx.util import Pt

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
PPTX = os.path.join(ROOT, "production", "slides", "01-ch1-energy-causes-change-slides.pptx")
PDF = os.path.join(ROOT, "production", "slides", "01-ch1-energy-causes-change-slides.pdf")
REPORT = os.path.join(ROOT, "qa", "reports", "PHASE_4_CH1_SLIDESHOW_QA.md")
if "--report" in sys.argv:
    REPORT = os.path.join(ROOT, sys.argv[sys.argv.index("--report") + 1])

# ---- FROZEN tokens (Ch1 + unit neutrals) -----------------------------------
ALLOWED_FILLS = {
    "1A1D26", "4A5268", "FFFFFF", "F6F7F9", "D8DCE4",   # unit neutrals
    "1F2A5A", "A67C16", "7A5620", "F5EEDC",             # Ch1 Discovery/Evidence
    "F3E8CF", "E3E7F2",                                 # Ch1 gold/navy tints (10%)
}

BQ_EXACT = "Where can we observe evidence of energy causing change?"
ANCHORS = ["BQ", "V1", "N1", "N2", "V2", "N3", "N4"]
ANSWER_MAP = {
    "V1": ["change", "different"],
    "V2": ["later"],
    "N1": ["evidence"],
    "N3": ["position"],
}
NOTE_WORDING = [
    "energy — the ability to cause",          # V1
    "change — to become",                      # V1
    "stored energy — energy that has the ability to cause change at a",  # V2
    "Energy has the ability to cause change.",  # N1 (printed)
    "When you see a change, you are seeing",    # N1 (one blank)
    "of energy.",                               # N1
    "Write one example that shows a change for each form.",  # N2
    "Objects also have stored energy because of their",  # N3
    "Draw a change you observed. Show what it looked like BEFORE and AFTER.",  # N4
    "What changed?",                           # N4
    "How do we know energy is involved if we cannot see energy itself?",  # N4 check
    "Start with: “We know because we see",     # N4 check starter
    "Now I Think",                             # BQ return
]
READER_REFS = ["Reader p. 1", "Reader p. 3", "Reader p. 4", "pp. 1–2", "pp. 4–5", "pp. 1–5"]
MODEL_CHAINS = [("ENERGY", "CHANGE", "EVIDENCE"), ("STORED ENERGY", "RELEASE", "CHANGE")]
N2_FORMS = ["motion", "sound", "light", "heat", "electrical energy"]
N3_EXAMPLES = ["battery (in a flashlight)", "stretched rubber band", "raised object (apple)"]

BANNED = [
    "26 items", "26 questions", "26Q", "26-item", "INB",
    "kinetic", "potential", "mechanical", "radiant", "momentum", "inertia",
    "compression wave", "gravitational", "proof", "proves",
]

# Font tiers (Phase 6.1 / F14). Chrome is exempt; everything else >= 24pt; sentence-length text
# below the title/kicker band (top >= 2.0in) >= 28pt; table column headers >= 20pt.
CHROME_PREFIXES = ("CHAPTER 1", "Reader p", "Guided Notes", "Unit 1 ·")
ROW_LABELS = {"battery (in a flashlight)", "stretched rubber band", "raised object (apple)"}
TABLE_HEADERS = {"FORM", "ONE EXAMPLE THAT SHOWS A CHANGE", "EXAMPLE", "WHAT CHANGE COULD HAPPEN LATER?"}

# ---------------------------------------------------------------- helpers
def slide_texts(prs):
    out = []
    for i, s in enumerate(prs.slides, 1):
        parts = []
        for sh in s.shapes:
            if sh.has_text_frame:
                t = " ".join(r.text for p in sh.text_frame.paragraphs for r in p.runs)
                if t.strip():
                    parts.append(t)
        out.append((i, " ".join(parts)))
    return out

def all_runs(prs):
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        yield s, r

def shape_fills(prs):
    fills = set()
    for s in prs.slides:
        for sh in s.shapes:
            try:
                if sh.fill.type is not None and sh.fill.type == 1:  # solid
                    c = sh.fill.fore_color.rgb
                    fills.add(str(c))
            except Exception:
                pass
    return fills

def slide_dims(prs):
    w = prs.slide_width / 914400
    h = prs.slide_height / 914400
    return w, h

def overflow_shapes(prs):
    """Return list of (slide_idx, kind) for shapes extending beyond slide bounds."""
    W = prs.slide_width; H = prs.slide_height
    bad = []
    tol = Emu_val(1.0)  # ~1pt tolerance
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            try:
                l, t, w, h = sh.left, sh.top, sh.width, sh.height
                if l is None or t is None or w is None or h is None:
                    continue
                if l < -tol or t < -tol or l + w > W + tol or t + h > H + tol:
                    bad.append((i, sh.shape_type, round(l/914400,2), round(t/914400,2),
                                round((l+w)/914400,2), round((t+h)/914400,2)))
            except Exception:
                continue
    return bad

def Emu_val(pt):
    return int(pt * 12700)

def pdf_pages(path):
    if not os.path.exists(path):
        return -1
    info = subprocess.run(["pdfinfo", path], capture_output=True, text=True).stdout
    m = re.search(r"Pages:\s+(\d+)", info)
    return int(m.group(1)) if m else -1

# ---------------------------------------------------------------- gates
def gate_dims(prs):
    w, h = slide_dims(prs)
    ok = abs(w - 13.333) < 0.02 and abs(h - 7.5) < 0.02
    return ("PASS" if ok else "FAIL", f"{w:.3f} x {h:.3f} in (expect 13.333 x 7.5 / 16:9)")

def gate_font_floor(prs):
    sizes = [r.font.size.pt for _, r in all_runs(prs) if r.font.size and (r.text or "").strip()]
    if not sizes:
        return ("FAIL", "no sized runs")
    lo = min(sizes); hi = max(sizes)
    return ("PASS" if lo >= 12.0 else "FAIL",
            f"min {lo:.1f} pt (floor 12, chrome) · max {hi:.1f} pt")

def gate_font_tiers(prs):
    bad = []
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            top = (sh.top or 0) / 914400
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    t = (r.text or "").strip()
                    if not t or not r.font.size:
                        continue
                    sz = r.font.size.pt
                    if t.isdigit() or t.startswith(CHROME_PREFIXES) or t in ("·",):
                        continue
                    if "Reader p" in t[:12] or t.startswith("Guided Notes"):
                        continue
                    floor = 20 if t in TABLE_HEADERS else 24
                    if len(t) >= 25 and top >= 2.0 and t not in TABLE_HEADERS and t not in ROW_LABELS:
                        floor = 28
                    # italic helper captions on the title/kicker band are covered by T_KICK=24
                    if sz < floor:
                        bad.append((i, t[:28], sz, floor))
    return ("PASS" if not bad else "FAIL", f"{len(bad)} run(s) under tier floor (24 labels / 28 sentences / 20 table headers; chrome exempt): {bad[:8] or 'none'}")

def gate_widescreen_meta(prs):
    sz = prs.part._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}sldSz')
    t = sz.get('type') if sz is not None else None
    return ("PASS" if t != "screen4x3" else "FAIL", f"sldSz type attr = {t!r} (must not be screen4x3 on a 16:9 deck)")

def note_lines(slide):
    return slide.notes_slide.notes_text_frame.text.split("\n")

def gate_pacing_notes(prs):
    """Speaker-note TIME/PLAN lines must be derived from tools/ch1_pacing.py (one authoritative system)."""
    problems = list(PACE.check_plan())
    for i, sl in enumerate(prs.slides, 1):
        ls = note_lines(sl)
        time_ln = next((l for l in ls if l.startswith("TIME ")), None)
        plan_ln = next((l for l in ls if l.startswith("PLAN ")), None)
        if i not in PACE.S:
            problems.append(f"slide {i} not in pacing plan"); continue
        want = "TIME %g min" % PACE.S[i]["minutes"]
        if time_ln != want:
            problems.append(f"slide {i}: {time_ln!r} != {want!r}")
        if not plan_ln:
            problems.append(f"slide {i}: no PLAN line")
        else:
            info = PACE.S[i]
            if info["session"] and f"Session {info['session']}" not in plan_ln:
                problems.append(f"slide {i}: PLAN session mismatch")
            if info["status"] == "OPT" and "OPTIONAL" not in plan_ln:
                problems.append(f"slide {i}: optional not tagged")
            if info["status"] == "REQ" and "REQUIRED" not in plan_ln:
                problems.append(f"slide {i}: required not tagged")
    tot = " ; ".join("S%d %.1f+buffer %.1f" % (k, PACE.session_required_minutes(k), PACE.session_buffer(k)) for k in PACE.SESSIONS)
    return ("PASS" if not problems else "FAIL", f"{tot} · problems: {problems[:6] or 'none'}")

def gate_session_openers(prs):
    """No session may open on a reveal/answer slide; re-entry slide 2 is a BQ slide."""
    bad = []
    for k, sess in PACE.SESSIONS.items():
        first = sess["order"][0]
        mode = note_lines(prs.slides[first - 1])[0]
        if "reveal" in mode.lower() or "check" in mode.lower() or "answer" in mode.lower():
            bad.append((k, first, mode))
    return ("PASS" if not bad else "FAIL", f"session opener slides: {[(k, v['order'][0]) for k, v in PACE.SESSIONS.items()]} · reveal-openers: {bad or 'none'}")

def gate_open_task_choreography(prs):
    """F7: open-response tables must not be pre-answered before the write slide.
       N2 (slides 17-19): none of the 4 non-motion example phrases; N3 (23-25): no 'flashlight lights'/'falls'."""
    texts = dict(slide_texts(prs))
    bad = []
    pre = {
        17: ["your hands", "you hear a clap", "warms", "plugs in"],
        18: ["your hands move", "you hear a clap", "may feel warm"],
        23: ["is released", "flashlight lights", "lights up", "chemical"],
        24: ["it falls", "will fall", "Let it go"],
        25: ["it moves and makes a sound", "makes a sound: the change"],
        28: ["We recognize energy by the change we observe", "Evidence = the change you see"],
    }
    for sl, phrases in pre.items():
        for ph in phrases:
            if ph.lower() in texts[sl].lower():
                bad.append((sl, ph))
    return ("PASS" if not bad else "FAIL", f"pre-answered open tasks: {bad or 'none'}")

def _all_note_text(prs):
    return "\n".join("\n".join(note_lines(sl)) for sl in prs.slides)

def gate_phase_6_1b(prs):
    """Phase 6.1B final Grok remediation gates (M1 M2 M3 L1 L4 L5 L6)."""
    texts = dict(slide_texts(prs)); notes = {i: "\n".join(note_lines(sl)) for i, sl in enumerate(prs.slides, 1)}
    bad = []
    # M2 (reviewer follow-up): no slide/speaker note before slide 23 may ask or answer the still-battery question
    for i in range(1, 23):
        if re.search(r"battery", notes[i], re.I) and re.search(r"ASK[^\n]*battery|LISTEN FOR[^\n]*battery", notes[i], re.I):
            bad.append("slide %d speaker note pre-asks the slide 23 battery prediction" % i)
    # 6.1C / L-C: slide 1 teaser must not show the completed ENERGY -> CHANGE -> EVIDENCE chain (N1 reveals it)
    if re.search(r"\bEVIDENCE\b", texts[1]) or "?" not in texts[1].replace("Where can we observe evidence of energy causing change?", ""):
        bad.append("slide 1 shows the completed evidence chain (must be ENERGY -> CHANGE -> open box)")
    # M1: N1 write slide must not show the model's answer; reveal slide completes it and tells students to copy the word
    if "EVIDENCE" in texts[15] or "evidence" in texts[15].lower():
        bad.append("slide 15 shows the N1 answer")
    if "last box" not in texts[15]:
        bad.append("slide 15 does not explain the open model box")
    if not all(w in texts[16] for w in ("ENERGY", "CHANGE", "EVIDENCE")) or "last box" not in texts[16]:
        bad.append("slide 16 must show the full model + copy-into-last-box instruction")
    if "?" not in texts[14]:
        bad.append("slide 14 model must keep its '?' box")
    # M2: slide 23 PREDICT must not answer itself
    for w in ("chemical", "stores energy", "store energy", "stored"):
        if w in texts[23].lower():
            bad.append("slide 23 contains %r (answers the prediction)" % w)
    if "chemicals" not in texts[25].lower() or "chemicals" not in notes[23].lower():
        bad.append("battery/chemicals fact must appear in the slide 23 reveal note and on slide 25")
    # M3: Friday clock in speaker notes of the Friday slides
    for sl in (28, 30, 33):
        if not re.search(r"minute (6|9\.5|12)|MINUTE 12", notes[sl]):
            bad.append("slide %d has no FRIDAY CLOCK note" % sl)
    # L1: N4 Reader range
    for sl in (28, 29, 30):
        if "pp. 1–5" not in texts[sl] and "pp.1–5" not in notes[sl]:
            bad.append("slide %d does not cite pp. 1–5" % sl)
    if re.search(r"pp\. ?1[–-]6", "\n".join(texts.values()) + _all_note_text(prs)):
        bad.append("a 'pp. 1–6' citation remains in the deck")
    # L4: slide 32 question state must not give the evidence rule
    if re.search(r"evidence of energy|that is the evidence", texts[32], re.I):
        bad.append("slide 32 question state leaks the evidence rule")
    # L5: slide 45 battery wording is source-faithful and has no 'released'
    if "released" in texts[45].lower() or "A battery stores energy" not in texts[45] or "later" not in texts[45]:
        bad.append("slide 45 wording: %r" % texts[45][-90:])
    # L6: no hyphenated range that silently includes optional slide 18
    allnotes = _all_note_text(prs)
    if re.search(r"\b17\s*[–-]\s*27\b", allnotes):
        bad.append("speaker notes still contain the range 17-27")
    if "18 is OPTIONAL" not in notes[1]:
        bad.append("slide 1 schedule note does not mark slide 18 OPTIONAL")
    return ("PASS" if not bad else "FAIL", "Phase 6.1B gates (M1 M2 M3 L1 L4 L5 L6): %s" % (bad or "all clear"))

def gate_review_letters(prs):
    texts = dict(slide_texts(prs))
    letters = []
    for q in range(6):
        qs, as_ = 34 + 2 * q, 35 + 2 * q
        note = note_lines(prs.slides[as_ - 1])
        ans = next(l for l in note if l.startswith("REVEAL "))[len("REVEAL "):]
        # option order = order of option texts on question slide shapes (after the question)
        opts = []
        for sh in prs.slides[qs - 1].shapes:
            if sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if t and t not in ("A", "B", "C", "D") and not t.startswith(("Review", "CHAPTER")) and not t.isdigit() and t != "C" :
                    opts.append(t)
        # drop question text (first long text) then match
        idx = None
        for k, t in enumerate(opts):
            if t == ans:
                idx = k
        letters.append(idx)
    order = []
    for sh in prs.slides[33].shapes:
        pass
    # recompute letter via position among the 4 option boxes
    res = []
    for q in range(6):
        qs, as_ = 34 + 2 * q, 35 + 2 * q
        ans = next(l for l in note_lines(prs.slides[as_ - 1]) if l.startswith("REVEAL "))[len("REVEAL "):]
        opt_shapes = [sh.text_frame.text.strip() for sh in prs.slides[qs - 1].shapes
                      if sh.has_text_frame and sh.text_frame.text.strip() and sh.top / 914400 >= 2.7
                      and sh.text_frame.text.strip() not in ("A", "B", "C", "D") and not sh.text_frame.text.strip().isdigit()]
        res.append("ABCD"[opt_shapes.index(ans)] if ans in opt_shapes else "?")
    seq = "".join(res)
    sequential = any(seq[i:i+3] in ("ABC", "BCD", "CBA", "DCB") for i in range(len(seq) - 2))
    repeats = any(seq[i] == seq[i+1] == seq[i+2] for i in range(len(seq) - 2))
    ok = "?" not in seq and not sequential and not repeats and max(seq.count(c) for c in "ABCD") <= 2 and len(set(seq)) == 4
    return ("PASS" if ok else "FAIL", f"correct letters {seq} (no ascending/descending run, no triple repeat, <=2 per letter, all four used)")

def gate_render_geometry(prs):
    rep = GEO.check(PDF, edge_pt=6.0)
    oob, ov, empty = GEO.summarize(rep)
    return ("PASS" if not oob and not ov and not empty else "FAIL",
            f"rendered text: out-of-bounds pages {sorted(oob)[:8] or 'none'} · overlapping-word pages {({p: v[:3] for p, v in ov.items()}) or 'none'} · empty pages {empty or 'none'}")

def gate_body_floor(prs):
    """Body/instructional floor: largest-font-per-slide sanity + vocab term presence."""
    vocab_big = any(r.font.size and r.font.size.pt >= 48 for _, r in all_runs(prs))
    return ("PASS" if vocab_big else "FAIL", "vocab term >= 48pt present")

def gate_overflow(prs):
    bad = overflow_shapes(prs)
    return ("PASS" if not bad else "FAIL", f"{len(bad)} shape(s) overflow slide: {bad or 'none'}")

def gate_tokens(prs):
    fills = shape_fills(prs)
    unknown = fills - ALLOWED_FILLS
    return ("PASS" if not unknown else "WARN", f"unknown fills: {sorted(unknown) or 'none'}")

def gate_anchors(prs):
    texts = slide_texts(prs)
    blob = "\n".join(t for _, t in texts)
    missing = [c for c in ANCHORS if c not in blob]
    # order: BQ < V1 < N1 < N2 < V2 < N3 < N4
    idx = [blob.find(c) for c in ANCHORS]
    order_ok = all(a >= 0 for a in idx) and all(a < b for a, b in zip(idx, idx[1:]))
    if not order_ok:
        missing.append("anchor order")
    return ("PASS" if not missing else "FAIL", f"missing/out-of-order anchors: {missing or 'none'}")

def gate_bq_exact(prs):
    blob = "\n".join(t for _, t in slide_texts(prs))
    return ("PASS" if BQ_EXACT in blob else "FAIL", "BQ exact wording present")

def gate_note_wording(prs):
    blob = "\n".join(t for _, t in slide_texts(prs))
    missing = [w for w in NOTE_WORDING if w not in blob]
    return ("PASS" if not missing else "FAIL", f"missing note wording: {missing or 'none'}")

def gate_answers_revealed(prs):
    """Each anchor's answers must appear somewhere in the deck (reveal/check slides)."""
    blob = "\n".join(t for _, t in slide_texts(prs))
    missing = []
    for anchor, answers in ANSWER_MAP.items():
        for a in answers:
            if a not in blob:
                missing.append(f"{anchor}:{a}")
    return ("PASS" if not missing else "FAIL", f"answers missing: {missing or 'none'}")

def gate_no_front_load(prs):
    """Write slides must NOT render the answer in the blank (answer styled bold+accent = revealed).
       We check that the answer word, when it appears bold+gold, is on a reveal slide only.
       Simpler: detect bold+accent runs equal to an answer and list their slides (informational)."""
    hits = []
    for s, r in all_runs(prs):
        if not r.text.strip():
            continue
        t = r.text.strip().lower().rstrip(".")
        if t in {"change", "different", "later", "cause", "evidence", "position"}:
            try:
                bold = r.font.bold
                color = str(r.font.color.rgb) if r.font.color and r.font.color.type is not None else None
                if bold and color in ("A67C16",):
                    # slide index
                    idx = list(prs.slides).index(s) + 1
                    hits.append((idx, t))
            except Exception:
                pass
    # Reveal slides: 12(V1) 16(N1) 22(V2) 27(N3). Allow hits only there.
    allowed = {12, 16, 22, 27}
    bad = [(i, t) for i, t in hits if i not in allowed]
    return ("PASS" if not bad else "FAIL", f"front-loaded (bold+gold) answers on wrong slides: {bad or 'none'}")

def gate_reader_refs(prs):
    blob = "\n".join(t for _, t in slide_texts(prs))
    present = [r for r in READER_REFS if r in blob]
    return ("PASS" if len(present) >= 6 else "WARN", f"reader refs found: {present}")

def gate_vocab_grouping(prs):
    texts = slide_texts(prs)
    by_slide = {i: t for i, t in texts}
    # V1 card slide must contain both energy + change definitions; V2 slide stored energy
    v1_slide = next((t for i, t in texts if "to become different" in t and "the ability to cause change" in t), None)
    v2_slide = next((t for i, t in texts if "at a later time" in t), None)
    ok = v1_slide is not None and v2_slide is not None
    return ("PASS" if ok else "FAIL", "V1 (energy+change) and V2 (stored energy) groupings present")

def gate_n2_forms(prs):
    blob = "\n".join(t for _, t in slide_texts(prs))
    missing = [f for f in N2_FORMS if f not in blob]
    return ("PASS" if not missing else "FAIL", f"N2 forms missing: {missing or 'none'}")

def gate_n3_examples(prs):
    blob = "\n".join(t for _, t in slide_texts(prs))
    missing = [e for e in N3_EXAMPLES if e not in blob]
    return ("PASS" if not missing else "FAIL", f"N3 examples missing: {missing or 'none'}")

def gate_model_chains(prs):
    blob = "\n".join(t for _, t in slide_texts(prs))
    missing = []
    for chain in MODEL_CHAINS:
        if not all(c in blob for c in chain):
            missing.append("→".join(chain))
    return ("PASS" if not missing else "FAIL", f"model chains missing: {missing or 'none'}")

def gate_review_structure(prs):
    texts = slide_texts(prs)
    q_slides = []
    for i, t in texts:
        if re.search(r"Review \d+ —", t) and "answer" not in t:
            q_slides.append((i, t))
    bad = []
    for i, t in q_slides:
        for letter in ("A", "B", "C", "D"):
            if letter not in t:
                bad.append(f"slide {i} missing option {letter}")
    n = len(q_slides)
    ok = 4 <= n <= 6 and not bad
    return ("PASS" if ok else "FAIL", f"{n} review questions; {bad or 'all have A/B/C/D'}")

def gate_banned(prs):
    blob = "\n".join(t for _, t in slide_texts(prs)).lower()
    hits = [b for b in BANNED if b in blob]
    return ("PASS" if not hits else "FAIL", f"banned language: {sorted(set(hits)) or 'none'}")

def gate_transfer_transform_scope(prs):
    """transfer/transform may appear ONLY in the official unit title 'Energy Transfer & Transformation'."""
    blob = "\n".join(t for _, t in slide_texts(prs))
    offenders = []
    for m in re.finditer(r"[Tt]ransfer|[Tt]ransform(ation)?", blob):
        start = max(0, m.start() - 30); end = min(len(blob), m.end() + 30)
        ctx = blob[start:end]
        if "Energy Transfer & Transformation" not in ctx:
            offenders.append(ctx.replace("\n", " "))
    return ("PASS" if not offenders else "WARN",
            f"transfer/transform outside unit title: {offenders or 'none'}")

def gate_render(prs):
    pages = pdf_pages(PDF)
    n_slides = len(prs.slides._sldIdLst)
    ok = pages == n_slides
    return ("PASS" if ok else "FAIL", f"PDF pages {pages} == slides {n_slides}")

def gate_slide_count(prs):
    n = len(prs.slides._sldIdLst)
    return ("PASS" if n >= 30 else "WARN", f"{n} slides")

def main():
    prs = Presentation(PPTX)
    gates = [
        ("slide dimensions (16:9)", gate_dims(prs)),
        ("font floor >= 12pt (chrome)", gate_font_floor(prs)),
        ("font tiers (24 labels / 28 sentences)", gate_font_tiers(prs)),
        ("widescreen metadata (no screen4x3)", gate_widescreen_meta(prs)),
        ("pacing notes = single plan", gate_pacing_notes(prs)),
        ("no session opens on a reveal", gate_session_openers(prs)),
        ("open tasks not pre-answered (F7)", gate_open_task_choreography(prs)),
        ("Phase 6.1B remediation (M1 M2 M3 L1 L4 L5 L6)", gate_phase_6_1b(prs)),
        ("review correct letters not patterned", gate_review_letters(prs)),
        ("rendered geometry (PDF words)", gate_render_geometry(prs)),
        ("vocab term >= 48pt", gate_body_floor(prs)),
        ("no shape overflow", gate_overflow(prs)),
        ("visual-token compliance", gate_tokens(prs)),
        ("BQ/V/N anchors + order", gate_anchors(prs)),
        ("BQ exact wording", gate_bq_exact(prs)),
        ("note wording mirror", gate_note_wording(prs)),
        ("answers revealed (check slides)", gate_answers_revealed(prs)),
        ("no front-loaded answers", gate_no_front_load(prs)),
        ("Reader page refs", gate_reader_refs(prs)),
        ("vocab grouping (V1/V2)", gate_vocab_grouping(prs)),
        ("N2 five forms", gate_n2_forms(prs)),
        ("N3 three examples", gate_n3_examples(prs)),
        ("model chains (N1/N3)", gate_model_chains(prs)),
        ("review A/B/C/D structure", gate_review_structure(prs)),
        ("no banned/stale language", gate_banned(prs)),
        ("transfer/transform scope", gate_transfer_transform_scope(prs)),
        ("render QA (PDF pages)", gate_render(prs)),
        ("slide count", gate_slide_count(prs)),
    ]
    print("=" * 76)
    print("CH1 SLIDESHOW — AUTOMATED QA (Phase 4)")
    print("=" * 76)
    overall = "PASS"
    for name, (s, note) in gates:
        if s == "FAIL":
            overall = "FAIL"
        elif s == "WARN" and overall == "PASS":
            overall = "WARN"
        print(f"[{s:4}] {name:34} {note}")
    print("=" * 76)
    print(f"OVERALL: {overall}")

    os.makedirs(os.path.dirname(REPORT), exist_ok=True)
    with open(REPORT, "w") as f:
        f.write("# Chapter 1 Slideshow Automated QA\n\nGenerated by tools/qa_ch1_slides.py (v1.1 gates).\n\n")
        f.write(f"Overall: **{overall}**\n\n")
        f.write(f"Artifact: `production/slides/01-ch1-energy-causes-change-slides.pptx`\n\n")
        for name, (s, note) in gates:
            f.write(f"- [{s}] {name} — {note}\n")
    print(f"report: {REPORT}")

if __name__ == "__main__":
    main()
