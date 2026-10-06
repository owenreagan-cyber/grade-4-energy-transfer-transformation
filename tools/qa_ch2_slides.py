"""Phase 7 — Chapter 2 Slideshow automated QA (deterministic, structural/text).

Headless gates against the FINAL Ch2 slideshow (PPTX + rendered PDF):
  slide dims · font floors · safe-margin/overflow · visual-token compliance ·
  BQ/V/N anchor sync + order · exact note wording · answer reveal (no front-load) ·
  Reader page refs · vocab grouping · review A/B/C/D structure ·
  no stale INB/26-item language · no banned advanced vocabulary ·
  no Ch3 speed-energy leakage · render QA (PDF page count).

Run: .venv/bin/python tools/qa_ch2_slides.py
"""
import os, re, subprocess, sys
from pptx import Presentation
from pptx.util import Pt

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
PPTX = os.path.join(ROOT, "production", "slides", "02-ch2-moving-objects-have-energy-slides.pptx")
PDF = os.path.join(ROOT, "production", "slides", "02-ch2-moving-objects-have-energy-slides.pdf")
REPORT = os.path.join(ROOT, "qa", "reports", "PHASE_7_CH2_SLIDESHOW_QA.md")

# ---- FROZEN tokens (Ch2 + unit neutrals) -----------------------------------
ALLOWED_FILLS = {
    "1A1D26", "4A5268", "FFFFFF", "F6F7F9", "D8DCE4",   # unit neutrals
    "0F5A5C", "0E8A97", "5B6770", "FFFFFF",             # Ch2 Motion
    "E7F0F0", "E7F3F5",                                 # Ch2 teal/aqua tints (10%)
}

BQ_EXACT = "How are energy, change, and movement of objects related?"
ANCHORS = ["BQ", "V1", "N1", "V2", "N2", "N3"]
ANSWER_MAP = {
    "V1": ["position"],
    "V2": ["moving"],
    "N1": ["change", "outside"],
    "N2": ["motion"],
}
NOTE_WORDING = [
    "motion — the process of an object changing",                      # V1
    "energy of motion — the energy an object possesses while it is",    # V2
    "A moving object can cause a",                                      # N1
    "Players cause their own motion. A ball does not. Something else puts the ball into motion.",  # N1 caption (does not reveal the blank "outside")
    "Some objects cause their own motion. Other objects are put into motion by",  # N1
    "All moving objects have energy of",                                # N2
    "Energy is the ability to cause change. A moving object causes change, so it has energy.",  # N2 caption
    "Write or draw the change each moving object causes.",              # N3 caption
    "What changed?",                                                    # N2 observe
    "How do we know a moving object has energy?",                       # CHECK
    "Start with: “We know because the moving object causes",            # CHECK starter
    "Now I Think",                                                      # BQ return
]
READER_REFS = ["Reader p. 7", "Reader p. 8", "Reader p. 9", "Reader pp. 10–12", "Reader pp. 7–12"]
MODEL_CHAINS = [("MOVING OBJECT", "CHANGE", "EVIDENCE")]
N3_EXAMPLES = ["hammer (swinging)", "pendulum (swinging)", "wind (moving air)"]
N3_ANSWERS = ["drives a nail into the wood", "turns the clock gears", "spins the turbine blades"]

BANNED = [
    "26 items", "26 questions", "26Q", "26-item", "INB",
    "kinetic", "potential", "mechanical", "radiant", "momentum", "inertia",
    "compression wave", "gravitational", "proof", "proves",
]
# Ch3 scope terms that must not be taught in Ch2 (informational FAIL if used as content)
CH3_LEAK = ["speed", "faster", "fastest", "slower"]

# Reveal slides where bold+accent answers legitimately appear (V1/N1/V2/N2 check slides)
REVEAL_SLIDES = {10, 13, 16, 20}
ANSWER_WORDS = {"position", "moving", "change", "outside", "motion"}

# ---------------------------------------------------------------- helpers
def slide_texts(prs):
    out = []
    for i, s in enumerate(prs.slides, 1):
        parts = []
        for sh in s.shapes:
            if sh.has_text_frame:
                t = " ".join(r.text for p in sh.text_frame.paragraphs for r in p.runs)
                if t.strip():
                    parts.append(t)
        out.append((i, " ".join(parts)))
    return out

def all_runs(prs):
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        yield s, r

def shape_fills(prs):
    fills = set()
    for s in prs.slides:
        for sh in s.shapes:
            try:
                if sh.fill.type is not None and sh.fill.type == 1:  # solid
                    c = sh.fill.fore_color.rgb
                    fills.add(str(c))
            except Exception:
                pass
    return fills

def slide_dims(prs):
    w = prs.slide_width / 914400
    h = prs.slide_height / 914400
    return w, h

def overflow_shapes(prs):
    W = prs.slide_width; H = prs.slide_height
    bad = []
    tol = int(1.0 * 12700)
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            try:
                l, t, w, h = sh.left, sh.top, sh.width, sh.height
                if l is None or t is None or w is None or h is None:
                    continue
                if l < -tol or t < -tol or l + w > W + tol or t + h > H + tol:
                    bad.append((i, str(sh.shape_type), round(l/914400,2), round(t/914400,2),
                                round((l+w)/914400,2), round((t+h)/914400,2)))
            except Exception:
                continue
    return bad

def pdf_pages(path):
    if not os.path.exists(path):
        return -1
    info = subprocess.run(["pdfinfo", path], capture_output=True, text=True).stdout
    m = re.search(r"Pages:\s+(\d+)", info)
    return int(m.group(1)) if m else -1

# ---------------------------------------------------------------- gates
def gate_dims(prs):
    w, h = slide_dims(prs)
    ok = abs(w - 13.333) < 0.02 and abs(h - 7.5) < 0.02
    return ("PASS" if ok else "FAIL", f"{w:.3f} x {h:.3f} in (expect 13.333 x 7.5 / 16:9)")

def gate_font_floor(prs):
    sizes = [r.font.size.pt for _, r in all_runs(prs) if r.font.size and (r.text or "").strip()]
    if not sizes:
        return ("FAIL", "no sized runs")
    lo = min(sizes); hi = max(sizes)
    return ("PASS" if lo >= 12.0 else "FAIL",
            f"min {lo:.1f} pt (floor 12, chrome) · max {hi:.1f} pt")

def gate_body_floor(prs):
    vocab_big = any(r.font.size and r.font.size.pt >= 48 for _, r in all_runs(prs))
    return ("PASS" if vocab_big else "FAIL", "vocab term >= 48pt present")

def gate_overflow(prs):
    bad = overflow_shapes(prs)
    return ("PASS" if not bad else "FAIL", f"{len(bad)} shape(s) overflow slide: {bad or 'none'}")

def gate_tokens(prs):
    fills = shape_fills(prs)
    unknown = fills - ALLOWED_FILLS
    return ("PASS" if not unknown else "WARN", f"unknown fills: {sorted(unknown) or 'none'}")

def gate_anchors(prs):
    texts = slide_texts(prs)
    blob = "\n".join(t for _, t in texts)
    missing = [c for c in ANCHORS if c not in blob]
    idx = [blob.find(c) for c in ANCHORS]
    order_ok = all(a >= 0 for a in idx) and all(a < b for a, b in zip(idx, idx[1:]))
    if not order_ok:
        missing.append("anchor order")
    return ("PASS" if not missing else "FAIL", f"missing/out-of-order anchors: {missing or 'none'}")

def gate_bq_exact(prs):
    blob = "\n".join(t for _, t in slide_texts(prs))
    return ("PASS" if BQ_EXACT in blob else "FAIL", "BQ exact wording present")

def gate_note_wording(prs):
    blob = "\n".join(t for _, t in slide_texts(prs))
    missing = [w for w in NOTE_WORDING if w not in blob]
    return ("PASS" if not missing else "FAIL", f"missing note wording: {missing or 'none'}")

def gate_answers_revealed(prs):
    blob = "\n".join(t for _, t in slide_texts(prs))
    missing = []
    for anchor, answers in ANSWER_MAP.items():
        for a in answers:
            if a not in blob:
                missing.append(f"{anchor}:{a}")
    return ("PASS" if not missing else "FAIL", f"answers missing: {missing or 'none'}")

def gate_no_front_load(prs):
    hits = []
    for s, r in all_runs(prs):
        if not r.text.strip():
            continue
        t = r.text.strip().lower().rstrip(".")
        if t in ANSWER_WORDS:
            idx = list(prs.slides).index(s) + 1
            try:
                bold = r.font.bold
                color = str(r.font.color.rgb) if r.font.color and r.font.color.type is not None else None
                if bold and color == "0E8A97":
                    hits.append((idx, t))
            except Exception:
                hits.append((idx, t))
    bad = [(i, t) for i, t in hits if i not in REVEAL_SLIDES]
    return ("PASS" if not bad else "FAIL", f"front-loaded (bold+accent) answers on wrong slides: {bad or 'none'}")

def gate_reader_refs(prs):
    blob = "\n".join(t for _, t in slide_texts(prs))
    present = [r for r in READER_REFS if r in blob]
    return ("PASS" if len(present) >= 5 else "WARN", f"reader refs found: {present}")

def gate_vocab_grouping(prs):
    texts = slide_texts(prs)
    v1 = next((t for i, t in texts if "the process of an object changing position" in t), None)
    v2 = next((t for i, t in texts if "the energy an object possesses while it is moving" in t), None)
    ok = v1 is not None and v2 is not None
    return ("PASS" if ok else "FAIL", "V1 (motion) and V2 (energy of motion) groupings present")

def gate_model_chains(prs):
    blob = "\n".join(t for _, t in slide_texts(prs))
    missing = []
    for chain in MODEL_CHAINS:
        if not all(c in blob for c in chain):
            missing.append("→".join(chain))
    return ("PASS" if not missing else "FAIL", f"model chains missing: {missing or 'none'}")

def gate_n3_examples(prs):
    blob = "\n".join(t for _, t in slide_texts(prs))
    missing = [e for e in N3_EXAMPLES if e not in blob]
    return ("PASS" if not missing else "FAIL", f"N3 examples missing: {missing or 'none'}")

def gate_n3_answers(prs):
    blob = "\n".join(t for _, t in slide_texts(prs))
    missing = [a for a in N3_ANSWERS if a not in blob]
    return ("PASS" if not missing else "FAIL", f"N3 sample answers missing: {missing or 'none'}")

def gate_review_structure(prs):
    texts = slide_texts(prs)
    q_slides = []
    for i, t in texts:
        if re.search(r"Review \d+ —", t) and "answer" not in t:
            q_slides.append((i, t))
    bad = []
    for i, t in q_slides:
        for letter in ("A", "B", "C", "D"):
            if letter not in t:
                bad.append(f"slide {i} missing option {letter}")
    n = len(q_slides)
    ok = 4 <= n <= 6 and not bad
    return ("PASS" if ok else "FAIL", f"{n} review questions; {bad or 'all have A/B/C/D'}")

def gate_banned(prs):
    blob = "\n".join(t for _, t in slide_texts(prs)).lower()
    hits = [b for b in BANNED if b in blob]
    return ("PASS" if not hits else "FAIL", f"banned language: {sorted(set(hits)) or 'none'}")

def gate_ch3_leak(prs):
    blob = "\n".join(t for _, t in slide_texts(prs)).lower()
    hits = [w for w in CH3_LEAK if re.search(r"\b" + w + r"\b", blob)]
    return ("PASS" if not hits else "FAIL", f"Ch3 speed-energy leakage: {sorted(set(hits)) or 'none'}")

def gate_transfer_transform_scope(prs):
    blob = "\n".join(t for _, t in slide_texts(prs))
    offenders = []
    for m in re.finditer(r"[Tt]ransfer|[Tt]ransform(ation)?", blob):
        start = max(0, m.start() - 30); end = min(len(blob), m.end() + 30)
        ctx = blob[start:end]
        if "Energy Transfer & Transformation" not in ctx:
            offenders.append(ctx.replace("\n", " "))
    return ("PASS" if not offenders else "WARN",
            f"transfer/transform outside unit title: {offenders or 'none'}")

def gate_render(prs):
    pages = pdf_pages(PDF)
    n_slides = len(prs.slides._sldIdLst)
    ok = pages == n_slides
    return ("PASS" if ok else "FAIL", f"PDF pages {pages} == slides {n_slides}")

def gate_pacing(prs):
    import ch2_pacing as P
    probs = P.check_plan()
    missing = []
    for i, s in enumerate(prs.slides, 1):
        note = ""
        if s.has_notes_slide and s.notes_slide.notes_text_frame:
            note = s.notes_slide.notes_text_frame.text
        if "PLAN " not in note:
            missing.append(i)
        elif i in P.S and ("TIME %g min" % P.S[i]["minutes"]) not in note:
            missing.append(i)
    ok = not probs and not missing
    return ("PASS" if ok else "FAIL", "plan problems: %s · slides missing PLAN/TIME: %s" % (probs or "none", missing or "none"))

def gate_pre_reveal(prs):
    import qa_ch2_guards as G
    faces = []
    for s in prs.slides:
        bits = []
        for sh in s.shapes:
            if sh.has_text_frame:
                bits.append(sh.text_frame.text)
        faces.append("\n".join(bits))
    hits = G.leak_outside(faces[:12])
    return ("PASS" if not hits else "FAIL", "outside before reveal on slides %s" % (hits or "none"))

def gate_week(prs):
    blob = "\n".join(t for _, t in slide_texts(prs))
    bad = "WEEK 2" in blob or "Block 1" in blob or "Model It" in blob
    return ("PASS" if (not bad and "WEEK 1" in blob) else "FAIL", "Week 1 on the title; no Week 2 / Block 1 / Model It")

def gate_slide_count(prs):
    n = len(prs.slides._sldIdLst)
    return ("PASS" if n >= 30 else "WARN", f"{n} slides")

def main():
    import argparse
    global ROOT, PPTX, PDF, REPORT
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=None)
    a, _unknown = ap.parse_known_args()
    if a.root:
        ROOT = os.path.abspath(a.root)
        PPTX = os.path.join(ROOT, "production", "slides", "02-ch2-moving-objects-have-energy-slides.pptx")
        PDF = os.path.join(ROOT, "production", "slides", "02-ch2-moving-objects-have-energy-slides.pdf")
        REPORT = os.path.join(ROOT, "qa", "reports", "PHASE_7_CH2_SLIDESHOW_QA.md")
    prs = Presentation(PPTX)
    gates = [
        ("slide dimensions (16:9)", gate_dims(prs)),
        ("font floor >= 12pt (chrome)", gate_font_floor(prs)),
        ("vocab term >= 48pt", gate_body_floor(prs)),
        ("no shape overflow", gate_overflow(prs)),
        ("visual-token compliance", gate_tokens(prs)),
        ("BQ/V/N anchors + order", gate_anchors(prs)),
        ("BQ exact wording", gate_bq_exact(prs)),
        ("note wording mirror", gate_note_wording(prs)),
        ("answers revealed (check slides)", gate_answers_revealed(prs)),
        ("no front-loaded answers", gate_no_front_load(prs)),
        ("Reader page refs", gate_reader_refs(prs)),
        ("vocab grouping (V1/V2)", gate_vocab_grouping(prs)),
        ("model chain (N2)", gate_model_chains(prs)),
        ("N3 three examples", gate_n3_examples(prs)),
        ("N3 sample answers", gate_n3_answers(prs)),
        ("review A/B/C/D structure", gate_review_structure(prs)),
        ("no banned/stale language", gate_banned(prs)),
        ("no Ch3 speed-energy leakage", gate_ch3_leak(prs)),
        ("transfer/transform scope", gate_transfer_transform_scope(prs)),
        ("render QA (PDF pages)", gate_render(prs)),
        ("slide count", gate_slide_count(prs)),
        ("pacing plan + speaker PLAN lines", gate_pacing(prs)),
        ("title says Week 1", gate_week(prs)),
        ("no outside before N1 reveal", gate_pre_reveal(prs)),
    ]
    print("=" * 76)
    print("CH2 SLIDESHOW — AUTOMATED QA (Phase 7)")
    print("=" * 76)
    overall = "PASS"
    for name, (s, note) in gates:
        if s == "FAIL":
            overall = "FAIL"
        elif s == "WARN" and overall == "PASS":
            overall = "WARN"
        print(f"[{s:4}] {name:34} {note}")
    print("=" * 76)
    print(f"OVERALL: {overall}")

    os.makedirs(os.path.dirname(REPORT), exist_ok=True)
    with open(REPORT, "w") as f:
        f.write("# Phase 7 — Chapter 2 Slideshow Automated QA\n\n")
        f.write(f"Overall: **{overall}**\n\n")
        f.write(f"Artifact: `production/slides/02-ch2-moving-objects-have-energy-slides.pptx`\n\n")
        for name, (s, note) in gates:
            f.write(f"- [{s}] {name} — {note}\n")
    print(f"report: {REPORT}")
    if overall == "FAIL":
        sys.exit(1)

if __name__ == "__main__":
    main()
