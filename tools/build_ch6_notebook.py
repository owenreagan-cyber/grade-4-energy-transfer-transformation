"""Build Chapter 6 Interactive Science Notebook (5 inserts)."""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from ch4_common import *
from ch4_common import _set_run

OUT = os.path.join(os.path.dirname(__file__), '..', 'production', 'student-notebook', 'chapter-06')


def _border(cell, color="666666", dashed=True):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'dashed' if dashed else 'single')
        el.set(qn('w:sz'), '12')
        el.set(qn('w:color'), color)
        borders.append(el)
    tcPr.append(borders)


def insert_card(d, instruction):
    add_par(d, instruction, size=9, bold=True, color=GRAY, space_after=2)
    t = d.add_table(rows=1, cols=1)
    t.alignment = 1
    cell = t.cell(0, 0)
    cell.width = Inches(7.0)
    _border(cell)
    return cell


def cell_add(cell, text="", size=11, bold=False, italic=False, color=INK, align=None,
             before=0, after=4):
    p = cell.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if text:
        _set_run(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return p


def cell_title(cell, text, sub):
    cell_add(cell, text, size=13, bold=True, color=HEADER, after=0)
    cell_add(cell, sub, size=10, italic=True, color=GRAY, after=4)


def cell_section(cell, text):
    cell_add(cell, text, size=12, bold=True, color=ACCENT, before=6, after=2)


def build():
    d = Document()
    setup_portrait(d)
    add_par(d, "Interactive Science Notebook — Chapter 6 · Energy Transformations + Engineering Design",
            size=16, bold=True, color=HEADER, space_after=2)
    add_par(d, "Trim each insert on the dashed line and glue into your notebook. "
              "Every interactive part has a flat, low-motor alternative.",
            size=10, italic=True, color=GRAY, space_after=10)

    # ---------------- INSERT 1: Transfer vs Transformation ----------------
    cell = insert_card(d, "INSERT 1 — Transfer vs Transformation    ·    trim on the dashed line    ·    glue the BACK")
    cell_title(cell, "Transfer vs Transformation", "Chapter 6 · Week 7")
    cell_section(cell, "TWO WAYS ENERGY CHANGES THINGS")
    cell_add(cell, "TRANSFER = energy MOVES from one object or place to another.", size=12, bold=True, after=0)
    cell_add(cell, "  example: warm water warms a cool spoon (heat moves).", size=11, after=4)
    cell_add(cell, "TRANSFORMATION = energy CHANGES FORM.", size=12, bold=True, after=0)
    cell_add(cell, "  example: a flashlight changes stored energy into light.", size=11, after=4)
    cell_add(cell, "BOTH = energy MOVES AND CHANGES FORM (most systems).", size=12, bold=True, color=ACCENT, after=0)
    cell_add(cell, "  example: falling dominoes — motion moves AND some motion becomes sound.", size=11, after=6)
    cell_section(cell, "MY OWN EXAMPLES")
    cell_add(cell, "TRANSFER: ____________________", size=12, after=4)
    cell_add(cell, "TRANSFORMATION: ____________________", size=12, after=4)
    cell_add(cell, "BOTH: ____________________", size=12, after=4)

    # ---------------- INSERT 2: Energy System Model ----------------
    cell = insert_card(d, "INSERT 2 — Energy System Model    ·    trim on the dashed line    ·    glue the BACK")
    cell_title(cell, "Energy System Model", "Chapter 6 · Week 7")
    cell_section(cell, "THE MODEL")
    cell_add(cell, "INPUT", size=12, bold=True, after=0)
    cell_add(cell, "→ CHANGE IN FORM (TRANSFORMATION)", size=12, bold=True, after=0)
    cell_add(cell, "→ OUTPUT / OBSERVABLE CHANGE", size=12, bold=True, color=ACCENT, after=6)
    cell_section(cell, "DRAW ONE DEVICE")
    cell_add(cell, "Pick one device (flashlight, buzzer, motor, wind-up toy). Draw it and label the chain.",
             size=11, after=2)
    p = cell.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(p.add_run("[  draw your device + label INPUT → TRANSFORMATION → OUTPUT  ]"),
             size=10, italic=True, color=GRAY)
    cell_section(cell, "MY CHAIN")
    cell_add(cell, "INPUT: ____________________", size=12, after=2)
    cell_add(cell, "TRANSFORMATION: ____________________", size=12, after=2)
    cell_add(cell, "OUTPUT: ____________________", size=12, after=4)

    # ---------------- INSERT 3: Device Detective ----------------
    cell = insert_card(d, "INSERT 3 — Device Detective    ·    trim on the dashed line    ·    glue the BACK")
    cell_title(cell, "Device Detective", "Chapter 6 · Week 7 · analyze one device")
    cell_section(cell, "ANALYZE ONE DEVICE")
    cell_add(cell, "1. What goes IN? ____________________", size=12, after=2)
    cell_add(cell, "2. What CHANGES? ____________________", size=12, after=2)
    cell_add(cell, "3. What comes OUT? ____________________", size=12, after=4)
    cell_section(cell, "VERDICT (circle one)")
    cell_add(cell, "TRANSFER     TRANSFORMATION     BOTH", size=12, bold=True, after=6)
    cell_section(cell, "WHY?")
    cell_add(cell, "I circled this because ____________________", size=11, after=4)
    cell_add(cell, "My evidence: I observed ____________________", size=11, after=4)

    # ---------------- INSERT 4: Engineering Design ----------------
    cell = insert_card(d, "INSERT 4 — Engineering Design    ·    trim on the dashed line    ·    glue the BACK")
    cell_title(cell, "Engineering Design", "Chapter 6 · Week 8")
    cell_section(cell, "THE CYCLE")
    cell_add(cell, "DESIGN → BUILD → TEST → REFINE", size=12, bold=True, color=ACCENT, after=6)
    cell_section(cell, "MY PROBLEM")
    cell_add(cell, "Problem: ____________________", size=12, after=4)
    cell_section(cell, "CRITERIA (what it must do)")
    cell_add(cell, "1. ____________________", size=12, after=2)
    cell_add(cell, "2. ____________________", size=12, after=2)
    cell_add(cell, "3. ____________________", size=12, after=4)
    cell_section(cell, "CONSTRAINTS (limits)")
    cell_add(cell, "1. ____________________", size=12, after=2)
    cell_add(cell, "2. ____________________", size=12, after=4)
    cell_section(cell, "BLUEPRINT")
    cell_add(cell, "Draw + label your device (trigger, output, energy change).", size=11, after=2)
    p = cell.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(p.add_run("[  draw + label your blueprint here  ]"), size=10, italic=True, color=GRAY)

    # ---------------- INSERT 5: Test -> Change -> Retest (flap) ----------------
    cell = insert_card(d, "INSERT 5 — Test → Change → Retest    ·    trim on the dashed line    ·    "
                          "fold the flap DOWN and glue the BASE    ·    flat option: write under the flap")
    cell_title(cell, "Test → Change → Retest", "Chapter 6 · Week 8 · lift the flap to check")
    cell_section(cell, "TEST 1 (on the flap)")
    cell_add(cell, "Did it trigger? ____  What output? ____  What failed? ____", size=11, after=6)
    cell_section(cell, "WHAT WE CHANGED + WHY (under the flap)")
    cell_add(cell, "We changed ONLY ____________________ because ____________________", size=11, after=4)
    cell_section(cell, "TEST 2 RESULT")
    cell_add(cell, "Did the change help? ____  Evidence: ____________________", size=11, after=6)
    cell_section(cell, "MY ENERGY PATHWAY")
    cell_add(cell, "INPUT → TRANSFORMATION → OUTPUT:", size=11, bold=True, after=0)
    cell_add(cell, "____________________ → ____________________ → ____________________", size=12, after=4)

    d.save(os.path.join(OUT, "06-ch6-interactive-notebook.docx"))
    print("built notebook")


if __name__ == "__main__":
    build()
