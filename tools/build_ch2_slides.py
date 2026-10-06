"""Build Chapter 2 "Moving Objects Have Energy" — GOLD-STANDARD slide deck (PPTX, 16:9).

Phase 7. Consumes the FROZEN Unit 1 Visual System v1.0 tokens and mirrors the
FINALIZED Chapter 2 Guided Notes (immutable BQ/V/N anchors) exactly.
Ch2 theme = Motion (primary #0F5A5C · accent #0E8A97 · neutral #5B6770 · surface #FFFFFF).

Authoritative source: CKSci Student Reader Ch.2 "Moving Objects Have Energy" (pp. 7-12)
+ Teacher Guide Lesson 3. See docs/chapter-02/CH2_SLIDESHOW_HANDOFF.md.

Run: .venv/bin/python tools/build_ch2_slides.py
Then: soffice --headless --convert-to pdf  (see render step)
"""
import json, os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import ch2_pacing as PACE
_ANCHORS = json.load(open(os.path.join(os.path.dirname(__file__), "..", "curriculum/chapter-02/ch2_anchor_map.json"), encoding="utf-8"))["anchors"]

def _lead(code, index=0):
    return _ANCHORS[code]["lines"][index].split("____", 1)[0].rstrip()

def _trail(code, index=0):
    parts = _ANCHORS[code]["lines"][index].split("____", 1)
    return parts[1].lstrip() if len(parts) > 1 else ""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
OUT = os.path.join(ROOT, "production", "slides")

# ---------------------------------------------------------------- FROZEN Ch2 tokens
PRIMARY = RGBColor(0x0F, 0x5A, 0x5C)   # deep teal
ACCENT  = RGBColor(0x0E, 0x8A, 0x97)   # aqua
NEUTRAL = RGBColor(0x5B, 0x67, 0x70)   # slate
SURFACE = RGBColor(0xFF, 0xFF, 0xFF)   # white
INK     = RGBColor(0x1A, 0x1D, 0x26)
INK_SOFT= RGBColor(0x4A, 0x52, 0x68)
PAPER   = RGBColor(0xFF, 0xFF, 0xFF)
MIST    = RGBColor(0xF6, 0xF7, 0xF9)
RULE    = RGBColor(0xD8, 0xDC, 0xE4)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)

# 10% tints (light fills) — teal primary tint + aqua accent tint
TEAL_TINT = RGBColor(0xE7, 0xF0, 0xF0)   # primary 10% tint
AQUA_TINT = RGBColor(0xE7, 0xF3, 0xF5)   # accent 10% tint

FONT = "Avenir Next"

# Type scale (frozen slide tokens + Ch1 exemplar)
T_MAJOR  = 48   # chapter title
T_TITLE  = 40   # slide title
T_HEAD   = 32   # notes section heading
T_BODY   = 30   # instructional body
T_BODY_S = 28   # body (denser lines)
T_OPT    = 28   # MCQ options
T_DEF    = 26   # vocab definition on card
T_KICK   = 24   # kicker / subtitle (matches Chapter 1; 24 pt floor)
T_PAGE   = 16   # Reader page refs / slide number (chrome)
T_TAB    = 12   # chapter tab / chrome
T_HUD    = 24   # HUD bubble text students must read

CHAPTER_TAB = "CHAPTER 2 · MOVING OBJECTS HAVE ENERGY"

# ---------------------------------------------------------------- presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

SLIDE_NO = [0]  # mutable counter

def new_slide(bg=SURFACE):
    SLIDE_NO[0] += 1
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background(); r.shadow.inherit = False
    return s

# ---------------------------------------------------------------- low-level helpers
def rect(s, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line_color=None, line_w=None, rounded=False):
    sp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else shape,
                            Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        try:
            sp.adjustments[0] = 0.08
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line_color is not None:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(line_w or 1.25)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp

def tb(s, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    return box, tf

def run(p, text, size=T_BODY, color=INK, bold=False, italic=False, underline=False, font=FONT, caps=False):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color; r.font.name = font
    if underline:
        r.font.underline = True
    if caps:
        r.font.caps = True
    return r

def text(s, x, y, w, h, lines, size=T_BODY, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False, gap=10, caps=False):
    if isinstance(lines, str):
        lines = [lines]
    box, tf = tb(s, x, y, w, h, anchor=anchor)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(gap)
        run(p, ln, size=size, color=color, bold=bold, italic=italic, caps=caps)
    return box

def rich(s, x, y, w, h, segs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, gap=12):
    box, tf = tb(s, x, y, w, h, anchor=anchor)
    for i, para_segs in enumerate(segs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(gap)
        for (txt, fmt) in para_segs:
            run(p, txt, size=fmt.get('size', T_BODY), color=fmt.get('color', INK),
                bold=fmt.get('bold', False), italic=fmt.get('italic', False),
                underline=fmt.get('underline', False))
    return box

def blank_str(n):
    return "_" * n

# ---------------------------------------------------------------- chrome / header
def chapter_tab(s, x=0.55, y=0.40, w=3.85, h=0.44):
    sp = rect(s, x, y, w, h, PRIMARY, rounded=True)
    tf = sp.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.12); tf.margin_right = Inches(0.12)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, CHAPTER_TAB, size=T_TAB, color=WHITE, bold=True)
    return sp

def mode_badge(s, mode, x=12.5, y=0.36, w=0.6, h=0.52, fill=PRIMARY, color=WHITE):
    sp = rect(s, x, y, w, h, fill, rounded=True)
    tf = sp.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, mode, size=T_HUD, color=color, bold=True)
    return sp

HUD_W = {  # 24pt bold widths, right-aligned clear of the mode badge
    "TURN & TALK": 2.75, "OBSERVE": 2.15, "PREDICT": 2.0, "EVIDENCE": 2.25,
    "WRITE": 1.6, "CHECK": 1.7, "RECITE": 1.75, "TRACE": 1.65, "DRAW": 1.5,
    "EXPLAIN": 1.95, "VOTE": 1.4, "WATCH": 1.7, "DATA": 1.55, "ENGINEER": 2.25,
}
def hud_bubble(s, text_label, x=None, y=0.36, h=0.52):
    w = HUD_W.get(text_label, 2.2)
    if x is None:
        x = 12.4 - w
    sp = rect(s, x, y, w, h, ACCENT, rounded=True)
    tf = sp.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.08)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, text_label, size=T_HUD, color=WHITE, bold=True)
    return sp

def slide_number(s, n=None):
    n = n or SLIDE_NO[0]
    text(s, 12.2, 7.02, 0.9, 0.4, str(n), size=T_PAGE, color=INK_SOFT, align=PP_ALIGN.RIGHT)

def anchor_chip(s, code, x=0.55, y=1.02, w=0.78, h=0.5):
    sp = rect(s, x, y, w, h, ACCENT, rounded=True)
    tf = sp.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, code, size=24, color=WHITE, bold=True)
    return sp

def header(s, title, kicker=None, mode=None, hud=None, anchor=None, title_size=T_TITLE):
    chapter_tab(s)
    if mode:
        mode_badge(s, mode)
    if hud:
        hud_bubble(s, hud)
    tx = 0.55
    if anchor:
        anchor_chip(s, anchor)
        tx = 1.5
    text(s, tx, 1.02, 12.75 - tx, 0.7, title, size=title_size, color=PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if kicker:
        text(s, 0.55, 1.72, 11.7, 0.4, kicker, size=T_KICK, color=INK_SOFT, italic=True)
    slide_number(s)

# ---------------------------------------------------------------- model / diagrams
def model_strip(s, items, x=0.75, y=None, w=11.8, box_h=0.82, gap=0.5, size=26, fills=None):
    n = len(items)
    box_w = (w - (n - 1) * gap) / n
    cx = x
    for i, it in enumerate(items):
        f = (fills[i] if fills else (PRIMARY if i % 2 == 0 else NEUTRAL))
        sp = rect(s, cx, y, box_w, box_h, f, rounded=True)
        tf = sp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.08)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        run(p, it, size=size, color=WHITE, bold=True)
        if i < n - 1:
            ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(cx + box_w + 0.02),
                                    Inches(y + box_h/2 - 0.16), Inches(gap - 0.04), Inches(0.32))
            ar.fill.solid(); ar.fill.fore_color.rgb = ACCENT
            ar.line.fill.background(); ar.shadow.inherit = False
        cx += box_w + gap
    return box_h

def before_after(s, x=0.75, y=2.5, w=5.6, h=2.6, left="BEFORE", right="AFTER", left_hint="", right_hint=""):
    gap = 0.35
    bw = (w - gap) / 2
    for idx, (lab, hint) in enumerate([(left, left_hint), (right, right_hint)]):
        cx = x + idx * (bw + gap)
        rect(s, cx, y, bw, h, PAPER, line_color=PRIMARY, line_w=1.25, rounded=True)
        text(s, cx + 0.2, y + 0.15, bw - 0.4, 0.4, lab, size=24, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
        if hint:
            text(s, cx + 0.2, y + 0.62, bw - 0.4, h - 0.8, hint, size=T_BODY_S, color=INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- vocab card
def vocab_card(s, x, y, w, h, term, definition, page, visual=None):
    rect(s, x, y, w, h, PAPER, line_color=RULE, line_w=1.25, rounded=True)
    rect(s, x, y, w, 0.16, PRIMARY)                 # teal top band
    pill = rect(s, x + w - 1.15, y + 0.5, 0.9, 0.4, TEAL_TINT, rounded=True)
    tf = pill.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, "TEACH", size=14, color=PRIMARY, bold=True)
    text(s, x + 0.35, y + 0.42, w - 1.7, 0.4, "VOCABULARY", size=18, color=NEUTRAL, bold=True, caps=True)
    text(s, x + 0.35, y + 0.78, w - 0.7, 1.05, term, size=54, color=PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.35, y + 1.95, w - 0.7, 0.95, definition, size=T_DEF, color=INK)
    if visual:
        rect(s, x + 0.35, y + 3.0, w - 0.7, h - 3.7, AQUA_TINT, line_color=ACCENT, line_w=0.75, rounded=True)
        visual(s, x + 0.35, y + 3.0, w - 0.7, h - 3.7)
    text(s, x + 0.35, y + h - 0.55, w - 0.7, 0.4, page, size=T_PAGE, color=INK_SOFT, bold=True)

def card_visual_chain(panel_words):
    """Return a visual-draw function that renders a small arrow chain inside a panel."""
    def draw(s, x, y, w, h):
        n = len(panel_words)
        cw = (w - 0.5 * (n - 1) - 0.3) / n
        cx = x + 0.15
        cy = y + h / 2 - 0.28
        for i, word in enumerate(panel_words):
            sp = rect(s, cx, cy, cw, 0.56, PRIMARY if i == 0 else NEUTRAL, rounded=True)
            tf = sp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = Inches(0.05); tf.margin_right = Inches(0.05)
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            run(p, word, size=15, color=WHITE, bold=True)
            if i < n - 1:
                ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(cx + cw + 0.02),
                                        Inches(cy + 0.14), Inches(0.46), Inches(0.28))
                ar.fill.solid(); ar.fill.fore_color.rgb = ACCENT
                ar.line.fill.background(); ar.shadow.inherit = False
            cx += cw + 0.5
    return draw

# ---------------------------------------------------------------- notes mirror
def notes_heading(s, code, title_text, reader_ref):
    wide = len(code) > 2
    anchor_chip(s, code, w=1.55 if wide else 0.78)
    tx = 2.25 if wide else 1.5
    text(s, tx, 1.02, 12.2 - tx, 0.6, title_text, size=T_HEAD, color=PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, tx, 1.62, 8.0, 0.35, reader_ref, size=T_KICK, color=INK_SOFT, italic=True)

def write_line(s, x, y, w, color=INK_SOFT):
    rect(s, x, y, w, 0.035, color)

def _plan_line(n):
    info = PACE.S[n]
    mins = "%g" % info["minutes"]
    if info["session"] == 0:
        return "PLAN OPTIONAL · review only, never inside a required block · not Week 1 Friday (that day is Chapter 1) · %s min" % mins
    sess = PACE.SESSIONS[info["session"]]
    tag = "REQUIRED" if info["status"] == "REQ" else "OPTIONAL"
    extra = ""
    if n == PACE.REENTRY_SLIDE:
        extra = " · RE-ENTRY: Session 2 also opens on this slide (%g min)" % PACE.REENTRY_MIN[2]
    return "PLAN Session %d (%s) · %s · %s min%s" % (info["session"], sess["day"], tag, mins, extra)

def set_note(slide, lines):
    n = SLIDE_NO[0]
    out = [ln for ln in lines if not ln.startswith("TIME ") and not ln.startswith("PLAN ")]
    out.insert(1, "TIME %g min" % PACE.S[n]["minutes"])
    out.insert(2, _plan_line(n))
    slide.notes_slide.notes_text_frame.text = "\n".join(out)

# ================================================================================
# SLIDE 1 — TITLE
# ================================================================================
s = new_slide(PRIMARY)
sp = rect(s, 0.9, 0.75, 4.9, 0.55, ACCENT, rounded=True)
tf = sp.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = Inches(0.15); tf.margin_right = Inches(0.15)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
run(p, "CHAPTER 2 · GRADE 4 · WEEK 1", size=14, color=WHITE, bold=True)
text(s, 0.9, 1.75, 11.5, 1.0, "Moving Objects Have Energy", size=T_MAJOR, color=WHITE, bold=True)
text(s, 0.9, 2.85, 11.5, 0.5, "Motion  ·  How do we know a moving object has energy?", size=T_KICK, color=RGBColor(0xD8, 0xDC, 0xE4))
model_strip(s, ["MOVING OBJECT", "CHANGE", "EVIDENCE"], x=0.9, y=4.2, w=9.2, box_h=0.95, gap=0.45, size=24)
text(s, 0.9, 5.9, 11.5, 0.5, "Unit 1 · Energy Transfer & Transformation", size=T_KICK, color=RGBColor(0xD8, 0xDC, 0xE4), italic=True)
slide_number(s)
set_note(s, ["MODE Title · orient", "TIME 2 min", "STUDENT RESOURCE Guided Notes (not yet open)",
             "ASK What can a moving object do to other things?", "LISTEN FOR words like move, hit, change",
             "NEXT Open the Big Question."])

# ================================================================================
# SLIDE 2 — BIG QUESTION (intro)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s)
mode_badge(s, "BQ")
hud_bubble(s, "TURN & TALK")
text(s, 0.55, 1.1, 11.7, 0.5, "Big Question", size=T_KICK, color=NEUTRAL, bold=True, caps=True)
rect(s, 0.55, 1.7, 12.23, 1.9, PRIMARY, rounded=True)
text(s, 0.9, 1.95, 11.5, 1.4, "How are energy, change, and movement of objects related?",
     size=34, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.9, 4.0, 11.5, 0.5, "Keep it in mind all chapter. We will come back to it.", size=T_BODY_S, color=INK_SOFT, align=PP_ALIGN.CENTER)
text(s, 0.55, 4.9, 11.7, 0.6, "Think: What is ONE change a moving object can cause?", size=T_BODY, color=INK, align=PP_ALIGN.CENTER, bold=True)
slide_number(s)
set_note(s, ["MODE BQ · introduce", "TIME 3 min", "STUDENT RESOURCE Guided Notes p.1 (BQ box)",
             "ASK What change have you seen a moving object cause?", "LISTEN FOR hit, knock over, push, move",
             "REVEAL Do not answer the BQ yet — we collect evidence first.", "NEXT My first thinking."])

# ================================================================================
# SLIDE 3 — BQ · MY FIRST THINKING (notes mirror)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "WRITE")
notes_heading(s, "BQ", "Big Question — My first thinking", "Guided Notes p. 1")
text(s, 1.5, 2.3, 11.0, 1.0, "How are energy, change, and movement of objects related?", size=T_BODY, color=PRIMARY, bold=True)
text(s, 1.5, 3.5, 6.0, 0.55, "My first thinking", size=T_BODY, color=INK, bold=True)
write_line(s, 1.5, 4.1, 10.5)
write_line(s, 1.5, 4.8, 10.5)
text(s, 1.5, 5.55, 10.5, 0.5, "Write 1–2 quick ideas. No wrong answer yet.", size=T_BODY_S, color=INK_SOFT, italic=True)
slide_number(s)
set_note(s, ["MODE Notes · BQ anchor", "TIME 2 min", "STUDENT RESOURCE Guided Notes p.1 top",
             "ASK What change can a moving object cause? Write your first idea.", "LISTEN FOR students naming a change (not 'energy')",
             "MISCONCEPTION Accept any change; do not correct yet.", "NEXT Observe the marble collision."])

# ================================================================================
# SLIDE 4 — DEMO OBSERVE: marble before
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "D"); hud_bubble(s, "OBSERVE")
text(s, 0.55, 1.02, 11.7, 0.7, "Marble Collision — before", size=T_TITLE, color=PRIMARY, bold=True)
text(s, 0.55, 1.75, 11.7, 0.4, "One marble sits still. Another marble rolls toward it.", size=T_KICK, color=INK_SOFT, italic=True)
rect(s, 2.0, 2.5, 9.3, 3.3, PAPER, line_color=PRIMARY, line_w=1.25, rounded=True)
text(s, 2.3, 2.8, 8.7, 0.5, "BEFORE", size=T_HEAD, color=ACCENT, bold=True, caps=True)
text(s, 2.3, 3.5, 8.7, 1.6, "A target marble is still.\nA moving marble rolls toward it.",
     size=T_BODY, color=INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.55, 6.15, 11.7, 0.5, "Observe closely. What do you see BEFORE the marbles touch?", size=T_BODY, color=INK, align=PP_ALIGN.CENTER)
slide_number(s)
set_note(s, ["MODE Demo · observe", "TIME 2 min", "STUDENT RESOURCE none (live demo)",
             "ASK What does each marble look like right now?", "LISTEN FOR still, rolling, moving",
             "NEXT Predict what will happen."])

# ================================================================================
# SLIDE 5 — DEMO PREDICT
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "D"); hud_bubble(s, "PREDICT")
text(s, 0.55, 1.02, 11.7, 0.7, "Predict: what will happen?", size=T_TITLE, color=PRIMARY, bold=True)
text(s, 0.55, 1.75, 11.7, 0.4, "Think before the moving marble reaches the still marble.", size=T_KICK, color=INK_SOFT, italic=True)
rect(s, 1.5, 2.45, 10.3, 3.15, PAPER, line_color=PRIMARY, line_w=1.25, rounded=True)
text(s, 1.8, 2.65, 9.7, 1.15, "When the moving marble hits the still marble, the still marble will…", size=T_BODY, color=INK, bold=True)
text(s, 1.8, 4.0, 9.7, 0.7, "get  ____________________", size=T_BODY, color=INK)
text(s, 0.55, 6.05, 11.7, 0.5, "Make a prediction. Say it to your partner.", size=T_BODY, color=INK, align=PP_ALIGN.CENTER)
slide_number(s)
set_note(s, ["MODE Demo · predict", "TIME 2 min", "STUDENT RESOURCE none",
             "ASK What will happen to the still marble?", "LISTEN FOR 'it will start moving'",
             "REVEAL Do not confirm yet — run the demo.", "NEXT Observe the change."])

# ================================================================================
# SLIDE 6 — DEMO OBSERVE: what changed
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "D"); hud_bubble(s, "OBSERVE")
text(s, 0.55, 1.02, 11.7, 0.7, "What changed?", size=T_TITLE, color=PRIMARY, bold=True)
before_after(s, x=0.75, y=2.4, w=7.4, h=3.0,
             left="BEFORE", right="AFTER",
             left_hint="target marble still", right_hint="target marble now moving")
rect(s, 8.7, 2.4, 3.8, 3.0, AQUA_TINT, line_color=ACCENT, line_w=0.75, rounded=True)
text(s, 8.95, 2.6, 3.3, 0.5, "CHANGES", size=24, color=NEUTRAL, bold=True, caps=True)
text(s, 8.95, 3.2, 3.3, 2.0, "• the still marble starts to move", size=T_BODY_S, color=INK)
slide_number(s)
set_note(s, ["MODE Demo · observe", "TIME 3 min", "STUDENT RESOURCE none",
             "ASK Name every change you saw.", "LISTEN FOR 'the still marble moved'",
             "REVEAL The target marble changed from still to moving.", "NEXT Turn the change into evidence."])

# ================================================================================
# SLIDE 7 — DEMO EVIDENCE
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "D"); hud_bubble(s, "EVIDENCE")
text(s, 0.55, 1.02, 11.7, 0.7, "How do we know the marble had energy?", size=T_TITLE, color=PRIMARY, bold=True)
text(s, 0.55, 1.75, 11.7, 0.4, "Energy can be invisible. We look for the change.", size=T_KICK, color=INK_SOFT, italic=True)
model_strip(s, ["MOVING OBJECT", "CHANGE", "EVIDENCE"], x=0.75, y=2.7, w=11.8, box_h=1.0, gap=0.5, size=28)
text(s, 0.55, 4.3, 11.7, 1.4,
     "The still marble started to move.\nThat change is evidence that the moving marble had energy.",
     size=T_BODY, color=INK, align=PP_ALIGN.CENTER)
slide_number(s)
set_note(s, ["MODE Demo · evidence", "TIME 3 min", "STUDENT RESOURCE Reader p.9",
             "ASK How do we know the moving marble had energy if we can't see energy?", "LISTEN FOR 'we saw the change'",
             "MISCONCEPTION We see the change, not the energy itself.", "NEXT Vocabulary: motion."])

# ================================================================================
# SLIDE 8 — V1 VOCABULARY CARD (motion)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "V")
text(s, 0.55, 1.02, 11.7, 0.7, "Vocabulary — Motion", size=T_TITLE, color=PRIMARY, bold=True)
text(s, 0.55, 1.72, 11.7, 0.4, "Reader p. 7", size=T_KICK, color=INK_SOFT, italic=True)
vocab_card(s, 3.0, 2.25, 7.3, 4.5, "motion",
           "the process of an object changing position",
           "Reader p. 7",
           visual=card_visual_chain(["STILL", "MOVING"]))
slide_number(s)
set_note(s, ["MODE Vocabulary · TEACH", "TIME 3 min", "STUDENT RESOURCE Guided Notes V1 (Reader p.7)",
             "ASK When is an object in motion?", "LISTEN FOR 'when it changes position / moves'",
             "REVEAL motion = the process of an object changing position.", "NEXT Record V1 on notes."])

# ================================================================================
# SLIDE 9 — V1 NOTES (write)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "WRITE")
notes_heading(s, "V1", "Vocabulary — Motion", "Guided Notes · Reader p. 7")
rich(s, 1.5, 2.6, 11.0, 1.0, [
    [("motion — the process of an object changing ", {}), (blank_str(8), {"underline": True, "color": INK}), (" .", {})],
])
text(s, 1.5, 4.0, 11.0, 0.5, "Write the missing word in the blank.", size=T_BODY_S, color=INK_SOFT, italic=True)
slide_number(s)
set_note(s, ["MODE Notes · V1 anchor", "TIME 2 min", "STUDENT RESOURCE Guided Notes V1 (Reader p.7)",
             "ASK Fill in the blank.", "LISTEN FOR 'position'", "REVEAL Reveal after writing.", "NEXT Check answer."])

# ================================================================================
# SLIDE 10 — V1 NOTES (check/reveal)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "CHECK")
notes_heading(s, "V1", "Vocabulary — Motion", "Guided Notes · Reader p. 7")
rich(s, 1.5, 2.6, 11.0, 1.0, [
    [("motion — the process of an object changing ", {}), ("position", {"bold": True, "color": ACCENT}), (" .", {})],
])
text(s, 1.5, 4.0, 11.0, 0.5, "position", size=T_BODY_S, color=INK_SOFT, italic=True)
slide_number(s)
set_note(s, ["MODE Notes · V1 reveal", "TIME 1 min", "STUDENT RESOURCE Guided Notes V1",
             "REVEAL position", "NEXT Moving Objects Can Cause Changes (N1)."])

# ================================================================================
# SLIDE 11 — N1 CONCEPT: moving objects can cause changes
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "R"); hud_bubble(s, "RECITE")
text(s, 0.55, 1.02, 11.7, 0.7, "What can a moving object do?", size=T_TITLE, color=PRIMARY, bold=True)
text(s, 0.55, 1.75, 11.7, 0.4, "Reader p. 8 — think before you write.", size=T_KICK, color=INK_SOFT, italic=True)
rect(s, 0.75, 2.5, 11.8, 2.2, PRIMARY, rounded=True)
text(s, 1.2, 2.9, 10.9, 1.4, "Some objects start their own motion.\nA ball does not.", size=32, color=WHITE, bold=True,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.75, 5.0, 11.8, 1.0, _ANCHORS["N1"]["lines"][1],
     size=T_BODY_S, color=INK, align=PP_ALIGN.CENTER)
slide_number(s)
set_note(s, ["MODE Read · concept", "TIME 3 min", "STUDENT RESOURCE Reader p.8",
             "ASK How does a ball get put into motion?", "LISTEN FOR a kick, a throw, or another push",
             "REVEAL Do not say the notes words yet. Students write first.", "NEXT Record N1 on notes."])

# ================================================================================
# SLIDE 12 — N1 NOTES (write)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "WRITE")
notes_heading(s, "N1", "Write the Missing Words", "Guided Notes · Reader p. 8")
rich(s, 1.5, 2.35, 11.0, 0.8, [
    [(_lead("N1", 0) + " ", {}), (blank_str(6), {"underline": True, "color": INK}), (" " + _trail("N1", 0), {})],
])
text(s, 1.5, 3.25, 11.0, 0.5, _ANCHORS["N1"]["lines"][1],
     size=T_BODY_S, color=INK, italic=True)
rich(s, 1.5, 4.15, 11.0, 0.8, [
    [(_lead("N1", 2) + " ", {}),
     (blank_str(7), {"underline": True, "color": INK}), (" " + _trail("N1", 2), {})],
])
slide_number(s)
set_note(s, ["MODE Notes · N1 anchor", "TIME 4 min", "STUDENT RESOURCE Guided Notes N1 (Reader p.8)",
             "ASK Fill in the two blanks.", "LISTEN FOR the two missing words", "REVEAL Reveal after writing.", "NEXT Check answers."])

# ================================================================================
# SLIDE 13 — N1 NOTES (check/reveal)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "CHECK")
notes_heading(s, "N1", "Moving Objects Can Cause Changes", "Guided Notes · Reader p. 8")
rich(s, 1.5, 2.35, 11.0, 0.8, [
    [("A moving object can cause a ", {}), ("change", {"bold": True, "color": ACCENT}), (" .", {})],
])
text(s, 1.5, 3.25, 11.0, 0.5, _ANCHORS["N1"]["lines"][1],
     size=T_BODY_S, color=INK, italic=True)
rich(s, 1.5, 4.15, 11.0, 0.8, [
    [("Some objects cause their own motion. Other objects are put into motion by ", {}),
     ("outside", {"bold": True, "color": ACCENT}), (" motion.", {})],
])
text(s, 1.5, 5.15, 11.0, 0.4, "change · outside", size=T_BODY_S, color=INK_SOFT, italic=True)
slide_number(s)
set_note(s, ["MODE Notes · N1 reveal", "TIME 1 min", "STUDENT RESOURCE Guided Notes N1",
             "REVEAL change · outside", "NEXT Vocabulary: energy of motion (V2)."])

# ================================================================================
# SLIDE 14 — V2 VOCABULARY CARD (energy of motion)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "V")
text(s, 0.55, 1.02, 11.7, 0.7, "Vocabulary — Energy of Motion", size=T_TITLE, color=PRIMARY, bold=True)
text(s, 0.55, 1.72, 11.7, 0.4, "Reader p. 9", size=T_KICK, color=INK_SOFT, italic=True)
vocab_card(s, 3.0, 2.25, 7.3, 4.5, "energy of motion",
           "the energy an object possesses while it is moving",
           "Reader p. 9",
           visual=card_visual_chain(["MOVING", "CHANGE"]))
slide_number(s)
set_note(s, ["MODE Vocabulary · TEACH", "TIME 3 min", "STUDENT RESOURCE Guided Notes V2 (Reader p.9)",
             "ASK Does a rolling marble have energy? How do you know?", "LISTEN FOR 'yes — it can cause a change'",
             "REVEAL energy of motion = energy an object has while moving.", "NEXT Record V2."])

# ================================================================================
# SLIDE 15 — V2 NOTES (write)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "WRITE")
notes_heading(s, "V2", "Vocabulary — Energy of Motion", "Guided Notes · Reader p. 9")
rich(s, 1.5, 2.65, 11.0, 1.15, [
    [("energy of motion — the energy an object possesses while it is ", {}),
     (blank_str(6), {"underline": True, "color": INK}), (" .", {})],
])
text(s, 1.5, 4.15, 11.0, 0.5, "Write the missing word in the blank.", size=T_BODY_S, color=INK_SOFT, italic=True)
slide_number(s)
set_note(s, ["MODE Notes · V2 anchor", "TIME 2 min", "STUDENT RESOURCE Guided Notes V2 (Reader p.9)",
             "ASK Fill in the blank.", "LISTEN FOR 'moving'", "REVEAL Reveal after writing.", "NEXT Check answer."])

# ================================================================================
# SLIDE 16 — V2 NOTES (check/reveal)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "CHECK")
notes_heading(s, "V2", "Vocabulary — Energy of Motion", "Guided Notes · Reader p. 9")
rich(s, 1.5, 2.65, 11.0, 1.15, [
    [("energy of motion — the energy an object possesses while it is ", {}),
     ("moving", {"bold": True, "color": ACCENT}), (" .", {})],
])
text(s, 1.5, 4.15, 11.0, 0.5, "moving", size=T_BODY_S, color=INK_SOFT, italic=True)
slide_number(s)
set_note(s, ["MODE Notes · V2 reveal", "TIME 1 min", "STUDENT RESOURCE Guided Notes V2",
             "REVEAL moving", "NEXT Moving Objects Have Energy (N2)."])

# ================================================================================
# SLIDE 17 — N2 CONCEPT: moving objects have energy
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "R"); hud_bubble(s, "TRACE")
text(s, 0.55, 1.02, 11.7, 0.7, "Moving objects have energy", size=T_TITLE, color=PRIMARY, bold=True)
text(s, 0.55, 1.75, 11.7, 0.4, "Reader p. 9 — energy is the ability to cause change.", size=T_KICK, color=INK_SOFT, italic=True)
model_strip(s, ["MOVING OBJECT", "CHANGE", "EVIDENCE"], x=0.75, y=2.7, w=11.8, box_h=1.1, gap=0.5, size=30)
text(s, 0.75, 4.3, 11.8, 1.6,
     "Energy is the ability to cause change.\nA moving object causes change, so it has energy.",
     size=T_BODY, color=INK, align=PP_ALIGN.CENTER)
slide_number(s)
set_note(s, ["MODE Read · model", "TIME 3 min", "STUDENT RESOURCE Reader p.9",
             "ASK Trace: moving object → change → evidence.", "LISTEN FOR students using 'evidence' not 'proof'",
             "NEXT Record N2 on notes."])

# ================================================================================
# SLIDE 18 — N2 NOTES (write)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "WRITE")
notes_heading(s, "N2", "Moving Objects Have Energy", "Guided Notes · Reader p. 9")
rich(s, 1.5, 2.15, 11.0, 0.7, [
    [("All moving objects have energy of ", {}), (blank_str(6), {"underline": True, "color": INK}), (" .", {})],
])
text(s, 1.5, 2.9, 11.0, 0.85, "Energy is the ability to cause change. A moving object causes change, so it has energy.",
     size=T_BODY_S, color=INK, italic=True)
model_strip(s, ["MOVING OBJECT", "CHANGE", "EVIDENCE"], x=1.5, y=4.0, w=10.4, box_h=0.9, gap=0.42, size=24)
slide_number(s)
set_note(s, ["MODE Notes · N2 anchor", "TIME 3 min", "STUDENT RESOURCE Guided Notes N2 (Reader p.9)",
             "ASK Fill in the blank.", "LISTEN FOR 'motion'", "REVEAL Reveal after writing.", "NEXT Observe: marble collision (OBSERVE box)."])

# ================================================================================
# SLIDE 19 — N2 OBSERVE (notes mirror)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "DRAW")
notes_heading(s, "N2", "Marble Collision — Observe", "Guided Notes · Reader p. 9")
rect(s, 1.5, 2.15, 11.0, 0.7, MIST, line_color=ACCENT, line_w=0.75, rounded=True)
text(s, 1.7, 2.25, 10.6, 0.5, "OBSERVE — one moving marble rolls into a target marble.", size=24, color=NEUTRAL, bold=True, anchor=MSO_ANCHOR.MIDDLE)
before_after(s, x=1.5, y=3.05, w=11.0, h=2.15, left="BEFORE", right="AFTER", left_hint="draw", right_hint="draw")
text(s, 1.5, 5.35, 11.0, 0.4, "What changed?", size=24, color=INK, bold=True)
write_line(s, 1.5, 5.85, 10.5)
write_line(s, 1.5, 6.35, 10.5)
slide_number(s)
set_note(s, ["MODE Notes · N2 OBSERVE", "TIME 2 min", "STUDENT RESOURCE Guided Notes N2 (OBSERVE box)",
             "RECALL 30 seconds: What did the target marble do yesterday? If students do not remember, roll the marble once. A reminder, not a new investigation.",
             "ASK Draw the two marbles before and after.", "LISTEN FOR 'the target marble moved'",
             "REVEAL Sample: the target marble changed from sitting still to moving.", "NEXT Check N2 answer."])

# ================================================================================
# SLIDE 20 — N2 NOTES (check/reveal)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "CHECK")
notes_heading(s, "N2", "Moving Objects Have Energy", "Guided Notes · Reader p. 9")
rich(s, 1.5, 2.15, 11.0, 0.7, [
    [("All moving objects have energy of ", {}), ("motion", {"bold": True, "color": ACCENT}), (" .", {})],
])
text(s, 1.5, 2.9, 11.0, 0.85, "Energy is the ability to cause change. A moving object causes change, so it has energy.",
     size=T_BODY_S, color=INK, italic=True)
model_strip(s, ["MOVING OBJECT", "CHANGE", "EVIDENCE"], x=1.5, y=4.0, w=10.4, box_h=0.9, gap=0.42, size=24)
text(s, 1.5, 5.1, 11.0, 0.7, "motion\n“The target marble changed from sitting still to moving.”", size=24, color=INK_SOFT, italic=True)
slide_number(s)
set_note(s, ["MODE Notes · N2 reveal", "TIME 1 min", "STUDENT RESOURCE Guided Notes N2",
             "REVEAL motion + sample change response", "NEXT People Use the Energy of Motion (N3)."])

# ================================================================================
# SLIDE 21 — N3 CONCEPT: people use the energy of motion
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "R")
text(s, 0.55, 1.02, 11.7, 0.7, "People use the energy of motion", size=T_TITLE, color=PRIMARY, bold=True)
text(s, 0.55, 1.75, 11.7, 0.4, "Reader pp. 10–12 — moving objects do useful work.", size=T_KICK, color=INK_SOFT, italic=True)
examples = ["hammer (swinging)", "pendulum (swinging)", "wind (moving air)"]
cw = 3.6; gap = 0.35; total = 3*cw + 2*gap
x0 = 0.55 + (12.23 - total)/2
for i, ex in enumerate(examples):
    cx = x0 + i*(cw+gap)
    sp = rect(s, cx, 2.7, cw, 1.6, PRIMARY, rounded=True)
    tf = sp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.1); tf.margin_right = Inches(0.1)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, ex, size=24, color=WHITE, bold=True)
text(s, 0.55, 4.7, 11.7, 1.2,
     "A swinging hammer, a swinging pendulum, and moving air (wind)\nall use the energy of motion to cause a change.",
     size=T_BODY_S, color=INK, align=PP_ALIGN.CENTER)
slide_number(s)
set_note(s, ["MODE Read · concept", "TIME 3 min", "STUDENT RESOURCE Reader pp.10–12",
             "ASK What change does each moving object cause?", "LISTEN FOR the change each one causes",
             "NEXT Record N3 table on notes."])

# ================================================================================
# SLIDE 22 — N3 NOTES (table)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "WRITE")
notes_heading(s, "N3", "People Use the Energy of Motion", "Guided Notes · Reader pp. 10–12")
text(s, 1.5, 2.3, 11.0, 0.5, "Write or draw the change each moving object causes.", size=T_BODY_S, color=INK, italic=True)
hx, hy, hw, hh = 1.5, 2.9, 4.8, 0.62
ex, ew = hx + hw + 0.15, 11.0 - hw - 0.15
rect(s, hx, hy, hw, hh, PRIMARY, rounded=False)
text(s, hx+0.15, hy+0.12, hw-0.3, 0.4, "MOVING OBJECT", size=24, color=WHITE, bold=True)
rect(s, ex, hy, ew, hh, PRIMARY, rounded=False)
text(s, ex+0.15, hy+0.12, ew-0.3, 0.4, "WHAT CHANGE DOES IT CAUSE?", size=24, color=WHITE, bold=True)
rows = ["hammer (swinging)", "pendulum (swinging)", "wind (moving air)"]
ry = hy + hh
for i, rname in enumerate(rows):
    rh = 0.95
    cy = ry + i*rh
    rect(s, hx, cy, hw, rh, PAPER, line_color=RULE, line_w=0.75)
    text(s, hx+0.15, cy+0.18, hw-0.3, 0.6, rname, size=24, color=PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, ex, cy, ew, rh, PAPER, line_color=RULE, line_w=0.75)
slide_number(s)
set_note(s, ["MODE Notes · N3 anchor", "TIME 5 min", "STUDENT RESOURCE Guided Notes N3 (Reader pp.10–12)",
             "ASK Write or draw the change each moving object causes.", "LISTEN FOR the change each object causes",
             "NEXT Check answers."])

# ================================================================================
# SLIDE 23 — N3 NOTES (check/reveal)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "CHECK")
notes_heading(s, "N3", "People Use the Energy of Motion", "Guided Notes · Reader pp. 10–12")
hx, hy, hw, hh = 1.5, 2.35, 4.8, 0.62
ex, ew = hx + hw + 0.15, 11.0 - hw - 0.15
rect(s, hx, hy, hw, hh, PRIMARY, rounded=False)
text(s, hx+0.15, hy+0.12, hw-0.3, 0.4, "MOVING OBJECT", size=24, color=WHITE, bold=True)
rect(s, ex, hy, ew, hh, PRIMARY, rounded=False)
text(s, ex+0.15, hy+0.12, ew-0.3, 0.4, "WHAT CHANGE DOES IT CAUSE?", size=24, color=WHITE, bold=True)
answers = [
    ("hammer (swinging)", "drives a nail into the wood"),
    ("pendulum (swinging)", "turns the clock gears"),
    ("wind (moving air)", "spins the turbine blades"),
]
ry = hy + hh
for i, (rname, ans) in enumerate(answers):
    rh = 0.95
    cy = ry + i*rh
    rect(s, hx, cy, hw, rh, PAPER, line_color=RULE, line_w=0.75)
    text(s, hx+0.15, cy+0.18, hw-0.3, 0.6, rname, size=24, color=PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, ex, cy, ew, rh, AQUA_TINT, line_color=ACCENT, line_w=0.75)
    text(s, ex+0.15, cy+0.18, ew-0.3, 0.6, ans, size=24, color=INK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
slide_number(s)
set_note(s, ["MODE Notes · N3 reveal", "TIME 2 min", "STUDENT RESOURCE Guided Notes N3",
             "REVEAL hammer → nail · pendulum → clock gears · wind → turbine blades",
             "NEXT Check Your Thinking."])

# ================================================================================
# SLIDE 24 — CHECK YOUR THINKING (write)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "WRITE")
notes_heading(s, "CHECK", "Check Your Thinking", "Guided Notes · Reader pp. 7–12")
text(s, 1.5, 2.3, 11.0, 0.8, "How do we know a moving object has energy?", size=T_BODY, color=INK, bold=True)
text(s, 1.5, 3.15, 11.0, 0.55, "Start with: “We know because the moving object causes …”", size=24, color=INK, italic=True)
write_line(s, 1.5, 4.15, 10.5)
write_line(s, 1.5, 4.9, 10.5)
slide_number(s)
set_note(s, ["MODE Notes · Check Your Thinking", "TIME 3 min", "STUDENT RESOURCE Guided Notes (CHECK)",
             "ASK Complete the sentence starter.", "LISTEN FOR 'a change, like a marble starting to move'",
             "NEXT Reveal a sample response."])

# ================================================================================
# SLIDE 25 — CHECK YOUR THINKING (sample)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "C"); hud_bubble(s, "CHECK")
text(s, 0.55, 1.02, 11.7, 0.7, "We know because it causes a change", size=T_TITLE, color=PRIMARY, bold=True)
rect(s, 1.5, 2.4, 10.3, 2.2, AQUA_TINT, line_color=ACCENT, line_w=0.75, rounded=True)
text(s, 1.8, 2.6, 9.7, 0.5, "SAMPLE", size=18, color=NEUTRAL, bold=True, caps=True)
text(s, 1.8, 3.15, 9.7, 1.3,
     "“We know because the moving object causes a change — like a moving marble making another marble move.”",
     size=T_BODY_S, color=INK, italic=True)
text(s, 1.5, 4.9, 10.3, 0.6, "Any change a moving object causes is evidence that it has energy.", size=T_BODY_S, color=INK, align=PP_ALIGN.CENTER)
slide_number(s)
set_note(s, ["MODE Check · evidence", "TIME 2 min", "STUDENT RESOURCE Guided Notes (CHECK)",
             "REVEAL Sample response (accept any change-as-evidence).", "NEXT Apply: which object shows the energy of motion?"])

# ================================================================================
# SLIDE 26 — APPLICATION: which change shows the energy of motion?
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "C"); hud_bubble(s, "EVIDENCE")
text(s, 0.55, 1.02, 11.7, 0.7, "Which change shows the energy of motion?", size=T_TITLE, color=PRIMARY, bold=True)
opts = [
    ("A moving marble makes a still marble move.", True),
    ("A rock keeps sitting still on the ground.", False),
    ("A pencil rests on a desk.", False),
]
y = 2.5
for i, (txt, is_yes) in enumerate(opts):
    cy = y + i * 1.25
    rect(s, 1.5, cy, 10.3, 1.0, PAPER, line_color=PRIMARY, line_w=1.25, rounded=True)
    text(s, 1.8, cy + 0.22, 9.7, 0.6, txt, size=T_BODY, color=INK, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.55, 6.4, 11.7, 0.5, "Look for a change — that is the evidence of the energy of motion.", size=T_BODY_S, color=INK_SOFT, align=PP_ALIGN.CENTER, italic=True)
slide_number(s)
set_note(s, ["MODE Apply · evidence", "TIME 3 min", "STUDENT RESOURCE none",
             "ASK Which one shows a moving object causing a change?", "LISTEN FOR 'the moving marble'",
             "REVEAL The moving marble makes the still marble move — a change.", "NEXT Return to the Big Question."])

# ================================================================================
# SLIDE 27 — BQ · NOW I THINK (return)
# ================================================================================
s = new_slide(SURFACE)
chapter_tab(s); mode_badge(s, "N"); hud_bubble(s, "WRITE")
notes_heading(s, "BQ", "BQ · Now I Think", "Guided Notes p. 3")
rect(s, 1.5, 2.3, 11.0, 1.75, PRIMARY, rounded=True)
text(s, 1.8, 2.45, 10.4, 0.4, "BIG QUESTION", size=18, color=ACCENT, bold=True, caps=True)
text(s, 1.8, 2.9, 10.4, 1.05, "How are energy, change, and movement of objects related?", size=T_BODY, color=WHITE, bold=True)
text(s, 1.5, 4.3, 11.0, 0.5, "Now I think…", size=T_BODY, color=INK, bold=True)
write_line(s, 1.5, 5.0, 10.5)
write_line(s, 1.5, 5.7, 10.5)
write_line(s, 1.5, 6.4, 10.5)
slide_number(s)
set_note(s, ["MODE Notes · BQ return", "TIME 3 min", "STUDENT RESOURCE Guided Notes p.3",
             "ASK Answer the Big Question now.", "LISTEN FOR a moving object + a change it causes",
             "NEXT End-of-chapter review."])

# ================================================================================
# END-OF-CHAPTER REVIEW (5 MCQ, question state + answer state)
# ================================================================================
REVIEW = [
    {"q": "A moving object has energy. What is this energy called?",
     "opts": ["stored energy", "energy of motion", "energy of light", "energy of sound"],
     "correct": 1, "type": "concept",
     "explain": "Energy of motion (Reader p. 9)."},
    {"q": "A marble rolls and hits a still marble. What change do you observe?",
     "opts": ["Both marbles stop moving", "The still marble starts to move", "The marbles change color", "Nothing happens"],
     "correct": 1, "type": "visual",
     "explain": "The still marble starts to move (Reader p. 8)."},
    {"q": "A swinging hammer hits a nail. What change does the hammer cause?",
     "opts": ["The hammer turns to liquid", "The hammer floats away", "The nail is driven into the wood", "The nail melts"],
     "correct": 2, "type": "scenario",
     "explain": "The hammer drives the nail in (Reader p. 10)."},
    {"q": "How do we know a moving object has energy if we cannot see the energy itself?",
     "opts": ["We can always see energy directly", "We guess without any evidence", "Moving objects never have energy", "We observe the change it causes"],
     "correct": 3, "type": "application",
     "explain": "We know because it causes a change."},
    {"q": "Wind (moving air) spins windmill blades. What does this show?",
     "opts": ["Moving air has energy of motion", "Still air is stronger", "Windmills work without energy", "Blades move on their own"],
     "correct": 0, "type": "scenario",
     "explain": "Moving air has energy of motion (Reader p. 12)."},
]

def review_question(slide_q, num, item):
    chapter_tab(slide_q); mode_badge(slide_q, "C")
    text(slide_q, 0.55, 1.02, 11.7, 0.7, f"Review {num} — {item['type']}", size=T_HEAD, color=PRIMARY, bold=True)
    text(slide_q, 0.55, 1.66, 11.7, 1.05, item["q"], size=T_BODY, color=INK, bold=True)
    y = 2.78
    letters = ["A", "B", "C", "D"]
    for i, opt in enumerate(item["opts"]):
        cy = y + i * 0.88
        rect(slide_q, 1.0, cy, 11.3, 0.74, PAPER, line_color=RULE, line_w=1.0, rounded=True)
        text(slide_q, 1.25, cy + 0.13, 0.6, 0.5, letters[i], size=T_OPT, color=ACCENT, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(slide_q, 1.95, cy + 0.11, 10.1, 0.54, opt, size=T_OPT, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    slide_number(slide_q)

def review_answer(slide_a, num, item):
    chapter_tab(slide_a); mode_badge(slide_a, "C"); hud_bubble(slide_a, "CHECK")
    text(slide_a, 0.55, 1.02, 11.7, 0.7, f"Review {num} — answer", size=T_HEAD, color=PRIMARY, bold=True)
    text(slide_a, 0.55, 1.66, 11.7, 1.05, item["q"], size=T_BODY, color=INK, bold=True)
    y = 2.78
    letters = ["A", "B", "C", "D"]
    for i, opt in enumerate(item["opts"]):
        cy = y + i * 0.88
        is_correct = (i == item["correct"])
        fill = AQUA_TINT if is_correct else PAPER
        line_c = ACCENT if is_correct else RULE
        rect(slide_a, 1.0, cy, 11.3, 0.74, fill, line_color=line_c, line_w=1.5 if is_correct else 1.0, rounded=True)
        text(slide_a, 1.25, cy + 0.13, 0.6, 0.5, letters[i], size=T_OPT, color=ACCENT, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(slide_a, 1.95, cy + 0.11, 10.1, 0.54, opt, size=T_OPT, color=INK, bold=is_correct, anchor=MSO_ANCHOR.MIDDLE)
    rect(slide_a, 0.55, 6.4, 12.23, 0.58, MIST, line_color=RULE, line_w=1.0, rounded=True)
    text(slide_a, 0.75, 6.48, 11.85, 0.42, item["explain"], size=24, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    slide_number(slide_a)

for idx, item in enumerate(REVIEW, 1):
    sq = new_slide(SURFACE)
    review_question(sq, idx, item)
    set_note(sq, [f"MODE Check · review Q{idx}", "STUDENT RESOURCE none",
                  "ASK Read aloud; students vote A/B/C/D.", "REVEAL Do not reveal yet.",
                  "NEXT Reveal the answer."])
    sa = new_slide(SURFACE)
    review_answer(sa, idx, item)
    set_note(sa, [f"MODE Check · review Q{idx} answer", "REVEAL " + REVIEW[idx-1]["opts"][item["correct"]],
                  "NEXT Next question (or close)."])

# ================================================================================
# save
# ================================================================================
out = os.path.join(OUT, "02-ch2-moving-objects-have-energy-slides.pptx")
prs.save(out)
# python-pptx writes type="screen4x3" even when cx/cy are 16:9. Chapter 1 omits the type.
# The parity gate treats screen4x3 as a 4:3 deck, so strip it.
import zipfile
buf = __import__("io").BytesIO()
with zipfile.ZipFile(out) as zin, zipfile.ZipFile(buf, "w") as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "ppt/presentation.xml":
            data = data.replace(b' type="screen4x3"', b"")
        zout.writestr(item.filename, data)
with open(out, "wb") as f:
    f.write(buf.getvalue())
print("built slides:", len(prs.slides._sldIdLst))
print("saved:", out)
