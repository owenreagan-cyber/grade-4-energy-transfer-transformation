"""Build Chapter 4 Energy Transfer slide deck (PPTX, 16:9)."""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = os.path.join(os.path.dirname(__file__), '..', 'production', 'slides')

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x2E, 0x6E, 0x8E)
AMBER = RGBColor(0xD9, 0xA4, 0x41)
LIGHT = RGBColor(0xF5, 0xF3, 0xEC)
INK = RGBColor(0x1A, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x55, 0x55, 0x55)
PALE = RGBColor(0xD5, 0xDE, 0xEA)

CUE_COLORS = {
    "NOTEBOOK": RGBColor(0x6B, 0x4F, 0xA0),
    "WRITE": RGBColor(0x2E, 0x6E, 0x8E),
    "DRAW": RGBColor(0x3E, 0x8E, 0x5A),
    "LABEL": RGBColor(0xD9, 0xA4, 0x41),
    "REVEAL": RGBColor(0xC0, 0x4A, 0x4A),
}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def add_slide():
    return prs.slides.add_slide(BLANK)

def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp

def textbox(slide, x, y, w, h, text, size=24, color=INK, bold=False, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = color; r.font.name = "Arial"
    return tb

def cue_badge(slide, cue):
    if not cue:
        return
    c = CUE_COLORS.get(cue, GRAY)
    sp = rect(slide, 12.25, 0.25, 0.9, 0.5, c, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = sp.text_frame; tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = cue
    r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = "Arial"

def model_strip(slide, items, y=5.6, x=0.8, w=11.7, box_h=0.7, gap=0.35):
    n = len(items)
    box_w = (w - (n - 1) * gap) / n
    cx = x
    for i, it in enumerate(items):
        sp = rect(slide, cx, y, box_w, box_h, NAVY if i == 0 else (AMBER if i == n - 1 else TEAL),
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        tf = sp.text_frame; tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = it
        r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = "Arial"
        if i < n - 1:
            ar = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(cx + box_w + 0.03),
                                        Inches(y + box_h/2 - 0.14), Inches(gap - 0.06), Inches(0.28))
            ar.fill.solid(); ar.fill.fore_color.rgb = GRAY; ar.line.fill.background(); ar.shadow.inherit = False
        cx += box_w + gap

def header(slide, title, subtitle, cue=None, num=None):
    rect(slide, 0, 0, 13.333, 1.0, NAVY)
    textbox(slide, 0.5, 0.12, 11.6, 0.75, title, size=30, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        textbox(slide, 0.5, 0.86, 11.6, 0.3, subtitle, size=13, color=PALE)
    if num:
        textbox(slide, 11.4, 0.32, 0.7, 0.4, str(num), size=14, color=PALE, align=PP_ALIGN.RIGHT)
    cue_badge(slide, cue)

def body(slide, lines, x=0.8, y=1.5, w=11.7, h=3.9, size=24, gap=8):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.color.rgb = INK; r.font.name = "Arial"
    return tb

def concept(title, subtitle, lines, cue=None, num=None, strip=None, size=24):
    s = add_slide()
    rect(s, 0, 0, 13.333, 7.5, LIGHT)
    header(s, title, subtitle, cue, num)
    body(s, lines, size=size)
    if strip:
        model_strip(s, strip)
    return s

# ---------------- Week 4 ----------------
s = add_slide(); rect(s, 0, 0, 13.333, 7.5, NAVY)
textbox(s, 0.8, 2.0, 11.7, 1.2, "Energy Transfer", size=46, color=WHITE, bold=True)
textbox(s, 0.8, 3.2, 11.7, 0.6, "Chapter 4 · Grade 4 · Weeks 4–5", size=20, color=PALE)
model_strip(s, ["SOURCE", "ENERGY MOVES", "TARGET CHANGES"], y=5.0)

concept("Something moved… no one touched it", "Opening phenomenon",
        ["A cup is tapped.", "A paper target a few inches away moves.", "Nothing touched the paper."], num=2)
concept("I notice… I wonder…", "Opening phenomenon",
        ["I notice: the paper moved.", "I wonder: how did it move without being touched?"], num=3)

s = add_slide(); rect(s, 0, 0, 13.333, 7.5, LIGHT)
rect(s, 0, 2.6, 13.333, 2.2, NAVY)
textbox(s, 0.8, 2.75, 11.7, 1.9, "What evidence shows that energy transfers from place to place?",
        size=34, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
textbox(s, 0.8, 2.0, 11.7, 0.5, "Big Question", size=16, color=GRAY, align=PP_ALIGN.CENTER)
textbox(s, 11.4, 0.32, 0.7, 0.4, "4", size=14, color=GRAY, align=PP_ALIGN.RIGHT)

concept("Energy transfer", "Definition",
        ["Energy transfer = energy moves from one place to another,", "or from one object to another."],
        cue="WRITE", num=5, strip=["SOURCE", "ENERGY MOVES", "TARGET"])
concept("Notebook — energy transfer", "Insert 1",
        ["Fill in the blanks:", "\u201cEnergy transfer happens when energy moves from one ______",
         "or ______ to another.\u201d", "Draw one example."],
        cue="NOTEBOOK", num=6, strip=["SOURCE", "ENERGY MOVES", "TARGET"])
concept("Sound prediction", "Week 4 · Day 1–2",
        ["What do you think the paper will do", "when the cup is tapped?", "Why?"], num=7)
concept("Sound setup", "Investigation",
        ["Cup + tapping object + paper target.", "The cup must NOT touch the paper."],
        cue="DRAW", num=8, strip=["SOUND SOURCE", "ENERGY MOVES", "TARGET CHANGES"])
concept("Sound safety", "Investigation",
        ["Keep your body very still.", "No bumping desks.", "No blowing air at the paper."], num=9)
concept("Sound evidence", "Investigation",
        ["The paper moved — and nothing touched it.", "The movement is the CHANGE.",
         "That change is evidence sound transferred energy."],
        cue="REVEAL", num=10)
concept("Sound explain", "Investigation",
        ["Did anything touch the paper?  (No.)", "So how did the energy get from the cup to the paper?",
         "Sound moved through the air and caused the change."], num=11)
concept("Notebook — sound", "Insert 2 (sound column)",
        ["Draw the sound setup.", "Label the source and the change you saw."],
        cue="NOTEBOOK", num=12)
concept("Light phenomenon", "Week 4 · Day 3",
        ["A flashlight lights up a card across the room.", "The card changes — it lights up."], num=13)
concept("Light prediction", "Investigation",
        ["What will change on the target", "when the flashlight is turned on?",
         "What if a barrier blocks the path?"], num=14)
concept("Light path model", "Investigation",
        ["Trace the path of the light."],
        cue="LABEL", num=15, strip=["SOURCE", "PATH", "TARGET", "OBSERVED CHANGE"])
concept("Light test", "Investigation",
        ["Shine the light on the target.", "Place a barrier in the path.", "What happens to the light spot?"], num=16)
concept("Light evidence", "Investigation",
        ["The target lights up when the beam reaches it.", "A barrier blocks the path — and the change disappears.",
         "We infer that light transferred energy."],
        cue="REVEAL", num=17)
concept("Notebook — light", "Insert 2 (light column)",
        ["Draw the light path.", "Label the source, path, and target."],
        cue="NOTEBOOK", num=18)
concept("The transfer model", "Generalize",
        ["Every energy transfer has three parts:", "a SOURCE, energy MOVES, and a TARGET CHANGES."],
        num=19, strip=["SOURCE", "ENERGY MOVES", "TARGET CHANGES"])
concept("Reader — Energy Transfer", "Chapter 4, pp. 17–18",
        ["A push on a skateboard transfers energy from you to your friend.",
         "A kicked ball transfers energy to a second ball."], num=20)
concept("Reader — Sound & Light", "Chapter 4, p. 19",
        ["Sound: a horn can make you feel a push.", "Light: the sun transfers light across space to Earth."], num=21)
concept("Transfer vs transformation — preview", "Getting ready",
        ["TRANSFER — energy moves.", "TRANSFORMATION — energy changes form.",
         "(We will sort these later this chapter.)"],
        cue="WRITE", num=22)
concept("Examples / non-examples", "Check understanding",
        ["Which show energy TRANSFER?",
         "bell → paper moves · flashlight → card lights · warm water → spoon warms"], num=23)

# ---------------- Week 5 ----------------
concept("Heat phenomenon", "Week 5 · Day 1",
        ["Warm water is placed next to a cool cup.", "The water cools and the cup warms."], num=24)
concept("Heat prediction", "Investigation",
        ["What will happen to the temperature of the warm water?", "Of the cooler object?"], num=25)
concept("Heat setup + safety", "Investigation",
        ["Warm water (teacher-controlled, not hot) + thermometer + cooler object.",
         "NO hot plates. NO open flames.", "Record temperature in °C."], num=26)
concept("Heat data", "Investigation",
        ["Record the temperature before and after.",
         "Warm water: start ___ °C → later ___ °C", "Cooler object: start ___ °C → later ___ °C"],
        cue="WRITE", num=27)
concept("Heat evidence", "Investigation",
        ["The warmer material cooled.", "The cooler material warmed.",
         "The temperature change is evidence heat transferred energy."],
        cue="REVEAL", num=28, strip=["WARMER", "ENERGY TRANSFERS", "COOLER"])
concept("Notebook — heat", "Insert 3 (heat column)",
        ["Fill in: warmer → energy transfer → cooler.", "Add your evidence: the temperature changed."],
        cue="NOTEBOOK", num=29)
concept("Electricity phenomenon", "Week 5 · Day 2",
        ["A battery is connected to a bulb by wires.", "The bulb lights up — a change far from the battery."], num=30)
concept("Electricity prediction", "Investigation",
        ["What will the device do when the circuit is complete?", "What happens if you break the circuit?"], num=31)
concept("Circuit setup + safety", "Investigation",
        ["Low-voltage battery (AA/AAA/D cell) + wires + bulb/buzzer/motor.",
         "NO wall outlets. NO coin cells."], num=32)
concept("Circuit model", "Investigation",
        ["Trace the path of the electric current."],
        cue="LABEL", num=33, strip=["BATTERY", "CIRCUIT", "DEVICE", "CHANGE"])
concept("Electricity evidence", "Investigation",
        ["Complete circuit → the device changes (light / sound / motion).",
         "Broken circuit → the change stops.",
         "That change is evidence electric current transferred energy."],
        cue="REVEAL", num=34)
concept("Notebook — electricity", "Insert 3 (electricity column)",
        ["Fill in: battery → circuit → device.", "Add your evidence: light / sound / motion."],
        cue="NOTEBOOK", num=35)
concept("Transfer vs transformation", "Synthesis",
        ["TRANSFER — energy moves from place to place / object to object.",
         "TRANSFORMATION — energy changes form."],
        cue="WRITE", num=36)
concept("TRANSFER ≠ TRANSFORMATION", "Synthesis",
        ["Transfer = MOVES.", "Transformation = CHANGES FORM.", "They are different ideas."], num=37)
concept("BOTH", "Synthesis",
        ["Energy can move AND change form — at the same time.",
         "Flashlight: energy moves to the bulb AND changes to light.",
         "Buzzer: energy moves through wires AND changes to sound."],
        num=38, strip=["TRANSFER", "+", "TRANSFORMATION", "= BOTH"])
concept("Sort — Transfer / Transformation / BOTH", "Synthesis",
        ["Sort: flashlight · buzzer · motor · warm water warms a spoon · sound source.",
         "BOTH is allowed — do not force either/or."],
        cue="NOTEBOOK", num=39)
concept("Notebook — sort", "Insert 4",
        ["Cut out the system cards.", "Sort them under TRANSFER, TRANSFORMATION, or BOTH.",
         "Flat option: draw a line from each system to its column."],
        cue="NOTEBOOK", num=40)
concept("Evidence + CER", "Week 5 · Day 4",
        ["NOTICE — what did I observe?", "EVIDENCE — what changed?",
         "EXPLAIN — why did it change?", "CONNECT — how does this show transfer?"], num=41)
concept("CER model", "Writing",
        ["\u201cI know ______ because ______.\u201d", "Claim + Evidence + Reasoning."],
        cue="WRITE", num=42)
concept("Notebook — CER", "Insert 5",
        ["Plan your claim:", "\u201cI know energy transferred because ______.\u201d",
         "Name the change you observed."],
        cue="NOTEBOOK", num=43)
concept("Synthesis", "All four forms",
        ["Sound, light, heat, and electric current all transfer energy.",
         "In every one: a SOURCE, energy MOVES, and a TARGET CHANGES."],
        num=44, strip=["SOURCE", "ENERGY MOVES", "TARGET CHANGES"])
concept("Misconception check", "Synthesis",
        ["\u201cDid the energy disappear?\u201d", "NO — energy transfers or transforms.", "It is never gone."],
        cue="REVEAL", num=45)
concept("Review / exit", "Chapter 4",
        ["Name a source, a change, and the evidence", "for ONE form of energy transfer."], num=46)

prs.save(os.path.join(OUT, "04-ch4-energy-transfer-slides.pptx"))
print("built slides:", len(prs.slides._sldIdLst))
