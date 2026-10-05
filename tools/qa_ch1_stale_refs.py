"""Chapter 1 stale-reference scan (Phase 6.1). Target: 0 hits in ACTIVE files.

ACTIVE = files a teacher/agent would act on. Historical records (v1.0 freeze/certification, retired archive,
qa/reports of earlier phases, curriculum/chapter-01 files carrying a SUPERSEDED banner, the Phase 6.1 audit docs
that quote the stale phrases) are excluded by design.

Usage: .venv/bin/python tools/qa_ch1_stale_refs.py [--report qa/reports/X.md]
"""
import argparse, glob, os, re, sys

REPO = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))

# (regex, why)
STALE = [
    (r"Block 1:? ?[—-]? ?Observe", "legacy two-block Ch1 plan"),
    (r"Block 2:? ?[—-]? ?Forms", "legacy two-block Ch1 plan"),
    (r"cup/balloon/salt|cup \+ balloon \+ salt", "retired Ch1 secondary demo (belongs to Ch4)"),
    (r"Flashlight \+ salt", "retired Ch1 demo"),
    (r"01-ch1[^|\n]{0,30}S1–S9|01-ch1[^|\n]{0,30}S12–S23", "legacy slide ranges"),
    (r"\(prototype\)", "prototype language"),
    (r"letters shuffled|Correct letters: Q1=A", "old review-letter pattern"),
    (r"Q1=A · Q2=B · Q3=C · Q4=D", "old review-letter pattern"),
    (r"Draw/Label", "legacy Ch1 note-section name"),
    (r"N1 Model It|Model It · N4", "legacy N1 name"),
    (r"(Ch1|Chapter 1)[^\n]*PLANNED — Phase 4", "stale planned-status on a Ch1 row"),
    (r"01-ch1-interactive-science-notebook\.pdf` \| RELEASED", "retired INB listed as released"),
    (r"key embedded in notes", "Ch1 keys are separate files since v1.1"),
    (r"Energy Detectives / Moving Marbles", "Detectives is optional, not Thursday"),
    (r"BQ:\*\* What are some forms of energy\? What is stored energy\?", "Ch1 has one BQ (evidence of energy causing change)"),
    (r"Energy Detectives lab \(AP 1\.1\)|Pencil sharpener demo \| Insert 1|Stored Energy \(CORE\)", "legacy Ch1 INB / Tuesday-lab plan"),
    (r"Detectives catch-up \(STANDARD\)|forms-of-energy sort \(FLEX\)", "legacy Ch1 Friday/Tuesday FLEX content"),
    (r"Deck:\*\* 25 slides|Energy Model · Stored Energy\)", "legacy Ch1 deck/INB counts"),
    (r"01-ch1[^|\n]{0,40}S12–S2\d", "legacy slide ranges"),
    (r"6 items / 12 pts\b(?![^\n]*Unit 1 final)", "trace total must say it is the Unit 1 final contribution"),
    # ---- Phase 6.1B (final Grok remediation)
    (r"(?<![\d–-])17\s*[–-]\s*S?27\b", "hyphenated schedule range 17–27 silently includes OPTIONAL slide 18 (write: 2, 17, 19–27; 18 optional)"),
    (r"(later than|at) minute 15\b", "Friday quiz start is minute 12 at the latest (M3)"),
    (r"N4[^|\n]{0,80}\bpp?\.? ?1[–-]6\b", "N4 cites core Reader pp. 1–5 (p. 6 = teacher-facing forward preview) (L1)"),
    (r"Our class looked around the room|what some classmates found", "recovery must not present Reader examples as class observations (M4)"),
    (r"chain printed on the page \(nothing to re-copy\)", "N1 model's last box is empty on the student page (M1)"),
    (r"(?i)batteries store energy in chemicals[^\n]*PREDICT|PREDICT[^\n]*batteries store energy in chemicals", "slide 23 prediction must not answer itself (M2)"),
    (r"A battery stores energy; it changes when released", "battery is not 'released' (L5)"),
]

# Retired / superseded Chapter 1 artifacts: a line that names one of these must say RETIRED or SUPERSEDED
# (Phase 6.1B / H1: PRODUCTION_MANIFEST.md pointed RELEASED rows at them).
RETIRED_CH1 = [
    (r"production/student-notes/01-ch1", "obsolete 'WHAT WE NOTICE' student-notes DOCX"),
    (r"curriculum/chapter-01/teacher/01-ch1-energy-causes-change-teacher-map", "superseded historical Chapter 1 teacher map"),
    (r"production/student-notebook/chapter-01|01-ch1-interactive-notebook|01-ch1-interactive-science-notebook", "retired Chapter 1 Interactive Science Notebook"),
]
RETIRED_OK = re.compile(r"RETIRED|SUPERSEDED|superseded by|archived|archive/retired", re.I)
MANIFEST = "production/PRODUCTION_MANIFEST.md"

ACTIVE_GLOBS = [
    "docs/chapter-01/CH1_PACKAGE_MANIFEST.md", "docs/chapter-01/CH1_SLIDESHOW_MANIFEST.md",
    "docs/chapter-01/CH1_ASSESSMENT_TRACE.md", "docs/chapter-01/CH1_STUDY_GUIDE_CONTRIBUTION.md",
    "docs/chapter-01/CH1_PACING_PLAN.md", "docs/playbook/UNIT1_RESOURCE_MAP.md", "docs/playbook/UNIT1_PLAYBOOK_ARCHITECTURE.md",
    "production/final-8.5-week/WEEK1_2_TEACHER_QUICKSTART.md",
    "production/answer-keys/01-ch1-*.md", "production/quizzes/01-ch1-*.md", "production/labs/01-ch1-*.md",
    "curriculum/chapter-01/student/*.md", "curriculum/chapter-01/student/recovery/*.md",
    "production/final-package/ENERGY_UNIT_MASTER_INDEX.md",
]
# Shared multi-chapter files: only lines that mention Chapter 1 are scanned.
SHARED_LINE_FILTER = {
    "docs/playbook/UNIT1_PACING_CALENDAR.md": r"Ch1|Chapter 1",
    "production/final-8.5-week/MASTER_CURRICULUM_MAP.md": r"01-ch1|Ch1 ",
    "production/final-8.5-week/FINAL_RELEASE_MANIFEST.md": r"ch1|Ch1|Chapter 1",
    "production/final-8.5-week/TEACHER_PREP_GUIDE.md": r"01-ch1",
    "production/final-8.5-week/MASTER_RESOURCE_INDEX.md": r"01-ch1|U1 Ch1|Ch1–2",
    "production/final-8.5-week/ENERGY_AND_NATURAL_RESOURCES_MASTER_PACING.md": r"01-ch1|Ch1|Detectives|Energy Detectives",
    "production/final-8.5-week/MASTER_RESOURCE_INVENTORY.md": r"Unit 1 keys|01-ch1",
    "production/final-8.5-week/VIDEO_PRESENTATION_INTEGRATION_MATRIX.md": r"\| Ch1 \|",
    "production/friday-flex/FRIDAY_FLEX_BANK.md": r"Ch1|Detectives catch-up",
    "production/final-package/ENERGY_UNIT_WEEK_BY_WEEK_TEACHER_MAP.md": r"Chapter 1:|Energy Detectives|Pencil sharpener|Stored Energy \(CORE\)|Deck:\*\* 25|Energy Model",
}


def scan_text(txt, line_filter=None):
    hits = []
    for i, line in enumerate(txt.splitlines(), 1):
        if line_filter and not re.search(line_filter, line):
            continue
        for rx, why in STALE:
            if re.search(rx, line):
                hits.append((i, why, line.strip()[:140]))
        # retired/superseded Ch1 artifacts may only be mentioned when the line says so (reviewer follow-up, 6.1B)
        if not RETIRED_OK.search(line):
            for rx, why in RETIRED_CH1:
                if re.search(rx, line):
                    hits.append((i, "retired Ch1 artifact referenced as current: " + why, line.strip()[:140]))
    return hits


# Note: notebook/glue-in/INB architecture reactivation is now guarded by the FAIL-CLOSED model in
# `tools/qa_ch1_note_architecture.py` (authority files + retired-standard leading-banner gate +
# reviewed allowlist). This tool retains only the legacy Ch1 CONTENT staleness patterns and retired
# Ch1 artifact-reference checks.

def scan_manifest(txt):
    """PRODUCTION_MANIFEST.md (H1): scan every Chapter 1 row. Fail if a retired/superseded Ch1 artifact is listed
    as RELEASED/canonical/current, if WHAT WE NOTICE is presented as current, or if P3 does not point at premium-notes."""
    hits = []
    lines = txt.splitlines()
    section = ""
    p3_ok = False
    for i, line in enumerate(lines, 1):
        if line.startswith("## "):
            section = line
        in_ch1 = section.startswith("## Chapter 1")
        mentions_ch1 = in_ch1 or re.search(r"01-ch1|chapter-01\b|premium-notes/01", line)
        if not mentions_ch1:
            continue
        ok_line = bool(RETIRED_OK.search(line))
        for rx, why in RETIRED_CH1:
            if re.search(rx, line) and not ok_line:
                hits.append((i, "retired Ch1 artifact referenced as canonical/current: " + why, line.strip()[:140]))
            if re.search(rx, line) and re.search(r"\|\s*RELEASED\b", line):
                hits.append((i, "retired Ch1 artifact marked RELEASED: " + why, line.strip()[:140]))
        if re.search(r"WHAT WE NOTICE", line) and not ok_line:
            hits.append((i, "WHAT WE NOTICE architecture presented as current", line.strip()[:140]))
        if in_ch1 and line.startswith("| P3 |") and "production/premium-notes/01-ch1-energy-causes-change-guided-notes" in line and "student-notes" not in line:
            p3_ok = True
        if in_ch1:
            for rx, why in STALE:
                if re.search(rx, line):
                    hits.append((i, why, line.strip()[:140]))
    if not p3_ok:
        hits.append((0, "P3 Guided Notes master must be production/premium-notes/01-ch1-energy-causes-change-guided-notes.docx", ""))
    return hits


def docx_text(path):
    from docx import Document
    d = Document(path)
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            out.append(" | ".join(c.text for c in r.cells))
    return "\n".join(out)


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--report", default=None); a = ap.parse_args()
    results = []   # (file, hits)
    files = []
    for g in ACTIVE_GLOBS:
        files += sorted(glob.glob(os.path.join(REPO, g)))
    for f in files:
        rel = os.path.relpath(f, REPO)
        txt = open(f, encoding="utf-8").read()
        h = scan_text(txt)
        # Chapter 1 pages of the unit calendar file are handled below; skip generic docs that are chapter-agnostic
        results.append((rel, h))
    # H1: production/PRODUCTION_MANIFEST.md is explicitly in scope (Chapter 1 rows + any 01-ch1 line)
    results.append((MANIFEST + "  [Ch1 rows]", scan_manifest(open(os.path.join(REPO, MANIFEST), encoding="utf-8").read())))
    for rel, flt in SHARED_LINE_FILTER.items():
        txt = open(os.path.join(REPO, rel), encoding="utf-8").read()
        results.append((rel + "  [Ch1 lines only]", scan_text(txt, flt)))
    # Playbook DOCX, Ch1 scope (reuse the QA scope helper)
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from docx import Document
    import qa_teacher_playbook as qp
    doc = Document(qp.DOCX)
    results.append(("production/teacher-playbook/unit1-teacher-playbook.docx  [Ch1 scope]", scan_text(qp._ch1_scope(doc))))

    assert any(r.startswith(MANIFEST) for r, _ in results), "PRODUCTION_MANIFEST.md must be inside stale-reference scope"
    total = sum(len(h) for _, h in results)
    print("Scanned %d active files/scopes · stale hits: %d · PRODUCTION_MANIFEST.md in scope: YES" % (len(results), total))
    for rel, h in results:
        for ln, why, text in h:
            print("  HIT %s:%s [%s] %s" % (rel, ln, why, text))
    if a.report:
        with open(os.path.join(REPO, a.report), "w") as f:
            f.write("# Chapter 1 Stale-Reference Scan\n\nTool: `tools/qa_ch1_stale_refs.py`. Active scopes scanned: **%d** (includes `production/PRODUCTION_MANIFEST.md` Ch1 rows). Stale hits: **%d** (target 0). Result: **%s**\n\n" %
                    (len(results), total, "PASS" if total == 0 else "FAIL"))
            f.write("Excluded by design: v1.0 freeze/certification history, retired archive, earlier-phase QA reports, SUPERSEDED-bannered legacy `curriculum/chapter-01/*`. Note-architecture guardrails (authority files + retired-standard banners + reviewed allowlist) are enforced by `tools/qa_ch1_note_architecture.py`; this tool covers only legacy Ch1 CONTENT staleness and retired Ch1 artifact references.\n\n")
            f.write("| Scope | Hits |\n|---|---|\n")
            for rel, h in results:
                f.write("| `%s` | %d |\n" % (rel, len(h)))
            for rel, h in results:
                for ln, why, text in h:
                    f.write("\n- HIT `%s:%s` — %s — `%s`" % (rel, ln, why, text))
    print("RESULT:", "PASS" if total == 0 else "FAIL")
    sys.exit(0 if total == 0 else 1)


if __name__ == "__main__":
    main()
