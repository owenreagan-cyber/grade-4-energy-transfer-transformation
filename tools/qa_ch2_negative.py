"""Real-file negative tests for Chapter 2.

Each case copies the tree inputs, mutates one real file, runs the release builder,
re-plants the mutation if the builder repaired it, refreshes hashes, then runs the
live QA tools. The canonical working tree is not modified.

Run: .venv/bin/python tools/qa_ch2_negative.py
"""
import json, os, shutil, subprocess, sys, tempfile

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
PY = os.path.join(ROOT, ".venv/bin/python")
FILES = [
    "production/premium-notes/02-ch2-moving-objects-have-energy-guided-notes.docx",
    "production/premium-notes/02-ch2-moving-objects-have-energy-guided-notes.pdf",
    "production/premium-notes/02-ch2-moving-objects-have-energy-guided-notes-key.docx",
    "production/premium-notes/02-ch2-moving-objects-have-energy-guided-notes-key.pdf",
    "production/slides/02-ch2-moving-objects-have-energy-slides.pptx",
    "production/slides/02-ch2-moving-objects-have-energy-slides.pdf",
    "production/teacher-playbook/unit1-teacher-playbook.docx",
    "production/teacher-playbook/unit1-teacher-playbook.pdf",
    "production/PRODUCTION_MANIFEST.md",
    "production/final-delivery/RELEASE_HASH_MANIFEST.csv",
    "tools/ch2_pacing.py",
    "docs/playbook/UNIT1_PACING_CALENDAR.md",
    "production/final-8.5-week/MASTER_CURRICULUM_MAP.md",
    "production/final-8.5-week/MASTER_RESOURCE_INDEX.md",
    "production/final-8.5-week/WEEK1_2_TEACHER_QUICKSTART.md",
    "production/final-8.5-week/VIDEO_PRESENTATION_INTEGRATION_MATRIX.md",
    "production/final-package/ENERGY_UNIT_WEEK_BY_WEEK_TEACHER_MAP.md",
    "production/final-package/ENERGY_UNIT_MEDIA_PLAYBACK_GUIDE.md",
    "production/worksheets/02-ch2-moving-objects-have-energy-practice.docx",
    "production/quizzes/02-ch2-moving-objects-have-energy-quiz.docx",
    "production/worksheets/02-ch2-moving-marbles-student.docx",
    "production/worksheets/02-ch2-absent-student-recovery.docx",
    "production/worksheets/02-ch2-friday-flex.docx",
    "production/answer-keys/02-ch2-moving-objects-have-energy-quiz-key.md",
    "production/answer-keys/02-ch2-moving-objects-have-energy-practice-key.md",
    "production/answer-keys/02-ch2-moving-marbles-key.md",
    "production/answer-keys/02-ch2-friday-flex-key.md",
    "production/archive/retired/chapter-02-pre-7.1/02-ch2-interactive-science-notebook.pdf",
    "production/archive/retired/chapter-02-pre-7.1/02-ch2-moving-objects-have-energy-guided-notes.pdf",
    "curriculum/chapter-02/ch2_anchor_map.json",
    "curriculum/chapter-02/teacher/02-ch2-teacher-notes.md",
    "production/final-delivery/Energy/Unit 1 Teacher Guide/02-ch2-teacher-notes.md",
]
DIRS = ["production/releases/chapter-02"]
EXPECT = {
    "A": "SYNC", "A2": "SYNC", "B1": "LEAK", "B1p": "LEAK", "B2": "LEAK",
    "C": "VOCAB", "D1": "PACING", "D2": "PACING", "E1": "OPTIONAL", "E1p": "OPTIONAL",
    "E2": "OPTIONAL", "F": "RETIRED", "F2": "RETIRED", "G": "MANIFEST",
    "H1": "FRIDAY", "H1p": "FRIDAY", "H2": "FRIDAY", "I": "PLAYBOOK",
    "J1": "RELEASE", "J2": "RELEASE", "M1": "MEDIA", "R": "SYNC",
    "A3": "SYNC", "A4": "SYNC", "A5": "SYNC", "A6": "SYNC", "AM": "SYNC",
    "B1q": "LEAK", "B3": "LEAK", "B4": "LEAK", "S2": "LEAK", "S3": "LEAK",
    "E1q": "OPTIONAL", "E1r": "OPTIONAL", "E1n": "OPTIONAL", "E1t": "OPTIONAL",
    "H1h": "FRIDAY", "H1n": "FRIDAY", "H1title": "FRIDAY",
    "F4": "RETIRED", "F5": "STALE_EXPORT", "M2": "MEDIA",
    "CONTROL_GOOD": "PASS", "H1fp": "PASS", "OPTIONAL_GOOD": "PASS",
}
CONTROLS = {"CONTROL_GOOD", "H1fp", "OPTIONAL_GOOD"}


def fresh():
    tmp = tempfile.mkdtemp(prefix="ch2neg_")
    for rel in FILES:
        dest = os.path.join(tmp, rel)
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        shutil.copy2(os.path.join(ROOT, rel), dest)
    for rel in DIRS:
        shutil.copytree(os.path.join(ROOT, rel), os.path.join(tmp, rel))
    os.makedirs(os.path.join(tmp, "qa", "reports"), exist_ok=True)
    return tmp


def run(cmd):
    proc = subprocess.run(cmd, cwd=ROOT, capture_output=True, text=True)
    return proc.returncode, proc.stdout + proc.stderr


def sub(path, old, new):
    text = open(path, encoding="utf-8").read()
    if old not in text:
        if new in text:
            return
        raise SystemExit("missing plant target in %s" % path)
    open(path, "w", encoding="utf-8").write(text.replace(old, new, 1))


def _replace_runs(paragraphs, old, new):
    hit = False
    for p in paragraphs:
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, new, 1)
                hit = True
    return hit


def notes_replace(tmp, old, new):
    from docx import Document
    path = os.path.join(tmp, FILES[0])
    d = Document(path)
    hit = _replace_runs(d.paragraphs, old, new)
    if not hit:
        for table in d.tables:
            for row in table.rows:
                for cell in row.cells:
                    if _replace_runs(cell.paragraphs, old, new):
                        hit = True
    if not hit and new not in "\n".join(p.text for p in d.paragraphs):
        raise SystemExit("notes text not found: " + old)
    d.save(path)


def slide_replace(tmp, slide_no, old, new):
    from pptx import Presentation
    path = os.path.join(tmp, "production/slides/02-ch2-moving-objects-have-energy-slides.pptx")
    prs = Presentation(path)
    hit = False
    targets = prs.slides if slide_no is None else [prs.slides[slide_no - 1]]
    for slide in targets:
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            if _replace_runs(sh.text_frame.paragraphs, old, new):
                hit = True
    if not hit:
        raise SystemExit("slide text not found: " + old)
    prs.save(path)


def fill_blank(tmp, slide_no, word):
    from pptx import Presentation
    path = os.path.join(tmp, "production/slides/02-ch2-moving-objects-have-energy-slides.pptx")
    prs = Presentation(path)
    hit = False
    for sh in prs.slides[slide_no - 1].shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if "___" in r.text:
                    r.text = word
                    hit = True
    if not hit:
        raise SystemExit("blank not found on slide %d" % slide_no)
    prs.save(path)


def speaker_replace(tmp, slide_no, old, new):
    from pptx import Presentation
    path = os.path.join(tmp, "production/slides/02-ch2-moving-objects-have-energy-slides.pptx")
    prs = Presentation(path)
    box = prs.slides[slide_no - 1].notes_slide.notes_text_frame
    if old not in box.text:
        if new in box.text:
            return
        raise SystemExit("speaker text missing on %d" % slide_no)
    box.text = box.text.replace(old, new, 1)
    prs.save(path)


def playbook_day(tmp):
    from docx import Document
    path = os.path.join(tmp, "production/teacher-playbook/unit1-teacher-playbook.docx")
    d = Document(path)
    hit = False
    for table in d.tables:
        for row in table.rows:
            blob = " ".join(c.text for c in row.cells)
            if "Ch2 Session 2" not in blob or "Thu 10/15" not in blob:
                continue
            for cell in row.cells:
                if "Thu 10/15" in cell.text and _replace_runs(cell.paragraphs, "Thu 10/15", "Fri 10/16"):
                    hit = True
    if not hit:
        raise SystemExit("playbook day not found")
    d.save(path)


def launder(tmp):
    import csv, hashlib
    csv_path = os.path.join(tmp, "production/final-delivery/RELEASE_HASH_MANIFEST.csv")
    rows = list(csv.reader(open(csv_path, newline="", encoding="utf-8")))
    for r in rows[1:]:
        path = os.path.join(tmp, r[3])
        if os.path.exists(path):
            h = hashlib.sha256(open(path, "rb").read()).hexdigest()
            r[4] = h
            r[5] = str(os.path.getsize(path))
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)
    man_path = os.path.join(tmp, "production/releases/chapter-02/RELEASE_MANIFEST.json")
    if os.path.exists(man_path):
        man = json.load(open(man_path))
        kept = []
        for item in man.get("files", []):
            path = os.path.join(tmp, "production/releases/chapter-02", item["name"])
            if os.path.exists(path):
                item["sha256"] = hashlib.sha256(open(path, "rb").read()).hexdigest()
                kept.append(item)
        man["files"] = kept
        json.dump(man, open(man_path, "w"), indent=2)


def prefixes(out):
    found = []
    for line in out.splitlines():
        if line.startswith("[FAIL]"):
            token = line.split()[1] if len(line.split()) > 1 else ""
            found.append(token)
    return found


def main():
    cases = []

    def case(name, fn):
        cases.append((name, fn))

    case("A", lambda t: notes_replace(t, "puts the ball", "pushed the ball"))
    case("A2", lambda t: slide_replace(t, None, "puts the ball", "pushed the ball"))
    case("B1", lambda t: slide_replace(t, 11, "A ball does not.", "A moving object can cause a change."))
    case("B1p", lambda t: slide_replace(t, 11, "A ball does not.", "Moving objects can cause a change."))
    case("B2", lambda t: fill_blank(t, 15, "moving"))
    case("C", lambda t: notes_replace(t, "the process of an object changing", "the energy an object possesses while it is"))
    case("D1", lambda t: sub(os.path.join(t, "tools/ch2_pacing.py"),
                             '_add(1, 1, "REQ", 1.0, "open", "Title")',
                             '_add(1, 1, "REQ", 30.0, "open", "Title")'))
    case("D2", lambda t: speaker_replace(t, 1, "TIME 1 min", "TIME 9 min"))
    case("E1", lambda t: speaker_replace(t, 12, "REQUIRED", "OPTIONAL · skip if short"))
    case("E1p", lambda t: speaker_replace(t, 12, "REQUIRED", "optional if time"))
    case("E2", lambda t: (
        sub(os.path.join(t, "tools/ch2_pacing.py"),
            '_add(12, 1, "REQ", 2.5, "write", "N1 notes - write")',
            '_add(12, 1, "OPT", 2.5, "write", "N1 notes - write")'),
        speaker_replace(t, 12, "REQUIRED", "optional if time")))
    case("F", lambda t: shutil.copy(
        os.path.join(t, "production/releases/chapter-02/02-ch2-friday-flex.pdf"),
        os.path.join(t, "production/releases/chapter-02/02-ch2-class-notebook.pdf")))
    case("F2", lambda t: shutil.copy(
        os.path.join(t, "production/archive/retired/chapter-02-pre-7.1/02-ch2-interactive-science-notebook.pdf"),
        os.path.join(t, "production/releases/chapter-02/02-ch2-reading-packet.pdf")))
    case("G", lambda t: sub(os.path.join(t, "production/PRODUCTION_MANIFEST.md"),
                            "production/premium-notes/02-ch2-moving-objects-have-energy-guided-notes.docx",
                            "production/student-notes/02-ch2-moving-objects-have-energy-guided-notes.docx"))
    case("H1", lambda t: open(os.path.join(t, "docs/playbook/UNIT1_PACING_CALENDAR.md"), "a").write(
        "\nFriday 10/16 Chapter 2 Session 1 required slides.\n"))
    case("H1p", lambda t: open(os.path.join(t, "docs/playbook/UNIT1_PACING_CALENDAR.md"), "a").write(
        "\nFri 10/16: Moving Objects Have Energy, Session 3\n"))
    case("H2", lambda t: speaker_replace(t, 12, "PLAN Session 1 (Wed 10/14)", "PLAN Session 3 (Fri 10/16)"))
    case("I", playbook_day)
    case("J1", lambda t: os.path.exists(os.path.join(t, "production/releases/chapter-02/02-ch2-absent-student-recovery.pdf")) and os.remove(
        os.path.join(t, "production/releases/chapter-02/02-ch2-absent-student-recovery.pdf")))
    case("J2", lambda t: (
        os.path.exists(os.path.join(t, "production/releases/chapter-02/02-ch2-absent-student-recovery.pdf")) and os.remove(
            os.path.join(t, "production/releases/chapter-02/02-ch2-absent-student-recovery.pdf")),
        sub(os.path.join(t, "production/releases/chapter-02/RELEASE_MANIFEST.json"),
            "02-ch2-absent-student-recovery.pdf", "removed-recovery.pdf")))
    def reorder_notes(tmp):
        from docx import Document
        path = os.path.join(tmp, FILES[0])
        d = Document(path)
        body = d.element.body
        v1 = v2 = None
        for el in list(body.iterchildren()):
            t = "".join(el.itertext()).strip().upper()
            if t.startswith("V1") and v1 is None:
                v1 = el
            if t.startswith("V2") and v2 is None:
                v2 = el
        if v1 is None or v2 is None:
            raise SystemExit("anchor headings not found")
        v1.addprevious(v2)
        d.save(path)

    def notes_add(tmp, marker, extra):
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.text.paragraph import Paragraph
        path = os.path.join(tmp, FILES[0])
        d = Document(path)
        for p in d.paragraphs:
            if extra in p.text:
                return
            if marker in p.text:
                el = OxmlElement("w:p")
                p._p.addnext(el)
                Paragraph(el, p._parent).add_run(extra)
                d.save(path)
                return
        raise SystemExit("notes marker missing: " + marker)

    def add_box(tmp, slide_no, sentence):
        from pptx import Presentation
        from pptx.util import Inches
        path = os.path.join(tmp, "production/slides/02-ch2-moving-objects-have-energy-slides.pptx")
        prs = Presentation(path)
        slide = prs.slides[slide_no - 1]
        for sh in slide.shapes:
            if sh.has_text_frame and sentence in sh.text_frame.text:
                return
        box = slide.shapes.add_textbox(Inches(0.6), Inches(6.4), Inches(8), Inches(0.4))
        box.text_frame.text = sentence
        prs.save(path)

    def splice_key(tmp):
        rel = os.path.join(tmp, "production/releases/chapter-02")
        key = os.path.join(rel, "02-ch2-moving-objects-have-energy-teacher-key.pdf")
        nb = os.path.join(tmp, "production/archive/retired/chapter-02-pre-7.1/02-ch2-interactive-science-notebook.pdf")
        out = key + ".tmp"
        subprocess.check_call(["qpdf", "--empty", "--pages", key, "1", nb, "1", key, "2-z", "--", out],
                              stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        os.replace(out, key)

    case("A3", lambda t: notes_add(t, "Energy is the ability to cause change.", "The object got heavier when it moved."))
    case("A4", lambda t: notes_replace(t, "Write or draw the change each moving object causes.", "List the objects and then stop."))
    case("A5", lambda t: notes_replace(t, "We know because the moving object causes", "We guess that the object might"))
    case("A6", lambda t: add_box(t, 18, "The marble got faster, so it has more energy."))
    case("AM", lambda t: (
        sub(os.path.join(t, "curriculum/chapter-02/ch2_anchor_map.json"), "puts the ball into motion", "shoves the ball along"),
        notes_replace(t, "puts the ball into motion", "shoves the ball along"),
        slide_replace(t, None, "puts the ball into motion", "shoves the ball along")))
    case("B1q", lambda t: slide_replace(t, 11, "A ball does not.", "Any moving object causes a change."))
    case("B3", lambda t: add_box(t, 17, "All moving objects have energy of motion."))
    case("B4", lambda t: add_box(t, 22, "A swinging hammer drives a nail into the wood."))
    case("S2", lambda t: speaker_replace(t, 18, "REVEAL Reveal after writing.", "REVEAL motion"))
    case("S3", lambda t: speaker_replace(t, 22, "LISTEN FOR the change each object causes", "LISTEN FOR nail, gears, and blades"))
    case("E1q", lambda t: speaker_replace(t, 12, "REQUIRED", "Enrichment, can be skipped"))
    case("E1r", lambda t: speaker_replace(t, 12, "REQUIRED", "Bonus"))
    case("E1n", lambda t: speaker_replace(t, 12, "REQUIRED", "Not required"))
    case("E1t", lambda t: speaker_replace(t, 12, "REQUIRED", "Only if there is time"))
    case("H1h", lambda t: open(os.path.join(t, "docs/playbook/UNIT1_PACING_CALENDAR.md"), "a").write(
        "\n### Friday 10/16\n- Chapter 2 Session 1 required slides.\n"))
    case("H1n", lambda t: open(os.path.join(t, "docs/playbook/UNIT1_PACING_CALENDAR.md"), "a").write(
        "\nDo not skip the quiz: Friday — Chapter 2 N3.\n"))
    case("H1fp", lambda t: open(os.path.join(t, "docs/playbook/UNIT1_PACING_CALENDAR.md"), "a").write(
        "\nFriday: Chapter 1 quiz, not Chapter 2.\n"))
    case("H1title", lambda t: open(os.path.join(t, "docs/playbook/UNIT1_PACING_CALENDAR.md"), "a").write(
        "\n### Friday\n- Moving Objects Have Energy — Session 3\n"))
    case("F4", splice_key)
    case("F5", lambda t: shutil.copy(
        os.path.join(t, "production/archive/retired/chapter-02-pre-7.1/02-ch2-moving-objects-have-energy-guided-notes.pdf"),
        os.path.join(t, "production/premium-notes/02-ch2-moving-objects-have-energy-guided-notes.pdf")))
    case("M2", lambda t: sub(
        os.path.join(t, "production/final-package/ENERGY_UNIT_MEDIA_PLAYBACK_GUIDE.md"),
        "| `Moving_Objects_Have_Energy.mp4` | W1 Wed | SUPPORT |",
        "| `Moving_Objects_Have_Energy.mp4` | W1 Wed | CORE |"))
    case("CONTROL_GOOD", lambda t: None)
    def optional_good(tmp):
        speaker_replace(tmp, 26, "OPTIONAL - skip freely", "optional if time")
        src = os.path.join(tmp, "production/slides/02-ch2-moving-objects-have-energy-slides.pptx")
        dst = os.path.join(tmp, "production/releases/chapter-02/02-ch2-moving-objects-have-energy-slides.pptx")
        shutil.copy2(src, dst)
        launder(tmp)
    case("OPTIONAL_GOOD", optional_good)
    case("R", reorder_notes)
    case("M1", lambda t: sub(
        os.path.join(t, "production/final-8.5-week/VIDEO_PRESENTATION_INTEGRATION_MATRIX.md"),
        "| Wk 1 Wed Session 1 | 5:59 | SUPPORT |",
        "| Wk 1 Thu | 5:59 | CORE |"))

    print("=" * 72)
    print("CH2 REAL-FILE NEGATIVE TESTS")
    print("=" * 72)
    caught = wrong = 0
    tools = [
        [PY, os.path.join(ROOT, "tools/qa_ch2_semantic.py")],
        [PY, os.path.join(ROOT, "tools/qa_ch2_slides.py")],
        [PY, os.path.join(ROOT, "tools/qa_ch2_guided_notes.py")],
        [PY, os.path.join(ROOT, "tools/qa_teacher_playbook.py")],
        [PY, os.path.join(ROOT, "tools/qa_release_parity.py"), "--chapter", "2"],
    ]
    for name, fn in cases:
        tmp = fresh()
        try:
            try:
                fn(tmp)
            except Exception as exc:
                print("[MISSED] %s — plant error %s" % (name, exc))
                wrong += 1
                continue
            if name in CONTROLS:
                build_code = "skip"
            else:
                build_code, build_out = run([PY, os.path.join(ROOT, "tools/build_ch2_release.py"), "--root", tmp])
                if build_code == 0:
                    fn(tmp)
                launder(tmp)
            sem_code, sem_out = run(tools[0] + ["--root", tmp])
            other = []
            for cmd in tools[1:]:
                code, _out = run(cmd + ["--root", tmp])
                other.append(code)
            found = prefixes(sem_out)
            if name in CONTROLS:
                ok = sem_code == 0 and all(c == 0 for c in other)
                wrong += int(not ok)
            else:
                ok = sem_code != 0 and EXPECT[name] in found
                wrong += int(sem_code != 0 and EXPECT[name] not in found)
            caught += int(ok)
            print("[%s] %s expect %s got %s build %s exits %s" % (
                "CAUGHT" if ok else "MISSED", name, EXPECT[name],
                ",".join(found[:4]) or "none", build_code, [sem_code] + other))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
    print("=" * 72)
    print("CAUGHT %d/%d  WRONG-GATE %d" % (caught, len(cases), wrong))
    if caught != len(cases) or wrong:
        sys.exit(1)


if __name__ == "__main__":
    main()
