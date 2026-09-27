"""Build Chapter 4 investigation student worksheets (sound/light/heat/electricity)."""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from ch4_common import *

OUT = os.path.join(os.path.dirname(__file__), '..', 'production', 'worksheets')

def build_sound():
    d = Document()
    setup_landscape(d)
    title(d, "Sound Transfer — Investigation",
          "Chapter 4 · Energy Transfer · Week 4 (Days 1–2)")
    name_date(d)

    section(d, "BIG QUESTION")
    prompt(d, "Can sound transfer energy and cause a change without direct contact?")

    section(d, "PREDICT")
    prompt(d, "What do you think will happen to the paper when the cup is tapped?")
    answer_lines(d, 2)

    section(d, "SETUP")
    prompt(d, "Draw your setup. Label the SOUND SOURCE and the PAPER TARGET.")
    drawing_box(d, height_in=1.6, label="Draw here: cup + tapping object + paper target")
    add_par(d, "The cup must NOT touch the paper.", size=10, italic=True, color=GRAY, space_after=6)

    section(d, "OBSERVE")
    prompt(d, "Tap the bottom of the cup. Watch the paper. Record what you see. "
              "Keep your body very still — no bumping the table.")
    t = d.add_table(rows=5, cols=4)
    t.style = 'Table Grid'
    headers = ["Distance (qualitative)", "What did the paper do?", "Clear / Slight / None", "What is your evidence?"]
    for i, h in enumerate(headers):
        set_cell_text(t.cell(0, i), h, size=11, bold=True, color=RGBColor(255,255,255))
    rows = ["Near", "Middle", "Farther", "Notes"]
    for r, label in enumerate(rows, start=1):
        set_cell_text(t.cell(r, 0), label, size=11, bold=True)
    style_table(t)
    add_par(d, "Distance is qualitative — just near / middle / farther. You do NOT need exact numbers.",
            size=10, italic=True, color=GRAY, space_before=4, space_after=8)

    section(d, "EXPLAIN")
    prompt(d, "How does the paper's movement provide evidence that sound transferred energy?")
    answer_lines(d, 3)

    d.save(os.path.join(OUT, "04-ch4-sound-transfer-student.docx"))
    print("built sound")

def build_light():
    d = Document()
    setup_landscape(d)
    title(d, "Light Transfer — Investigation",
          "Chapter 4 · Energy Transfer · Week 4 (Day 3)")
    name_date(d)

    section(d, "BIG QUESTION")
    prompt(d, "How can light transfer energy from one place to another?")

    section(d, "PREDICT")
    prompt(d, "What will change on the target when the flashlight is turned on?")
    answer_lines(d, 2)

    section(d, "TRACE THE PATH")
    prompt(d, "Fill in the model for your light investigation:")
    t = d.add_table(rows=2, cols=4)
    t.style = 'Table Grid'
    for i, h in enumerate(["SOURCE", "PATH", "TARGET", "OBSERVED CHANGE"]):
        set_cell_text(t.cell(0, i), h, size=11, bold=True, color=RGBColor(255,255,255))
    for i in range(4):
        set_cell_text(t.cell(1, i), " ", size=11)
    style_table(t)
    add_par(d, "We observe light reaching the target — then we INFER that energy transferred. "
              "We do not 'see energy.'", size=10, italic=True, color=GRAY, space_before=4, space_after=8)

    section(d, "TEST")
    prompt(d, "1. Shine the flashlight on the target. What changed?")
    answer_lines(d, 2)
    prompt(d, "2. (Barrier test) Put a barrier between the flashlight and the target. What happens to the light spot?")
    answer_lines(d, 2)
    prompt(d, "3. Remove the barrier. What returns?")
    answer_lines(d, 2)

    section(d, "DRAW")
    prompt(d, "Draw the light path from the flashlight to the target. Add the barrier.")
    drawing_box(d, height_in=1.4, label="Draw here: flashlight → beam → target (with barrier)")

    section(d, "EXPLAIN")
    prompt(d, "What evidence shows that light transferred energy from the flashlight to the target?")
    answer_lines(d, 3)

    d.save(os.path.join(OUT, "04-ch4-light-transfer-student.docx"))
    print("built light")

def build_heat():
    d = Document()
    setup_landscape(d)
    title(d, "Heat Transfer — Investigation",
          "Chapter 4 · Energy Transfer · Week 5 (Day 1)")
    name_date(d)

    section(d, "BIG QUESTION")
    prompt(d, "What evidence shows that energy transferred between warmer and cooler materials?")

    section(d, "PREDICT")
    prompt(d, "What do you think will happen to the temperature of the warm water? Of the cooler object?")
    answer_lines(d, 2)

    section(d, "SAFETY")
    add_par(d, "Water is warm (not hot), controlled by the teacher. No hot plates. Handle the thermometer carefully.",
            size=11, bold=True)

    section(d, "MEASURE")
    prompt(d, "Record the temperature in °C. Repeat the reading if you can.")
    t = d.add_table(rows=4, cols=3)
    t.style = 'Table Grid'
    for i, h in enumerate(["Material", "Temperature at start (°C)", "Temperature later (°C)"]):
        set_cell_text(t.cell(0, i), h, size=11, bold=True, color=RGBColor(255,255,255))
    set_cell_text(t.cell(1, 0), "Warm water", size=11, bold=True)
    set_cell_text(t.cell(2, 0), "Cooler object/container", size=11, bold=True)
    set_cell_text(t.cell(3, 0), "Reading 2 (optional)", size=11, bold=True)
    style_table(t)
    add_par(d, "Warm water = about 40 °C or less. Cooler object = room temperature or cooler.",
            size=10, italic=True, color=GRAY, space_before=4, space_after=8)

    section(d, "OBSERVE")
    prompt(d, "Which material got cooler? Which material got warmer?")
    answer_lines(d, 2)

    section(d, "EXPLAIN")
    prompt(d, "How is the temperature change evidence that heat transferred energy?")
    answer_lines(d, 3)

    d.save(os.path.join(OUT, "04-ch4-heat-transfer-student.docx"))
    print("built heat")

def build_electricity():
    d = Document()
    setup_landscape(d)
    title(d, "Electricity Transfer — Investigation",
          "Chapter 4 · Energy Transfer · Week 5 (Day 2)")
    name_date(d)

    section(d, "BIG QUESTION")
    prompt(d, "What evidence shows that electric current can transfer energy from place to place?")

    section(d, "PREDICT")
    prompt(d, "What will the device do when the circuit is complete?")
    answer_lines(d, 2)

    section(d, "SAFETY")
    add_par(d, "Low-voltage battery only (AA/AAA/D cell). No wall outlets. No coin cells.", size=11, bold=True)

    section(d, "MODEL")
    prompt(d, "Fill in the model for your circuit:")
    t = d.add_table(rows=2, cols=4)
    t.style = 'Table Grid'
    for i, h in enumerate(["BATTERY", "CIRCUIT", "DEVICE", "CHANGE"]):
        set_cell_text(t.cell(0, i), h, size=11, bold=True, color=RGBColor(255,255,255))
    for i in range(4):
        set_cell_text(t.cell(1, i), " ", size=11)
    style_table(t)
    add_par(d, "The device can light up (bulb/LED), make sound (buzzer), or move (motor).",
            size=10, italic=True, color=GRAY, space_before=4, space_after=8)

    section(d, "OBSERVE")
    prompt(d, "1. Complete the circuit. What changed in the device?")
    answer_lines(d, 2)
    prompt(d, "2. Break the circuit (disconnect one wire). What stops?")
    answer_lines(d, 2)

    section(d, "CLASSIFY")
    prompt(d, "Did the energy only MOVE, only CHANGE FORM, or BOTH? Circle one and explain.")
    add_par(d, "TRANSFER    TRANSFORMATION    BOTH", size=13, bold=True, space_before=2, space_after=2)
    answer_lines(d, 2)

    section(d, "EXPLAIN")
    prompt(d, "What evidence shows that electric current transferred energy from the battery to the device?")
    answer_lines(d, 3)

    d.save(os.path.join(OUT, "04-ch4-electricity-transfer-student.docx"))
    print("built electricity")

if __name__ == "__main__":
    build_sound()
    build_light()
    build_heat()
    build_electricity()
    print("done")
