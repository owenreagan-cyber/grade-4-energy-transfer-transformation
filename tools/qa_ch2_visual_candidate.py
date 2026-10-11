"""Deterministic content and structure QA for the 49-slide Chapter 2 visual candidate.

Usage:
  python tools/qa_ch2_visual_candidate.py DECK.pptx REPORT.md
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from pptx import Presentation


EXPECTED_SIZE = (12192000, 6858000)
NOTE_FIELDS = (
    "MODE ",
    "TIME ",
    "SESSION ",
    "PLAN ",
    "STUDENT RESOURCE ",
    "READER PAGE ",
    "SAY / ASK ",
    "DO ",
    "LISTEN FOR ",
    "REVEAL ",
    "MISCONCEPTION ",
    "NEXT ",
    "IF SHORT ",
    "MATERIALS ",
)

SLIDE_NAMES = [
    "Chapter opener", "Big Question", "BQ first thinking", "Marble before",
    "Predict", "Optional media", "What changed", "Evidence model",
    "V1 teach", "V1 write", "V1 check", "Own/outside bridge",
    "N1 write", "N1 check", "Optional basketball", "V2 teach",
    "V2 write", "V2 check", "N2 model", "N2 write", "N2 draw",
    "N2 check", "N3 overview", "N3 table write", "N3 table check",
    "CHECK write", "CHECK reveal", "Optional application",
    "Optional observation/inference", "BQ now", "Quiz directions",
    "Review 1 question", "Review 1 answer", "Review 2 question",
    "Review 2 answer", "Review 3 question", "Review 3 answer",
    "Review 4 question", "Review 4 answer", "Review 5 question",
    "Review 5 answer", "Moving Marbles opener", "Materials and safety",
    "Plan and predict", "Test 1", "Record evidence", "Explain evidence",
    "Reset and repeat", "Extension reflection",
]

BQ = "How are energy, change, and movement of objects related?"

ANCHORS = {
    2: (BQ,),
    3: (BQ, "My first thinking"),
    10: ("motion — the process of an object changing ______________ .",),
    11: ("motion — the process of an object changing position .",),
    13: (
        "A moving object can cause a ______________ .",
        "Some objects cause their own motion. Other objects are put into motion by",
        "____________ motion.",
    ),
    14: (
        "A moving object can cause a change .",
        "Some objects cause their own motion. Other objects are put into motion by",
        "outside motion.",
    ),
    17: ("energy of motion — the energy an object possesses while it is ______________ .",),
    18: ("energy of motion — the energy an object possesses while it is moving .",),
    20: ("All moving objects have energy of ______________ .",),
    21: ("OBSERVE: one moving marble rolls into a target marble.", "What changed?"),
    22: ("All moving objects have energy of motion .",),
    24: ("Write or draw the change each moving object causes.",),
    25: ("drives a nail into the wood", "turns the clock gears", "spins the turbine blades"),
    26: ("How do we know a moving object has energy?", "We know because the moving object causes"),
    27: ("We know because the moving object causes a change",),
    30: (BQ, "NOW I THINK"),
    31: ("12 POINTS", "about 12 minutes"),
}

REVIEW = [
    (
        "A moving object has energy. What is this energy called?",
        ("stored energy", "energy of motion", "energy of light", "energy of sound"),
        1,
        "Energy of motion (Reader p. 9).",
    ),
    (
        "A marble rolls and hits a still marble. What change do you observe?",
        ("Both marbles stop moving", "The still marble starts to move", "The marbles change color", "Nothing happens"),
        1,
        "The still marble starts to move (Reader p. 8).",
    ),
    (
        "A swinging hammer hits a nail. What change does the hammer cause?",
        ("The hammer turns to liquid", "The hammer floats away", "The nail is driven into the wood", "The nail melts"),
        2,
        "The hammer drives the nail in (Reader p. 10).",
    ),
    (
        "How do we know a moving object has energy if we cannot see the energy itself?",
        ("We can always see energy directly", "We guess without any evidence", "Moving objects never have energy", "We observe the change it causes"),
        3,
        "We know because it causes a change.",
    ),
    (
        "Wind (moving air) spins windmill blades. What does this show?",
        ("Moving air has energy of motion", "Still air is stronger", "Windmills work without energy", "Blades move on their own"),
        0,
        "Moving air has energy of motion (Reader p. 12).",
    ),
]


def slide_text(slide) -> str:
    chunks: list[str] = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            value = shape.text.strip()
            if value:
                chunks.append(value)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    value = cell.text.strip()
                    if value:
                        chunks.append(value)
    return "\n".join(chunks)


def normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def main() -> int:
    if len(sys.argv) != 3:
        raise SystemExit("usage: qa_ch2_visual_candidate.py DECK.pptx REPORT.md")

    deck = Path(sys.argv[1]).resolve()
    report = Path(sys.argv[2]).resolve()
    prs = Presentation(deck)
    per_slide: dict[int, list[str]] = {n: [] for n in range(1, 50)}
    global_issues: list[str] = []

    if len(prs.slides) != 49:
        global_issues.append(f"slide count is {len(prs.slides)}, expected 49")
    if (prs.slide_width, prs.slide_height) != EXPECTED_SIZE:
        global_issues.append(f"slide size is {(prs.slide_width, prs.slide_height)}, expected {EXPECTED_SIZE}")

    texts: dict[int, str] = {}
    notes_by_slide: dict[int, str] = {}
    for n, slide in enumerate(prs.slides, 1):
        texts[n] = slide_text(slide)
        notes = slide.notes_slide.notes_text_frame.text
        notes_by_slide[n] = notes
        lines = notes.splitlines()
        for field in NOTE_FIELDS:
            if not any(line.startswith(field) for line in lines):
                per_slide[n].append(f"speaker notes missing `{field.strip()}` field")
        plan_lines = [line for line in lines if line.startswith("PLAN ")]
        if len(plan_lines) != 1 or plan_lines[0] not in ("PLAN REQUIRED", "PLAN OPTIONAL"):
            per_slide[n].append("speaker notes PLAN must be exactly REQUIRED or OPTIONAL")

    for n, expected in ANCHORS.items():
        for phrase in expected:
            if normalize(phrase) not in normalize(texts[n]):
                per_slide[n].append(f"missing required anchor text: {phrase}")

    # The exact answer must remain absent from each write/prompt slide.
    leakage = {
        10: ("motion — the process of an object changing position .",),
        13: ("A moving object can cause a change .", "outside motion."),
        17: ("while it is moving .",),
        20: ("All moving objects have energy of motion .",),
        24: ("drives a nail into the wood", "turns the clock gears", "spins the turbine blades"),
        26: ("We know because the moving object causes a change.",),
        30: ("Moving objects have energy because", "causes a change, so"),
    }
    for n, forbidden in leakage.items():
        lower = normalize(texts[n]).casefold()
        for phrase in forbidden:
            if normalize(phrase).casefold() in lower:
                per_slide[n].append(f"answer leakage before reveal: {phrase}")

    for write_n, reveal_n in ((10, 11), (13, 14), (17, 18), (20, 22), (21, 22), (24, 25), (26, 27)):
        if write_n >= reveal_n:
            per_slide[write_n].append("reveal order is invalid")

    letters = "ABCD"
    for idx, (question, options, correct, rationale) in enumerate(REVIEW):
        q_slide = 32 + idx * 2
        a_slide = q_slide + 1
        for n in (q_slide, a_slide):
            if question not in texts[n]:
                per_slide[n].append("canonical review question differs")
            for option in options:
                if option not in texts[n]:
                    per_slide[n].append(f"canonical option missing: {option}")
        if rationale in texts[q_slide]:
            per_slide[q_slide].append("review rationale leaked onto question slide")
        if rationale not in texts[a_slide]:
            per_slide[a_slide].append("canonical review rationale missing")
        expected_reveal = f"{letters[correct]}: {options[correct]}"
        if expected_reveal not in notes_by_slide[a_slide]:
            per_slide[a_slide].append(f"speaker-note answer key differs; expected {expected_reveal}")

    banned = (
        "kinetic", "potential", "gravitational", "mechanical", "radiant",
        "inertia", "momentum", "compression wave", "energy transferred",
        "energy transfer", "energy transformation", "transforms energy",
    )
    for n, text in texts.items():
        lower = text.casefold()
        # The unit footer/title is the only allowed transfer/transformation mention.
        scope_text = lower.replace("energy transfer & transformation", "")
        for term in banned:
            if term in scope_text:
                per_slide[n].append(f"out-of-scope or advanced vocabulary detected: {term}")
        if re.search(r"\bproof\b|\bproves\b", lower):
            per_slide[n].append("overclaim uses proof/proves instead of evidence")
        if "faster = more energy" in lower or "more energy than" in lower:
            per_slide[n].append("Chapter 3 speed–energy relationship leaked into Chapter 2")

    required_picture_slides = (1, 4, 7, 9, 12, 15, 16, 23, 42, 43, 48)
    for n in required_picture_slides:
        pictures = sum(1 for shape in prs.slides[n - 1].shapes if shape.shape_type == 13)
        if pictures < 1:
            per_slide[n].append("expected visual asset is missing")

    for n in (24, 25):
        if not any(getattr(shape, "has_table", False) for shape in prs.slides[n - 1].shapes):
            per_slide[n].append("required native editable table is missing")

    media = texts[6].casefold()
    if "unavailable" not in media or not any(word in media for word in ("fallback", "instead", "use", "continue")):
        per_slide[6].append("media shell lacks a visible unavailable-media fallback")

    for n in range(42, 50):
        visible = texts[n].casefold()
        if "optional" not in visible or "not graded" not in visible:
            per_slide[n].append("extension slide must visibly say OPTIONAL and NOT GRADED")

    version_match = re.search(r"-v(\d+)$", deck.stem)
    candidate_version = f"v{version_match.group(1)}" if version_match else "v1"
    build_dir = deck.parents[2] / ".." / "work" / f"ch2-visual-candidate-{candidate_version}"
    for name in ("CH2_SLIDE_MANIFEST.json", "CH2_INSTRUCTIONAL_ALIGNMENT_CONTRACT.json"):
        path = build_dir.resolve() / name
        if not path.exists():
            global_issues.append(f"missing companion {name}")
            continue
        data = json.loads(path.read_text(encoding="utf-8"))
        count = data.get("slide_count") or len(data.get("slides", []))
        if count != 49:
            global_issues.append(f"{name} reports {count} slides, expected 49")

    slide_failures = sum(bool(issues) for issues in per_slide.values())
    fail_count = slide_failures + len(global_issues)
    report.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# Chapter 2 Visual Candidate V1 — Deterministic Content QA",
        "",
        f"Artifact: `{deck}`",
        "",
        f"Overall: **{'PASS' if fail_count == 0 else 'FAIL'}**",
        "",
        "Checks: 49-slide count, 16:9 size, exact Big Question and Guided Notes anchor text, write-before-reveal sequencing, answer leakage, canonical review wording and key, 14-field speaker-note schema, scope and vocabulary guardrails, media fallback, required visuals, optional-extension labeling, native editable tables, and companion manifest counts.",
        "",
        "| Slide | Purpose | Result | Detail |",
        "|---:|---|---|---|",
    ]
    for n in range(1, 50):
        issues = per_slide[n]
        lines.append(
            f"| {n} | {SLIDE_NAMES[n - 1]} | **{'FAIL' if issues else 'PASS'}** | "
            f"{'<br>'.join(issues) if issues else 'All deterministic checks passed.'} |"
        )
    lines.extend(["", "## Global checks", ""])
    if global_issues:
        lines.extend(f"- FAIL: {issue}" for issue in global_issues)
    else:
        lines.append("- PASS: deck dimensions/count and both 49-slide companion manifests match.")
    lines.extend([
        "", "## Answer-leakage summary", "",
        f"- Write/prompt slides checked: {', '.join(map(str, leakage))}.",
        f"- Result: **{'PASS' if not any(per_slide[n] for n in leakage) else 'FAIL'}**.",
        "", "## Summary", "",
        f"- PASS slides: {sum(not issues for issues in per_slide.values())}",
        f"- FAIL slides: {slide_failures}",
        f"- Global issues: {len(global_issues)}", "",
    ])
    report.write_text("\n".join(lines), encoding="utf-8")
    print(f"{'PASS' if fail_count == 0 else 'FAIL'} {deck.name}: slide_failures={slide_failures} global_issues={len(global_issues)}")
    print(f"report: {report}")
    return 1 if fail_count else 0


if __name__ == "__main__":
    raise SystemExit(main())
