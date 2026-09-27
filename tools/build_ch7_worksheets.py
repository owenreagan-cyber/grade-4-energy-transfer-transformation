"""Build Chapter 7 unit synthesis review + Energy's Big Questions worksheets (DOCX)."""
import os, sys
sys.path.insert(0, os.path.dirname(__file__))
from ch4_common import *
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

WS = "production/worksheets"

# ================= UNIT SYNTHESIS REVIEW =================
doc = Document()
setup_portrait(doc)
title(doc, "Unit Synthesis Review", "Grade 4 Energy Transfer & Transformation — Week 9")
name_date(doc)
add_par(doc, "Use varied retrieval — not just memory. Show evidence for your ideas.", size=11, italic=True, color=GRAY, space_after=8)

# A. ENERGY -> CHANGE -> EVIDENCE
section(doc, "A. ENERGY → CHANGE → EVIDENCE")
prompt(doc, "1. Energy is the ability to ______________ .")
answer_lines(doc, 1)
prompt(doc, "2. Look at each change. Circle the BEST evidence that energy is involved.")
prompt(doc, "   a) A ball rolls across the floor.")
prompt(doc, "   b) A bell rings across the playground.")
prompt(doc, "   c) A plant grows taller.")
prompt(doc, "3. Pick ONE change above. Write the evidence you would observe:")
answer_lines(doc, 2)

# B. MOTION + ENERGY OF MOTION
section(doc, "B. MOTION + ENERGY OF MOTION")
prompt(doc, "4. What does a moving object have that a still object does not?")
answer_lines(doc, 1)
prompt(doc, "5. Give ONE example of a moving object causing a change:")
answer_lines(doc, 2)

# C. SAME OBJECT / DIFFERENT SPEED
section(doc, "C. SAME OBJECT / DIFFERENT SPEED")
prompt(doc, "6. Read the table. The SAME toy car rolled down a ramp at different heights and pushed a cup.")
t = doc.add_table(rows=4, cols=3)
style_table(t)
hdr = ["Ramp height", "Cup movement (measured)", "Energy of motion (inferred)"]
for j, h in enumerate(hdr):
    set_cell_text(t.rows[0].cells[j], h, size=10, bold=True, color=RGBColor(0xFF,0xFF,0xFF))
data = [["Low", "moved a little", "less"], ["Medium", "moved more", "more"], ["High", "moved the most", "most"]]
for i, row in enumerate(data):
    for j, v in enumerate(row):
        set_cell_text(t.rows[i+1].cells[j], v, size=11)
prompt(doc, "   What do we MEASURE? What do we INFER? Circle one pair:")
prompt(doc, "   (A) measure cup movement, infer speed   (B) measure speed, infer cup movement")
answer_lines(doc, 1)
prompt(doc, "7. For the SAME object, as speed goes up, energy of motion goes ________ .")
answer_lines(doc, 1)

# D. TRANSFER PATHS
section(doc, "D. TRANSFER PATHS")
prompt(doc, "8. Name the FOUR ways energy can transfer from place to place:")
prompt(doc, "   ____________   ____________   ____________   ____________")
prompt(doc, "9. Warm water warms a cool spoon. What kind of energy transfers to the spoon?")
answer_lines(doc, 1)

# E. TRANSFER / TRANSFORMATION / BOTH
section(doc, "E. TRANSFER / TRANSFORMATION / BOTH")
prompt(doc, "10. Label each as TRANSFER, TRANSFORMATION, or BOTH.")
prompt(doc, "    a) A flashlight turns on (battery → light): ________________")
prompt(doc, "    b) Warm water warms a spoon (heat moves, stays heat): ________________")
prompt(doc, "    c) A domino falls and clicks (motion passes + sound): ________________")
prompt(doc, "    d) A buzzer buzzes (electrical → sound): ________________")

# F. COLLISION PREDICTIONS
section(doc, "F. COLLISION PREDICTIONS")
prompt(doc, "11. A toy car rolls into a stack of blocks. The top block falls off.")
prompt(doc, "    You add weight to the car (change ONLY weight) and test again.")
prompt(doc, "    Predict what happens, and explain using energy:")
answer_lines(doc, 2)
prompt(doc, "12. During a collision, energy of motion transfers to the target, and some becomes ________ .")
answer_lines(doc, 1)

# G. DATA / GRAPH READING
section(doc, "G. DATA / GRAPH READING")
prompt(doc, "13. The graph shows cup movement at different ramp heights for the SAME car.")
prompt(doc, "    As ramp height increases, cup movement ________________ .")
prompt(doc, "    This is evidence that ________________ .")
answer_lines(doc, 2)

# H. ENGINEERING
section(doc, "H. ENGINEERING — criteria / constraints / test / refine")
prompt(doc, "14. Match the word to its meaning by drawing a line:")
prompt(doc, "    criteria        the limits (materials, time, safety)")
prompt(doc, "    constraints     what a good solution must do")
prompt(doc, "    refine          to improve based on test evidence")
prompt(doc, "15. A fair test changes ________ thing(s) and keeps everything else the same.")
answer_lines(doc, 1)

# I. EDISON / PERSISTENCE
section(doc, "I. EDISON / ENGINEERING PERSISTENCE")
prompt(doc, "16. Did Edison invent the FIRST light bulb? Circle: YES / NO")
prompt(doc, "17. What did Edison actually do? (one sentence)")
answer_lines(doc, 2)
prompt(doc, "18. How many materials did Edison's team test for the filament? (about) ________")
prompt(doc, "19. What does persistence mean in engineering?")
answer_lines(doc, 2)

# J. SHORT EXPLANATION
section(doc, "J. SHORT EXPLANATION")
prompt(doc, "20. Pick ONE device (flashlight, buzzer, ramp toy, or warm-water spoon).")
prompt(doc, "    Describe the energy change from beginning to end. Use TRANSFER or TRANSFORMATION.")
prompt(doc, "    Claim: ________________________________________________")
prompt(doc, "    Evidence: _____________________________________________")
prompt(doc, "    Reasoning (why the evidence supports your claim): ________________________")
answer_lines(doc, 3)

add_par(doc, "Provenance: ADAPTED (synthesis aligned to TG Unit Review pp.94-98 + official assessment concepts, re-authored).",
        size=9, italic=True, color=GRAY)
doc.save(os.path.join(WS, "07-unit-synthesis-review.docx"))
print("wrote", os.path.join(WS, "07-unit-synthesis-review.docx"))

# ================= ENERGY'S BIG QUESTIONS =================
doc = Document()
setup_portrait(doc)
title(doc, "Energy's Big Questions", "Unit Synthesis — aligned to AP UR.1")
name_date(doc)
add_par(doc, "Answer in complete sentences. Underline any Core Vocabulary you use. Use EVIDENCE, not proof.",
        size=11, italic=True, color=GRAY, space_after=8)

qs = [
    ("1. What is energy?", "Energy is the ability to cause change."),
    ("2. How do observable changes provide evidence that energy is involved?", "A change we can see, hear, or measure shows energy is involved."),
    ("3. How does a moving object show it has energy?", "A moving object has energy of motion; it can cause a change."),
    ("4. How does speed affect energy of motion for the SAME object?", "Faster = more energy of motion = bigger change (same object)."),
    ("5. How can energy transfer from place to place?", "By sound, light, heat, and electric current."),
    ("6. What happens to energy in a collision?", "Energy of motion transfers to the target; some becomes sound. It does NOT disappear."),
    ("7. What is transformation?", "Energy changes from one form to another."),
    ("8. When can BOTH transfer and transformation occur?", "When energy moves AND changes form (e.g., dominoes, flashlight)."),
    ("9. How do engineers use testing and evidence?", "They test, record evidence, and change one thing to improve."),
    ("10. What did Edison's work show about engineering persistence?", "Testing and redesign — many tries — can improve a design over time."),
]
for q, hint in qs:
    section(doc, q)
    answer_lines(doc, 2)

add_par(doc, "Provenance: ADAPTED (aligned to AP UR.1 Energy's Big Questions; re-authored).",
        size=9, italic=True, color=GRAY)
doc.save(os.path.join(WS, "07-energy-big-questions.docx"))
print("wrote", os.path.join(WS, "07-energy-big-questions.docx"))
