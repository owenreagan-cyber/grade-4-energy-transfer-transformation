"""Chapter 3 guided notes and key. Notes-first. Reads the chapter spec.

Theme: Measurement / Data. Cobalt #1B4FC2, chartreuse #5E8A18, charcoal #2B2F36, light gray #E9EBEF.
Run: .venv/bin/python tools/build_ch3_guided_notes.py
"""
import json
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
SPEC = json.load(open(os.path.join(ROOT, "docs/chapter-03/CH3_CHAPTER_SPEC.json"), encoding="utf-8"))
OUT_DIR = os.path.join(ROOT, "production/premium-notes")

INK = RGBColor(0x1A, 0x1D, 0x26)
INK_SOFT = RGBColor(0x4A, 0x52, 0x68)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PRIMARY = RGBColor(0x1B, 0x4F, 0xC2)
ACCENT = RGBColor(0x5E, 0x8A, 0x18)
NEUTRAL = RGBColor(0x2B, 0x2F, 0x36)
PRIMARY_HEX = "1B4FC2"
ACCENT_HEX = "5E8A18"
NEUTRAL_HEX = "2B2F36"
INK_HEX = "1A1D26"
MIST_HEX = "E9EBEF"
FONT = "Avenir Next"
BODY = 14.0
PROMPT = 14.5
SECTION = 16.0
TITLE = 22.0
CAPTION = 12.0
ATTR = 12.0
MARGIN = Inches(0.55)
USABLE = 8.5 - 1.1
FULL_W = Inches(USABLE)


def blank_count(target_in, size_pt):
    return max(2, int(round(target_in * 144.0 / size_pt)))


def _shade(cell, hexstr):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexstr)
    tcPr.append(shd)


def _margins(cell, top=60, left=90, bottom=60, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        edge = OxmlElement("w:%s" % tag)
        edge.set(qn("w:w"), str(value))
        edge.set(qn("w:type"), "dxa")
        mar.append(edge)
    tcPr.append(mar)


def _border(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge, spec in edges.items():
        node = OxmlElement("w:%s" % edge)
        node.set(qn("w:val"), spec.get("val", "single"))
        node.set(qn("w:sz"), str(spec.get("sz", 6)))
        node.set(qn("w:color"), spec.get("color", INK_HEX))
        borders.append(node)
    tcPr.append(borders)


def _no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement("w:%s" % edge)
        node.set(qn("w:val"), "none")
        borders.append(node)
    tblPr.append(borders)


def _row_height(row, dxa, rule="atLeast"):
    trPr = row._tr.get_or_add_trPr()
    tr_h = OxmlElement("w:trHeight")
    tr_h.set(qn("w:val"), str(dxa))
    tr_h.set(qn("w:hRule"), rule)
    trPr.append(tr_h)


def _vcenter(cell):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _run(p, text, size=BODY, bold=False, italic=False, color=INK, underline=False):
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    run.font.name = FONT
    run.font.underline = underline
    rPr = run._element.get_or_add_rPr()
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:cs"), FONT)
    rPr.append(fonts)
    return run


def _spacing(p, before=0, after=4, line_pt=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line_pt is not None:
        pf.line_spacing = Pt(line_pt)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    return p


def _para(doc, text="", size=BODY, bold=False, italic=False, color=INK, before=0, after=4, line_pt=None):
    p = doc.add_paragraph()
    _spacing(p, before, after, line_pt)
    if text:
        _run(p, text, size=size, bold=bold, italic=italic, color=color)
    return p


def _writing_lines(doc, n, gap_pt=28):
    table = doc.add_table(rows=n, cols=1)
    _no_borders(table)
    for row in table.rows:
        cell = row.cells[0]
        cell.width = FULL_W
        _margins(cell, top=0, left=0, bottom=0, right=0)
        _border(cell, bottom=dict(sz=8, color=INK_HEX))
        _spacing(cell.paragraphs[0], 0, 0, line_pt=2)
        _run(cell.paragraphs[0], "", size=2)
        _row_height(row, int(gap_pt * 20), rule="exact")
    return table


def _fill(doc, line, answer, is_key, width=1.6):
    lead, trail = line.split("____", 1)
    p = doc.add_paragraph()
    _spacing(p, 0, 6, line_pt=26)
    _run(p, lead, size=BODY, color=INK)
    if is_key:
        _run(p, answer, size=BODY, bold=True, color=NEUTRAL)
    else:
        _run(p, "_" * blank_count(width, BODY), size=BODY, underline=True, color=INK)
    if trail:
        _run(p, trail, size=BODY, color=INK)
    return p


def _band(doc, text):
    table = doc.add_table(rows=1, cols=1)
    _no_borders(table)
    cell = table.cell(0, 0)
    cell.width = FULL_W
    _margins(cell, top=70, left=110, bottom=70, right=110)
    _shade(cell, PRIMARY_HEX)
    _vcenter(cell)
    _run(cell.paragraphs[0], text, size=15, bold=True, color=WHITE)
    return table


def _section(doc, code, title):
    p = doc.add_paragraph()
    _spacing(p, 8, 2)
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:color"), PRIMARY_HEX)
    borders.append(bottom)
    pPr.append(borders)
    _run(p, code + "   ", size=SECTION, bold=True, color=ACCENT)
    _run(p, title, size=SECTION, bold=True, color=PRIMARY)
    return p


def _model(doc):
    words = ["BALL", "OBSERVE", "EVIDENCE"]
    table = doc.add_table(rows=1, cols=5)
    _no_borders(table)
    for index, word in enumerate(words):
        cell = table.cell(0, index * 2)
        _shade(cell, PRIMARY_HEX if index != 1 else NEUTRAL_HEX)
        _margins(cell, top=40, left=40, bottom=40, right=40)
        _vcenter(cell)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(cell.paragraphs[0], word, size=12, bold=True, color=WHITE)
        if index < 2:
            arrow = table.cell(0, index * 2 + 1)
            _vcenter(arrow)
            arrow.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(arrow.paragraphs[0], "→", size=16, bold=True, color=ACCENT)
    _row_height(table.rows[0], 520, rule="exact")
    return table


def _header(doc, is_key):
    _para(doc, "UNIT 1  ·  ENERGY TRANSFER & TRANSFORMATION", size=ATTR, bold=True, color=INK_SOFT, after=2)
    _para(doc, "Chapter 3  ·  Energy and Speed Are Related", size=TITLE, bold=True, color=PRIMARY, after=2)
    label = "Teacher Key  ·  Guided Notes" if is_key else "GUIDED NOTES"
    _para(doc, label, size=CAPTION, bold=True, color=NEUTRAL, after=4)
    p = doc.add_paragraph()
    _spacing(p, 0, 6)
    _run(p, "Name: ", size=CAPTION, bold=True, color=INK)
    _run(p, "_" * 28, size=CAPTION, color=INK)
    _run(p, "    Date: ", size=CAPTION, bold=True, color=INK)
    _run(p, "_" * 16, size=CAPTION, color=INK)


def _footer(doc, is_key):
    footer = doc.sections[0].footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label = "Teacher Key" if is_key else "Guided Notes"
    _run(p, "Unit 1  ·  Chapter 3  ·  %s  ·  Energy and Speed Are Related" % label, size=12, color=INK_SOFT)


def _new_doc():
    doc = Document()
    for paragraph in list(doc.paragraphs):
        paragraph._p.getparent().remove(paragraph._p)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = MARGIN
    section.right_margin = MARGIN
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(BODY)
    style.font.color.rgb = INK
    return doc


def _page_ref(doc, pages):
    if pages:
        if pages.lower().startswith("tg"):
            label = "Teacher Guide " + pages[2:].lstrip(" :·")
        else:
            label = "Reader p. " + pages
        _para(doc, label, size=ATTR, italic=True, color=INK_SOFT, after=2)


def _anchor_map():
    return {row["code"]: row for row in SPEC["anchors"]}


def _render_cloze(doc, anchor, is_key):
    _section(doc, anchor["code"], anchor["title"])
    _page_ref(doc, anchor.get("readerPages", ""))
    answers = list(anchor.get("answers") or [])
    used = 0
    for line in anchor["lines"]:
        if "____" in line:
            _fill(doc, line, answers[used], is_key)
            used += 1
        else:
            _para(doc, line, size=BODY, after=4)
    return used


def build(is_key):
    doc = _new_doc()
    _footer(doc, is_key)
    anchors = _anchor_map()
    _header(doc, is_key)
    bq = anchors["BQ"]
    _band(doc, "BQ    " + bq["lines"][0])
    _para(doc, "", after=4)
    _para(doc, "My first thinking", size=PROMPT, bold=True, color=PRIMARY, after=2)
    if is_key:
        _para(doc, "Any first idea is acceptable. Do not require the relationship yet.", size=CAPTION, italic=True, color=INK_SOFT)
    _writing_lines(doc, 2)
    blanks = 0
    blanks += _render_cloze(doc, anchors["V1"], is_key)
    blanks += _render_cloze(doc, anchors["N1"], is_key)
    _para(doc, "", after=4)
    _model(doc)
    _para(doc, "Preview: one ball, two setups. Which cup move was bigger?", size=CAPTION, italic=True, color=INK_SOFT, before=6)
    _writing_lines(doc, 2)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    blanks += _render_cloze(doc, anchors["N2"], is_key)
    _para(doc, "Record the nine measured trials on the Ramp It Up sheet. The first roll is a calibration trial. Do not record it.", size=CAPTION, italic=True, color=INK_SOFT, before=6)
    _para(doc, "Which setup do you predict will move the cup farther? Why?", size=PROMPT, bold=True, color=PRIMARY, before=8, after=2)
    _writing_lines(doc, 2)
    _para(doc, "What change did you measure?", size=PROMPT, bold=True, color=PRIMARY, before=8, after=2)
    if is_key:
        _para(doc, "How far the cup moved, in centimeters. Not a speed number.", size=CAPTION, italic=True, color=INK_SOFT)
    _writing_lines(doc, 2)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    blanks += _render_cloze(doc, anchors["N3"], is_key)
    blanks += _render_cloze(doc, anchors["N4"], is_key)
    _para(doc, "A louder sound is also a bigger change you can observe.", size=CAPTION, italic=True, color=INK_SOFT, before=4)
    check = anchors["CHECK"]
    _section(doc, "CHECK", check["title"])
    _para(doc, check["lines"][0], size=PROMPT, bold=True, after=2)
    if is_key:
        _para(doc, check["sampleAnswer"], size=CAPTION, italic=True, color=INK_SOFT)
    _writing_lines(doc, 2)
    now = anchors["BQ_NOW"]
    _para(doc, "", after=6)
    _band(doc, "BQ  ·  NOW I THINK    " + now["lines"][0])
    _para(doc, "", after=4)
    if is_key:
        _para(doc, now["sampleAnswer"], size=CAPTION, italic=True, color=INK_SOFT)
    _writing_lines(doc, 3)
    name = "03-ch3-energy-speed-guided-notes-key.docx" if is_key else "03-ch3-energy-speed-guided-notes.docx"
    path = os.path.join(OUT_DIR, name)
    doc.save(path)
    return path, blanks


def main():
    student, student_blanks = build(False)
    key, key_blanks = build(True)
    expected = SPEC["notesExpectation"]["blankCount"]
    if student_blanks != expected or key_blanks != expected:
        sys.exit("blank count %s/%s, expected %s" % (student_blanks, key_blanks, expected))
    print(student)
    print(key)
    print("blanks", student_blanks)


if __name__ == "__main__":
    main()
