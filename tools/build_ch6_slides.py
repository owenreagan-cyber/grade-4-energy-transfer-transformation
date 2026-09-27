"""Build Chapter 6 Energy Transformations + Engineering Design slide deck (PPTX, 16:9)."""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = os.path.join(os.path.dirname(__file__), '..', 'production', 'slides')

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x2E, 0x6E, 0x8E)
AMBER = RGBColor(0xD9, 0xA4, 0x41)
GREEN = RGBColor(0x3E, 0x8E, 0x5A)
LIGHT = RGBColor(0xF5, 0xF3, 0xEC)
INK = RGBColor(0x1A, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x55, 0x55, 0x55)
PALE = RGBColor(0xD5, 0xDE, 0xEA)

CUE_COLORS = {
    "NOTEBOOK": RGBColor(0x6B, 0x4F, 0xA0),
    "WRITE": RGBColor(0x2E, 0x6E, 0x8E),
    "DRAW": RGBColor(0x3E, 0x8E, 0x5A),
    "LABEL": RGBColor(0xD9, 0xA4, 0x41),
    "TEST": RGBColor(0xC0, 0x4A, 0x4A),
    "REVEAL": RGBColor(0xC0, 0x4A, 0x4A),
}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def textbox(slide, x, y, w, h, text, size=24, color=INK, bold=False, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = color; r.font.name = "Arial"
    return tb


def cue_badge(slide, cue):
    if not cue:
        return
    c = CUE_COLORS.get(cue, GRAY)
    sp = rect(slide, 12.25, 0.25, 0.9, 0.5, c, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = sp.text_frame; tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = cue
    r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = "Arial"


def model_strip(slide, items, y=5.6, x=0.8, w=11.7, box_h=0.7, gap=0.3):
    n = len(items)
    box_w = (w - (n - 1) * gap) / n
    cx = x
    for i, it in enumerate(items):
        sp = rect(slide, cx, y, box_w, box_h, NAVY if i == 0 else (AMBER if i == n - 1 else TEAL),
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        tf = sp.text_frame; tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = it
        r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = "Arial"
        if i < n - 1:
            ar = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(cx + box_w + 0.02),
                                        Inches(y + box_h / 2 - 0.12), Inches(gap - 0.04), Inches(0.24))
            ar.fill.solid(); ar.fill.fore_color.rgb = GRAY; ar.line.fill.background(); ar.shadow.inherit = False
        cx += box_w + gap


def header(slide, title, subtitle, cue=None, num=None):
    rect(slide, 0, 0, 13.333, 1.0, NAVY)
    textbox(slide, 0.5, 0.12, 11.6, 0.75, title, size=30, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        textbox(slide, 0.5, 0.86, 11.6, 0.3, subtitle, size=13, color=PALE)
    if num:
        textbox(slide, 11.4, 0.32, 0.7, 0.4, str(num), size=14, color=PALE, align=PP_ALIGN.RIGHT)
    cue_badge(slide, cue)


def body(slide, lines, x=0.8, y=1.5, w=11.7, h=3.9, size=24, gap=8):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.color.rgb = INK; r.font.name = "Arial"
    return tb


def concept(title, subtitle, lines, cue=None, num=None, strip=None, size=24):
    s = add_slide()
    rect(s, 0, 0, 13.333, 7.5, LIGHT)
    header(s, title, subtitle, cue, num)
    body(s, lines, size=size)
    if strip:
        model_strip(s, strip)
    return s


MODEL = ["INPUT", "CHANGE IN FORM", "OUTPUT / CHANGE"]
THREE = ["TRANSFER", "TRANSFORMATION", "BOTH"]

# ================= Week 7 =================
s = add_slide(); rect(s, 0, 0, 13.333, 7.5, NAVY)
textbox(s, 0.8, 2.0, 11.7, 1.2, "Energy Transformations", size=46, color=WHITE, bold=True)
textbox(s, 0.8, 3.2, 11.7, 0.6, "Chapter 6 · Grade 4 · Week 7", size=20, color=PALE)
model_strip(s, MODEL, y=5.0)

s = add_slide(); rect(s, 0, 0, 13.333, 7.5, LIGHT)
rect(s, 0, 2.6, 13.333, 2.2, NAVY)
textbox(s, 0.8, 2.75, 11.7, 1.9, "How do energy transformations help people?",
        size=34, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
textbox(s, 0.8, 2.0, 11.7, 0.5, "Big Question", size=16, color=GRAY, align=PP_ALIGN.CENTER)
textbox(s, 11.4, 0.32, 0.7, 0.4, "2", size=14, color=GRAY, align=PP_ALIGN.RIGHT)

concept("Wind it up…", "Opening phenomenon",
        ["Wind the toy. Let it go. It zips across the table.",
         "The energy did not appear. It was STORED — then it CHANGED FORM."], num=3)
concept("Energy can CHANGE FORM", "Define transformation",
        ["TRANSFORMATION = energy changes from one form to another.",
         "stored energy → energy of motion (wind-up toy)",
         "electrical energy → light (flashlight)"],
        num=4, strip=["ONE FORM", "→", "ANOTHER FORM"])
concept("Review: energy can MOVE", "Define transfer",
        ["TRANSFER = energy moves from one object or place to another.",
         "Warm water warms a cool spoon: heat moves."],
        num=5, strip=["HEAT", "MOVES", "TO THE SPOON"])
concept("Two ways energy changes things", "Transfer vs transformation",
        ["TRANSFER = energy MOVES.", "TRANSFORMATION = energy CHANGES FORM.",
         "Many systems do BOTH."],
        num=6, strip=THREE)
concept("A flashlight", "INPUT → TRANSFORMATION → OUTPUT",
        ["INPUT: stored energy (battery)",
         "TRANSFORMATION: stored → electrical → light (+ a little heat)",
         "OUTPUT: light we see"],
        cue="LABEL", num=7, strip=MODEL)
concept("Notebook — Insert 1", "Transfer vs Transformation",
        ["Write the two definitions and one example each.",
         "Circle whether your example is TRANSFER, TRANSFORMATION, or BOTH."],
        cue="NOTEBOOK", num=8)

concept("Energy System Model", "Everyday systems",
        ["For each device ask: 1) What energy is at the beginning?  2) What changes?",
         "3) What observable output occurs?  4) Transfer?  5) Transformation?  6) BOTH?"],
        num=9)
concept("Buzzer", "System chain",
        ["INPUT: electrical energy → TRANSFORMATION: electrical → sound → OUTPUT: we hear a buzz.",
         "Transfer? yes (energy moves to the buzzer). Transformation? yes (→ sound). BOTH."],
        num=10, strip=["ELECTRICAL", "→ SOUND", "BUZZ"])
concept("Motor / fan", "System chain",
        ["INPUT: electrical energy → TRANSFORMATION: electrical → energy of motion → OUTPUT: the fan spins.",
         "Transfer AND transformation — BOTH."],
        num=11, strip=["ELECTRICAL", "→ MOTION", "FAN SPINS"])
concept("Wind-up toy", "System chain",
        ["INPUT: stored energy → TRANSFORMATION: stored → energy of motion → OUTPUT: the toy moves.",
         "BOTH (energy moves through the toy AND changes form)."],
        num=12, strip=["STORED", "→ MOTION", "TOY MOVES"])
concept("Notebook — Insert 2", "Energy System Model",
        ["Draw ONE device. Label INPUT → TRANSFORMATION → OUTPUT / CHANGE.",
         "Write your own chain."],
        cue="DRAW", num=13, strip=MODEL)

concept("Transfer only", "Warm water → cooler spoon",
        ["Heat MOVES from the warm water to the cooler spoon.",
         "The form (heat) stays the same. TRANSFER."],
        num=14, strip=["HEAT", "MOVES", "SAME FORM"])
concept("Transformation only", "Battery → light (simple view)",
        ["Stored energy CHANGES FORM to light.",
         "If we look only at the form change: TRANSFORMATION."],
        num=15, strip=["STORED", "→ LIGHT", "FORM CHANGED"])
concept("BOTH", "Falling dominoes",
        ["Motion TRANSFERS down the line (domino to domino).",
         "AND some motion TRANSFORMS to sound (the click)."],
        num=16, strip=THREE)
concept("BOTH", "Toy-car collision",
        ["Motion TRANSFERS to the cup (the cup moves).",
         "AND some motion TRANSFORMS to sound (we hear it)."],
        num=17, strip=THREE)
concept("Sort it out", "Transfer / Transformation / BOTH",
        ["Classify each card: warm spoon · flashlight · dominoes · buzzer · motor · toy-car collision.",
         "State the system boundary first. Don't force either/or."],
        cue="WRITE", num=18)
concept("Notebook — Insert 3", "Device Detective",
        ["Analyze ONE device: what goes in? what changes? what comes out?",
         "Circle TRANSFER / TRANSFORMATION / BOTH and give your evidence."],
        cue="NOTEBOOK", num=19)

concept("Reader p. 31", "The toaster",
        ["Problem: heat bread to just the right temperature.",
         "Solution: electrical energy → heat + motion + sound. DING!"],
        num=20, strip=["ELECTRICAL", "→ HEAT + MOTION + SOUND", "TOAST"])
concept("Reader p. 33", "Generators",
        ["Faraday (1831): moving magnets past copper wire → electrical energy.",
         "Generators convert motion → electricity (dams, wind turbines)."],
        num=21, strip=["MOTION", "→ ELECTRICAL", "POWER"])
concept("Reader p. 34", "Batteries",
        ["Batteries store energy (chemicals with stored energy).",
         "Stored energy → electrical energy → sound, light, motion (cell phone)."],
        num=22, strip=["STORED", "→ ELECTRICAL", "SOUND / LIGHT / MOTION"])
concept("Reader p. 35", "Solar cells",
        ["Solar cells transform sunlight → electrical energy.",
         "Stored in batteries to power devices at night."],
        num=23, strip=["LIGHT", "→ ELECTRICAL", "POWER"])
concept("Reader p. 36", "Engineering design",
        ["A process to develop a solution to a problem.",
         "Identify the problem, then judge solutions by practicality (cost, weight)."],
        num=24, strip=["PROBLEM", "DESIGN", "SOLUTION"])
concept("Vocabulary", "Core vocabulary",
        ["convert = to change form · conversion = the process of changing form",
         "energy transformation = one form of energy changes to another",
         "engineering design = a process to develop a solution to a problem"],
        cue="WRITE", num=25)
concept("Design-thinking bridge", "A signal without wall electricity",
        ["A classroom needs a signal without using wall electricity.",
         "What energy change could help? (brainstorm only — no building yet)"],
        cue="WRITE", num=26)
concept("Check — Lesson 12", "AP 12.1",
        ["Answer the 6 questions: Big Question, stove, blender, wind-up timer, two transformations, "
         "electrical → light.", "Use your Core Vocabulary cards."],
        cue="WRITE", num=27)
concept("Week 7 exit", "Synthesis",
        ["One sentence: How can changing energy from one form to another solve a problem?"],
        cue="WRITE", num=28)

# ================= Week 8 =================
s = add_slide(); rect(s, 0, 0, 13.333, 7.5, NAVY)
textbox(s, 0.8, 2.0, 11.7, 1.2, "Engineering Performance Task", size=46, color=WHITE, bold=True)
textbox(s, 0.8, 3.2, 11.7, 0.6, "Chapter 6 · Grade 4 · Week 8", size=20, color=PALE)
model_strip(s, ["DESIGN", "BUILD", "TEST", "REFINE"], y=5.0)

s = add_slide(); rect(s, 0, 0, 13.333, 7.5, LIGHT)
rect(s, 0, 2.6, 13.333, 2.2, NAVY)
textbox(s, 0.8, 2.75, 11.7, 1.9, "How can I design a device that transforms energy to solve a problem?",
        size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
textbox(s, 0.8, 2.0, 11.7, 0.5, "Big Question", size=16, color=GRAY, align=PP_ALIGN.CENTER)
textbox(s, 11.4, 0.32, 0.7, 0.4, "30", size=14, color=GRAY, align=PP_ALIGN.RIGHT)

concept("Your challenge", "The problem",
        ["Design a NON-ELECTRICAL tabletop device that makes a clear signal",
         "when a model door/gate/trigger moves.",
         "Your device must show at least one energy transformation."],
        num=31, strip=["PROBLEM", "→", "DEVICE"])
concept("Criteria", "What a good device must do",
        ["respond to the trigger · make an observable signal (sound / motion / visible)",
         "show one energy transformation · work repeatedly · let you explain the pathway",
         "use supplied materials · stay in the tabletop test area"],
        cue="WRITE", num=32)
concept("Constraints", "The limits",
        ["only provided materials · no wall electricity · no flame · no hot plate · no launching",
         "limited build time · limited material quantity · tabletop footprint"],
        cue="WRITE", num=33)
concept("Materials", "Low-cost, reusable-heavy",
        ["cardboard tubes · paper cups · craft sticks · rubber bands · string · paper clips · washers",
         "bells / metal cup · index cards / paper · tape · glue · blocks · foil"],
        num=34)
concept("Notebook — Insert 4", "Engineering Design",
        ["Record: the problem, the criteria, the constraints, and a labeled blueprint.",
         "The cycle: DESIGN → BUILD → TEST → REFINE."],
        cue="NOTEBOOK", num=35, strip=["DESIGN", "BUILD", "TEST", "REFINE"])
concept("Blueprint", "Day 1 — design (AP 13.1)",
        ["Include: materials · labels · trigger · output · energy-transformation explanation.",
         "NO building until the blueprint check is complete."],
        cue="DRAW", num=36)
concept("Build day", "Day 2 — build + Test 1 (AP 13.2)",
        ["Build your device. Test it 3 times.",
         "Record: did it trigger? what output? clear/reliable? what failed? what evidence?"],
        cue="TEST", num=37)
concept("Notebook — Insert 5", "Test → Change → Retest",
        ["Test 1 result → what we changed → Test 2 result.",
         "Lift the flap to check your reasoning."],
        cue="NOTEBOOK", num=38)
concept("Failure is data", "Engineering mindset",
        ["A failed test is NOT a bad engineer. It is a CLUE.",
         "Initial failure is not graded negatively."],
        cue="REVEAL", num=39)
concept("Redesign", "Day 3 — change ONE thing",
        ["Choose ONE change based on your Test 1 evidence.",
         "Record: WHAT WE CHANGED · WHY · PREDICTION."],
        cue="WRITE", num=40, strip=["EVIDENCE", "→", "ONE CHANGE"])
concept("Test 2", "Day 3 — retest",
        ["Test again under the SAME conditions.",
         "Compare Test 1 vs Test 2. Did the change help? What is the evidence?"],
        cue="TEST", num=41)
concept("Final explanation", "Day 4 — claim · evidence · reasoning",
        ["CLAIM: what my device does.", "EVIDENCE: what I observed.",
         "REASONING: why the evidence supports my claim.",
         "Use \u201csupports\u201d — NOT \u201cproves.\u201d"],
        cue="WRITE", num=42)
concept("Present", "Day 4 — AP 13.3 + AP 13.4",
        ["Share: problem · criteria · constraints · blueprint · Test 1 · redesign · Test 2 · energy pathway.",
         "Peers score with the Presentation Scoring Guide."],
        num=43, strip=["PROBLEM", "DESIGN", "TEST", "EXPLAIN"])
concept("Reflect", "Week 8 exit",
        ["What is ONE thing you would change next time, and why?",
         "What did you learn about energy transformation and engineering design?"],
        cue="WRITE", num=44)

prs.save(os.path.join(OUT, "06-ch6-energy-transformations-slides.pptx"))
print("built slides:", len(prs.slides._sldIdLst))
