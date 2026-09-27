"""Build the premium full-size Chapter 1 Guided Notes comparison prototype (US Letter, 3 pages)."""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from ch1_notes_common import *

OUT = os.path.abspath(os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "..", "production",
    "comparisons", "chapter-01-notes", "guided-notes",
    "01-ch1-guided-notes-comparison-master.docx"))

doc = Document()
set_page(doc, left=0.75, right=0.75, top=0.7, bottom=0.7)
for _p in list(doc.paragraphs):
    if not _p.text.strip():
        _p._p.getparent().remove(_p._p)

FULL_W = 1440 * 7  # 7.0" usable width


def title(text, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(text), size=22, bold=True, color=NAVY, font=HEAD_FONT)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p2.add_run(subtitle), size=12, italic=True, color=GRAY)
    # name/date line
    n = doc.add_paragraph()
    set_run(n.add_run("Name: "), size=12, bold=True, color=INK)
    r = n.add_run("____________________________")
    r.underline = True
    set_run(r, size=12, color=INK)
    set_run(n.add_run("    Date: "), size=12, bold=True, color=INK)
    r2 = n.add_run("________________")
    r2.underline = True
    set_run(r2, size=12, color=INK)


def heading(text, size=16):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    sp = OxmlElement('w:spacing'); sp.set(qn('w:before'), '120'); sp.set(qn('w:after'), '40'); pPr.append(sp)
    set_run(p.add_run(text), size=size, bold=True, color=NAVY, caps=True)
    return p


def body(text, size=13, bold=False, italic=False, color=INK, after=80, before=0):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    sp = OxmlElement('w:spacing'); sp.set(qn('w:before'), str(int(before))); sp.set(qn('w:after'), str(int(after))); pPr.append(sp)
    set_run(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return p


def wline(count=1, gap=26, after=20):
    """Writing line with ~0.33" spacing (gap pt)."""
    for _ in range(count):
        p = doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        sp = OxmlElement('w:spacing'); sp.set(qn('w:after'), str(int(after * 20)))
        sp.set(qn('w:line'), str(int(gap * 20))); sp.set(qn('w:lineRule'), 'exact')
        pPr.append(sp)
        add_paragraph_bottom_border(p, sz=8, val='single', color=INK_HEX)
    return p


def fill(lead, blank="________________________", trail="", size=13):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    sp = OxmlElement('w:spacing'); sp.set(qn('w:after'), '50'); pPr.append(sp)
    set_run(p.add_run(lead + " "), size=size, color=INK)
    r = p.add_run(blank)
    r.underline = True
    set_run(r, size=size, color=INK)
    if trail:
        set_run(p.add_run(" " + trail), size=size, color=INK)
    return p


def page_break():
    from docx.enum.text import WD_BREAK
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def box(width_dxa, height_dxa, label):
    t = doc.add_table(rows=1, cols=1)
    no_table_borders(t)
    c = t.cell(0, 0)
    set_cell_width(c, width_dxa)
    set_cell_margins(c, top=20, left=60, bottom=20, right=60)
    set_cell_border(c, top=dict(sz=8, val='single', color=NAVY_HEX),
                    left=dict(sz=8, val='single', color=NAVY_HEX),
                    bottom=dict(sz=8, val='single', color=NAVY_HEX),
                    right=dict(sz=8, val='single', color=NAVY_HEX))
    set_row_height(t.rows[0], height_dxa, rule='exact')
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(label), size=12, italic=True, color=GRAY)
    return t


# ============================================================
# PAGE 1 — Energy Causes Change
# ============================================================
title("Energy Causes Change", "Chapter 1 Guided Notes · Grade 4 Energy Transfer & Transformation")

heading("Big Question", 17)
body("Where can we observe evidence of energy causing change?", size=14, bold=True, italic=True)

heading("Key Words", 17)
fill("energy — Energy is the ability to cause", trail=".")
fill("change — Change means something becomes", trail=".")
fill("evidence — Evidence is something we observe that helps us", "________________", "what happened.")

heading("The Model", 17)
m = doc.add_table(rows=1, cols=5)
no_table_borders(m)
box_w = 2620; arr_w = 700
for idx in range(5):
    mc = m.cell(0, idx)
    set_cell_width(mc, box_w if idx % 2 == 0 else arr_w)
    set_cell_margins(mc, top=40, left=10, bottom=40, right=10)
    vcenter(mc)
    mp = mc.paragraphs[0]
    mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if idx % 2 == 0:
        shade_cell(mc, NAVY_HEX)
        set_run(mp.add_run(["ENERGY", "CHANGE", "EVIDENCE"][idx // 2]), size=13, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    else:
        set_run(mp.add_run("→"), size=15, bold=True, color=CORAL)
body("In my own words: energy causes a change, and the change is our evidence.", size=12, italic=True, color=GRAY)

heading("One Example — Before & After", 17)
body("Pick one change from class. Draw what it looked like before and after.", size=12, color=GRAY)
ba = doc.add_table(rows=1, cols=2)
no_table_borders(ba)
for idx, lab in enumerate(["BEFORE", "AFTER"]):
    bc = ba.cell(0, idx)
    set_cell_width(bc, (FULL_W - 100) // 2)
    set_cell_margins(bc, top=20, left=60, bottom=20, right=60)
    set_cell_border(bc, top=dict(sz=8, val='single', color=NAVY_HEX),
                    left=dict(sz=8, val='single', color=NAVY_HEX),
                    bottom=dict(sz=8, val='single', color=NAVY_HEX),
                    right=dict(sz=8, val='single', color=NAVY_HEX))
    set_row_height(ba.rows[0], 1900, rule='exact')
    bp = bc.paragraphs[0]
    bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(bp.add_run(lab), size=12, bold=True, color=CORAL, caps=True)
fill("What changed?", "________________________________________________", size=12)

heading("Forms of Energy", 17)
f = doc.add_table(rows=1, cols=5)
no_table_borders(f)
cues = [("MOTION", "rolling"), ("LIGHT", "bright"), ("SOUND", "loud"), ("HEAT", "warm"), ("ELECTRICAL", "powers")]
for idx, (form, cue) in enumerate(cues):
    fc = f.cell(0, idx)
    set_cell_width(fc, FULL_W // 5)
    set_cell_margins(fc, top=30, left=20, bottom=30, right=20)
    shade_cell(fc, LIGHT_HEX)
    fp = fc.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(fp.add_run(form), size=11, bold=True, color=NAVY)
    fp2 = fc.add_paragraph(); fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(fp2.add_run(cue), size=11, italic=True, color=GRAY)

page_break()

# ============================================================
# PAGE 2 — Stored Energy
# ============================================================
title("Stored Energy", "Chapter 1 Guided Notes · Page 2")

heading("Key Idea", 17)
body("Stored energy can cause a change later.", size=14, bold=True)

heading("Three Examples", 17)
body("For each, write the change that could happen later.", size=12, color=GRAY)
ex = doc.add_table(rows=3, cols=2)
no_table_borders(ex)
examples = [("battery", "in a flashlight"), ("stretched rubber band", "held still"), ("raised object", "held up high")]
for i, (obj, state) in enumerate(examples):
    ec = ex.cell(i, 0)
    set_cell_width(ec, 3400)
    set_cell_margins(ec, top=60, left=80, bottom=60, right=80)
    ep = ec.paragraphs[0]
    set_run(ep.add_run(obj), size=14, bold=True, color=NAVY)
    ep2 = ec.add_paragraph()
    set_run(ep2.add_run(state), size=11, italic=True, color=GRAY)
    wc = ex.cell(i, 1)
    set_cell_width(wc, FULL_W - 3400)
    set_cell_margins(wc, top=40, left=80, bottom=40, right=80)
    wp = wc.paragraphs[0]
    set_run(wp.add_run("What change could happen later?"), size=11, bold=True, color=INK)
    lp = wc.add_paragraph()
    lpPr = lp._p.get_or_add_pPr()
    sp = OxmlElement('w:spacing'); sp.set(qn('w:after'), '10'); sp.set(qn('w:line'), '360'); lpPr.append(sp)
    add_paragraph_bottom_border(lp, sz=8, val='single', color=INK_HEX)
    set_row_height(ex.rows[i], 950, rule='atLeast')

heading("Observation vs Explanation", 17)
body("Observation = what I see.        Explanation = what it means.", size=13)
body("Example: “The light turned on” is an observation. “Energy caused the change” is an explanation.", size=12, italic=True, color=GRAY)

heading("Draw One Stored-Energy Example", 17)
box(FULL_W, 2300, "draw your stored-energy example here")

page_break()

# ============================================================
# PAGE 3 — Energy Detectives / Synthesis
# ============================================================
title("Energy Detectives — Evidence", "Chapter 1 Guided Notes · Page 3")

heading("Evidence Table", 17)
body("Record one change you found as an Energy Detective.", size=12, color=GRAY)
et = doc.add_table(rows=4, cols=2)
no_table_borders(et)
labels = ["Before", "After", "What changed?", "What is the evidence?"]
for i, lab in enumerate(labels):
    lc = et.cell(i, 0)
    set_cell_width(lc, 2600)
    set_cell_margins(lc, top=40, left=80, bottom=40, right=80)
    lp = lc.paragraphs[0]
    set_run(lp.add_run(lab), size=13, bold=True, color=NAVY)
    bc = et.cell(i, 1)
    set_cell_width(bc, FULL_W - 2600)
    set_cell_margins(bc, top=40, left=80, bottom=40, right=80)
    bp = bc.paragraphs[0]
    add_paragraph_bottom_border(bp, sz=8, val='single', color=INK_HEX)
    set_row_height(et.rows[i], 800, rule='atLeast')

heading("Synthesis", 17)
body("How do scientists know energy is involved if energy itself cannot be seen directly?", size=13, bold=True)
wline(4, gap=30, after=20)

heading("Remember", 17)
rm = doc.add_table(rows=1, cols=3)
no_table_borders(rm)
for idx, w in enumerate(["ENERGY", "CHANGE", "EVIDENCE"]):
    rc = rm.cell(0, idx)
    set_cell_width(rc, FULL_W // 3)
    set_cell_margins(rc, top=30, left=10, bottom=30, right=10)
    shade_cell(rc, NAVY_HEX)
    rp = rc.paragraphs[0]; rp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(rp.add_run(w), size=12, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

doc.save(OUT)
print("saved", OUT)
