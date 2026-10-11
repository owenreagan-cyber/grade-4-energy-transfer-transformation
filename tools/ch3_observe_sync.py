"""Print whether each spec anchor line is present in the notes and the deck."""
import json
import re
import zipfile
from xml.etree import ElementTree as ET

from pptx import Presentation

spec = json.load(open("docs/chapter-03/CH3_CHAPTER_SPEC.json", encoding="utf-8"))
ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
root = ET.fromstring(zipfile.ZipFile("production/premium-notes/03-ch3-energy-speed-guided-notes.docx").read("word/document.xml"))
notes = re.sub(r"\s+", " ", " ".join((node.text or "") for node in root.iter(ns + "t")))
deck = Presentation("production/slides/03-ch3-energy-speed-slides.pptx")
slides = []
for slide in deck.slides:
    chunks = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            chunks.append(shape.text_frame.text)
    slides.append(re.sub(r"\s+", " ", "\n".join(chunks)))
blob = "\n".join(slides)
for anchor in spec["anchors"]:
    print("==", anchor["code"])
    for line in anchor.get("lines") or []:
        bare = re.sub(r"\s+", " ", line.replace("____", "")).strip()
        print(" notes", bare[:60] in notes, "slides", bare[:60] in blob, "|", bare[:80])
