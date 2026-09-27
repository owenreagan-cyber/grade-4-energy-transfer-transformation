"""Premium guided-note design system (Phase 12A).

Shared components so Chapters 1–7 share a recognizable design family:
BIG QUESTION · KEY IDEA · NOTES · MODEL IT · DRAW/LABEL · CHECK YOUR THINKING · CONNECT.
Reuses the palette + low-level helpers in ch1_notes_common.py.
"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from ch1_notes_common import (
    NAVY, NAVY_HEX, CORAL, CORAL_HEX, GOLD, INK, INK_HEX, GRAY,
    LIGHT_HEX, MID_HEX, BODY_FONT, HEAD_FONT,
    set_run, shade_cell, set_cell_border, set_cell_margins, set_cell_width,
    set_row_height, set_table_borders, no_table_borders, add_paragraph_bottom_border,
    para, cell_para, vcenter, set_page,
)
from ch1_notes_common import OxmlElement, qn

FULL_W = 1440 * 7  # 7.0" usable width at 0.75" margins

# ---------------------------------------------------------------- document
def new_doc():
    doc = Document()
    set_page(doc, left=0.75, right=0.75, top=0.7, bottom=0.7)
    for _p in list(doc.paragraphs):
        if not _p.text.strip():
            _p._p.getparent().remove(_p._p)
    return doc

# ---------------------------------------------------------------- title / banner
def title_block(doc, title, subtitle, new_page=False):
    p = doc.add_paragraph()
    if new_page:
        p.paragraph_format.page_break_before = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(title), size=23, bold=True, color=NAVY, font=HEAD_FONT)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p2.add_run(subtitle), size=12, italic=True, color=GRAY)
    n = doc.add_paragraph()
    pPr = n._p.get_or_add_pPr()
    sp = OxmlElement('w:spacing'); sp.set(qn('w:after'), '40'); pPr.append(sp)
    set_run(n.add_run("Name: "), size=12, bold=True, color=INK)
    r = n.add_run("____________________________"); r.underline = True
    set_run(r, size=12, color=INK)
    set_run(n.add_run("      Date: "), size=12, bold=True, color=INK)
    r2 = n.add_run("________________"); r2.underline = True
    set_run(r2, size=12, color=INK)

def big_question(doc, text):
    t = doc.add_table(rows=1, cols=1)
    no_table_borders(t)
    c = t.cell(0, 0)
    set_cell_width(c, FULL_W)
    set_cell_margins(c, top=60, left=120, bottom=60, right=120)
    shade_cell(c, NAVY_HEX)
    p = c.paragraphs[0]
    set_run(p.add_run("BIG QUESTION   "), size=13, bold=True, color=GOLD, caps=True)
    set_run(p.add_run(text), size=15, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    # spacing after
    sp = doc.add_paragraph()
    spPr = sp._p.get_or_add_pPr()
    spelm = OxmlElement('w:spacing'); spelm.set(qn('w:after'), '80'); spPr.append(spelm)
    return t

def heading(doc, text, size=16, before=100):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    sp = OxmlElement('w:spacing'); sp.set(qn('w:before'), str(before)); sp.set(qn('w:after'), '40'); pPr.append(sp)
    set_run(p.add_run(text), size=size, bold=True, color=NAVY, caps=True)
    return p

def key_idea(doc, text, size=14):
    t = doc.add_table(rows=1, cols=1)
    no_table_borders(t)
    c = t.cell(0, 0)
    set_cell_width(c, FULL_W)
    set_cell_margins(c, top=50, left=100, bottom=50, right=100)
    shade_cell(c, MID_HEX)
    p = c.paragraphs[0]
    set_run(p.add_run("KEY IDEA   "), size=12, bold=True, color=CORAL, caps=True)
    set_run(p.add_run(text), size=size, bold=True, color=INK)
    sp = doc.add_paragraph()
    spPr = sp._p.get_or_add_pPr()
    spelm = OxmlElement('w:spacing'); spelm.set(qn('w:after'), '40'); spPr.append(spelm)
    return t

# ---------------------------------------------------------------- body / writing
def body(doc, text, size=13, bold=False, italic=False, color=INK, after=70, before=0):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    sp = OxmlElement('w:spacing'); sp.set(qn('w:before'), str(before)); sp.set(qn('w:after'), str(after)); pPr.append(sp)
    set_run(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return p

def fill(doc, lead, blank="________________________", trail="", size=13, after=60, bold_lead=False):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    sp = OxmlElement('w:spacing'); sp.set(qn('w:after'), str(after)); pPr.append(sp)
    set_run(p.add_run(lead + " "), size=size, bold=bold_lead, color=INK)
    r = p.add_run(blank)
    r.underline = True
    set_run(r, size=size, color=INK)
    if trail:
        set_run(p.add_run(" " + trail), size=size, color=INK)
    return p

def wline(doc, count=1, gap=28, after=22):
    """Writing line: gap pt = line spacing (28pt ≈ 0.39")."""
    last = None
    for _ in range(count):
        p = doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        sp = OxmlElement('w:spacing'); sp.set(qn('w:after'), str(after * 20))
        sp.set(qn('w:line'), str(int(gap * 20))); sp.set(qn('w:lineRule'), 'exact')
        pPr.append(sp)
        add_paragraph_bottom_border(p, sz=8, val='single', color=INK_HEX)
        last = p
    return last

# ---------------------------------------------------------------- model / boxes
def model_boxes(doc, words, box_w=2620, arr_w=700, size=13, fill_hex=NAVY_HEX, txt_color=None):
    """ENERGY → CHANGE → EVIDENCE style arrow chain."""
    n = len(words)
    cols = n * 2 - 1
    t = doc.add_table(rows=1, cols=cols)
    no_table_borders(t)
    for idx in range(cols):
        mc = t.cell(0, idx)
        set_cell_width(mc, box_w if idx % 2 == 0 else arr_w)
        set_cell_margins(mc, top=40, left=8, bottom=40, right=8)
        vcenter(mc)
        mp = mc.paragraphs[0]
        mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if idx % 2 == 0:
            shade_cell(mc, fill_hex)
            set_run(mp.add_run(words[idx // 2]), size=size, bold=True,
                    color=txt_color or RGBColor(0xFF, 0xFF, 0xFF))
        else:
            set_run(mp.add_run("→"), size=size + 2, bold=True, color=CORAL)
    sp = doc.add_paragraph()
    spPr = sp._p.get_or_add_pPr()
    spelm = OxmlElement('w:spacing'); spelm.set(qn('w:after'), '40'); spPr.append(spelm)
    return t

def model_stack(doc, words, size=13, box_w=None, arrow="↓"):
    """Vertical model: full-width shaded boxes with down-arrows (for long 4–5 step chains at ≥12pt)."""
    box_w = box_w or (FULL_W - 1400)
    first = True
    for w in words:
        if not first:
            ap = doc.add_paragraph()
            ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            apPr = ap._p.get_or_add_pPr()
            sp = OxmlElement('w:spacing'); sp.set(qn('w:before'), '20'); sp.set(qn('w:after'), '20'); apPr.append(sp)
            set_run(ap.add_run(arrow), size=size + 2, bold=True, color=CORAL)
        t = doc.add_table(rows=1, cols=1)
        no_table_borders(t)
        t.alignment = 1
        c = t.cell(0, 0)
        set_cell_width(c, box_w)
        set_cell_margins(c, top=40, left=60, bottom=40, right=60)
        shade_cell(c, NAVY_HEX)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(w), size=size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        set_row_height(t.rows[0], 440, rule='atLeast')
        first = False
    sp = doc.add_paragraph()
    spPr = sp._p.get_or_add_pPr()
    spelm = OxmlElement('w:spacing'); spelm.set(qn('w:after'), '40'); spPr.append(spelm)
    return

def labeled_pair(doc, left_label, right_label, height_dxa=1900, left_hint="", right_hint="", label_color=CORAL):
    """BEFORE / AFTER style two-box drawing area."""
    t = doc.add_table(rows=1, cols=2)
    no_table_borders(t)
    for idx, (lab, hint) in enumerate([(left_label, left_hint), (right_label, right_hint)]):
        bc = t.cell(0, idx)
        set_cell_width(bc, (FULL_W - 120) // 2)
        set_cell_margins(bc, top=20, left=60, bottom=20, right=60)
        set_cell_border(bc, top=dict(sz=8, val='single', color=NAVY_HEX),
                        left=dict(sz=8, val='single', color=NAVY_HEX),
                        bottom=dict(sz=8, val='single', color=NAVY_HEX),
                        right=dict(sz=8, val='single', color=NAVY_HEX))
        bp = bc.paragraphs[0]
        bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(bp.add_run(lab), size=13, bold=True, color=label_color, caps=True)
        if hint:
            hp = bc.add_paragraph()
            hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run(hp.add_run(hint), size=12, italic=True, color=GRAY)
    set_row_height(t.rows[0], height_dxa, rule='exact')
    return t

def draw_box(doc, width_dxa, height_dxa, label, label_size=12):
    t = doc.add_table(rows=1, cols=1)
    no_table_borders(t)
    c = t.cell(0, 0)
    set_cell_width(c, width_dxa)
    set_cell_margins(c, top=20, left=60, bottom=20, right=60)
    set_cell_border(c, top=dict(sz=8, val='single', color=NAVY_HEX),
                    left=dict(sz=8, val='single', color=NAVY_HEX),
                    bottom=dict(sz=8, val='single', color=NAVY_HEX),
                    right=dict(sz=8, val='single', color=NAVY_HEX))
    set_row_height(t.rows[0], height_dxa, rule='exact')
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(label), size=label_size, italic=True, color=GRAY)
    return t

def two_col_table(doc, headers, rows, col_widths=(3400, FULL_W - 3400), row_h=850, header_fill=NAVY_HEX):
    """Generic labeled two-column table with a bold header row + bordered body rows."""
    t = doc.add_table(rows=1 + len(rows), cols=2)
    no_table_borders(t)
    for j, h in enumerate(headers):
        hc = t.cell(0, j)
        set_cell_width(hc, col_widths[j])
        set_cell_margins(hc, top=40, left=80, bottom=40, right=80)
        shade_cell(hc, header_fill)
        hp = hc.paragraphs[0]
        set_run(hp.add_run(h), size=12, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), caps=True)
    set_row_height(t.rows[0], 420, rule='exact')
    for i, (a, b) in enumerate(rows):
        ac = t.cell(i + 1, 0)
        set_cell_width(ac, col_widths[0])
        set_cell_margins(ac, top=40, left=80, bottom=40, right=80)
        ap = ac.paragraphs[0]
        set_run(ap.add_run(a), size=13, bold=True, color=NAVY)
        bc = t.cell(i + 1, 1)
        set_cell_width(bc, col_widths[1])
        set_cell_margins(bc, top=40, left=80, bottom=40, right=80)
        bp = bc.paragraphs[0]
        if b:
            set_run(bp.add_run(b), size=12, color=GRAY)
        add_paragraph_bottom_border(bp, sz=8, val='single', color=INK_HEX)
        set_row_height(t.rows[i + 1], row_h, rule='atLeast')
    return t

def data_table(doc, headers, rows, col_widths, row_h=720, header_fill=NAVY_HEX, first_col_bold=True):
    """Generic N-column table: shaded header row + bordered body rows (blank = writing line)."""
    ncols = len(headers)
    t = doc.add_table(rows=1 + len(rows), cols=ncols)
    no_table_borders(t)
    for j, h in enumerate(headers):
        hc = t.cell(0, j)
        set_cell_width(hc, col_widths[j])
        set_cell_margins(hc, top=40, left=60, bottom=40, right=60)
        shade_cell(hc, header_fill)
        hp = hc.paragraphs[0]
        set_run(hp.add_run(h), size=12, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), caps=True)
    set_row_height(t.rows[0], 420, rule='exact')
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i + 1, j)
            set_cell_width(c, col_widths[j])
            set_cell_margins(c, top=40, left=60, bottom=40, right=60)
            cp = c.paragraphs[0]
            if val:
                set_run(cp.add_run(val), size=12, bold=(first_col_bold and j == 0), color=NAVY if j == 0 else GRAY)
            add_paragraph_bottom_border(cp, sz=8, val='single', color=INK_HEX)
        set_row_height(t.rows[i + 1], row_h, rule='atLeast')
    sp = doc.add_paragraph()
    spPr = sp._p.get_or_add_pPr()
    spelm = OxmlElement('w:spacing'); spelm.set(qn('w:after'), '40'); spPr.append(spelm)
    return t

# ---------------------------------------------------------------- footer strips
def check_thinking(doc, prompt, lines=2):
    heading(doc, "CHECK YOUR THINKING", 15)
    body(doc, prompt, size=13, bold=True, after=30)
    wline(doc, lines, gap=30, after=18)

def connect(doc, text):
    t = doc.add_table(rows=1, cols=1)
    no_table_borders(t)
    c = t.cell(0, 0)
    set_cell_width(c, FULL_W)
    set_cell_margins(c, top=50, left=100, bottom=50, right=100)
    shade_cell(c, MID_HEX)
    p = c.paragraphs[0]
    set_run(p.add_run("CONNECT   "), size=12, bold=True, color=NAVY, caps=True)
    set_run(p.add_run(text), size=12, italic=True, color=INK)
    return t

def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
