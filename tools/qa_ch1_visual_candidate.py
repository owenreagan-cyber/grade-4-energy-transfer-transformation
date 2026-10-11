"""Deterministic content/structure QA for the 45-slide Chapter 1 visual candidate.

Usage:
  python tools/qa_ch1_visual_candidate.py DECK.pptx REPORT.md
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation


EXPECTED_SIZE = (12192000, 6858000)
NOTE_FIELDS = (
    "MODE ", "TIME ", "PLAN ", "STUDENT RESOURCE ", "ASK ",
    "LISTEN FOR ", "REVEAL ", "MISCONCEPTION ", "NEXT ", "IF SHORT ",
)

SLIDE_NAMES = [
    "Chapter opener", "Big Question", "BQ first thinking", "Dark-to-bright phenomenon",
    "Notice and wonder", "Pencil before", "Predict", "Demo media shell", "Pencil evidence",
    "Vocabulary: energy", "Vocabulary: change", "V1 write", "V1 check", "Seashore bridge",
    "N1 write", "N1 check", "Five forms overview", "Five forms examples", "Five forms recap",
    "N2 write", "N2 check", "Vocabulary: stored energy", "V2 write", "V2 check",
    "Stored-energy examples", "N3 write", "N3 check", "Evidence", "N4 write/draw",
    "Check write", "Check reveal", "Optional review media", "BQ revisit",
    "Review 1 question", "Review 1 answer", "Review 2 question", "Review 2 answer",
    "Review 3 question", "Review 3 answer", "Review 4 question", "Review 4 answer",
    "Review 5 question", "Review 5 answer", "Review 6 question", "Review 6 answer and close",
]

ANCHORS = {
    2: ("Where can we observe evidence of energy causing change?",),
    3: ("Where can we observe evidence of energy causing change?", "My first thinking"),
    12: ("energy — the ability to cause ____________ .", "change — to become ________________ ."),
    13: ("energy — the ability to cause change .", "change — to become different ."),
    15: ("Energy has the ability to cause change.", "When you see a change, you are seeing ____________ of energy."),
    16: ("Energy has the ability to cause change.", "When you see a change, you are seeing evidence of energy."),
    20: ("Write one example that shows a change for each form.",),
    21: ("SAMPLE ANSWERS — YOUR EXAMPLES MAY DIFFER",),
    23: ("stored energy — energy that has the ability to cause change at a ___________ time.",),
    24: ("stored energy — energy that has the ability to cause change at a later time.",),
    26: ("Objects also have stored energy because of their ______________ .",),
    27: ("Objects also have stored energy because of their position .",),
    29: ("Draw a change you observed. Show what it looked like BEFORE and AFTER.",),
    30: ("How do we know energy is involved if we cannot see energy itself?",),
    31: ("We know because we see a change, like a ball starting to move or a light turning on.",),
    33: ("Where can we observe evidence of energy causing change?", "NOW I THINK"),
}

REVIEW = [
    ("Energy is best described as —",
     ["a visible glow", "something only batteries have", "the ability to cause change", "a kind of fuel"],
     "Energy is the ability to cause change (Reader p. 1)."),
    ("A pencil is dull, then you sharpen it. What change do you observe?",
     ["It gets shorter and sharper", "It gets longer", "It changes color", "It disappears"],
     "The pencil becomes shorter and sharper — a change."),
    ("Which object has stored energy that can cause a change LATER?",
     ["A ball already rolling", "A toy already moving", "A flashlight already on", "A stretched rubber band"],
     "The rubber band stores energy; it changes when released."),
    ("A carousel spins and plays music. Which two forms of energy come from the spinning and the music?",
     ["only light", "motion and sound", "motion and silence", "sound and stillness"],
     "The spinning is motion; the music is sound (Reader p. 3)."),
    ("How do we know energy is involved if we cannot see it?",
     ["We observe a change happen", "Energy can always be seen", "We guess without evidence", "Nothing can change"],
     "We see a change — like movement — as evidence."),
    ("A battery sits still in a flashlight. Does it have energy?",
     ["No — still objects have no energy", "Only when it is moving", "Yes — stored energy that can cause change later", "Only when the flashlight is on"],
     "A battery stores energy. It can cause a change later."),
]


def slide_text(slide) -> str:
    chunks: list[str] = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            value = shape.text.strip()
            if value:
                chunks.append(value)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    value = cell.text.strip()
                    if value:
                        chunks.append(value)
    return "\n".join(chunks)


def main() -> int:
    if len(sys.argv) != 3:
        raise SystemExit("usage: qa_ch1_visual_candidate.py DECK.pptx REPORT.md")
    deck = Path(sys.argv[1]).resolve()
    report = Path(sys.argv[2]).resolve()
    prs = Presentation(deck)
    per_slide: dict[int, list[str]] = {n: [] for n in range(1, 46)}
    global_issues: list[str] = []

    if len(prs.slides) != 45:
        global_issues.append(f"slide count is {len(prs.slides)}, expected 45")
    if (prs.slide_width, prs.slide_height) != EXPECTED_SIZE:
        global_issues.append(f"slide size is {(prs.slide_width, prs.slide_height)}, expected {EXPECTED_SIZE}")

    texts: dict[int, str] = {}
    for n, slide in enumerate(prs.slides, 1):
        text = slide_text(slide)
        texts[n] = text
        notes = slide.notes_slide.notes_text_frame.text
        for field in NOTE_FIELDS:
            if not any(line.startswith(field) for line in notes.splitlines()):
                per_slide[n].append(f"speaker notes missing `{field.strip()}` field")

    for n, expected in ANCHORS.items():
        for phrase in expected:
            if phrase not in texts[n]:
                per_slide[n].append(f"missing required anchor text: {phrase}")

    leakage = {
        12: ("energy — the ability to cause change .", "change — to become different ."),
        15: ("When you see a change, you are observing evidence that energy was involved.",),
        23: ("at a later time",),
        26: ("because of their position",),
        30: ("We know because we see a change",),
    }
    for n, forbidden in leakage.items():
        for phrase in forbidden:
            if phrase.casefold() in texts[n].casefold():
                per_slide[n].append(f"answer leakage before reveal: {phrase}")

    reveal_pairs = ((12, 13), (15, 16), (20, 21), (23, 24), (26, 27), (30, 31))
    for write_n, reveal_n in reveal_pairs:
        if write_n >= reveal_n:
            per_slide[write_n].append("reveal order is invalid")

    for idx, (question, options, rationale) in enumerate(REVIEW):
        q_slide = 34 + idx * 2
        a_slide = q_slide + 1
        for n in (q_slide, a_slide):
            if question not in texts[n]:
                per_slide[n].append("canonical review question differs")
        for option in options:
            if option not in texts[q_slide]:
                per_slide[q_slide].append(f"canonical option missing: {option}")
        answer_options = options if a_slide != 45 else (options[2],)
        for option in answer_options:
            if option not in texts[a_slide]:
                per_slide[a_slide].append(f"canonical option missing: {option}")
        if rationale in texts[q_slide]:
            per_slide[q_slide].append("review rationale leaked onto question slide")
        if rationale not in texts[a_slide]:
            per_slide[a_slide].append("canonical review rationale missing")

    banned = ("kinetic", "potential", "gravitational", "mechanical", "radiant", "inertia", "momentum", "compression wave")
    for n, text in texts.items():
        lower = text.casefold()
        for term in banned:
            if term in lower:
                per_slide[n].append(f"advanced required vocabulary detected: {term}")
        if "proves energy" in lower or "proof of energy" in lower:
            per_slide[n].append("overclaim uses proof/proves instead of evidence")

    required_picture_slides = (1, 4, 5, 6, 9, 10, 11, 14, 17, 19, 22, 25, 28, 33, 36, 37, 38, 39, 44, 45)
    for n in required_picture_slides:
        slide = prs.slides[n - 1]
        picture_count = sum(1 for shape in slide.shapes if shape.shape_type == 13)
        if picture_count < 1:
            per_slide[n].append("expected visual asset is missing")

    for n in (20, 21, 26, 27):
        if not any(getattr(shape, "has_table", False) for shape in prs.slides[n - 1].shapes):
            per_slide[n].append("required native editable table is missing")

    media_text = texts[8].casefold() + "\n" + texts[32].casefold()
    if "media slot" not in media_text or "unavailable" not in media_text:
        global_issues.append("media shells do not preserve a visible slot/fallback path")

    fail_count = sum(bool(v) for v in per_slide.values()) + len(global_issues)
    report.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# Chapter 1 Visual Candidate — Deterministic Content QA",
        "",
        f"Artifact: `{deck}`",
        "",
        f"Overall: **{'PASS' if fail_count == 0 else 'FAIL'}**",
        "",
        "Checks: 45-slide count, 16:9 size, exact BQ/V/N anchor text, write-before-reveal sequencing, answer leakage, canonical review wording, speaker-note schema, vocabulary guardrails, media fallback, visual placement, and native editable tables.",
        "",
        "| Slide | Purpose | Result | Detail |",
        "|---:|---|---|---|",
    ]
    for n in range(1, 46):
        issues = per_slide[n]
        lines.append(f"| {n} | {SLIDE_NAMES[n - 1]} | **{'FAIL' if issues else 'PASS'}** | {'; '.join(issues) if issues else 'All deterministic checks passed.'} |")
    lines.extend(["", "## Global checks", ""])
    if global_issues:
        lines.extend(f"- FAIL: {issue}" for issue in global_issues)
    else:
        lines.append("- PASS: 45 slides at 16:9; required media fallback language is present.")
    lines.extend(["", "## Summary", "", f"- PASS slides: {sum(not v for v in per_slide.values())}", f"- FAIL slides: {sum(bool(v) for v in per_slide.values())}", f"- Global issues: {len(global_issues)}", ""])
    report.write_text("\n".join(lines), encoding="utf-8")
    print(f"{'PASS' if fail_count == 0 else 'FAIL'} {deck.name}: slide_failures={sum(bool(v) for v in per_slide.values())} global_issues={len(global_issues)}")
    print(f"report: {report}")
    return 1 if fail_count else 0


if __name__ == "__main__":
    raise SystemExit(main())
