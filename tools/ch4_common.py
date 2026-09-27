"""Shared helpers for building Chapter 4 DOCX/PPTX artifacts (Phase 6)."""
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------- palette ----------
INK = RGBColor(0x1A, 0x1A, 0x1A)        # body text
HEADER = RGBColor(0x1F, 0x3A, 0x5F)     # dark blue headers
ACCENT = RGBColor(0x2E, 0x6E, 0x8E)     # teal accent
GRAY = RGBColor(0x55, 0x55, 0x55)
LINE = RGBColor(0x66, 0x66, 0x66)

# ---------- paragraph helpers ----------
def _set_run(run, size=11, bold=False, italic=False, color=INK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    return run


def add_par(doc, text="", size=11, bold=False, italic=False, color=INK,
            align=None, space_before=0, space_after=4, font="Calibri"):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if text:
        _set_run(p.add_run(text), size=size, bold=bold, italic=italic, color=color, font=font)
    return p


def add_runs(doc, runs, align=None, space_before=0, space_after=4):
    """runs = list of (text, dict) tuples."""
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    for text, kw in runs:
        _set_run(p.add_run(text), **kw)
    return p


def title(doc, text, subtitle=None):
    add_par(doc, text, size=22, bold=True, color=INK, space_after=0)
    if subtitle:
        add_par(doc, subtitle, size=11, color=GRAY, space_after=2)


def name_date(doc):
    add_par(doc, "Name: ____________________________    Date: ________________",
            size=12, space_after=8)


def section(doc, text):
    add_par(doc, text, size=13, bold=True, color=HEADER, space_before=8, space_after=2)


def prompt(doc, text, size=11):
    add_par(doc, text, size=size, space_after=2)


def answer_lines(doc, n=2):
    for _ in range(n):
        add_par(doc, " " * 60, size=12, space_after=8)


def drawing_box(doc, height_in=1.3, label=""):
    """A bordered box for student drawings."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.4)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'dashed')
        el.set(qn('w:sz'), '8')
        el.set(qn('w:color'), '999999')
        borders.append(el)
    tcPr.append(borders)
    # set row height
    tr = table.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    trHeight = OxmlElement('w:trHeight')
    trHeight.set(qn('w:val'), str(int(height_in * 1440)))
    trHeight.set(qn('w:hRule'), 'atLeast')
    trPr.append(trHeight)
    if label:
        p = cell.paragraphs[0]
        _set_run(p.add_run(label), size=10, italic=True, color=GRAY)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return table


def shade_cell(cell, hex_color="1F3A5F"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def style_table(table, header_shade="1F3A5F", header_color=RGBColor(0xFF, 0xFF, 0xFF)):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # header row
    if len(table.rows) > 0:
        for cell in table.rows[0].cells:
            shade_cell(cell, header_shade)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = header_color
                    r.font.size = Pt(11)


def set_cell_text(cell, text, size=11, bold=False, color=INK, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    _set_run(p.add_run(text), size=size, bold=bold, color=color)


def setup_portrait(doc):
    s = doc.sections[0]
    s.page_width = Inches(8.5)
    s.page_height = Inches(11)
    s.top_margin = Inches(0.5)
    s.bottom_margin = Inches(0.5)
    s.left_margin = Inches(0.6)
    s.right_margin = Inches(0.6)
    return s


def setup_landscape(doc):
    s = doc.sections[0]
    s.page_width = Inches(11)
    s.page_height = Inches(8.5)
    s.top_margin = Inches(0.5)
    s.bottom_margin = Inches(0.5)
    s.left_margin = Inches(0.6)
    s.right_margin = Inches(0.6)
    return s


def page_break(doc):
    from docx.enum.text import WD_BREAK
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
