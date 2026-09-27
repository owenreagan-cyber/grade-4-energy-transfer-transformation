"""Build Chapter 6 engineering performance task (7 pages) + rubric DOCX."""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from ch4_common import *

WS = os.path.join(os.path.dirname(__file__), '..', 'production', 'worksheets')
RB = os.path.join(os.path.dirname(__file__), '..', 'production', 'rubrics')


def _page(d, n, label):
    add_par(d, f"PAGE {n} of 7 — {label}", size=9, bold=True, color=GRAY, align=WD_ALIGN_PARAGRAPH.RIGHT,
            space_after=4)


def build_task():
    d = Document()
    setup_portrait(d)

    # ---- PAGE 1: Problem / Criteria / Constraints ----
    title(d, "Engineering Performance Task — Tabletop Alert / Signal Device",
          "Chapter 6 · Energy Transformations + Engineering Design · Week 8")
    name_date(d)
    _page(d, 1, "Problem · Criteria · Constraints")

    section(d, "THE PROBLEM")
    prompt(d, "Design a non-electrical tabletop device that produces a clear signal when a model "
              "door/gate/trigger moves. Your device must demonstrate at least one energy transformation.")

    section(d, "CRITERIA (what a good device must do)")
    for i, c in enumerate([
        "responds to the trigger (the model door/gate moves → the device reacts)",
        "produces an observable signal (sound, motion, or a visible signal)",
        "demonstrates at least one energy transformation",
        "works repeatedly (at least 2 of 3 trials)",
        "you can explain the energy pathway",
    ], 1):
        add_par(d, f"□  {i}. {c}", size=11, space_after=2)

    section(d, "CONSTRAINTS (limits)")
    for i, c in enumerate([
        "only the provided materials",
        "no wall electricity · no open flame · no hot plate · no projectile launching",
        "limited build time (one period per build/test)",
        "limited material quantity",
        "tabletop footprint (fits on the tray / sheet of paper)",
    ], 1):
        add_par(d, f"□  {i}. {c}", size=11, space_after=2)

    section(d, "MY TEAM'S PROBLEM (in our own words)")
    answer_lines(d, 2)

    page_break(d)

    # ---- PAGE 2: Blueprint + Materials ----
    _page(d, 2, "Blueprint + Materials")

    section(d, "DEVICE NAME")
    answer_lines(d, 1)

    section(d, "MATERIALS I WILL USE")
    add_par(d, "Circle or list: cardboard tubes · paper cups · craft sticks · rubber bands · string · "
              "paper clips · washers · bells / metal cup · index cards / paper · tape · glue · blocks · foil",
            size=11, space_after=2)
    answer_lines(d, 2)

    section(d, "BLUEPRINT (draw + label)")
    prompt(d, "Label: the TRIGGER, the OUTPUT, and where the ENERGY CHANGE happens.")
    drawing_box(d, height_in=2.4, label="[ draw + label your blueprint here ]")

    section(d, "HOW I WILL KNOW IT WORKS")
    answer_lines(d, 2)

    page_break(d)

    # ---- PAGE 3: Energy Path ----
    _page(d, 3, "Energy Path")

    section(d, "MY ENERGY PATHWAY")
    add_par(d, "INPUT  →  CHANGE IN FORM (TRANSFORMATION)  →  OUTPUT / OBSERVABLE CHANGE",
            size=12, bold=True, space_after=6)
    add_par(d, "INPUT (energy at the start):  ____________________", size=12, space_after=4)
    add_par(d, "TRANSFORMATION (what changes form):  ____________________", size=12, space_after=4)
    add_par(d, "OUTPUT (what I can observe):  ____________________", size=12, space_after=6)

    section(d, "TRANSFER / TRANSFORMATION / BOTH?")
    prompt(d, "Does my device transfer energy, transform energy, or BOTH? Circle one and explain.")
    add_par(d, "TRANSFER      TRANSFORMATION      BOTH", size=12, bold=True, space_after=4)
    answer_lines(d, 2)

    page_break(d)

    # ---- PAGE 4: Test 1 ----
    _page(d, 4, "Test 1")

    section(d, "TEST 1 — record your evidence (3 trials)")
    t = d.add_table(rows=4, cols=4)
    t.style = 'Table Grid'
    headers = ["Trial", "Did it trigger? (yes/no)", "What output occurred?", "Evidence (what I observed)"]
    for j, h in enumerate(headers):
        set_cell_text(t.cell(0, j), h, size=11, bold=True, color=RGBColor(255, 255, 255))
    for r in range(1, 4):
        set_cell_text(t.cell(r, 0), str(r), size=11)
        for c in (1, 2, 3):
            set_cell_text(t.cell(r, c), "", size=11)
    style_table(t)

    section(d, "WHAT WORKED / WHAT FAILED")
    add_par(d, "What worked: ____________________", size=11, space_after=2)
    add_par(d, "What failed (this is DATA, not a penalty): ____________________", size=11, space_after=4)
    add_par(d, "Evidence I collected: ____________________", size=11, space_after=4)

    page_break(d)

    # ---- PAGE 5: Redesign ----
    _page(d, 5, "Redesign")

    section(d, "REDESIGN — change ONE thing, based on evidence")
    add_par(d, "WHAT WE CHANGED (only one thing):  ____________________", size=12, space_after=4)
    add_par(d, "WHY (what evidence from Test 1 supports this change):  ____________________", size=12, space_after=4)
    add_par(d, "PREDICTION (I think Test 2 will show):  ____________________", size=12, space_after=6)

    section(d, "NEW BLUEPRINT SKETCH")
    drawing_box(d, height_in=1.6, label="[ sketch the change ]")

    page_break(d)

    # ---- PAGE 6: Test 2 ----
    _page(d, 6, "Test 2")

    section(d, "TEST 2 — same conditions as Test 1 (3 trials)")
    t = d.add_table(rows=4, cols=4)
    t.style = 'Table Grid'
    for j, h in enumerate(headers):
        set_cell_text(t.cell(0, j), h, size=11, bold=True, color=RGBColor(255, 255, 255))
    for r in range(1, 4):
        set_cell_text(t.cell(r, 0), str(r), size=11)
        for c in (1, 2, 3):
            set_cell_text(t.cell(r, c), "", size=11)
    style_table(t)

    section(d, "TEST 1 vs TEST 2 — compare")
    add_par(d, "Test 1 result: ____________________", size=11, space_after=2)
    add_par(d, "Test 2 result: ____________________", size=11, space_after=4)
    add_par(d, "Did the change help? (yes / no / same)", size=11, bold=True, space_after=2)
    add_par(d, "Evidence that supports my answer: ____________________", size=11, space_after=4)

    page_break(d)

    # ---- PAGE 7: Final Explanation / Presentation Plan ----
    _page(d, 7, "Final Explanation · Presentation Plan")

    section(d, "FINAL EXPLANATION (claim · evidence · reasoning)")
    add_par(d, "CLAIM (what my device does):  ____________________", size=11, space_after=4)
    add_par(d, "EVIDENCE (what I observed):  ____________________", size=11, space_after=4)
    add_par(d, "REASONING (why the evidence supports my claim):  ____________________", size=11, space_after=6)

    section(d, "MY ENERGY PATHWAY (final)")
    add_par(d, "INPUT → TRANSFORMATION → OUTPUT:", size=11, bold=True, space_after=0)
    add_par(d, "____________________ → ____________________ → ____________________", size=12, space_after=6)

    section(d, "PRESENTATION PLAN")
    prompt(d, "Be ready to say: the problem · the criteria · the constraints · the blueprint · "
              "Test 1 · the redesign · Test 2 · the energy pathway · the final evaluation.")
    add_par(d, "My one-sentence summary: ____________________", size=11, space_after=4)

    section(d, "WHAT I LEARNED")
    answer_lines(d, 2)

    d.save(os.path.join(WS, "06-ch6-engineering-performance-task.docx"))
    print("built engineering task")


def build_rubric():
    d = Document()
    setup_portrait(d)
    title(d, "Engineering Rubric — Tabletop Alert / Signal Device",
          "Chapter 6 · Grade 4 · 15 points (aligned to AP 13.4)")
    name_date(d)

    def band(name, pts, rows):
        section(d, f"{name} — {pts} pts")
        t = d.add_table(rows=len(rows) + 1, cols=3)
        t.style = 'Table Grid'
        for j, h in enumerate(["What to look for", "Full credit", "Partial credit"]):
            set_cell_text(t.cell(0, j), h, size=10, bold=True, color=RGBColor(255, 255, 255))
        for r, (crit, full, partial) in enumerate(rows, 1):
            set_cell_text(t.cell(r, 0), crit, size=10)
            set_cell_text(t.cell(r, 1), full, size=10)
            set_cell_text(t.cell(r, 2), partial, size=10)
        style_table(t)

    band("DESIGN", "5", [
        ("Blueprint", "clear, labeled (trigger, output, energy change)", "some labels missing"),
        ("Criteria", "all criteria listed and addressed", "some criteria missing"),
        ("Constraints", "all constraints listed and respected", "constraints partly respected"),
        ("Materials", "materials listed and used as planned", "materials listed but not matched to design"),
        ("Problem", "problem restated clearly", "problem unclear"),
    ])

    band("ENERGY MODEL", "3", [
        ("Input/output", "correct initial AND final forms of energy", "one form incorrect"),
        ("Transformation", "correctly identifies the energy transformation", "transformation missing"),
        ("Observable change", "explains the observable output/change", "output unclear"),
    ])

    band("TESTING / REDESIGN", "3", [
        ("Test 1", "records Test 1 results (3 trials)", "Test 1 partly recorded"),
        ("Evidence-based change", "changes ONE thing with a reason", "change without a reason"),
        ("Test 2", "records Test 2 and compares to Test 1", "Test 2 incomplete"),
    ])

    band("FINAL EXPLANATION", "4", [
        ("Claim/conclusion", "clear claim about the device", "claim unclear"),
        ("Evidence", "uses observed evidence (not \u201cproves\u201d)", "little or no evidence"),
        ("Reasoning", "links evidence to the claim", "reasoning weak"),
        ("Engineering vocabulary", "uses criteria/constraints/test/refine/energy correctly", "vocabulary missing"),
    ])

    add_par(d, "Total: ____ / 15", size=12, bold=True, space_before=8, space_after=4)
    add_par(d, "Note: a failed prototype is NOT a penalty — score the reasoning and evidence, not the decoration.",
            size=10, italic=True, color=GRAY)

    d.save(os.path.join(RB, "06-ch6-engineering-rubric.docx"))
    print("built rubric")


if __name__ == "__main__":
    os.makedirs(RB, exist_ok=True)
    build_task()
    build_rubric()
    print("done")
