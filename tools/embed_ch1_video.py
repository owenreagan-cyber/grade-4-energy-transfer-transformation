#!/usr/bin/env python3
"""Phase 5H — genuinely embed the approved Stored_Energy.mp4 into the premium hybrid deck.

Loads the protected baseline v4 PPTX, adds the movie to slide 25 (STORED→RELEASE→CHANGE) as a
small click-to-play player with a poster frame, appends a clearly-delineated supplemental media
note (canonical note text preserved as-is), and saves a NEW candidate v6. Does NOT modify v4.
"""
import os, shutil, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
BASELINE = os.path.join(ROOT, "production", "candidates", "chapter-01",
                        "ch1-premium-hybrid-v4", "01-ch1-energy-causes-change-premium-hybrid-v4.pptx")
OUTDIR = os.path.join(ROOT, "production", "candidates", "chapter-01", "ch1-premium-hybrid-v6-embedded-media")
MOVIE = os.path.join(ROOT, "source-intake", "2026-10-09-chapter-1", "media", "Stored_Energy.mp4")
POSTER = os.path.join(ROOT, "qa", "reports", "ch1-phase5g", "_scripts", "poster-stored-energy.png")

TARGET_SLIDE = 25   # "Stored energy can cause a change later" (STORED→RELEASE→CHANGE)
LEFT, TOP, W, H = 0.75, 5.6, 3.2, 1.8   # 16:9, bottom-left open region under the rubber-band photo

MEDIA_NOTE = (
    "\n----\n"
    "MEDIA (Phase 5H supplemental — NOT canonical text):\n"
    "OPTIONAL VIDEO embedded on this slide: Stored_Energy.mp4 (4:37, click-to-play, no autoplay).\n"
    "Play after the STORED→RELEASE→CHANGE trace. Pre-view once before class.\n"
    "FALLBACK: if the video does not play, skip it — the STORED→RELEASE→CHANGE model plus the live "
    "battery / rubber-band / raised-apple examples fully cover N3. No required learning depends on the video."
)


def make_poster(src, dst, w=960):
    import subprocess
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-ss", "3", "-i", src,
                    "-frames:v", "1", "-vf", "scale=%d:-1" % w, dst], check=True)
    return dst


def main():
    os.makedirs(OUTDIR, exist_ok=True)
    if not os.path.exists(POSTER):
        make_poster(MOVIE, POSTER)

    prs = Presentation(BASELINE)
    assert len(prs.slides) == 45, "baseline must be 45 slides"

    slide = prs.slides[TARGET_SLIDE - 1]
    slide.shapes.add_movie(MOVIE, Inches(LEFT), Inches(TOP), Inches(W), Inches(H),
                           poster_frame_image=POSTER, mime_type="video/mp4")

    # preserve canonical note text; append delineated supplemental media block
    notes_frame = slide.notes_slide.notes_text_frame
    existing = notes_frame.text
    if "MEDIA (Phase 5H supplemental" not in existing:
        notes_frame.text = existing.rstrip("\n") + MEDIA_NOTE

    out = os.path.join(OUTDIR, "01-ch1-energy-causes-change-premium-hybrid-v6-embedded-media.pptx")
    prs.save(out)
    print("embedded Stored_Energy.mp4 on slide", TARGET_SLIDE)
    print("saved:", out)


if __name__ == "__main__":
    main()
