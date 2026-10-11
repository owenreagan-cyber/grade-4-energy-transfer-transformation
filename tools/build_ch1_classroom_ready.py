#!/usr/bin/env python3
"""Build the Chapter 1 classroom-ready wrapper without touching frozen files.

This is intentionally separate from the gold-standard release builder. It copies
accepted frozen artifacts, creates classroom convenience documents, and records
their source hashes. It never writes into production/releases/chapter-01.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


REPO = Path(__file__).resolve().parents[1]
OUT = REPO / "production/classroom-ready/week-2026-10-12/chapter-01"
RELEASE = REPO / "production/releases/chapter-01"

NAVY = "1F2A5A"
GOLD = "A67C16"
CREAM = "F5EEDC"
PALE = "F6F7F9"
RULE = "D9D9D9"
INK = "1A1D26"
SOFT = "4A5268"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()


def pdf_pages(path: Path) -> int:
    return len(PdfReader(str(path)).pages)


def ensure_dirs() -> None:
    for rel in (
        "01_TEACHER",
        "02_STUDENT_PRINT",
        "03_SLIDES",
        "04_EXTENSIONS",
        "05_OPTIONAL_MEDIA",
        "06_KEYS",
    ):
        (OUT / rel).mkdir(parents=True, exist_ok=True)


def copy_frozen() -> None:
    mapping = {
        RELEASE / "01-ch1-energy-causes-change-guided-notes.pdf": OUT / "02_STUDENT_PRINT/Chapter_1_Guided_Notes.pdf",
        RELEASE / "01-ch1-energy-causes-change-practice.pdf": OUT / "02_STUDENT_PRINT/Chapter_1_Practice.pdf",
        RELEASE / "01-ch1-energy-causes-change-quiz.pdf": OUT / "02_STUDENT_PRINT/Chapter_1_Check.pdf",
        RELEASE / "01-ch1-energy-causes-change-slides.pptx": OUT / "03_SLIDES/Chapter_1_Slides.pptx",
        RELEASE / "01-ch1-energy-causes-change-slides.pdf": OUT / "03_SLIDES/Chapter_1_Slides.pdf",
        RELEASE / "01-ch1-energy-detectives-student.pdf": OUT / "04_EXTENSIONS/Energy_Detectives_Optional.pdf",
        RELEASE / "01-ch1-energy-causes-change-guided-notes-key.pdf": OUT / "06_KEYS/Guided_Notes_Key.pdf",
    }
    for src, dst in mapping.items():
        shutil.copy2(src, dst)


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_cell_fill(cell, hex_value: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_value)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, left=110, bottom=90, right=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for edge, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def set_table_borders(table, color=RULE, size=5) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)
        borders.append(node)
    tbl_pr.append(borders)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_font(run, size=10.5, bold=False, color=INK, italic=False) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = rgb(color)


def add_text(doc, text="", *, size=10.5, bold=False, color=INK, italic=False,
             before=0, after=5, align=None, keep=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.keep_with_next = keep
    if text:
        r = p.add_run(text)
        set_font(r, size=size, bold=bold, color=color, italic=italic)
    return p


def add_bullets(doc, items, *, size=10.5, after=3) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.14)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(item)
        set_font(r, size=size)


def add_heading(doc, text, *, level=1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(4 if level == 1 else 2)
    p.paragraph_format.space_after = Pt(7 if level == 1 else 4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, size=17 if level == 1 else 12, bold=True, color="000000")


def add_table(doc, headers, rows, widths, *, font_size=9.1) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = 1
    set_table_borders(table)
    set_repeat_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        set_cell_fill(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_font(r, size=font_size, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = Inches(widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2:
                set_cell_fill(cell, PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_font(r, size=font_size, bold=(index == 0), color=INK if index == 0 else SOFT)
    add_text(doc, after=3)


def add_page_break(doc) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


def configure_doc(doc: Document, footer_label: str) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)

    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.color.rgb = rgb("000000")

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{footer_label}  |  Week of October 12-16, 2026  |  ")
    set_font(r, size=8.2, color=SOFT)
    field_run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run._r.extend([begin, instr, end])
    set_font(field_run, size=8.2, color=SOFT)


def cover(doc: Document) -> None:
    add_text(doc, "GRADE 4 SCIENCE", size=11, bold=True, color=GOLD, after=20,
             align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "Chapter 1", size=20, bold=True, color="000000", after=5,
             align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "Energy Causes Change", size=28, bold=True, color="000000", after=10,
             align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "Teacher Packet Classroom Ready", size=16, color=SOFT, after=28,
             align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "Week of October 12-16, 2026", size=13, bold=True, color=NAVY, after=26,
             align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "Big Question", size=11, bold=True, color=GOLD, after=4,
             align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "Where can we observe evidence of energy causing change?", size=17,
             bold=True, color="000000", after=28, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_table(doc, ["Teach", "Project", "Collect"], [[
        "Mon 10/12 and Tue 10/13; close Fri 10/16",
        "Chapter 1 Slides (45 slides)",
        "Guided Notes; Friday Check (16 points + percentage)",
    ]], [2.35, 2.15, 2.75], font_size=9.5)
    add_text(doc, "Chapter 2 begins Wednesday 10/14. Friday returns to Chapter 1 for recognition, reflection, and the scored check. Friday contains no first exposure to essential Chapter 1 science.",
             size=10.5, italic=True, color=SOFT, before=8, after=18,
             align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "Classroom-ready wrapper of frozen Chapter 1 gold-standard version 1.1",
             size=9.2, color=SOFT, align=WD_ALIGN_PARAGRAPH.CENTER)


def build_teacher_packet() -> Path:
    doc = Document()
    configure_doc(doc, "Chapter 1 Teacher Packet")
    cover(doc)
    add_page_break(doc)

    add_heading(doc, "Chapter Overview")
    add_text(doc, "Students learn that energy is the ability to cause change. They use observable before-and-after changes as evidence that energy was involved, identify five forms of energy, and explain that objects can have stored energy that can cause change later.", size=11, after=10)
    add_table(doc, ["Field", "Classroom Decision"], [
        ["Authoritative sources", "CKSci Student Reader, Chapter 1, printed pages 1-6; Teacher Guide, Lessons 1-2"],
        ["Core Reader pages", "Pages 1-5; page 6 is a forward preview and is not required in Guided Notes"],
        ["Big Question", "Where can we observe evidence of energy causing change?"],
        ["Core ideas", "Energy can cause change; change is evidence; energy occurs in different forms; stored energy can cause change later; position can relate to stored energy"],
        ["Vocabulary", "energy; change; stored energy"],
        ["Language to model", "evidence; motion; sound; light; heat; electrical energy; position"],
        ["Do not require", "kinetic energy; potential energy; mechanical energy; radiant energy; joules; force; work equations"],
    ], [1.55, 5.55], font_size=9.6)
    add_heading(doc, "Learning Objectives", level=2)
    add_bullets(doc, [
        "Give an example of energy causing an observable change.",
        "Identify motion, sound, light, heat, and electrical energy in familiar examples.",
        "Explain that an object can have stored energy even when it is not moving.",
        "Use a before-and-after change as evidence that energy was involved.",
    ])
    add_page_break(doc)

    add_heading(doc, "Week Placement and Pacing")
    add_table(doc, ["Day", "Chapter", "Required Sequence", "Student Work", "Time"], [
        ["Mon 10/12", "Ch. 1 Session 1", "Observe and Name: slides 1-3, 6-16", "Notes p. 1: BQ, V1, N1", "24.5 min"],
        ["Tue 10/13", "Ch. 1 Session 2", "Forms and Stored Energy: slide 2, 17, 19-27", "Notes p. 2: N2, V2, N3", "24.5 min"],
        ["Wed 10/14", "Ch. 2 begins", "Chapter 2 Session 1", "Chapter 2 notes", "24 min"],
        ["Thu 10/15", "Ch. 2", "Chapter 2 Session 2 and quiz", "Chapter 2 notes and quiz", "24 min"],
        ["Fri 10/16", "Ch. 1 Session 3", "Recognize, Reflect, Check: slide 2, 28-31, 33", "Notes p. 3 and 16-point check", "25 min"],
    ], [0.9, 1.25, 2.8, 1.75, 0.7], font_size=8.7)
    add_heading(doc, "Required and Optional", level=2)
    add_table(doc, ["Required", "Recommended", "Optional"], [[
        "Guided Notes; slides; electric pencil-sharpener demo; Friday check",
        "Practice after Tuesday, before Friday",
        "Energy Detectives after the check or during a later FLEX window; review slides 34-45",
    ]], [2.45, 2.0, 2.95], font_size=9.2)
    add_text(doc, "Friday rule: all vocabulary and core models are taught by Tuesday. Friday asks students to recognize and synthesize content already taught, then complete the scored check. Hand out the check no later than minute 12.", size=10.5, bold=True, before=8, after=8)
    add_page_break(doc)

    add_heading(doc, "Session 1 Observe and Name")
    add_text(doc, "Monday October 12  |  Goal: Students observe change, name it as evidence, and complete Guided Notes page 1.", size=10.6, bold=True, color=NAVY, after=8)
    add_table(doc, ["Min", "Slides", "Teacher Move", "Student Action"], [
        ["0-3", "1-3", "Introduce the Big Question. Ask for first thinking without correcting yet.", "Turn and talk; write BQ first thinking."],
        ["3-8", "6-9", "Run the electric sharpener sequence: BEFORE, PREDICT, OBSERVE, AFTER, EVIDENCE.", "Predict; watch; name visible and audible changes."],
        ["8-13", "10-12", "Teach V1 exactly: energy and change. Pause on the WRITE slide before reveal.", "Fill change and different in V1."],
        ["13-17", "13-14", "Read and discuss Reader pages 1-2. Ask how a change becomes evidence.", "Follow along; name one change."],
        ["17-21", "15-16", "Use N1 WRITE, then REVEAL. Complete ENERGY -> CHANGE -> EVIDENCE.", "Fill evidence and complete the model."],
        ["21-24.5", "-", "Quick oral check and collect a few responses. Preserve the buffer.", "Explain one observed change."],
    ], [0.55, 0.7, 3.85, 2.2], font_size=8.7)
    add_heading(doc, "Ask and Listen For", level=2)
    add_table(doc, ["Ask", "Listen For"], [
        ["What changed?", "The pencil became shorter and sharper; the sharpener changed from quiet to noisy."],
        ["What did you see or hear?", "Specific observations, not a claim that the student saw energy."],
        ["What evidence tells us energy was involved?", "The observable change is the evidence."],
    ], [3.25, 3.85], font_size=9.2)
    add_text(doc, "Optional slides 4-5 may be skipped freely. If slide 15 has not started by minute 24, move N1 to the first three minutes of Tuesday.", size=9.6, italic=True, color=SOFT, before=4)
    add_page_break(doc)

    add_heading(doc, "Session 2 Forms and Stored Energy")
    add_text(doc, "Tuesday October 13  |  Goal: Students identify five forms of energy, explain stored energy, and complete Guided Notes page 2.", size=10.6, bold=True, color=NAVY, after=8)
    add_table(doc, ["Min", "Slides", "Teacher Move", "Student Action"], [
        ["0-1", "2", "Re-read the Big Question and ask what evidence the class has collected.", "Recall Monday's evidence."],
        ["1-9", "17, 19", "Name the five forms exactly. Model one example, then release students to N2.", "Write one change example for each form."],
        ["9-13.5", "20-22", "Teach V2 exactly. Pause on WRITE before reveal.", "Fill later in the stored energy definition."],
        ["13.5-19.5", "23-25", "Show the still battery and a raised object. Use STORED ENERGY -> RELEASE -> CHANGE.", "Predict a later change; discuss position."],
        ["19.5-24.5", "26-27", "Guide N3, then reveal samples. Correct only science, not personal example choice.", "Fill position and record later changes."],
    ], [0.65, 0.75, 3.8, 2.15], font_size=8.7)
    add_heading(doc, "Ask and Listen For", level=2)
    add_table(doc, ["Ask", "Listen For"], [
        ["Can an object have energy while it is still?", "Yes. A battery, stretched rubber band, or raised object can have stored energy."],
        ["What could change later?", "A flashlight can light; a rubber band can snap back; a raised object can fall."],
        ["Why does position matter?", "A raised object can cause a change later when released."],
    ], [3.25, 3.85], font_size=9.2)
    add_text(doc, "Slide 18 is optional. If needed, complete only three N2 rows in class and finish the other two before Friday. All five rows are required by Friday.", size=9.6, italic=True, color=SOFT, before=4)
    add_page_break(doc)

    add_heading(doc, "Session 3 Recognize Reflect and Check")
    add_text(doc, "Friday October 16  |  Goal: Students synthesize evidence already taught, answer the Big Question, and complete the 16-point Chapter 1 check.", size=10.6, bold=True, color=NAVY, after=8)
    add_table(doc, ["Clock", "Slides", "Priority", "Action"], [
        ["0:00-0:30", "2", "Required", "Re-enter through the Big Question."],
        ["0:30-6:00", "28-29", "Priority 1", "Complete N4 BEFORE/AFTER drawing and answer What changed?"],
        ["6:00-9:30", "30-31", "Priority 3", "Complete CHECK YOUR THINKING if time remains."],
        ["9:30-12:00", "33", "Priority 2", "Write BQ NOW I THINK if it can finish before minute 12."],
        ["12:00-25:00", "Check", "Required", "Administer all 9 items. Do not shorten the check."],
        ["25:00-30:00", "-", "Buffer", "Collect, finish BQ/CHECK, or begin quiet optional work."],
    ], [1.05, 0.75, 1.15, 4.0], font_size=8.9)
    add_heading(doc, "Friday Fallback", level=2)
    add_bullets(doc, [
        "N4 always happens in class.",
        "If N4 ends after minute 6, skip slides 30-31. CHECK YOUR THINKING moves to catch-up or home.",
        "If BQ NOW I THINK has not started by minute 9.5, complete it after the check or during catch-up.",
        "Hand out the check at minute 12 regardless of unfinished reflection work.",
    ], size=10.2)
    add_text(doc, "Optional slide 32 and review slides 34-45 are not part of the required Friday block.", size=9.7, italic=True, color=SOFT, before=5)
    add_page_break(doc)

    add_heading(doc, "Electric Pencil Sharpener Demonstration")
    add_text(doc, "Science focus: ENERGY -> CHANGE -> EVIDENCE. Do not explain the internal machinery or require advanced energy terms.", size=11, bold=True, color=NAVY, after=10)
    add_table(doc, ["Stage", "Teacher Action", "Prompt", "Expected Evidence"], [
        ["BEFORE", "Show a dull pencil and a silent sharpener.", "What do you notice before we begin?", "Dull or rounded tip; pencil is longer; sharpener is quiet."],
        ["PREDICT", "Hold the pencil near the sharpener without starting it.", "What do you predict will change?", "The tip may become sharper or shorter; the sharpener may make sound."],
        ["OBSERVE", "Sharpen once while all students watch and listen.", "What did you see or hear?", "The pencil changed shape and length; the sharpener made noise."],
        ["AFTER", "Hold the sharpened pencil beside the before pencil if available.", "What is different now?", "The tip is sharp; the pencil is slightly shorter; shavings may be visible."],
        ["EVIDENCE", "Build the class model.", "What evidence tells us energy was involved?", "The observable changes are evidence that energy was involved."],
    ], [0.8, 2.25, 2.0, 2.1], font_size=8.8)
    add_heading(doc, "Materials and Safety", level=2)
    add_bullets(doc, [
        "One electric pencil sharpener on a stable, visible surface.",
        "One dull pencil; an optional second already sharpened pencil for comparison.",
        "Teacher operates or directly supervises the sharpener. Keep fingers, hair, clothing, and loose items away from the opening.",
        "Do not claim heat unless the teacher actually observes it. Heat is not required for the demonstration.",
    ], size=10.2)
    add_heading(doc, "If Students Ask", level=2)
    add_text(doc, "Where does the energy come from? Say: It comes through the cord or from batteries. We will study how energy moves and changes form in later chapters.", size=10.2)
    add_page_break(doc)

    add_heading(doc, "Guided Notes and Slide Anchor Map")
    add_text(doc, "When students write on Guided Notes, project the matching BQ, V, or N code. Slides without an anchor may support discussion but may not hide a required note answer.", size=10.6, after=8)
    add_table(doc, ["Notes Anchor", "Exact Student Task", "Slides", "Reveal"], [
        ["BQ", "My first thinking; later, BQ NOW I THINK", "2-3 and 33", "Teacher discussion; no single answer"],
        ["V1", "energy - ability to cause change; change - become different", "10-12", "Slide 12: change, different"],
        ["N1", "When you see a change, you are seeing evidence of energy; complete model", "13-16", "Slide 16: evidence; ENERGY -> CHANGE -> EVIDENCE"],
        ["N2", "One example showing a change for each of five forms", "17 and 19", "Teacher checks examples"],
        ["V2", "stored energy can cause change at a later time", "20-22", "Slide 22: later"],
        ["N3", "Stored energy examples and position", "23-27", "Slide 27: position plus samples"],
        ["N4", "Draw BEFORE and AFTER; answer What changed?", "28-31", "Slide 31: sample response"],
    ], [0.85, 3.25, 1.0, 2.05], font_size=8.7)
    add_heading(doc, "Required Blanks in Order", level=2)
    add_text(doc, "V1: change, different  |  N1: evidence  |  V2: later  |  N3: position", size=11, bold=True, color=NAVY, after=8)
    add_heading(doc, "HUD Labels", level=2)
    add_text(doc, "RECITE, WRITE, PREDICT, VOTE, TURN & TALK, CHECK, TRACE, DRAW, WATCH, OBSERVE, EVIDENCE, DATA, EXPLAIN are action labels. They remain separate from the semantic BQ/V/N anchor.", size=10.1)
    add_page_break(doc)

    add_heading(doc, "Teacher Questions Misconceptions and Responses")
    add_table(doc, ["Student Idea", "Teacher Response", "What to Reinforce"], [
        ["I saw energy.", "What did you actually see or hear change?", "We observe the change. The change is evidence of energy."],
        ["Only moving objects have energy.", "Can a battery cause a change later even while it is sitting still?", "Objects can have stored energy."],
        ["Stored energy means it is moving.", "Is the battery or raised book moving now? What could happen after release?", "Stored energy can cause change later."],
        ["The desk is brown, so that is evidence.", "Did the desk become brown during the observation?", "A description is not automatically an observed change."],
        ["Proof and evidence are the same word here.", "Accept proof once, then restate the class word evidence.", "Use evidence in the Big Question and the N1 model."],
        ["The sharpener uses mechanical or kinetic energy.", "Those terms are not needed in this chapter. What observable change did you notice?", "Use the accepted Grade 4 terms."],
    ], [1.65, 3.0, 2.45], font_size=8.7)
    add_heading(doc, "High Value Prompts", level=2)
    add_bullets(doc, [
        "What changed in this specific object or system?",
        "What did it look or sound like before and after?",
        "Which form of energy is connected to the change?",
        "Why is that change evidence that energy was involved?",
        "Could this object cause a change later? What would release the stored energy?",
    ], size=10.2)
    add_page_break(doc)

    add_heading(doc, "What to Collect and How to Grade")
    add_table(doc, ["Artifact", "When", "Status", "Teacher Action"], [
        ["Guided Notes", "Mon, Tue, Fri", "Required; not scored", "Check completion and science accuracy. Use the mirrored key."],
        ["Practice", "After Tue; before Fri preferred", "Recommended; ungraded", "Review quickly or use the key for self-correction."],
        ["Chapter 1 Check", "Fri minutes 12-25", "Required; 16 points", "Record raw score and whole-number percentage."],
        ["Energy Detectives", "After the check or later FLEX", "Optional; ungraded", "Use look-fors only. Do not assign points."],
    ], [1.45, 1.45, 1.65, 2.6], font_size=8.9)
    add_heading(doc, "Check Administration", level=2)
    add_bullets(doc, [
        "Choose OPEN NOTE or CLOSED NOTE before the check. If open note, allow only each student's own completed Guided Notes.",
        "Do not change the note policy during the assessment.",
        "Allow partial credit on items 3, 5, 6, 8, and 9 according to the key.",
        "Do not subtract for spelling or handwriting when the science idea is clear.",
    ], size=10.2)
    add_heading(doc, "Raw Score to Percentage", level=2)
    add_table(doc, ["Raw", "%", "Raw", "%", "Raw", "%"], [
        ["16", "100", "10", "63", "4", "25"],
        ["15", "94", "9", "56", "3", "19"],
        ["14", "88", "8", "50", "2", "13"],
        ["13", "81", "7", "44", "1", "6"],
        ["12", "75", "6", "38", "0", "0"],
        ["11", "69", "5", "31", "", ""],
    ], [0.75, 0.75, 0.75, 0.75, 0.75, 0.75], font_size=9.1)
    add_text(doc, "Key files: Guided_Notes_Key.pdf, Practice_Key.pdf, Check_Key.pdf. The optional Energy Detectives look-fors are in Energy_Detectives_Optional_Key.pdf.", size=9.6, italic=True, color=SOFT, before=5)
    add_page_break(doc)

    add_heading(doc, "Absent Student Procedure Extensions and Media")
    add_heading(doc, "Absent Student", level=2)
    add_bullets(doc, [
        "Have the student read Student Reader pages 1-5 and complete Guided Notes in order: page 1, page 2, page 3.",
        "Use Reader examples and the written pencil-sharpener description. Do not ask the student to claim a live observation they did not make.",
        "Require the same check policy used with the class. Open note is allowed only if the class used open note.",
        "Optional Energy Detectives does not need to be made up.",
    ], size=10.2)
    add_heading(doc, "Optional Energy Detectives", level=2)
    add_text(doc, "Use after the Friday check, on a review day, or during a later FLEX block. Students record an object or system, BEFORE, AFTER, the observable change, and the related form of energy. It is an ADAPTED extension from AP 1.1 and remains ungraded.", size=10.3)
    add_heading(doc, "Optional Media Decision", level=2)
    add_text(doc, "The specifically named uploaded video was not present in the supplied attachment set. The closest repository micro-video was reviewed and is not recommended because it uses unsupported wording including zero state, null result, verified proof, and all physical actions. The lesson does not depend on video. Use the live pencil-sharpener demonstration.", size=10.3)
    add_heading(doc, "Transition", level=2)
    add_text(doc, "Chapter 2 begins Wednesday October 14. Friday October 16 closes Chapter 1 through recognition, Big Question synthesis, and the scored check; it does not add new Chapter 1 science.", size=10.3)
    add_page_break(doc)

    add_heading(doc, "Classroom Ready Checklist and Provenance")
    add_table(doc, ["Before Monday", "Ready When"], [
        ["Print student packets", "20 copies of the 7-page Student Print Packet"],
        ["Print teacher materials", "Teacher Packet, Materials and Prep, Guided Notes Key, Practice Key, Check Key"],
        ["Project slides", "Chapter_1_Slides.pptx opens; title slide is ready"],
        ["Prepare demo", "Sharpener works; dull pencil is ready; optional comparison pencil is ready"],
        ["Choose check policy", "OPEN NOTE or CLOSED NOTE is decided before Friday copies are used"],
        ["Protect Friday timing", "Check will be handed out no later than minute 12"],
    ], [2.15, 4.95], font_size=9.4)
    add_heading(doc, "Provenance", level=2)
    add_table(doc, ["Tag", "Use in this packet"], [
        ["SOURCE", "Reader definitions, core science, Big Question, Teacher Guide demonstration and prompts"],
        ["ADAPTED", "Guided Notes models, practice, Chapter 1 check, Energy Detectives student page"],
        ["CUSTOM", "Classroom pacing wrapper, print directions, Friday fallback summary, score conversion table"],
        ["EXTENSION", "Energy Detectives as optional FLEX; everyday practice examples"],
    ], [1.15, 5.95], font_size=9.3)
    add_heading(doc, "Frozen Baseline", level=2)
    add_text(doc, "Gold-standard version 1.1 remains authoritative. This classroom-ready packet and bundle are derived wrappers. No frozen Chapter 1 file is modified in place.", size=10.4, bold=True)

    path = OUT / "01_TEACHER/Chapter_1_Teacher_Packet_Classroom_Ready.docx"
    doc.save(path)
    return path


def build_materials_doc() -> Path:
    doc = Document()
    configure_doc(doc, "Chapter 1 Materials and Prep")
    add_text(doc, "Chapter 1 Materials and Prep", size=22, bold=True, color="000000", after=4)
    add_text(doc, "Energy Causes Change  |  Week of October 12-16, 2026", size=11.5, color=SOFT, after=10)
    add_table(doc, ["Day", "Have Ready", "Print or Project"], [
        ["Mon 10/12", "Electric sharpener; one dull pencil; optional sharpened comparison pencil", "Project slides 1-3 and 6-16; Guided Notes p. 1"],
        ["Tue 10/13", "Optional rubber band; a battery or flashlight may be shown without making the lesson depend on it", "Project slide 2, 17, 19-27; Guided Notes p. 2"],
        ["Fri 10/16", "Timer; decided OPEN NOTE or CLOSED NOTE policy", "Project slide 2, 28-31, 33; Guided Notes p. 3; Chapter 1 Check"],
    ], [1.1, 3.05, 3.1], font_size=8.9)
    add_heading(doc, "Electric Sharpener Sequence", level=2)
    add_table(doc, ["1", "2", "3", "4", "5"], [["BEFORE", "PREDICT", "OBSERVE", "AFTER", "EVIDENCE"]], [1.4, 1.4, 1.4, 1.4, 1.4], font_size=9.4)
    add_text(doc, "Ask: What changed? What did you see or hear? What evidence tells us energy was involved?", size=10.2, bold=True, color=NAVY, after=8)
    add_heading(doc, "Print Today", level=2)
    add_bullets(doc, [
        "20 copies of Chapter_1_Student_Print_Packet.pdf (7 pages each; 140 printed page sides).",
        "1 copy each: Teacher Packet, Materials and Prep, Guided Notes Key, Practice Key, Check Key.",
        "Optional: 20 copies of Energy Detectives (2 pages each) only if planning FLEX or early-finisher use.",
    ], size=10.2)
    add_heading(doc, "Final Checks", level=2)
    add_bullets(doc, [
        "Test the sharpener before class.",
        "Open Chapter_1_Slides.pptx and confirm presenter display works.",
        "Keep the Friday check start fixed at minute 12.",
        "Do not use the repository substitute video listed in MEDIA_GUIDE.md.",
    ], size=10.2)
    path = OUT / "01_TEACHER/Materials_and_Prep.docx"
    doc.save(path)
    return path


def pdf_styles():
    styles = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("KeyTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=20,
                                leading=24, textColor=colors.HexColor("#000000"), alignment=TA_LEFT,
                                spaceAfter=10),
        "h1": ParagraphStyle("KeyH1", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=15,
                             leading=18, textColor=colors.HexColor("#000000"), spaceBefore=7, spaceAfter=6),
        "h2": ParagraphStyle("KeyH2", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=13,
                             leading=16, textColor=colors.HexColor("#1F2A5A"), spaceBefore=6, spaceAfter=4),
        "body": ParagraphStyle("KeyBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=12.5,
                               leading=15.5, textColor=colors.HexColor("#1A1D26"), spaceAfter=5),
        "small": ParagraphStyle("KeySmall", parent=styles["BodyText"], fontName="Helvetica", fontSize=11.5,
                                leading=14, textColor=colors.HexColor("#4A5268"), spaceAfter=4),
        "center": ParagraphStyle("KeyCenter", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=12,
                                 leading=15, alignment=TA_CENTER, textColor=colors.HexColor("#1F2A5A")),
    }


def key_doc(path: Path, title: str):
    return SimpleDocTemplate(str(path), pagesize=letter, rightMargin=0.55 * inch, leftMargin=0.55 * inch,
                             topMargin=0.52 * inch, bottomMargin=0.52 * inch, title=title, author="Grade 4 Science")


def key_footer(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 12)
    canvas.setFillColor(colors.HexColor("#4A5268"))
    canvas.drawCentredString(letter[0] / 2, 0.28 * inch, f"Chapter 1 Classroom Ready | {doc.page}")
    canvas.restoreState()


def key_header(title: str, subtitle: str, styles):
    return [Paragraph(title, styles["title"]), Paragraph(subtitle, styles["small"]), Spacer(1, 5)]


def styled_table(data, widths, *, header=True, font=8.2):
    table = Table(data, colWidths=[w * inch for w in widths], repeatRows=1 if header else 0, hAlign="LEFT")
    commands = [
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("FONTNAME", (0, 1 if header else 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), font),
        ("TEXTCOLOR", (0, 1 if header else 0), (-1, -1), colors.HexColor("#1A1D26")),
    ]
    if header:
        commands += [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2A5A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ]
    for row in range(2 if header else 1, len(data), 2):
        commands.append(("BACKGROUND", (0, row), (-1, row), colors.HexColor("#F6F7F9")))
    table.setStyle(TableStyle(commands))
    return table


def build_practice_key() -> Path:
    path = OUT / "06_KEYS/Practice_Key.pdf"
    styles = pdf_styles()
    story = key_header("Chapter 1 Practice Key", "Energy Causes Change | Ungraded practice | Reader Chapter 1, pages 1-6", styles)
    story += [Paragraph("Part A What Changed", styles["h1"])]
    rows = [["Scenario", "Observable change", "Form"]]
    for row in [
        ("Light switch", "Room changed from dark to bright", "light or electrical energy"),
        ("Toy car", "Car changed from still to moving", "motion"),
        ("Bell", "Quiet changed to a ringing sound", "sound"),
        ("Ice cube on warm hand", "Ice melted or became smaller", "heat"),
    ]:
        rows.append([Paragraph(x, styles["small"]) for x in row])
    story += [styled_table(rows, [2.0, 3.0, 2.0], font=11.5), Spacer(1, 7), Paragraph("Part B Evidence or Not Enough Evidence", styles["h1"])]
    rows = [["Observation", "Answer", "Because"]]
    for row in [
        ("Fan blades started spinning when plugged in", "EVIDENCE", "still to moving is an observed change"),
        ("Battery looks shiny", "NOT ENOUGH", "looking shiny does not show a change"),
        ("Room got brighter when the light came on", "EVIDENCE", "dark to bright is a change"),
        ("I think the motor is strong", "NOT ENOUGH", "this is an opinion; no change was observed"),
    ]:
        rows.append([Paragraph(x, styles["small"]) for x in row])
    story += [styled_table(rows, [3.0, 1.25, 2.75], font=11.5), PageBreak(), Paragraph("Part C Stored Energy", styles["h1"])]
    rows = [["Example", "Answer", "Expected explanation"]]
    for row in [
        ("Battery in switched-off flashlight", "YES", "When switched on, it can make the bulb light."),
        ("Apple held above the ground", "YES", "If released, it can fall because of its position."),
        ("Stretched rubber band", "YES", "If released, it can snap back and may make sound."),
        ("Toy car rolling", "NO - happening now", "Its energy is in its motion now, not stored for later."),
    ]:
        rows.append([Paragraph(x, styles["small"]) for x in row])
    story += [styled_table(rows, [2.35, 1.4, 3.25], font=11.5), Spacer(1, 9), Paragraph("Part D Explain", styles["h1"]),
              Paragraph("Sample: I flipped a switch and the room changed from dark to bright. A change is evidence of energy, so the light shows energy was at work.", styles["body"]),
              Paragraph("Accept any response that names an observed change and connects that change to evidence of energy. Do not require stored energy. Accept proof once, then restate evidence.", styles["body"]),
              Spacer(1, 10), Paragraph("Teacher Note", styles["h2"]),
              Paragraph("The bell and ice-cube rows are everyday EXTENSION examples. Accept any reasonable observable change and one of the five taught forms.", styles["small"])]
    key_doc(path, "Chapter 1 Practice Key").build(story, onFirstPage=key_footer, onLaterPages=key_footer)
    return path


def build_check_key() -> Path:
    path = OUT / "06_KEYS/Check_Key.pdf"
    styles = pdf_styles()
    story = key_header("Chapter 1 Check Key", "9 questions | 16 points | Record raw score and whole-number percentage", styles)
    story += [Paragraph("Administration", styles["h1"]),
              Paragraph("Teacher choice: OPEN NOTE or CLOSED NOTE. Decide before the check. Open note means the student's own completed Guided Notes only. Friday slot: minutes 12-25.", styles["body"])]
    items = [
        ("1 | 1 point", "c) cause change", "Energy is the ability to cause change."),
        ("2 | 1 point", "b) evidence", "A change is evidence that energy is present."),
        ("3 | 2 points", "1 point: car changed from still to moving. 1 point: motion.", "Accept it moved. Do not penalize spelling."),
        ("4 | 1 point", "d) motion, sound, light, heat, electrical energy", "The five forms taught in Reader page 3 and N2."),
        ("5 | 2 points", "ENERGY -> change -> evidence", "1 point per blank in order. Swapped words earn 0 for those blanks."),
        ("6 | 2 points", "Yes. It has stored energy that can cause a later change, such as lighting the bulb.", "1 point yes/stored energy; 1 point later change."),
        ("7 | 1 point", "a) has stored energy", "A raised object has stored energy because of its position."),
        ("8 | 3 points", "Yes; names off-to-on or dark-to-bright change; connects change to evidence of energy.", "1 point each. Student must use the word change."),
        ("9 | 3 points", "a YES; b NO; c YES", "1 point each."),
    ]
    for label, answer, scoring in items:
        story.append(KeepTogether([
            Paragraph(label, styles["h2"]),
            Paragraph(f"<b>Answer:</b> {answer}", styles["body"]),
            Paragraph(f"<b>Scoring:</b> {scoring}", styles["small"]),
        ]))
    story += [Spacer(1, 8), Paragraph("Score Conversion", styles["h1"])]
    score_rows = [["Raw", "%", "Raw", "%", "Raw", "%"]]
    for row in [(16, 100, 10, 63, 4, 25), (15, 94, 9, 56, 3, 19), (14, 88, 8, 50, 2, 13),
                (13, 81, 7, 44, 1, 6), (12, 75, 6, 38, 0, 0), (11, 69, 5, 31, "", "")]:
        score_rows.append([str(x) for x in row])
    story += [styled_table(score_rows, [0.8] * 6, font=11.5), Spacer(1, 10), Paragraph("Partial Credit and Boundaries", styles["h1"])]
    for text in [
        "Partial credit is allowed on items 3, 5, 6, 8, and 9 as described above.",
        "Do not subtract for spelling or handwriting when the science idea is clear.",
        "Accept proof for evidence once, then restate the class term evidence.",
        "The check does not require kinetic, potential, mechanical, radiant, or other advanced vocabulary.",
    ]:
        story.append(Paragraph(f"- {text}", styles["body"]))
    key_doc(path, "Chapter 1 Check Key").build(story, onFirstPage=key_footer, onLaterPages=key_footer)
    return path


def build_detectives_key() -> Path:
    path = OUT / "06_KEYS/Energy_Detectives_Optional_Key.pdf"
    styles = pdf_styles()
    story = key_header("Energy Detectives Optional Look Fors", "OPTIONAL FLEX | Ungraded | Adapted from Teacher Guide AP 1.1", styles)
    story += [Paragraph("Use", styles["h1"]), Paragraph("Use after the Friday check, on a review day, or during a later FLEX block. Students may work individually or in pairs. No points are assigned.", styles["body"]),
              Paragraph("Look Fors", styles["h1"])]
    for text in [
        "Each row names a specific classroom object or system.",
        "BEFORE and AFTER describe observable states.",
        "What changed identifies a difference, such as dark to light, still to moving, or quiet to noisy.",
        "The energy form is one of motion, sound, light, heat, or electrical energy.",
        "Students say they saw the change, not that they saw energy.",
    ]:
        story.append(Paragraph(f"- {text}", styles["body"]))
    story += [Paragraph("Possible Examples", styles["h1"])]
    rows = [["Object", "Before", "After", "Change / evidence"]]
    for row in [
        ("Classroom light", "dark", "bright", "the room became brighter"),
        ("Pencil sharpener", "dull pencil; quiet", "sharp pencil; sound", "shape and sound changed"),
        ("Rubber band", "stretched", "released", "shape and motion changed"),
        ("Bell", "quiet", "ringing", "sound began"),
        ("Flashlight", "off", "on", "light appeared"),
    ]:
        rows.append([Paragraph(x, styles["small"]) for x in row])
    story += [styled_table(rows, [1.45, 1.55, 1.55, 2.45], font=11.5), Spacer(1, 8),
              Paragraph("Redirect: If a student writes verified proof or says energy itself was visible, ask: What specific change did you observe? Restate the conclusion using evidence.", styles["body"])]
    key_doc(path, "Energy Detectives Optional Look Fors").build(story, onFirstPage=key_footer, onLaterPages=key_footer)
    return path


def merge_student_packet() -> Path:
    output = OUT / "02_STUDENT_PRINT/Chapter_1_Student_Print_Packet.pdf"
    writer = PdfWriter()
    for rel in (
        "02_STUDENT_PRINT/Chapter_1_Guided_Notes.pdf",
        "02_STUDENT_PRINT/Chapter_1_Practice.pdf",
        "02_STUDENT_PRINT/Chapter_1_Check.pdf",
    ):
        reader = PdfReader(str(OUT / rel))
        for page in reader.pages:
            writer.add_page(page)
    with output.open("wb") as stream:
        writer.write(stream)
    return output


def write_readmes() -> None:
    readme = """# Chapter 1 Print First

## Print today for 20 students

Print **20 copies** of `02_STUDENT_PRINT/Chapter_1_Student_Print_Packet.pdf`.

- Guided Notes: 3 pages per student = 60 printed page sides
- Practice: 2 pages per student = 40 printed page sides
- Chapter 1 Check: 2 pages per student = 40 printed page sides
- Required total: 7 pages per student = **140 printed page sides**
- Single-sided: 140 sheets
- Duplex: 4 sheets per student = 80 sheets total, with the final side blank

The packet order is Guided Notes, Practice, then Check. There is no cover sheet and no duplicate page.

## Print one teacher copy

- `01_TEACHER/Chapter_1_Teacher_Packet_Classroom_Ready.pdf`
- `01_TEACHER/Materials_and_Prep.pdf`
- `06_KEYS/Guided_Notes_Key.pdf`
- `06_KEYS/Practice_Key.pdf`
- `06_KEYS/Check_Key.pdf`

Optional: print `06_KEYS/Energy_Detectives_Optional_Key.pdf` if using the extension.

## Project

- Open `03_SLIDES/Chapter_1_Slides.pptx` for teaching.
- Use `03_SLIDES/Chapter_1_Slides.pdf` only as a backup or preview.

## Optional printing

- `04_EXTENSIONS/Energy_Detectives_Optional.pdf`: 2 pages per participating student.
- For all 20 students: 40 additional printed page sides, or 20 duplex sheets.
- Use only after the Friday check, on a review day, or during a later FLEX block.

## Before copying the check

Choose OPEN NOTE or CLOSED NOTE. If open note, students may use only their own completed Guided Notes. Record every scored result as a raw score out of 16 and a whole-number percentage.
"""
    (OUT / "README_PRINT_FIRST.md").write_text(readme, encoding="utf-8")

    media = """# Chapter 1 Optional Media Guide

## Decision

No video is required or recommended for the classroom-ready sequence. Use the live electric pencil-sharpener demonstration.

## Requested upload status

`How_Change_Proves_Energy_Exists.mp4` was not present in the supplied attachment folder or repository under that name, so its runtime and content could not be verified directly.

## Closest repository substitute reviewed

- File: `assets/supplemental/media/chapter-01/How_Scientists_Spot_Invisible_Energy.mp4`
- Runtime: 1 minute 15.95 seconds
- Classification: SUPPLEMENTAL
- Decision: **DO NOT USE**

The substitute contains unsupported or misleading phrases, including `zero state`, `null result`, `verified proof`, and `all physical actions`. Those phrases conflict with the accepted Chapter 1 evidence language and Grade 4 scope.

## If the missing video is supplied later

Review it before use for exact Reader alignment, grade-level language, evidence-not-proof wording, and whether it adds value beyond the live demo. A safe insertion point would be after slide 9 and before V1 only if it passes review. The lesson must remain independent of the video.
"""
    (OUT / "05_OPTIONAL_MEDIA/MEDIA_GUIDE.md").write_text(media, encoding="utf-8")


def finalize_manifest() -> None:
    files = []
    for path in sorted(OUT.rglob("*")):
        if not path.is_file() or path.name == "CLASSROOM_READY_MANIFEST.json":
            continue
        row = {
            "path": str(path.relative_to(OUT)),
            "bytes": path.stat().st_size,
            "sha256": sha256(path),
        }
        if path.suffix.lower() == ".pdf":
            row["pages"] = pdf_pages(path)
        files.append(row)
    manifest = {
        "release": "Chapter 1 Classroom Ready",
        "week": "2026-10-12 through 2026-10-16",
        "goldStandardVersion": "1.1",
        "goldStandardStatus": "UNCHANGED",
        "frozenSourceDirectory": "production/releases/chapter-01",
        "files": files,
    }
    (OUT / "CLASSROOM_READY_MANIFEST.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--finalize", action="store_true", help="Refresh the output manifest only")
    args = parser.parse_args()
    if args.finalize:
        finalize_manifest()
        return
    ensure_dirs()
    copy_frozen()
    build_teacher_packet()
    build_materials_doc()
    build_practice_key()
    build_check_key()
    build_detectives_key()
    merge_student_packet()
    write_readmes()


if __name__ == "__main__":
    main()
