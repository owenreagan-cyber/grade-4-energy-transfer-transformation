"""Build printable Unit 2 activity handouts (DOCX) using the shared design system.

Outputs (production/activities/unit-02/):
  1. energy-resource-detective-chart.docx   (2-page Grade 4 reference)
  2. transfer-transformation-both-sort.docx (mat + 15 cards + teacher key)
  3. energy-resource-scenario-challenge.docx(cards + answer frame + key)
  4. fast-and-curious.docx                  (worksheet + key)
  5. chocolate-chip-mining-flex-model.docx  (FLEX worksheet)
  6. solar-device-mini-investigation.docx    (Wk 8 Thu data sheet)
"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from ch1_notes_common import (
    NAVY, NAVY_HEX, CORAL, CORAL_HEX, GOLD, INK, INK_HEX, GRAY,
    LIGHT_HEX, MID_HEX, BODY_FONT, HEAD_FONT,
    set_run, shade_cell, set_cell_border, set_cell_margins, set_cell_width,
    set_row_height, no_table_borders, add_paragraph_bottom_border,
    para, cell_para, vcenter, set_page,
)
from gn_design import title_block, heading, key_idea, body, wline, two_col_table, connect

OUT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..",
    "production", "activities", "unit-02"))
os.makedirs(OUT, exist_ok=True)

FULL_W = 1440 * 7  # 7.0" usable at 0.75" margins


def new_doc():
    doc = Document()
    set_page(doc, left=0.75, right=0.75, top=0.7, bottom=0.7)
    for _p in list(doc.paragraphs):
        if not _p.text.strip():
            _p._p.getparent().remove(_p._p)
    return doc


def hbreak(doc, text, size=16, before=100):
    """Heading that starts a new page (avoids stray blank pages)."""
    h = heading(doc, text, size=size, before=before)
    h.paragraph_format.page_break_before = True
    return h


def resource_cards_row(doc, left, right):
    """Two cards side by side. Each = (name, tag, rows)."""
    from ch1_notes_common import OxmlElement, qn
    CARD_W = (FULL_W - 140) // 2
    t = doc.add_table(rows=1, cols=2)
    no_table_borders(t)
    for idx, (name, tag, rows) in enumerate([left, right]):
        cell = t.cell(0, idx)
        set_cell_width(cell, CARD_W)
        set_cell_margins(cell, top=0, left=0, bottom=0, right=0)
        # inner table for the card
        inner = cell.add_table(rows=1 + len(rows), cols=1)
        no_table_borders(inner)
        hc = inner.cell(0, 0)
        set_cell_width(hc, CARD_W)
        set_cell_margins(hc, top=30, left=80, bottom=30, right=80)
        shade_cell(hc, NAVY_HEX)
        hp = hc.paragraphs[0]
        set_run(hp.add_run(name), size=13, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        set_run(hp.add_run("  " + tag), size=10, italic=True, color=GOLD)
        set_row_height(inner.rows[0], 320, rule='exact')
        for i, (lab, val) in enumerate(rows):
            c = inner.cell(i + 1, 0)
            set_cell_width(c, CARD_W)
            set_cell_margins(c, top=20, left=80, bottom=20, right=80)
            set_cell_border(c, left=dict(sz=8, val='single', color=NAVY_HEX),
                            right=dict(sz=8, val='single', color=NAVY_HEX),
                            bottom=dict(sz=8, val='single', color=NAVY_HEX))
            p = c.paragraphs[0]
            set_run(p.add_run(lab + ": "), size=11, bold=True, color=NAVY)
            set_run(p.add_run(val), size=11, color=INK)
    # spacing after the row
    sp = doc.add_paragraph()
    spPr = sp._p.get_or_add_pPr()
    elm = OxmlElement('w:spacing'); elm.set(qn('w:after'), '80'); spPr.append(elm)
    return t


# ================================================================ 1. Detective chart
def build_detective_chart():
    doc = new_doc()
    title_block(doc, "Energy Resource Detective Chart", "Grade 4 reference · every resource has benefits AND limitations")
    heading(doc, "Page 1 — Nonrenewable Resources", 16)

    nonrenewable = [
        ("Coal", "NONRENEWABLE · solid", [
            ("Renewable / Nonrenewable", "Nonrenewable (solid)"),
            ("How it is used", "Burned to heat homes, cook, and make electricity"),
            ("Energy pathway", "burn coal → heat water → steam → turbine → generator → electricity"),
            ("Benefit", "Abundant and inexpensive"),
            ("Limitation", "Mining is dangerous; burning releases substances into the air"),
            ("Works well when…", "A lot of coal is available and inexpensive"),
        ]),
        ("Petroleum (crude oil)", "NONRENEWABLE · liquid", [
            ("Renewable / Nonrenewable", "Nonrenewable (liquid)"),
            ("How it is used", "Refined into gasoline, diesel, kerosene, plastics, and more"),
            ("Energy pathway", "burn fuel → heat → move engine parts (or make electricity)"),
            ("Benefit", "Powers vehicles; makes many products"),
            ("Limitation", "Drilling and spills can harm land, water, and wildlife"),
            ("Works well when…", "Liquid fuel is needed for engines and products"),
        ]),
        ("Natural gas", "NONRENEWABLE · gas", [
            ("Renewable / Nonrenewable", "Nonrenewable (gas)"),
            ("How it is used", "Burned for heating and cooking; piped into homes"),
            ("Energy pathway", "burn gas → heat"),
            ("Benefit", "Heats homes; cooks food; burns cleaner than coal"),
            ("Limitation", "Flammable; pipes can break and leak"),
            ("Works well when…", "Pipelines can reach homes and buildings"),
        ]),
        ("Uranium", "NONRENEWABLE · mined", [
            ("Renewable / Nonrenewable", "Nonrenewable (mined)"),
            ("How it is used", "Heats water in a nuclear power plant"),
            ("Energy pathway", "nuclear heat → steam → turbine → generator → electricity"),
            ("Benefit", "Reliable electricity with little air pollution"),
            ("Limitation", "Used fuel stays dangerous a long time and must be stored safely"),
            ("Works well when…", "Reliable electricity is needed"),
        ]),
    ]
    for i in range(0, len(nonrenewable), 2):
        resource_cards_row(doc, nonrenewable[i], nonrenewable[i + 1])

    hbreak(doc, "Page 2 — Renewable Resources", 16)
    renewable = [
        ("Wind", "RENEWABLE", [
            ("Renewable / Nonrenewable", "Renewable"),
            ("How it is used", "Turns turbine blades to make electricity"),
            ("Energy pathway", "moving air → blades → generator → electricity"),
            ("Benefit", "Renewable; no fuel burned"),
            ("Limitation", "Wind is not always blowing; needs a windy place"),
            ("Works well when…", "The wind blows steadily and strongly"),
        ]),
        ("Hydroelectric (moving water)", "RENEWABLE", [
            ("Renewable / Nonrenewable", "Renewable"),
            ("How it is used", "Water through a dam spins a turbine"),
            ("Energy pathway", "water in reservoir → dam → turbine → generator → electricity"),
            ("Benefit", "Renewable; water can be stored (predictable)"),
            ("Limitation", "Dam floods land; can block fish"),
            ("Works well when…", "There is a large, steady river"),
        ]),
        ("Solar", "RENEWABLE", [
            ("Renewable / Nonrenewable", "Renewable"),
            ("How it is used", "Sunlight makes electricity or heats water and food"),
            ("Energy pathway", "sunlight → solar cell → electricity"),
            ("Benefit", "Renewable; free sunlight; no fuel burned"),
            ("Limitation", "Depends on sunlight (location, weather, time of day)"),
            ("Works well when…", "A place gets steady, strong sunlight"),
        ]),
        ("Geothermal (Earth's heat)", "RENEWABLE", [
            ("Renewable / Nonrenewable", "Renewable"),
            ("How it is used", "Heat from underground warms buildings or makes steam for electricity"),
            ("Energy pathway", "Earth's heat → steam → turbine → generator → electricity"),
            ("Benefit", "Little pollution; heat is always underground"),
            ("Limitation", "Needs a special (volcanic) location; costly to build"),
            ("Works well when…", "There is hot rock near the surface"),
        ]),
    ]
    for i in range(0, len(renewable), 2):
        resource_cards_row(doc, renewable[i], renewable[i + 1])

    key_idea(doc, "Remember: abundant ≠ renewable · renewable ≠ available everywhere · every resource has benefits AND limitations.", size=13)
    doc.save(os.path.join(OUT, "energy-resource-detective-chart.docx"))
    print("saved energy-resource-detective-chart.docx")


# ================================================================ 2. Sort (mat + cards + key)
SORT_CARDS = [
    ("A ball rolls and hits another ball, which starts moving.", "TRANSFER"),
    ("A toaster turns electrical energy into heat and light.", "TRANSFORMATION"),
    ("Sunlight warms your skin.", "TRANSFER"),
    ("A battery powers a flashlight, and light shines on the wall.", "BOTH"),
    ("You shout, and a friend across the room hears you.", "TRANSFER"),
    ("A wind turbine spins and a generator makes electricity.", "TRANSFORMATION"),
    ("A pot on a stove gets hot, and the handle warms too.", "TRANSFER"),
    ("A car engine burns fuel to make the car move.", "TRANSFORMATION"),
    ("A solar panel turns sunlight into electricity that powers a fan.", "BOTH"),
    ("A candle burns and lights the room.", "BOTH"),
    ("A cup of warm water warms a cooler spoon placed in it.", "TRANSFER"),
    ("A hair dryer turns electric current into heat and moving air.", "TRANSFORMATION"),
    ("A stretched rubber band is let go and flies across the room.", "BOTH"),
    ("A loud speaker makes sound you can hear across the room.", "TRANSFER"),
    ("A hand pushes a swing, and the swing moves.", "TRANSFER"),
]

SORT_KEY = {
    "TRANSFER": [1, 3, 5, 7, 11, 14, 15],
    "TRANSFORMATION": [2, 6, 8, 12],
    "BOTH": [4, 9, 10, 13],
}


def build_sort():
    doc = new_doc()
    title_block(doc, "Transfer / Transformation / BOTH", "Sort the cards · Unit 1 Wk 3 · cut along the lines")
    # category mat
    t = doc.add_table(rows=2, cols=3)
    no_table_borders(t)
    cats = ["TRANSFER", "TRANSFORMATION", "BOTH"]
    hints = ["energy moves to a new place/object", "energy changes form", "energy moves AND changes form"]
    for j in range(3):
        hc = t.cell(0, j)
        set_cell_width(hc, FULL_W // 3)
        set_cell_margins(hc, top=40, left=60, bottom=40, right=60)
        shade_cell(hc, NAVY_HEX if j != 1 else CORAL_HEX)
        hp = hc.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(hp.add_run(cats[j]), size=14, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        bc = t.cell(1, j)
        set_cell_width(bc, FULL_W // 3)
        set_cell_margins(bc, top=20, left=60, bottom=20, right=60)
        set_cell_border(bc, top=dict(sz=8, val='single', color=NAVY_HEX),
                        left=dict(sz=8, val='single', color=NAVY_HEX),
                        bottom=dict(sz=8, val='single', color=NAVY_HEX),
                        right=dict(sz=8, val='single', color=NAVY_HEX))
        bp = bc.paragraphs[0]; bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(bp.add_run(hints[j]), size=11, italic=True, color=GRAY)
    set_row_height(t.rows[0], 420, rule='exact')
    set_row_height(t.rows[1], 2400, rule='exact')

    hbreak(doc, "Cards — cut out (15)", 16)
    t2 = doc.add_table(rows=15, cols=1)
    no_table_borders(t2)
    for i, (text, _k) in enumerate(SORT_CARDS):
        c = t2.cell(i, 0)
        set_cell_width(c, FULL_W)
        set_cell_margins(c, top=30, left=100, bottom=30, right=100)
        set_cell_border(c, left=dict(sz=8, val='single', color=INK_HEX),
                        right=dict(sz=8, val='single', color=INK_HEX),
                        top=dict(sz=8, val='single', color=INK_HEX),
                        bottom=dict(sz=8, val='single', color=INK_HEX))
        p = c.paragraphs[0]
        set_run(p.add_run(f"{i+1}. "), size=12, bold=True, color=NAVY)
        set_run(p.add_run(text), size=13, color=INK)
        set_row_height(t2.rows[i], 520, rule='exact')

    hbreak(doc, "Teacher Key", 16)
    for cat in ["TRANSFER", "TRANSFORMATION", "BOTH"]:
        body(doc, f"{cat}: cards {', '.join(map(str, SORT_KEY[cat]))}", size=13, bold=True, after=40)
    body(doc, "Transfer = energy moving place-to-place or object-to-object (form stays the same).", size=12, after=30)
    body(doc, "Transformation = energy changing form (motion → light, stored → motion).", size=12, after=30)
    body(doc, "Both = the system changes form AND the energy then travels.", size=12, after=30)
    body(doc, "Accept reasoned placements; discuss edges (a speaker = sound transfer after the transformation inside the speaker).", size=12, italic=True)
    doc.save(os.path.join(OUT, "transfer-transformation-both-sort.docx"))
    print("saved transfer-transformation-both-sort.docx")


# ================================================================ 3. Scenario challenge
SCENARIOS = [
    ("Sunny Valley", "A warm town that gets strong sunlight almost every day. Nights are long, and there is no large river nearby.",
     "solar", "Strong sunlight almost every day", "Long nights / clouds reduce output"),
    ("Windy Coast", "A seaside town with steady, strong ocean breezes most of the year. The water near shore is calm with no strong currents.",
     "wind", "Steady, strong breezes", "Wind can be calm some days"),
    ("River Gorge", "A community sits beside a large, fast river that flows all year. The area has steep hills above the river.",
     "hydroelectric", "Large, steady river", "Dam could flood land / block fish"),
    ("Volcano Ridge", "A town near a mountain with hot rocks and steam vents close to the surface.",
     "geothermal", "Hot rocks near surface", "Only works in that special location"),
    ("Coal Country", "A region with thick coal seams underground and a long history of mining.",
     "coal", "Thick coal seams", "Burning releases substances into air"),
    ("Oil Basin", "A flat region with large deposits of petroleum below the surface.",
     "petroleum", "Large oil deposits", "Spills can harm land/water"),
]


def build_scenario():
    doc = new_doc()
    title_block(doc, "Energy Resource Scenario Challenge", "Friday PRACTICE / EXTEND · fictional places · which resource is reasonable to investigate?")
    heading(doc, "Scenario cards — cut out (6)", 16)
    t = doc.add_table(rows=6, cols=1)
    no_table_borders(t)
    for i, (place, desc, _r, _f, _l) in enumerate(SCENARIOS):
        c = t.cell(i, 0)
        set_cell_width(c, FULL_W)
        set_cell_margins(c, top=40, left=100, bottom=40, right=100)
        set_cell_border(c, left=dict(sz=8, val='single', color=NAVY_HEX),
                        right=dict(sz=8, val='single', color=NAVY_HEX),
                        top=dict(sz=8, val='single', color=NAVY_HEX),
                        bottom=dict(sz=8, val='single', color=NAVY_HEX))
        p = c.paragraphs[0]
        set_run(p.add_run(place + "  "), size=14, bold=True, color=CORAL)
        set_run(p.add_run(desc), size=13, color=INK)
        set_row_height(t.rows[i], 700, rule='atLeast')

    hbreak(doc, "Student Answer Frame", 16)
    body(doc, "Choose ONE resource that would be reasonable to investigate for that place.", size=13, after=60)
    wline(doc, 1)
    body(doc, "Place: ______________________________________", size=13, after=80)
    body(doc, "Resource we would investigate: ________________________", size=13, after=80)
    body(doc, "One fact that supports it: ____________________________", size=13, after=80)
    body(doc, "One limitation to still consider: ______________________", size=13, after=80)

    hbreak(doc, "Teacher Key — reasonable answers (accept any defensible choice)", 16)
    for place, _desc, res, fact, lim in SCENARIOS:
        body(doc, f"{place} → {res}", size=13, bold=True, after=10)
        body(doc, f"   Supporting fact: {fact}", size=12, after=10)
        body(doc, f"   Limitation: {lim}", size=12, after=40)
    body(doc, "Students may choose a different resource as long as they give one supporting fact and one limitation. This is about reasoning, not one right answer.", size=12, italic=True)
    doc.save(os.path.join(OUT, "energy-resource-scenario-challenge.docx"))
    print("saved energy-resource-scenario-challenge.docx")


# ================================================================ 4. Fast & Curious
def build_fast_curious():
    doc = new_doc()
    title_block(doc, "Fast & Curious", "Ramp height, displacement, and inferred energy of motion · Friday PRACTICE")
    heading(doc, "The Data File", 16)
    body(doc, "A student recorded two trials for each ramp height. Complete the table.", size=13, after=60)
    t = doc.add_table(rows=5, cols=4)
    no_table_borders(t)
    headers = ["Ramp height (books)", "Trial 1 push (cm)", "Trial 2 push (cm)", "Average push (cm)"]
    for j, h in enumerate(headers):
        hc = t.cell(0, j)
        set_cell_width(hc, FULL_W // 4)
        set_cell_margins(hc, top=40, left=40, bottom=40, right=40)
        shade_cell(hc, NAVY_HEX)
        hp = hc.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(hp.add_run(h), size=11, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    set_row_height(t.rows[0], 420, rule='exact')
    data = [("1 book", "3", "5"), ("2 books", "7", "9"), ("3 books", "11", "13"), ("4 books", "15", "17")]
    for i, (a, b, c) in enumerate(data):
        for j, val in enumerate([a, b, c, ""]):
            c0 = t.cell(i + 1, j)
            set_cell_width(c0, FULL_W // 4)
            set_cell_margins(c0, top=30, left=40, bottom=30, right=40)
            p = c0.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if val:
                set_run(p.add_run(val), size=13, bold=(j == 0), color=NAVY if j == 0 else INK)
            add_paragraph_bottom_border(p, sz=8, val='single', color=INK_HEX)
        set_row_height(t.rows[i + 1], 460, rule='exact')

    heading(doc, "Measured vs. Inferred", 16)
    for line in ["Ramp height", "Cup push (displacement)", "Relative speed of the marble", "Relative energy of motion"]:
        body(doc, line + " — MEASURED / INFERRED — why? ____________________", size=13, after=40)

    heading(doc, "Make a Claim", 16)
    body(doc, "The SAME marble rolled from a HIGHER ramp made the cup move ______ (more / less).", size=13, after=30)
    body(doc, "A faster roll causes a ______ (greater / smaller) change, so it has ______ (more / less) energy of motion.", size=13, after=60)

    hbreak(doc, "Teacher Key", 16)
    body(doc, "Averages: 1 book = 4 · 2 books = 8 · 3 books = 12 · 4 books = 16", size=13, bold=True, after=40)
    body(doc, "Measured: ramp height, cup push (read with a ruler/tape).", size=13, after=30)
    body(doc, "Inferred: relative speed (from ramp height) and relative energy of motion (from the cup push — we infer energy from the change we see).", size=13, after=30)
    body(doc, "Claim: more / greater / more. (Same object, faster = more energy of motion = greater change.)", size=13, after=30)
    body(doc, "No speed is directly measured and no m/s value is invented.", size=12, italic=True)
    doc.save(os.path.join(OUT, "fast-and-curious.docx"))
    print("saved fast-and-curious.docx")


# ================================================================ 5. Cookie mining (FLEX)
def build_cookie_mining():
    doc = new_doc()
    title_block(doc, "Chocolate Chip Mining", "OPTIONAL FLEX MODEL · this is a MODEL of resource extraction")
    heading(doc, "What the Model Stands For", 16)
    two_col_table(doc, ["Model piece", "Stands for"], [
        ("Cookie", "Rock underground that holds a resource"),
        ("Chocolate chips", "The resource being mined (like coal or metal)"),
        ("Toothpick", "Mining tools"),
        ("Cookie crumbs", "Disturbance to the land"),
    ], row_h=520)

    heading(doc, "What to Do", 16)
    body(doc, "1.  Place the cookie on a napkin. DRAW it before mining (BEFORE box).", size=13, after=30)
    body(doc, "2.  Use a toothpick to dig out as many whole chocolate chips as you can.", size=13, after=30)
    body(doc, "3.  DRAW the cookie after mining (AFTER box).", size=13, after=30)
    body(doc, "4.  Count the chips you removed.", size=13, after=60)

    # BEFORE / AFTER boxes
    t = doc.add_table(rows=1, cols=2)
    no_table_borders(t)
    for idx, lab in enumerate(["BEFORE", "AFTER"]):
        bc = t.cell(0, idx)
        set_cell_width(bc, (FULL_W - 120) // 2)
        set_cell_margins(bc, top=20, left=60, bottom=20, right=60)
        set_cell_border(bc, top=dict(sz=8, val='single', color=NAVY_HEX),
                        left=dict(sz=8, val='single', color=NAVY_HEX),
                        bottom=dict(sz=8, val='single', color=NAVY_HEX),
                        right=dict(sz=8, val='single', color=NAVY_HEX))
        bp = bc.paragraphs[0]; bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(bp.add_run(lab), size=13, bold=True, color=CORAL, caps=True)
    set_row_height(t.rows[0], 1500, rule='exact')

    heading(doc, "Record", 16)
    body(doc, "Chips extracted: ______", size=13, after=60)

    heading(doc, "Think About It", 16)
    body(doc, "What does this model help us understand?  Mining takes a resource out of the ground and disturbs the land around it.", size=13, after=40)
    body(doc, "What does this model NOT show?  A real mine is much bigger, deeper, and affects more land, water, and air than a cookie can show.", size=13, after=40)
    body(doc, "DO NOT eat the materials.", size=13, bold=True, color=CORAL)
    doc.save(os.path.join(OUT, "chocolate-chip-mining-flex-model.docx"))
    print("saved chocolate-chip-mining-flex-model.docx")


# ================================================================ 6. Solar mini-investigation
def build_solar_mini():
    doc = new_doc()
    title_block(doc, "Available Light & a Solar Device", "Mini-investigation · Wk 8 Thu · how does available light affect a solar device?")
    body(doc, "Question:  How does the amount of available light change what our solar device does?", size=13, bold=True, after=40)
    heading(doc, "Prediction", 16)
    body(doc, "If we give the solar cell more light, then ____________________.", size=13, after=40)
    body(doc, "If we give the solar cell less light (shade it), then ____________________.", size=13, after=60)
    heading(doc, "Observe", 16)
    t = doc.add_table(rows=4, cols=3)
    no_table_borders(t)
    for j, h in enumerate(["Condition", "What I observe", "Change I see (evidence)"]):
        hc = t.cell(0, j)
        set_cell_width(hc, FULL_W // 3)
        set_cell_margins(hc, top=40, left=40, bottom=40, right=40)
        shade_cell(hc, NAVY_HEX)
        hp = hc.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(hp.add_run(h), size=11, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    set_row_height(t.rows[0], 420, rule='exact')
    for i, cond in enumerate(["Bright / close light", "Dim / far light", "Shaded"]):
        c0 = t.cell(i + 1, 0)
        set_cell_width(c0, FULL_W // 3); set_cell_margins(c0, top=30, left=40, bottom=30, right=40)
        p0 = c0.paragraphs[0]; set_run(p0.add_run(cond), size=12, bold=True, color=NAVY)
        for j in (1, 2):
            cj = t.cell(i + 1, j)
            set_cell_width(cj, FULL_W // 3); set_cell_margins(cj, top=30, left=40, bottom=30, right=40)
            add_paragraph_bottom_border(cj.paragraphs[0], sz=8, val='single', color=INK_HEX)
        set_row_height(t.rows[i + 1], 520, rule='exact')
    heading(doc, "Explain", 16)
    body(doc, "The change I saw is my evidence. The device worked best when ____________________.", size=13, after=40)
    body(doc, "This tells me solar devices depend on the light that is available.", size=13, italic=True)
    doc.save(os.path.join(OUT, "solar-device-mini-investigation.docx"))
    print("saved solar-device-mini-investigation.docx")


if __name__ == "__main__":
    build_detective_chart()
    build_sort()
    build_scenario()
    build_fast_curious()
    build_cookie_mining()
    build_solar_mini()
    print("done")
