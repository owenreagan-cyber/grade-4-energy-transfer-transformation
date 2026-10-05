"""Phase 3 — Chapter 1 Guided Notes automated student-artifact QA.

Headless gates against the FINAL Ch1 guided notes (student + teacher key):
  US Letter · margins · font floors · page count · blank widths · multi-word
  blank separation · handwriting vertical clearance · line spacing · visual-system
  token compliance · grayscale contrast · Reader page refs · BQ/V/N consistency ·
  no stale INB/26-item language · no unsupported advanced vocabulary ·
  key/student pagination parity · all blanks answered in key.

Run: .venv/bin/python tools/qa_ch1_guided_notes.py
"""
import os, re, subprocess, sys
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
DIR = os.path.join(ROOT, "production", "premium-notes")
STUDENT = os.path.join(DIR, "01-ch1-energy-causes-change-guided-notes")
KEY = os.path.join(DIR, "01-ch1-energy-causes-change-guided-notes-key")
REPORT_DIR = os.path.join(ROOT, "qa", "reports")

# ---- FROZEN tokens (Ch1 + unit neutrals) -----------------------------------
ALLOWED_FILLS = {
    "1A1D26", "4A5268", "FFFFFF", "F6F7F9", "D8DCE4",   # unit neutrals
    "1F2A5A", "A67C16", "7A5620", "F5EEDC",             # Ch1 Discovery/Evidence
}

# Exact authoritative wording (verified vs Student Reader pp.1-6 + TG Lessons 1-2)
BQ_EXACT = "Where can we observe evidence of energy causing change?"
V_GROUPS = ["V1", "V2"]
N_SECTIONS = ["N1", "N2", "N3", "N4"]
VOCAB_ANSWERS = ["change", "different", "later"]
CONCEPT_ANSWERS = ["evidence", "position"]   # v1.1: N1 has ONE blank (evidence); N3 position
KEY_SAMPLE_MARKERS = ["Sample:", "Accept any"]

BANNED = [
    "26 items", "26 questions", "26Q", "26 Q", "26-item",
    "Interactive Science Notebook is the primary", "INB-primary", "INB",
    "kinetic", "potential", "mechanical", "radiant", "momentum", "inertia",
    "compression wave", "gravitational potential", "proof", "proves",
]

def iter_paragraphs(container):
    if hasattr(container, "paragraphs"):
        for p in container.paragraphs:
            yield p
    if hasattr(container, "tables"):
        for t in container.tables:
            for row in t.rows:
                for c in row.cells:
                    yield from iter_paragraphs(c)

def all_runs(doc):
    for p in iter_paragraphs(doc):
        for r in p.runs:
            yield r, p

def full_text(doc):
    return "\n".join(p.text for p in iter_paragraphs(doc))

def pdf_pages(path):
    info = subprocess.run(["pdfinfo", path], capture_output=True, text=True).stdout
    m = re.search(r"Pages:\s+(\d+)", info)
    return int(m.group(1)) if m else -1

def pdf_size(path):
    info = subprocess.run(["pdfinfo", path], capture_output=True, text=True).stdout
    m = re.search(r"Page size:\s+([\d.]+) x ([\d.]+)", info)
    return (float(m.group(1)), float(m.group(2))) if m else (0, 0)

def underline_blank_runs(doc):
    """Return list of (width_in, size_pt) for underscore-only runs (content blanks)."""
    out = []
    for r, _p in all_runs(doc):
        t = r.text or ""
        if t and t.strip("_") == "" and "_" in t and r.font.underline:
            sz = r.font.size.pt if r.font.size else 0
            out.append((len(t) * sz / 144.0, sz))
    return out

# ---------------------------------------------------------------- gates
def gate_page_size(doc):
    sec = doc.sections[0]
    w, h = sec.page_width / 914400, sec.page_height / 914400
    ok = abs(w - 8.5) < 0.01 and abs(h - 11.0) < 0.01
    return ("PASS" if ok else "FAIL", f"{w:.2f} x {h:.2f} in (expect 8.5 x 11)")

def gate_margins(doc):
    sec = doc.sections[0]
    vals = {"left": sec.left_margin / 914400, "right": sec.right_margin / 914400,
            "top": sec.top_margin / 914400, "bottom": sec.bottom_margin / 914400}
    ok = all(0.50 <= vals[k] <= 0.60 for k in ("left", "right")) and \
         vals["top"] >= 0.49 and vals["bottom"] >= 0.49
    return ("PASS" if ok else "WARN", f"margins { {k: round(v, 2) for k, v in vals.items()} }")

def gate_font_floors(doc):
    sizes = [r.font.size.pt for r, _ in all_runs(doc) if r.font.size and (r.text or "").strip()]
    if not sizes:
        return ("FAIL", "no sized runs")
    lo = min(sizes)
    # body/caption/attribution floor: attribution is 9.5, footer 8.5 -> absolute floor 8
    return ("PASS" if lo >= 8.0 else "FAIL", f"min run size {lo:.1f} pt (floor 8)")

def gate_fill_line_clearance(doc):
    """Deterministically check fill-in lines reserve >= 24pt line height (handwriting clearance)."""
    low = []
    for p in iter_paragraphs(doc):
        if p.runs and any((r.text or "").strip("_") == "" and "_" in (r.text or "") for r in p.runs):
            ls = p.paragraph_format.line_spacing
            if ls is not None and ls.pt < 24:
                low.append(round(ls.pt, 1))
    return ("PASS" if not low else "WARN", f"fill lines below 24pt: {low or 'none'}")

def gate_blank_widths(blanks):
    bad = [round(w, 2) for w, s in blanks if w < 0.4 or w > 2.3]
    return ("PASS" if not bad else "WARN", f"blank widths out of 0.4-2.3in: {bad or 'none'}")

def gate_multiword_separation(doc):
    """No single blank may exceed 2.3in (multi-word answers must be split)."""
    return gate_blank_widths(underline_blank_runs(doc))

def gate_tokens(doc):
    fills = set()
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                tcPr = cell._tc.find(qn("w:tcPr"))
                if tcPr is not None:
                    for el in tcPr.findall(qn("w:shd")):
                        f = el.get(qn("w:fill"))
                        if f and f != "auto":
                            fills.add(f.lstrip("#").upper())
    unknown = fills - ALLOWED_FILLS
    return ("PASS" if not unknown else "WARN", f"fills: {sorted(fills)}; unknown: {sorted(unknown) or 'none'}")

def gate_grayscale_contrast():
    def lin(c):
        c /= 255
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    def lum(h):
        return 0.2126 * lin(int(h[0:2], 16)) + 0.7152 * lin(int(h[2:4], 16)) + 0.0722 * lin(int(h[4:6], 16))
    bad = []
    for h in ("1F2A5A", "A67C16", "7A5620"):
        ratio = 1.05 / (lum(h) + 0.05)
        if ratio < 3.0:
            bad.append(f"{h}:{ratio:.1f}")
    return ("PASS" if not bad else "WARN", f"chapter tokens vs white <3:1: {bad or 'none'}")

def gate_bq_vn(doc):
    text = full_text(doc)
    missing = []
    if BQ_EXACT not in text:
        missing.append("BQ exact wording")
    for code in V_GROUPS + N_SECTIONS:
        if code not in text:
            missing.append(code)
    # Notes-first architecture order: BQ → V1 → N1 → N2 → V2 → N3 → N4
    seq = ["BQ", "V1", "N1", "N2", "V2", "N3", "N4"]
    idx = [text.find(c) for c in seq]
    order_ok = all(a >= 0 for a in idx) and all(a < b for a, b in zip(idx, idx[1:]))
    if not order_ok:
        missing.append("BQ/V/N order")
    return ("PASS" if not missing else "FAIL", f"missing/out-of-order: {missing or 'none'}")

def gate_reader_refs(doc):
    text = full_text(doc)
    refs = ["Reader p. 1", "Reader pp. 1", "Reader p. 3", "Reader p. 4", "Reader pp. 4"]
    present = [r for r in refs if r in text]
    return ("PASS" if len(present) >= 5 else "WARN", f"Reader page refs found: {present}")

def gate_banned(doc):
    text = full_text(doc).lower()
    hits = [b for b in BANNED if b.lower() in text]
    return ("PASS" if not hits else "FAIL", f"banned/stale language: {sorted(set(hits)) or 'none'}")

def gate_answers_in_key(key_doc):
    text = full_text(key_doc)
    missing = [a for a in VOCAB_ANSWERS + CONCEPT_ANSWERS if a not in text]
    samples_ok = all(m in text for m in KEY_SAMPLE_MARKERS)
    return ("PASS" if not missing and samples_ok else "FAIL",
            f"missing answers: {missing or 'none'}; sample markers: {samples_ok}")

def gate_key_no_unanswered_blanks(key_doc):
    """Key must fill every content blank; only the Name line may remain underscored."""
    unders = [w for w, s in underline_blank_runs(key_doc)]
    # Name line = one 2.0-2.3in blank; content blanks are <=1.6in
    content_blanks = [round(w, 2) for w in unders if w < 1.9]
    return ("PASS" if not content_blanks else "FAIL", f"unfilled content blanks in key: {content_blanks or 'none'}")

EXPECTED_RULED_LINE_GROUPS = [2, 2, 3, 3]   # first thinking · What changed? · CHECK · BQ NOW I THINK
EXPECTED_CONTENT_BLANKS = 5                 # V1 x2 · N1 x1 · V2 x1 · N3 x1

def ruled_line_groups(doc):
    """Writing-line groups = single-column tables whose cells are empty and carry a bottom border."""
    groups = []
    for t in doc.tables:
        if len(t.columns) != 1:
            continue
        ok = True
        for row in t.rows:
            c = row.cells[0]
            if c.text.strip():
                ok = False; break
            b = c._tc.find(qn("w:tcPr"))
            bd = b.find(qn("w:tcBorders")) if b is not None else None
            if bd is None or bd.find(qn("w:bottom")) is None:
                ok = False; break
        if ok:
            groups.append(len(t.rows))
    return groups

def gate_ruled_lines(doc, expected=None):
    expected = expected or EXPECTED_RULED_LINE_GROUPS
    g = ruled_line_groups(doc)
    merged = [p for p in iter_paragraphs(doc)
              if p._p.pPr is not None and p._p.pPr.find(qn("w:pBdr")) is not None and not p.text.strip()]
    ok = g == expected and not merged
    return ("PASS" if ok else "FAIL", f"ruled-line row groups {g} (expect {expected}); empty bordered paragraphs (merge bug): {len(merged)}")

def gate_blank_budget(doc):
    n = len(underline_blank_runs(doc))
    # Name line is one extra underlined run (>=1.9in)
    content = [w for w, _ in underline_blank_runs(doc) if w < 1.9]
    return ("PASS" if len(content) == EXPECTED_CONTENT_BLANKS else "FAIL", f"content blanks {len(content)} (expect {EXPECTED_CONTENT_BLANKS}; v1.0 had 7)")

def ordered_text(doc):
    """Document-order text (paragraphs and tables interleaved)."""
    out = []
    for el in doc.element.body.iterchildren():
        out.append("".join(x.text or "" for x in el.iter(qn("w:t"))))
    return "\n".join(out)

def _n1_segment(doc):
    t = ordered_text(doc)
    a, b = t.find("N1   ENERGY CAUSES"), t.find("N2   ENERGY OCCURS")
    return t[a:b]

def _n1_chain_cells(doc):
    """Return the text of the 5 cells of the N1 model chain (ENERGY -> CHANGE -> last box), or None."""
    for t in doc.tables:
        if len(t.rows) == 1 and len(t.columns) == 5:
            cells = [c.text.strip() for c in t.rows[0].cells]
            if cells[0] == "ENERGY":
                return cells
    return None

def gate_n1_redundancy(doc, is_key=False):
    seg = _n1_segment(doc)
    bad = [x for x in ("A change is", "Energy has the ability to _") if x in seg]
    one = "_" in seg and len(re.findall(r"_{4,}", seg)) == 1
    cells = _n1_chain_cells(doc)
    chain_ok = cells is not None and cells[0] == "ENERGY" and cells[2] == "CHANGE"
    if is_key:
        # teacher key keeps the FULL model and the answer in the sentence
        ok = (not bad) and chain_ok and cells[4] == "EVIDENCE" and "evidence" in seg
        note = "key N1: full model %s; answer in sentence" % (cells,)
    else:
        # M1 (Phase 6.1B): the student page must NOT print the answer to its own blank -> no 'evidence' anywhere in N1,
        # and the model's last box is an EMPTY write-in box
        ok = (not bad) and one and chain_ok and cells[4] == "" and "evidence" not in seg.lower()
        note = "N1 redundant text %s; single blank: %s; model %s; 'evidence' printed on page: %s" % (
            bad or "none", one, cells, "evidence" in seg.lower())
    return ("PASS" if ok else "FAIL", note)

def gate_n4_reader_range(doc):
    """L1 (Phase 6.1B): N4 cites core pp. 1-5 (p. 6 is a teacher-facing forward preview)."""
    t = ordered_text(doc)
    seg = t[t.find("N4   RECOGNIZING"):]
    ok = "Reader pp. 1\u20135" in seg and "1\u20136" not in t and "pp. 1-6" not in t
    return ("PASS" if ok else "FAIL", "N4 cites 'Reader pp. 1\u20135'; no 'pp. 1\u20136' anywhere on the page")

def gate_first_thinking_rules(pdf_path):
    """L2 (Phase 6.1B): rendered check - TWO separated writing rules under 'My first thinking' (the 2nd must not touch V1)."""
    import tempfile, glob
    from PIL import Image
    tmp = tempfile.mkdtemp(prefix="ft_")
    subprocess.run(["pdftoppm", "-png", "-r", "100", "-f", "1", "-l", "1", pdf_path, os.path.join(tmp, "p")], check=True)
    img = Image.open(glob.glob(os.path.join(tmp, "p*.png"))[0]).convert("L"); px = img.load(); w, h = img.size
    bb = subprocess.run(["pdftotext", "-bbox", "-f", "1", "-l", "1", pdf_path, "-"], capture_output=True, text=True).stdout
    def ymin(word, nth=0):
        ys = [float(m.group(1)) for m in re.finditer(r'yMin="([\d.]+)"[^>]*>' + re.escape(word) + '<', bb)]
        return ys[nth] if len(ys) > nth else None
    y_lab = ymin("thinking"); y_v1 = ymin("V1")
    if y_lab is None or y_v1 is None:
        return ("FAIL", "could not locate 'My first thinking' / 'V1' in the rendered PDF")
    rows = [y for y in range(h) if sum(1 for x in range(w) if px[x, y] < 160) > w * 0.6]
    groups = []
    for y in rows:
        if groups and y - groups[-1][-1] <= 2: groups[-1].append(y)
        else: groups.append([y])
    rules = [g for g in groups if y_lab * 100 / 72 + 8 < g[0] < y_v1 * 100 / 72]
    gap_pt = (y_v1 - (rules[-1][-1] * 72 / 100)) if rules else -1
    spacing = [round((rules[i + 1][0] - rules[i][0]) * 72 / 100, 1) for i in range(len(rules) - 1)]
    ok = len(rules) == 2 and gap_pt >= 12 and spacing and spacing[0] >= 24
    return ("PASS" if ok else "FAIL", "rendered rules under 'My first thinking': %d (expect 2) · rule pitch %s pt (>=24) · gap rule2->V1 %.1f pt (>=12)" % (len(rules), spacing, gap_pt))

def gate_page2_row_rules(pdf_path):
    """L-A (Phase 6.1C): rendered check - on notes page 2 every N2 form row and N3 example row has its OWN writing rule
    directly under its label (<= 8 pt below the label's text bottom) and the previous rule is clearly above the label
    (>= 10 pt above its text top), so no rule can be mistaken for the next row's."""
    import tempfile, glob
    from PIL import Image
    tmp = tempfile.mkdtemp(prefix="p2r_")
    subprocess.run(["pdftoppm", "-png", "-r", "144", "-f", "2", "-l", "2", pdf_path, os.path.join(tmp, "p")], check=True)
    img = Image.open(glob.glob(os.path.join(tmp, "p*.png"))[0]).convert("L"); px = img.load(); w, h = img.size
    S = 2.0
    x0, x1 = int(330 * S), int(560 * S)
    ys = [y for y in range(h) if sum(1 for x in range(x0, x1, 2) if px[x, y] < 110) > 0.9 * ((x1 - x0) / 2)]
    groups = []
    for y in ys:
        if groups and y - groups[-1][-1] <= 2: groups[-1].append(y)
        else: groups.append([y])
    rules = [sum(g) / len(g) / S for g in groups]
    bb = subprocess.run(["pdftotext", "-bbox", "-f", "2", "-l", "2", pdf_path, "-"], capture_output=True, text=True).stdout
    bad, rows = [], 0
    for lab in ("motion", "sound", "light", "heat", "electrical", "battery", "stretched", "raised"):
        m = re.search(r'yMin="([\d.]+)" xMax="[\d.]+" yMax="([\d.]+)">' + lab + "<", bb)
        if not m:
            bad.append(lab + ": label not found"); continue
        ymin, ymax = float(m.group(1)), float(m.group(2)); rows += 1
        below = [r for r in rules if r >= ymin + 5]; above = [r for r in rules if r < ymin]
        if not below or not above:
            bad.append(lab + ": rule not found"); continue
        d_own, d_prev = below[0] - ymax, ymin - above[-1]
        if d_own > 8 or d_prev < 10:
            bad.append("%s (own rule %.1f pt below text, previous rule %.1f pt above)" % (lab, d_own, d_prev))
    return ("PASS" if not bad and rows == 8 else "FAIL", "page-2 rows checked: %d/8 · misaligned: %s" % (rows, bad or "none"))

def gate_n2_write_prompt(doc):
    t = full_text(doc)
    return ("PASS" if "Write one example that shows a change for each form." in t and "Write or draw" not in t else "FAIL",
            "N2 prompt says Write (cells are single ruled rows)")

REPORT_OVERRIDE = None
if "--report" in sys.argv:
    REPORT_OVERRIDE = os.path.join(ROOT, sys.argv[sys.argv.index("--report") + 1])

def main():
    s_doc = Document(STUDENT + ".docx")
    k_doc = Document(KEY + ".docx")
    s_pages = pdf_pages(STUDENT + ".pdf")
    k_pages = pdf_pages(KEY + ".pdf")
    s_size = pdf_size(STUDENT + ".pdf")
    k_size = pdf_size(KEY + ".pdf")

    gates = []
    def add(name, result, pdf_note=None):
        gates.append((name, result, pdf_note))

    add("US Letter (student DOCX)", gate_page_size(s_doc))
    add("US Letter (key DOCX)", gate_page_size(k_doc))
    add("Margins 0.50-0.60in", gate_margins(s_doc))
    add("Font floors >= 8pt", gate_font_floors(s_doc))
    add("Fill-line clearance >= 24pt", gate_fill_line_clearance(s_doc))
    add("Blank widths 0.4-2.3in", gate_blank_widths(underline_blank_runs(s_doc)))
    add("Multi-word blank separation", gate_multiword_separation(s_doc))
    add("Visual token compliance", gate_tokens(s_doc))
    add("Grayscale contrast", gate_grayscale_contrast())
    add("BQ/V/N consistency + order", gate_bq_vn(s_doc))
    add("Reader page refs", gate_reader_refs(s_doc))
    add("No banned/stale language", gate_banned(s_doc))
    add("Key/student pagination parity", ("PASS" if s_pages == k_pages == 3 else "FAIL",
        f"student={s_pages} key={k_pages} (expect 3/3)"))
    add("PDF page count = 3", ("PASS" if s_pages == 3 else "FAIL", f"student={s_pages}"))
    add("PDF US Letter size", ("PASS" if (abs(s_size[0] - 612) < 2 and abs(s_size[1] - 792) < 2) else "FAIL",
        f"student {s_size} (expect 612x792)"))
    add("Ruled writing lines (table rows, no merge bug)", gate_ruled_lines(s_doc))
    add("Content blank budget (<= v1.0)", gate_blank_budget(s_doc))
    add("N1 redundancy removed + model box open (M1)", gate_n1_redundancy(s_doc))
    add("N4 cites Reader pp. 1\u20135 (L1)", gate_n4_reader_range(s_doc))
    add("My first thinking: 2 usable rendered rules (L2)", gate_first_thinking_rules(STUDENT + ".pdf"))
    add("Student page 2: each writing rule belongs to its own row (L-A)", gate_page2_row_rules(STUDENT + ".pdf"))
    add("Key page 2: each writing rule belongs to its own row (L-A)", gate_page2_row_rules(KEY + ".pdf"))
    add("N2 prompt wording", gate_n2_write_prompt(s_doc))
    # key: the sample answer for "What changed?" takes the place of one ruled line (keeps the key at 3 pages)
    add("Key ruled lines parity", gate_ruled_lines(k_doc, [2, 1, 3, 3]))
    add("Key N1 keeps full model (M1)", gate_n1_redundancy(k_doc, is_key=True))
    add("Key N4 cites Reader pp. 1\u20135 (L1)", gate_n4_reader_range(k_doc))
    add("All blanks answered in key", gate_answers_in_key(k_doc))
    add("No unanswered blanks in key", gate_key_no_unanswered_blanks(k_doc))

    print("=" * 74)
    print("CH1 GUIDED NOTES — AUTOMATED STUDENT-ARTIFACT QA (Phase 3)")
    print("=" * 74)
    overall = "PASS"
    for name, (s, note), _ in gates:
        if s == "FAIL":
            overall = "FAIL"
        elif s == "WARN" and overall == "PASS":
            overall = "WARN"
        print(f"[{s:4}] {name:38} {note}")
    print("=" * 74)
    print(f"OVERALL: {overall}")

    os.makedirs(REPORT_DIR, exist_ok=True)
    out = REPORT_OVERRIDE or os.path.join(REPORT_DIR, "PHASE_3_CH1_GUIDED_NOTES_QA.md")
    with open(out, "w") as f:
        f.write("# Chapter 1 Guided Notes Automated QA\n\nGenerated by tools/qa_ch1_guided_notes.py (v1.1 gates).\n\n")
        f.write(f"Overall: **{overall}**\n\n")
        for name, (s, note), _ in gates:
            f.write(f"- [{s}] {name} — {note}\n")
    print(f"report: {out}")

if __name__ == "__main__":
    main()
