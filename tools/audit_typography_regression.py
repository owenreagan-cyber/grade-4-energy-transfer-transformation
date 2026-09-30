#!/usr/bin/env python3
"""
Phase 13S/13T — typography + science/climate regression audit.

Produces qa/temp/<phase>/audit.json with:
  - per-deck slide font-size stats (min/max, small-text count, text-heavy count)
  - per-docx font-size stats (min/max, small-text count)
  - science regression hits
  - climate regression hits
"""
import argparse
import glob
import json
import os
import re

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

SCIENCE = [
    "proof", "proves", "kinetic", "potential", "mechanical energy",
    "chemical energy", "thermal energy", "radiant", "momentum", "impulse",
    "inertia", "electron", "photon", "semiconductor", "compression wave",
    "fission", "U-235", "base-load", "baseload", "dispatch", "peak demand",
    "grid stabilization", "impact force",
]
CLIMATE = [
    "climate change", "global warming", "climate crisis", "carbon footprint",
    "zero carbon", "carbon-free", "carbon tax", "cap and trade", "activism",
    "save the planet", "energy policy",
]


def iter_decks():
    return sorted(glob.glob(os.path.join(ROOT, "production", "slides", "*.pptx")))


def iter_docx():
    dirs = [
        "production/premium-notes", "production/worksheets", "production/labs",
        "production/activities", "production/quizzes", "production/student-notebook",
        "production/review", "production/lab-setup-guides", "production/rubrics",
    ]
    out = []
    for d in dirs:
        out += glob.glob(os.path.join(ROOT, d, "**", "*.docx"), recursive=True)
    return sorted(out)


def deck_text_and_sizes(path):
    from pptx import Presentation
    prs = Presentation(path)
    slides = []
    for idx, slide in enumerate(prs.slides, 1):
        sizes, texts = [], []
        for shp in slide.shapes:
            if shp.has_text_frame:
                for para in shp.text_frame.paragraphs:
                    for run in para.runs:
                        if run.font.size is not None:
                            sizes.append(run.font.size.pt)
                        if run.text.strip():
                            texts.append(run.text.strip())
        slides.append({"slide": idx, "sizes": sizes, "text": " ".join(texts)})
    return slides


def docx_text_and_sizes(path):
    from docx import Document
    doc = Document(path)
    sizes, texts = [], []
    def walk_paras(paras):
        for p in paras:
            for r in p.runs:
                if r.font.size is not None:
                    sizes.append(r.font.size.pt)
                if r.text.strip():
                    texts.append(r.text.strip())
            if p.text.strip() and not p.runs:
                texts.append(p.text.strip())
    walk_paras(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                walk_paras(cell.paragraphs)
    return sizes, texts


def scan(text, terms):
    low = text.lower()
    hits = []
    for t in terms:
        if " " in t:
            if t in low:
                hits.append(t)
        else:
            if re.search(rf"\b{re.escape(t)}\b", low):
                hits.append(t)
    return hits


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--phase", default="phase13s")
    args = ap.parse_args()

    audit = {"decks": [], "docx": [], "science": {}, "climate": {}}

    sci_hits = {}
    cli_hits = {}

    for path in iter_decks():
        name = os.path.basename(path)
        slides = deck_text_and_sizes(path)
        all_sizes = [s for sl in slides for s in sl["sizes"]]
        small = [s for s in all_sizes if s < 20]
        text_chars = [sl["text"] for sl in slides]
        n_dense = sum(1 for t in text_chars if len(t) > 800)
        for sl in slides:
            for term in scan(sl["text"], SCIENCE):
                sci_hits.setdefault(term, []).append(f"{name} S{sl['slide']}")
            for term in scan(sl["text"], CLIMATE):
                cli_hits.setdefault(term, []).append(f"{name} S{sl['slide']}")
        audit["decks"].append({
            "file": name, "n_slides": len(slides),
            "min_pt": round(min(all_sizes), 1) if all_sizes else None,
            "max_pt": round(max(all_sizes), 1) if all_sizes else None,
            "n_text_lt20pt": len(small),
            "n_dense_slides": n_dense,
        })

    for path in iter_docx():
        name = os.path.basename(path)
        sizes, texts = docx_text_and_sizes(path)
        small = [s for s in sizes if s < 12]
        for t in texts:
            for term in scan(t, SCIENCE):
                sci_hits.setdefault(term, []).append(name)
            for term in scan(t, CLIMATE):
                cli_hits.setdefault(term, []).append(name)
        audit["docx"].append({
            "file": name,
            "min_pt": round(min(sizes), 1) if sizes else None,
            "max_pt": round(max(sizes), 1) if sizes else None,
            "n_text_lt12pt": len(small),
        })

    audit["science"] = sci_hits
    audit["climate"] = cli_hits

    out = os.path.join(ROOT, "qa", "temp", args.phase, "audit.json")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, "w") as f:
        json.dump(audit, f, indent=2)

    # Summary
    print(f"DECKS: {len(audit['decks'])}")
    for d in audit["decks"]:
        flag = " !!" if (d["n_text_lt20pt"] > 0 or d["n_dense_slides"] > 0) else ""
        print(f"  {d['file']}: slides={d['n_slides']} min={d['min_pt']}pt max={d['max_pt']}pt "
              f"small<20={d['n_text_lt20pt']} dense={d['n_dense_slides']}{flag}")
    print(f"\nDOCX: {len(audit['docx'])}")
    small_docx = [d for d in audit["docx"] if d["n_text_lt12pt"] > 0]
    for d in audit["docx"]:
        flag = " !!" if d["n_text_lt12pt"] > 0 else ""
        print(f"  {d['file']}: min={d['min_pt']}pt max={d['max_pt']}pt small<12={d['n_text_lt12pt']}{flag}")
    print(f"\nSCIENCE hits: {len(sci_hits)}")
    for k, v in sci_hits.items():
        print(f"  {k}: {v}")
    print(f"\nCLIMATE hits: {len(cli_hits)}")
    for k, v in cli_hits.items():
        print(f"  {k}: {v}")
    print(f"\nWrote {out}")


if __name__ == "__main__":
    main()
