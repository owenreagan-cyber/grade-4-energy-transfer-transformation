"""Build Chapter 5 Interactive Science Notebook (3 inserts)."""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from ch4_common import *
from ch4_common import _set_run

OUT = os.path.join(os.path.dirname(__file__), '..', 'production', 'student-notebook', 'chapter-05')


def _border(cell, color="666666", dashed=True):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'dashed' if dashed else 'single')
        el.set(qn('w:sz'), '12')
        el.set(qn('w:color'), color)
        borders.append(el)
    tcPr.append(borders)


def insert_card(d, instruction):
    add_par(d, instruction, size=9, bold=True, color=GRAY, space_after=2)
    t = d.add_table(rows=1, cols=1)
    t.alignment = 1
    cell = t.cell(0, 0)
    cell.width = Inches(7.0)
    _border(cell)
    return cell


def cell_add(cell, text="", size=11, bold=False, italic=False, color=INK, align=None,
             before=0, after=4):
    p = cell.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if text:
        _set_run(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return p


def cell_title(cell, text, sub):
    cell_add(cell, text, size=13, bold=True, color=HEADER, after=0)
    cell_add(cell, sub, size=10, italic=True, color=GRAY, after=4)


def cell_section(cell, text):
    cell_add(cell, text, size=12, bold=True, color=ACCENT, before=6, after=2)


def build():
    d = Document()
    setup_portrait(d)
    add_par(d, "Interactive Science Notebook — Chapter 5 · Collisions Transfer Energy",
            size=16, bold=True, color=HEADER, space_after=2)
    add_par(d, "Trim each insert on the dashed line and glue into your notebook. "
              "Every interactive part has a flat, low-motor alternative.",
            size=10, italic=True, color=GRAY, space_after=10)

    # ---------------- INSERT 1 ----------------
    cell = insert_card(d, "INSERT 1 — What Is a Collision?    ·    trim on the dashed line    ·    glue the BACK")
    cell_title(cell, "What Is a Collision?", "Chapter 5 · Week 6")
    cell_section(cell, "VOCABULARY")
    cell_add(cell, "collide = to come together with ________________", size=12, after=4)
    cell_add(cell, "collision = an instance of ________________", size=12, after=4)
    cell_add(cell, "sound = a form of energy that comes from a ________________ object", size=12, after=4)
    cell_section(cell, "WHAT A COLLISION CAN CHANGE")
    cell_add(cell, "1. MOTION (speed or direction)", size=11, bold=True, after=0)
    cell_add(cell, "2. SHAPE", size=11, bold=True, after=0)
    cell_add(cell, "3. SOUND (we hear it)", size=11, bold=True, after=6)
    cell_section(cell, "DRAW")
    cell_add(cell, "Draw one collision and label what changed (motion / shape / sound).", size=11, after=2)
    p = cell.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(p.add_run("[  draw your collision here  ]"), size=10, italic=True, color=GRAY)

    # ---------------- INSERT 2 ----------------
    cell = insert_card(d, "INSERT 2 — The Collision Model + Fair Test    ·    trim on the dashed line    ·    glue the BACK")
    cell_title(cell, "The Collision Model + Fair Test", "Chapter 5 · Week 6")
    cell_section(cell, "THE MODEL")
    cell_add(cell, "COLLISION", size=12, bold=True, after=0)
    cell_add(cell, "→ MOTION CAN CHANGE", size=12, bold=True, after=0)
    cell_add(cell, "→ ENERGY CAN TRANSFER", size=12, bold=True, after=0)
    cell_add(cell, "→ ENERGY DOES NOT DISAPPEAR", size=12, bold=True, color=ACCENT, after=6)
    cell_section(cell, "FAIR TEST RULE")
    cell_add(cell, "CHANGE ONE THING AT A TIME.", size=12, bold=True, after=0)
    cell_add(cell, "Part A (speed): same car, same target, same surface — change ONLY ramp height.", size=11, after=2)
    cell_add(cell, "Part B (mass): same ramp, same target, same surface — change ONLY mass/weight.", size=11, after=6)
    cell_section(cell, "MY PREDICTIONS")
    cell_add(cell, "Part A — I predict the target moves farthest at ____ books, because ____________________", size=12, after=4)
    cell_add(cell, "Part B — I predict the target moves farthest with ____ cars, because ____________________", size=12, after=4)

    # ---------------- INSERT 3 (interactive flap) ----------------
    cell = insert_card(d, "INSERT 3 — Energy Does NOT Disappear (lift-the-flap)    ·    trim on the dashed line    ·    "
                          "fold the flap DOWN and glue the BASE    ·    flat option: write the answer under the flap")
    cell_title(cell, "Energy Does NOT Disappear", "Chapter 5 · Week 6 · lift the flap to check")
    cell_section(cell, "THE QUESTION (on the flap)")
    cell_add(cell, "When a car hits a cup and the cup moves, DID THE ENERGY DISAPPEAR?", size=12, bold=True, after=6)
    cell_section(cell, "THE ANSWER (under the flap)")
    cell_add(cell, "NO — the energy did NOT disappear.", size=12, bold=True, color=ACCENT, after=2)
    cell_add(cell, "Some energy moved to the cup (it moved). Some became sound (we heard it). "
                   "A tiny bit became heat (hard to measure).", size=11, after=4)
    cell_section(cell, "MY EVIDENCE")
    cell_add(cell, "In our Crash Lab, I observed: the target moved ____________________ cm when ____________________",
             size=11, after=4)
    cell_add(cell, "The change I observed is evidence that energy ____________________.", size=12, after=4)

    d.save(os.path.join(OUT, "05-ch5-interactive-notebook.docx"))
    print("built notebook")


if __name__ == "__main__":
    build()
