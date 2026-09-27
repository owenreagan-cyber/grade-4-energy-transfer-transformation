"""Build the original Unit Final Assessment (26 items, 50 pts) as DOCX."""
import os, sys
sys.path.insert(0, os.path.dirname(__file__))
from ch4_common import *
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw, ImageFont

OUT = "production/quizzes/07-unit-final-assessment.docx"
GRAPH = "/tmp/ch7_ramp_graph.png"


def make_graph(path):
    """Simple grayscale-friendly bar graph: ramp height vs cup movement (same car)."""
    W, H = 760, 380
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    gray = (70, 70, 70)
    # axes
    d.line([(90, 40), (90, 300)], fill=gray, width=3)          # y axis
    d.line([(90, 300), (720, 300)], fill=gray, width=3)         # x axis
    # y-axis labels
    try:
        f = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 22)
        fs = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 20)
    except Exception:
        f = ImageFont.load_default(); fs = ImageFont.load_default()
    d.text((30, 120), "Cup\nmovement\n(more →)", fill=gray, font=fs)
    # bars: low/medium/high
    bars = [("Low", 40), ("Medium", 110), ("High", 210)]
    x = 160
    for label, h in bars:
        y0 = 300 - h
        d.rectangle([x, y0, x + 90, 300], fill=(120, 140, 160))
        d.text((x + 18, 310), label, fill=gray, font=f)
        x += 200
    d.text((320, 340), "Ramp height  (same toy car)", fill=gray, font=f)
    img.save(path)


make_graph(GRAPH)

doc = Document()
setup_portrait(doc)
title(doc, "Unit Final Assessment", "What Have I Learned About Energy?")
add_par(doc, "Grade 4 Energy Transfer & Transformation · 26 items · 50 points", size=11, italic=True, color=GRAY, space_after=4)
name_date(doc)
add_par(doc, "Directions: Read each item carefully. Show your evidence. You may use your Core Vocabulary cards.",
        size=11, space_after=8)

def item(n, text):
    add_par(doc, f"{n}. {text}", size=11, bold=True, space_before=6, space_after=2)

def opt(text):
    add_par(doc, text, size=11, space_after=1)

def lines(n=2):
    answer_lines(doc, n)

# ---- Section 1: Energy Causes Change (Ch.1) ----
section(doc, "Section 1 — Energy Causes Change")

item(1, "Which show evidence of energy causing a change? Circle ALL correct answers.")
opt("    a) A basketball hits the backboard.")
opt("    b) You get a small shock from a doorknob.")
opt("    c) The school bus engine makes a sound when it starts.")
opt("    d) A candle flame lights up a dark room.")

item(2, "Energy is the ability to ______.")
opt("    a) stay still")
opt("    b) cause change")
opt("    c) make heat only")
opt("    d) be invisible")

item(3, "Draw a model of energy causing a change. Show a BEFORE and an AFTER. "
        "Label the type of energy. Write one sentence describing the change.")
drawing_box(doc, 1.6, "BEFORE  →  AFTER")
prompt(doc, "Type of energy: ____________________________")
prompt(doc, "Describe the change: ____________________________")
lines(1)

item(4, "Explain how an observable change is EVIDENCE that energy is involved.")
lines(2)

# ---- Section 2: Moving Objects Have Energy (Ch.2) ----
section(doc, "Section 2 — Moving Objects Have Energy")

item(5, "A moving object has ______.")
opt("    a) stored energy only")
opt("    b) no energy")
opt("    c) energy of motion")
opt("    d) only sound energy")

item(6, "Which shows a moving object causing a change?")
opt("    a) A ball sitting on a shelf")
opt("    b) A rolling ball knocks over a block")
opt("    c) A pencil lying on a desk")
opt("    d) A book held still in your hand")

item(7, "Give ONE example of a moving object causing a change. Describe the change.")
lines(2)

# ---- Section 3: Speed & Energy (same object) (Ch.3) ----
section(doc, "Section 3 — Speed and Energy (Same Object)")

item(8, "Read the table. The SAME toy car rolled down a ramp at three heights and pushed a cup.")
t = doc.add_table(rows=4, cols=3)
style_table(t)
for j, h in enumerate(["Ramp height", "Cup movement (measured)", "Energy of motion (inferred)"]):
    set_cell_text(t.rows[0].cells[j], h, size=10, bold=True, color=RGBColor(0xFF,0xFF,0xFF))
for i, row in enumerate([["Low", "moved a little", "less"], ["Medium", "moved more", "more"], ["High", "moved the most", "most"]]):
    for j, v in enumerate(row):
        set_cell_text(t.rows[i+1].cells[j], v, size=11)
prompt(doc, "What pattern do you see?")
lines(2)

item(9, "In the ramp test, what do we MEASURE, and what do we INFER?")
opt("    a) measure cup movement, infer speed")
opt("    b) measure speed, infer cup movement")
opt("    c) measure energy, infer height")
opt("    d) measure height, infer the car's color")

item(10, "Look at the graph. It shows how far the cup moved at different ramp heights for the SAME car.")
doc.add_picture(GRAPH, width=Inches(4.6))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
prompt(doc, "As ramp height increases, cup movement ______. "
        "This is evidence that a faster car has ______ energy of motion.")
lines(2)

item(11, "For the SAME object, what is the relationship between speed and energy of motion? "
         "Use evidence.")
lines(3)

# ---- Section 4: Energy Transfer (Ch.4) ----
section(doc, "Section 4 — Energy Transfer")

item(12, "Which shows sound energy transferring from place to place?")
opt("    a) A bell rings and you hear it across the room")
opt("    b) A ball sits still on the ground")
opt("    c) A block stays on a table")
opt("    d) A ruler rests in a drawer")

item(13, "Warm water warms a cool spoon. The heat MOVES but stays heat. This is ______.")
opt("    a) TRANSFER")
opt("    b) TRANSFORMATION")
opt("    c) BOTH")
opt("    d) neither")

item(14, "Which are evidence that energy TRANSFERRED from place to place? Circle ALL correct.")
opt("    a) You hear a bell from far away")
opt("    b) A book stays closed on a desk")
opt("    c) You flip a switch and the room lights up")
opt("    d) A pencil does not move")
opt("    e) A backpack left in the sun gets warm")

item(15, "A flashlight turns on. Stored energy in the battery becomes light. This is ______.")
opt("    a) TRANSFER")
opt("    b) TRANSFORMATION")
opt("    c) neither")

item(16, "Label the energy change in a flashlight:")
prompt(doc, "    INPUT: ________________  →  OUTPUT: ________________")

# ---- Section 5: Collisions (Ch.5) ----
section(doc, "Section 5 — Collisions")

item(17, "During a collision, what happens to the energy of motion?")
opt("    a) It disappears")
opt("    b) It transfers to the target, and some becomes sound")
opt("    c) It becomes cold only")
opt("    d) Nothing happens")

item(18, "Read the table: a toy car rolls into a stack of blocks.")
t = doc.add_table(rows=3, cols=3)
style_table(t)
for j, h in enumerate(["Car", "Weight changed?", "Blocks knocked over"]):
    set_cell_text(t.rows[0].cells[j], h, size=10, bold=True, color=RGBColor(0xFF,0xFF,0xFF))
for i, row in enumerate([["Test 1", "no extra weight", "1 block"], ["Test 2", "added weight (ONLY change)", "2 blocks"]]):
    for j, v in enumerate(row):
        set_cell_text(t.rows[i+1].cells[j], v, size=11)
prompt(doc, "Explain the result using energy:")
lines(2)

item(19, "Is the test in item 18 a fair test? Why or why not? Predict what happens if you add even more weight.")
lines(3)

item(20, "After a collision, the energy ______.")
opt("    a) is gone forever")
opt("    b) turns into light only")
opt("    c) does not disappear — it transfers, and some becomes sound")
opt("    d) becomes nothing")

# ---- Section 6: Transformations + Engineering (Ch.6) ----
section(doc, "Section 6 — Transformations and Engineering")

item(21, "A domino falls and clicks. Motion passes to the next domino AND some becomes sound. This is ______.")
opt("    a) TRANSFER only")
opt("    b) TRANSFORMATION only")
opt("    c) BOTH")
opt("    d) neither")

item(22, "Draw a line to match each word to its meaning.")
prompt(doc, "    criteria        the limits (materials, time, safety)")
prompt(doc, "    constraints     what a good solution must do")
prompt(doc, "    refine          to improve using test evidence")

item(23, "Pick ONE device (flashlight, buzzer, or ramp toy). Write a claim about its energy change, "
         "give evidence, and explain your reasoning. Use \"supports,\" not \"proves.\"")
prompt(doc, "Claim: ________________________________________")
prompt(doc, "Evidence: ______________________________________")
prompt(doc, "Reasoning: ______________________________________")
lines(2)

item(24, "Label the energy chain for a fan:")
prompt(doc, "    INPUT: electrical energy  →  TRANSFORMATION: ____________  →  OUTPUT: the fan spins")

# ---- Section 7: Edison & Engineering Persistence (Ch.7) ----
section(doc, "Section 7 — Edison and Engineering Persistence")

item(25, "What is TRUE about Thomas Edison and the light bulb?")
opt("    a) He invented the first light bulb all by himself")
opt("    b) He improved earlier designs by testing and redesigning")
opt("    c) He never tested any materials")
opt("    d) He gave up after one try")

item(26, "How did Edison use TESTING and REDESIGN to solve problems? What does this show about persistence?")
lines(4)

add_par(doc, "End of assessment.", size=11, italic=True, color=GRAY, space_before=8)
add_par(doc, "Provenance: ADAPTED — original assessment aligned to the official Unit Assessment blueprint (50 pts), re-authored.",
        size=9, italic=True, color=GRAY)

doc.save(OUT)
print("wrote", OUT)
