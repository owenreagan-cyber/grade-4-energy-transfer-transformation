"""Structured Chapter 2 date/session check.

Canonical week:
  2026-10-14 Wednesday  Chapter 2 Session 1 required
  2026-10-15 Thursday   Chapter 2 Session 2 and the quiz required
  2026-10-16 Friday     Chapter 1 only

A required Chapter 2 session assigned to 2026-10-16 fails.
A sentence that says Chapter 2 never moves to Friday does not.
"""
import argparse, os, re, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

FRIDAY = "2026-10-16"
DATE_TOKEN = re.compile(
    r"2026-10-1[456]|10/1[456]|oct(?:ober)?\s+1[456]|\bfriday\b|\bfri\b",
    re.I,
)
CH2_SESSION = re.compile(
    r"chapter\s*2|chapter two|moving objects have energy|\bch\s*2\b|\bch2 session\b",
    re.I,
)
NEGATION = re.compile(
    r"\b(never moves|not friday|not week 1 friday|not chapter 2|no chapter 2|is chapter 1)\b",
    re.I,
)


def to_iso(token):
    t = token.lower().replace("october", "oct")
    if "2026-10-14" in t or "10/14" in t or t.strip() in ("oct 14",):
        return "2026-10-14"
    if "2026-10-15" in t or "10/15" in t or "oct 15" in t:
        return "2026-10-15"
    if "2026-10-16" in t or "10/16" in t or "oct 16" in t or t.strip() in ("friday", "fri"):
        return FRIDAY
    if "friday" in t or re.search(r"\bfri\b", t):
        return FRIDAY
    if "10/14" in t or "14" == t.strip():
        return "2026-10-14"
    return None


def _assigned(text):
    if NEGATION.search(text):
        return False
    if not CH2_SESSION.search(text):
        return False
    # A required instructional assignment, not a bare mention that Friday is Chapter 1.
    return bool(re.search(r"session\s*[123]|required|n[123]\b|quiz", text, re.I))


def hits(text):
    """Return lines that assign required Chapter 2 instruction to 2026-10-16."""
    bad = []
    stack = []
    for raw in text.splitlines():
        line = raw.strip()
        m = re.match(r"^(#+)\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, to_iso(m.group(2)) == FRIDAY))
        in_friday = any(flag for _, flag in stack)
        if in_friday and _assigned(line):
            bad.append(line[:160])
            continue
        cells = re.split(r"\|", line)
        friday_cell = any(to_iso(c) == FRIDAY and not CH2_SESSION.search(c) for c in cells)
        ch2_cell = any(_assigned(c) for c in cells)
        if friday_cell and ch2_cell:
            bad.append(line[:160])
            continue
        for clause in re.split(r"[.;|·]", line):
            if to_iso(clause) == FRIDAY and _assigned(clause):
                bad.append(line[:160])
                break
    return bad


def check_root(root):
    files = [
        "docs/playbook/UNIT1_PACING_CALENDAR.md",
        "production/final-8.5-week/MASTER_CURRICULUM_MAP.md",
        "production/final-8.5-week/MASTER_RESOURCE_INDEX.md",
        "production/final-8.5-week/WEEK1_2_TEACHER_QUICKSTART.md",
        "production/final-package/ENERGY_UNIT_WEEK_BY_WEEK_TEACHER_MAP.md",
        "production/final-package/ENERGY_UNIT_MEDIA_PLAYBACK_GUIDE.md",
        "production/final-8.5-week/VIDEO_PRESENTATION_INTEGRATION_MATRIX.md",
    ]
    fails = []
    for rel in files:
        path = os.path.join(root, rel)
        if not os.path.exists(path):
            fails.append("SCHEDULE missing " + rel)
            continue
        for hit in hits(open(path, encoding="utf-8").read()):
            fails.append("FRIDAY %s: %s" % (rel, hit))
    pb = os.path.join(root, "production/teacher-playbook/unit1-teacher-playbook.docx")
    if os.path.exists(pb):
        from docx import Document
        d = Document(pb)
        lines = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                lines.append(" | ".join(c.text for c in row.cells))
        for hit in hits("\n".join(lines)):
            fails.append("FRIDAY playbook: " + hit)
    return fails


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=None)
    a = ap.parse_args()
    root = os.path.abspath(a.root or os.path.join(os.path.dirname(__file__), ".."))
    fails = check_root(root)
    if fails:
        for f in fails:
            print("[FAIL]", f)
        sys.exit(1)
    print("[PASS] Chapter 2 required sessions are Wed 2026-10-14 and Thu 2026-10-15")
    sys.exit(0)


if __name__ == "__main__":
    main()
