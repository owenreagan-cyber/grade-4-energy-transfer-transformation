"""QA gates for the Unit 1 Teacher Playbook (Phase 2 foundation).

Headless checks: page size, margins, font floors, table widths (no clipping),
header/footer consistency, visual-system token use, grayscale contrast,
no unsupported-artifact claims, no stale INB-primary / 26-item language.

Run: .venv/bin/python tools/qa_teacher_playbook.py
"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import Pt

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
DOCX = os.path.join(ROOT, "production", "teacher-playbook", "unit1-teacher-playbook.docx")
PDF = os.path.join(ROOT, "production", "teacher-playbook", "unit1-teacher-playbook.pdf")
REPORT_DIR = os.path.join(ROOT, "qa", "reports")

# FROZEN tokens (design-tokens.json)
ALLOWED_FILLS = {
    "1A1D26", "4A5268", "FFFFFF", "F6F7F9", "D8DCE4",  # unit neutrals
    "1F2A5A", "A67C16", "7A5620", "F5EEDC",  # ch1
    "0F5A5C", "0E8A97", "5B6770",              # ch2
    "1B4FC2", "5E8A18", "2B2F36", "E9EBEF",    # ch3
    "6A3FB5", "0E93A6", "2A2350", "EFEDF6",    # ch4
    "A62626", "C25A1E", "2E2E30", "F7EFE2",    # ch5
    "1F5C3D", "8F6D24", "5A6B7A", "F4F5F3",    # ch6
    "3E2C1F", "A06E16", "2E3238", "EFE3CE",    # ch7
}

BANNED = [
    "26 items", "26 questions", "26Q", "26 Q",
    "Interactive Science Notebook is the primary", "INB-primary",
    "kinetic energy is required", "potential energy is required",
]

def hex_of_color(c):
    try:
        if c and c.rgb:
            return str(c.rgb)
    except Exception:
        pass
    return None

def _runs_in(doc):
    for p in doc.paragraphs:
        yield from p.runs
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield from p.runs
    sec = doc.sections[0]
    for p in list(sec.footer.paragraphs) + list(sec.header.paragraphs):
        yield from p.runs

def check_font_floors(doc):
    # text-free runs (the 1pt page-break carrier paragraph, Phase 6.1) are not readable text
    mins = [r.font.size.pt for r in _runs_in(doc) if r.font.size and r.text.strip()]
    if not mins:
        return ("FAIL", "no sized runs found")
    lo = min(mins)
    return ("PASS" if lo >= 8.0 else "WARN", f"min run size {lo:.1f} pt (floor 8pt, incl. footer)")

def check_table_widths(doc):
    bad = []
    for ti, t in enumerate(doc.tables):
        try:
            w = sum(c.width for c in t.rows[0].cells if c.width)
            if w > 0 and w > 6400800 + 20000:  # 7.0in + tolerance
                bad.append((ti, round(w / 914400, 2)))
        except Exception:
            pass
    return ("PASS" if not bad else "WARN", f"{len(bad)} table(s) exceed 7.0in usable width {bad}")

def check_token_use(doc):
    fills = set()
    from docx.oxml.ns import qn
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                shd = cell._tc.find(qn('w:tcPr'))
                if shd is not None:
                    for el in shd.findall(qn('w:shd')):
                        f = el.get(qn('w:fill'))
                        if f and f != 'auto':
                            fills.add(f.lstrip('#').upper())
    unknown = fills - ALLOWED_FILLS
    return ("PASS" if not unknown else "WARN", f"fills used: {sorted(fills)}; unknown: {sorted(unknown) or 'none'}")

def check_banned_text(doc):
    hits = []
    def scan(texts):
        for t in texts:
            for b in BANNED:
                if b.lower() in t.lower():
                    hits.append(b)
    scan([p.text for p in doc.paragraphs])
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                scan([p.text for p in c.paragraphs])
    return ("PASS" if not hits else "FAIL", f"banned language: {sorted(set(hits)) or 'none'}")

def check_required_terms(doc):
    """Verify locked assessment wording is present and correct."""
    blob = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                blob += "\n" + c.text
    checks = {
        "25 questions / 50 points": "25 questions" in blob and "50 points" in blob,
        "matching Q1–8": "Q1–8" in blob and "matching" in blob.lower(),
        "MC Q9–25": "Q9–25" in blob,
        "no short answer": "no short answer" in blob.lower(),
        "no CER": "no cer" in blob.lower(),
        "no fill-in-the-blank": "fill-in-the-blank" in blob.lower(),
        "raw score + percentage": "raw score" in blob.lower() and "%" in blob,
        "weighting not specified": "not specified" in blob,
    }
    missing = [k for k, v in checks.items() if not v]
    return ("PASS" if not missing else "WARN", f"missing: {missing or 'none'}")

def check_docx_page_size(doc):
    sec = doc.sections[0]
    w = sec.page_width; h = sec.page_height
    ok = abs(w - 7772400) < 10000 and abs(h - 10058400) < 10000  # 8.5x11 in EMU
    return ("PASS" if ok else "FAIL", f"section page = {w/914400:.2f} x {h/914400:.2f} in (expect 8.5 x 11)")

def check_grayscale_contrast():
    # WCAG relative luminance (linearized sRGB); desaturation preserves luminance, so this
    # also proves the grayscale (B/W print) survival of chapter primaries vs white.
    def _lin(c):
        c = c / 255
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    def lum(hexstr):
        r = _lin(int(hexstr[0:2], 16)); g = _lin(int(hexstr[2:4], 16)); b = _lin(int(hexstr[4:6], 16))
        return 0.2126 * r + 0.7152 * g + 0.0722 * b
    prims = ["1F2A5A", "0F5A5C", "1B4FC2", "6A3FB5", "A62626", "1F5C3D", "3E2C1F"]
    bad = []
    for h in prims:
        L = lum(h)
        ratio = (1.0 + 0.05) / (L + 0.05)
        if ratio < 3.0:
            bad.append(h)
    return ("PASS" if not bad else "WARN", f"chapter primaries vs white < 3:1: {bad or 'none'}")

def check_pdf():
    import subprocess
    if not os.path.exists(PDF):
        return ("FAIL", "PDF missing")
    info = subprocess.run(["pdfinfo", PDF], capture_output=True, text=True).stdout
    size = [l for l in info.splitlines() if "Page size" in l]
    pages = [l for l in info.splitlines() if l.startswith("Pages")]
    ok_size = "612 x 792" in (size[0] if size else "")
    # footer page number on multiple pages
    txt = subprocess.run(["pdftotext", PDF, "-"], capture_output=True, text=True).stdout
    footer_ok = txt.count("Unit 1 Teacher Playbook") > 10
    return ("PASS" if ok_size and footer_ok else "WARN",
            f"size={ok_size} ({size[0].strip() if size else '?'}); {pages[0].strip() if pages else '?'}; footer_pages={txt.count('Unit 1 Teacher Playbook')}")

STALE_CH1 = ["(prototype)", "Model It", "Draw/Label", "Block 1 \u2014 Observe", "Block 2 \u2014 Forms", "Lesson blocks (Block",
             "cup/balloon/salt", "letters shuffled", "Flashlight + salt", "PLANNED placeholders until Phases 3"]

def _blob(doc):
    b = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                b += "\n" + c.text
    return b

def _ordered_lines(doc):
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    out = []
    for el in doc.element.body.iterchildren():
        if el.tag.endswith("}p"):
            out.append(Paragraph(el, doc).text)
        elif el.tag.endswith("}tbl"):
            for row in Table(el, doc).rows:
                out.append(" | ".join(c.text for c in row.cells))
    return out

def _ch1_scope(doc):
    """Ch1 text only: Mon/Tue/Fri-wk1 pacing rows + everything from the Ch1 spread to the Ch2 spread."""
    lines = _ordered_lines(doc)
    keep, on = [], False
    for ln in lines:
        if ln.startswith(("Mon 10/12", "Tue 10/13", "Fri 10/16")):
            keep.append(ln)
        if "Chapter 1 \u2014 Teacher Spread" in ln:
            on = True
        if on and "Chapter 2 \u2014 Teacher Spread" in ln:
            on = False
        if on:
            keep.append(ln)
    return "\n".join(keep)

def check_ch1_pacing_sync(doc):
    """Phase 6.1: Playbook Ch1 timing must match tools/ch1_pacing.py (one authoritative plan)."""
    import ch1_pacing as P
    blob = _blob(doc)
    scope = _ch1_scope(doc)
    stale = sorted({t for t in STALE_CH1 if t in scope})
    need = ["Chapter 1 Pacing Plan", "Session 1", "Session 2", "Session 3", "9 items", "16 pts"]
    missing = [n for n in need if n not in blob]
    sess = [P.session_required_minutes(i) for i in (1, 2, 3)]
    ok = not stale and not missing and all(m <= 25.0 for m in sess)
    return ("PASS" if ok else "FAIL", f"stale Ch1 tokens: {stale or 'none'} · missing: {missing or 'none'} · session minutes {sess}")

def check_ch2_pacing_sync(doc):
    """Chapter 2 teacher text must match tools/ch2_pacing.py and drop legacy row language."""
    import re
    import ch2_pacing as P
    lines = _ordered_lines(doc)
    keep, on = [], False
    for ln in lines:
        if ln.startswith(("Wed 10/14", "Thu 10/15")):
            keep.append(ln)
        if "Chapter 2 \u2014 Teacher Spread" in ln:
            on = True
        if on and "Answer-Key Architecture" in ln:
            on = False
        if on:
            keep.append(ln)
    scope = "\n".join(keep)
    stale = [t for t in ("Model It", "Draw/Label", "Block 1", "Block 2", "Key Words") if t in scope]
    paths = re.findall(r"(docs/|curriculum/|production/|student/recovery|[\w./-]+\.md\b)", scope)
    need = ["Chapter 2 Pacing Plan", "24 min", "12 pts", "minute 12"]
    missing = [n for n in need if n not in scope and n not in _blob(doc)]
    # minutes themselves live in the pacing module; the spread must name the plan
    sess = [P.session_required_minutes(i) for i in (1, 2)]
    ok = not stale and not paths and not missing and all(m <= 25.0 for m in sess) and not P.check_plan()
    return ("PASS" if ok else "FAIL",
            "stale: %s · paths: %s · missing: %s · minutes %s" % (stale or "none", paths or "none", missing or "none", sess))

def check_no_repo_paths(doc, _doc_texts=None):
    """6.1C / L-D: Unit-level + Chapter 1 teacher-facing text must use classroom names, not repo paths.
    Out of scope until their own phases: Chapter 2-7 rows and the Ch3-7 SKELETON appendix register."""
    import re
    rx = re.compile(r"(docs/|curriculum/|production/|qa/|tools/|student/recovery|\bCH\d_[A-Z_]+|[\w-]+\.md\b)")
    later = re.compile(r"chapter-0[2-7N]|\bch[2-7]\b|\bCH[2-7]_|-ch[2-7]-|ch[2-7]-|/03-|/07-|FRIDAY_FLEX_FRAMEWORK|WHOLE_UNIT|production/review|MASTER_CURRICULUM_MAP|SKELETON|Chapter [2-7]\b|Ch[2-7]\b", re.I)
    texts = _doc_texts if _doc_texts is not None else [c.text for t in doc.tables for r in t.rows for c in r.cells] + [p.text for p in doc.paragraphs]
    hits = sorted({m.group(0) for t in texts if not later.search(t) for m in rx.finditer(t)})
    return ("PASS" if not hits else "FAIL", f"Ch1/unit-level repo paths in teacher text: {hits or 'none'}")

def main():
    import argparse
    global ROOT, DOCX, PDF, REPORT_DIR
    ap = argparse.ArgumentParser(); ap.add_argument("--report", default=None); ap.add_argument("--root", default=None); args = ap.parse_args()
    if args.root:
        ROOT = os.path.abspath(args.root)
        DOCX = os.path.join(ROOT, "production", "teacher-playbook", "unit1-teacher-playbook.docx")
        PDF = os.path.join(ROOT, "production", "teacher-playbook", "unit1-teacher-playbook.pdf")
        REPORT_DIR = os.path.join(ROOT, "qa", "reports")
    doc = Document(DOCX)
    gates = [
        ("page-size (US Letter)", check_docx_page_size(doc)),
        ("font floors", check_font_floors(doc)),
        ("table widths / clipping", check_table_widths(doc)),
        ("visual-system token use", check_token_use(doc)),
        ("grayscale contrast (primary vs white)", check_grayscale_contrast()),
        ("no banned/stale language", check_banned_text(doc)),
        ("required assessment terms", check_required_terms(doc)),
        ("PDF render (size + footer)", check_pdf()),
        ("Ch1 pacing sync / stale-language (6.1)", check_ch1_pacing_sync(doc)),
        ("Ch1/unit-level pages: no raw repo paths (6.1C L-D)", check_no_repo_paths(doc)),
        ("Ch2 pacing sync / no legacy rows", check_ch2_pacing_sync(doc)),
    ]
    print("=" * 70)
    print("UNIT 1 TEACHER PLAYBOOK — QA GATES (Phase 2)")
    print("=" * 70)
    status = "PASS"
    for name, (s, note) in gates:
        if s == "FAIL":
            status = "FAIL"
        elif s == "WARN" and status == "PASS":
            status = "WARN"
        print(f"[{s:4}] {name:38} {note}")
    print("=" * 70)
    print(f"OVERALL: {status}")
    os.makedirs(REPORT_DIR, exist_ok=True)
    rep = args.report or os.path.join(REPORT_DIR, "PHASE_2_PLAYBOOK_QA.md")
    with open(rep, "w") as f:
        f.write("# Phase 2 — Teacher Playbook Automated QA\n\n")
        f.write(f"Overall: **{status}**\n\n")
        for name, (s, note) in gates:
            f.write(f"- [{s}] {name} — {note}\n")
    print(f"report: {rep}")
    if status == "FAIL":
        sys.exit(1)

if __name__ == "__main__":
    main()
