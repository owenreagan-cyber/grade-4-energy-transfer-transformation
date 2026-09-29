"""Embed approved rebuilt infographics as full-screen REFERENCE slides (Phase 13R).

Appends ONE reference slide to the END of each target deck so existing slide
numbering (and all master-map references) stay intact. Each slide = the 16:9
infographic + a concise bottom prompt bar + teacher speaker notes.
"""
import os, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
SLIDES = os.path.join(ROOT, "production", "slides")
INFO = os.path.join(ROOT, "production", "infographics")

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# (deck filename, infographic filename, prompt badge, prompt question, purpose)
EMBEDS = [
    ("01-ch1-energy-causes-change-slides.pptx", "unit1-energy-detective-anchor.png",
     "REFERENCE", "What change is the evidence?", "Ch1 Energy Detective anchor (energy → change → evidence)."),
    ("04-ch4-energy-transfer-slides.pptx", "unit1-transfer-vs-transformation.png",
     "LOOK / NOTICE", "Did energy move, change form, or both?", "Ch4 transfer vs transformation reference."),
    ("05-ch5-collisions-slides.pptx", "unit1-collision-energy-transfer.png",
     "REFERENCE", "What transfers during a collision?", "Ch5 collision energy-transfer reference."),
    ("06-ch6-energy-transformations-slides.pptx", "unit1-engineering-design-cycle.png",
     "COMPARE", "What are the criteria and constraints?", "Ch6 engineering design cycle + Edison/Latimer reference."),
    ("01-u2-ch1-renewable-nonrenewable-slides.pptx", "unit2-renewable-nonrenewable-resources.png",
     "COMPARE", "Which resources renew quickly enough to be called renewable?", "U2 Ch1 renewable vs nonrenewable matrix."),
    ("03-u2-ch3-slides.pptx", "unit2-turbine-generator-pathways.png",
     "COMPARE", "What stays the same? What changes?", "U2 Ch3–6 turbine → generator pathways reference."),
]

def add_reference_slide(deck, img, badge, question, purpose):
    path = os.path.join(SLIDES, deck)
    prs = Presentation(path)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # image (16:9 = 2752x1536) centered, with room for the bottom prompt bar
    img_path = os.path.join(INFO, img)
    img_w = Inches(12.0)
    img_h = Inches(12.0 * 1536 / 2752)
    x = Inches((13.333 - 12.0) / 2)
    y = Inches(0.2)
    slide.shapes.add_picture(img_path, x, y, width=img_w, height=img_h)

    # bottom prompt bar
    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(7.02), img_w, Inches(0.42))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    bar.shadow.inherit = False
    tf = bar.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r1 = p.add_run(); r1.text = badge + "   ·   "
    r1.font.size = Pt(15); r1.font.bold = True; r1.font.color.rgb = RGBColor(0xD9, 0xA4, 0x41); r1.font.name = "Arial"
    r2 = p.add_run(); r2.text = question
    r2.font.size = Pt(15); r2.font.bold = True; r2.font.color.rgb = WHITE; r2.font.name = "Arial"

    # teacher speaker notes
    notes = slide.notes_slide.notes_text_frame
    notes.text = f"INFOGRAPHIC REFERENCE ({badge})\n{question}\n\n{purpose}\nApproved rebuilt visual — production/infographics/{img}. Locked Grade 4 vocabulary only."

    prs.save(path)
    print(f"embedded {img} -> {deck} (slide {len(prs.slides.__iter__.__self__._sldIdLst)})")

def main():
    for deck, img, badge, q, purpose in EMBEDS:
        add_reference_slide(deck, img, badge, q, purpose)

if __name__ == "__main__":
    main()
