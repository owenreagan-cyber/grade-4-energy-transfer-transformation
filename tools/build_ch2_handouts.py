"""Chapter 2 student handouts (Phase 7.4). Quiz, practice, Moving Marbles, recovery, flex.

Quiz matches the Chapter 1 header, points, and writing-line pattern.
Moving Marbles stays letter landscape and is labeled an after-quiz extension.

Run: .venv/bin/python tools/build_ch2_handouts.py
"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_ch2_guided_notes as g
from build_ch2_guided_notes import (Inches, Pt, WD_ALIGN_PARAGRAPH, FULL_W,
    PRIMARY_HEX, INK_HEX, CH2_PRIMARY, CH2_ACCENT, CH2_NEUTRAL, INK, INK_SOFT, WHITE)
from docx.enum.section import WD_ORIENT

ROOT = g.ROOT
BODY, PROMPT, CAP = 14.0, 14.5, 12.0

def header(doc, title, tag):
    g._para(doc, "UNIT 1 · ENERGY TRANSFER & TRANSFORMATION", size=g.ATTR_SZ, bold=True, color=INK_SOFT, after=2, caps=True)
    g._para(doc, title, size=22, bold=True, color=CH2_PRIMARY, after=2)
    g._para(doc, tag, size=CAP, bold=True, color=CH2_NEUTRAL, after=4)
    p = doc.add_paragraph(); g._spacing(p, 0, 6, line_pt=26)
    g._run(p, "Name: ", size=CAP, bold=True)
    g._run(p, "_" * 28, size=CAP, underline=True)
    g._run(p, "    Homeroom:  ", size=CAP, bold=True)
    for i, lbl in enumerate(("4A", "4B", "4C", "4D")):
        g._run(p, ("  " if i else "") + "○ " + lbl, size=CAP)

def score_box(doc, total):
    t = doc.add_table(rows=1, cols=1); g._no_borders(t)
    c = t.cell(0, 0); c.width = FULL_W
    g._margins(c, top=60, left=120, bottom=60, right=120); g._shade(c, "F6F7F9")
    g._border(c, left=dict(sz=16, color="0E8A97"))
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    g._run(p, "SCORE:  ", size=CAP, bold=True, color=CH2_NEUTRAL)
    g._run(p, "_____", size=PROMPT, underline=True)
    g._run(p, " / %d      " % total, size=PROMPT, bold=True)
    g._run(p, "_______", size=PROMPT, underline=True)
    g._run(p, " %", size=PROMPT, bold=True)

def q_head(doc, n, pts, text):
    p = doc.add_paragraph(); g._spacing(p, 8, 2, line_pt=20)
    g._run(p, "%d.  " % n, size=PROMPT, bold=True, color=CH2_ACCENT)
    g._run(p, text, size=PROMPT, bold=True, color=INK)
    g._run(p, "   (%d pt%s)" % (pts, "" if pts == 1 else "s"), size=11, italic=True, color=INK_SOFT)

def mc(doc, opts):
    for L, o in zip("abcd", opts):
        p = doc.add_paragraph(); g._spacing(p, 0, 1, line_pt=20)
        p.paragraph_format.left_indent = Inches(0.3)
        g._run(p, "%s)  " % L, size=BODY, bold=True, color=CH2_PRIMARY)
        g._run(p, o, size=BODY)

def lines(doc, n):
    g._writing_line(doc, size=BODY, n=n)

def page_break(doc):
    p = doc.add_paragraph(); p.paragraph_format.page_break_before = True
    g._spacing(p, 0, 0, line_pt=1)
    r = p.add_run(""); r.font.size = Pt(1)

def draw_pair(doc, height_twips=2880):
    t = doc.add_table(rows=2, cols=2); g._no_borders(t)
    half = Inches((g.USABLE_IN - 0.2) / 2)
    for i, lab in enumerate(("BEFORE", "AFTER")):
        c = t.cell(0, i); c.width = half
        g._margins(c, top=40, left=80, bottom=40, right=80)
        g._shade(c, PRIMARY_HEX)
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        g._run(p, lab, size=CAP, bold=True, color=WHITE)
    g._row_height(t.rows[0], 360, rule="exact")
    for i in range(2):
        c = t.cell(1, i); c.width = half
        g._margins(c, top=40, left=80, bottom=40, right=80)
        for edge in ("top", "left", "bottom", "right"):
            g._border(c, **{edge: dict(sz=8, color=PRIMARY_HEX)})
        g._run(c.paragraphs[0], "", size=2)
    g._row_height(t.rows[1], height_twips, rule="exact")

def grid(doc, headers, widths, rows, row_h=700):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers)); g._no_borders(t)
    for r in t.rows:
        for i, w in enumerate(widths):
            r.cells[i].width = Inches(w)
    for i, h in enumerate(headers):
        c = t.cell(0, i); g._margins(c, top=40, left=70, bottom=40, right=70)
        g._shade(c, PRIMARY_HEX)
        g._run(c.paragraphs[0], h, size=CAP, bold=True, color=WHITE)
    g._row_height(t.rows[0], 400, rule="exact")
    for ri, row in enumerate(rows, 1):
        for ci, val in enumerate(row):
            c = t.cell(ri, ci); g._margins(c, top=40, left=70, bottom=40, right=70)
            g._border(c, bottom=dict(sz=8, color=INK_HEX))
            if val:
                g._run(c.paragraphs[0], val, size=BODY, bold=(ci == 0), color=INK)
        g._row_height(t.rows[ri], row_h, rule="atLeast")

def finish(doc, path, footer):
    g._footer(doc, footer)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    doc.save(path)
    return path

def build_quiz():
    doc = g.new_doc()
    header(doc, "Chapter 2 Quiz — Moving Objects Have Energy",
           "Chapter 2 · Grade 4 · SCORED — 8 questions · 12 points")
    p = doc.add_paragraph(); g._spacing(p, 0, 4, line_pt=18)
    g._run(p, "Notes today:   ○ OPEN NOTE     ○ CLOSED NOTE", size=CAP, bold=True, color=CH2_NEUTRAL)
    g._run(p, "      (teacher circles one)", size=11, italic=True, color=INK_SOFT)
    score_box(doc, 12)
    q_head(doc, 1, 1, "Motion is …")
    mc(doc, ["the process of an object changing position", "how fast something goes",
             "a kind of stored energy", "a visible substance"])
    q_head(doc, 2, 1, "A moving object has …")
    mc(doc, ["stored energy only", "energy of motion", "no energy", "energy we can see"])
    q_head(doc, 3, 2, "A student rolls a marble into a still marble. The still marble starts to move.")
    g._para(doc, "What was the moving object?", size=BODY, bold=True, after=1)
    lines(doc, 1)
    g._para(doc, "What changed?", size=BODY, bold=True, after=1)
    lines(doc, 1)
    q_head(doc, 4, 1, "Energy is the ability to …")
    mc(doc, ["cause change", "be seen", "stay still", "disappear"])
    page_break(doc)
    q_head(doc, 5, 2, "Draw the marble collision. Show the target before and after.")
    draw_pair(doc, 2880)
    q_head(doc, 6, 2, "How do we know a moving object has energy if we cannot see the energy?")
    g._para(doc, "Start with: “We know because …”", size=BODY, italic=True, color=INK_SOFT, after=2)
    lines(doc, 3)
    q_head(doc, 7, 2, "A swinging hammer hits a nail. What change does it cause?")
    lines(doc, 2)
    q_head(doc, 8, 1, "A book sits still on a desk. Nothing changes. Does the book show energy of motion?")
    mc(doc, ["Yes — every object has energy of motion", "No — nothing is moving and nothing changes",
             "Yes — the desk is holding it up", "No — books cannot have energy of any kind"])
    return finish(doc, os.path.join(ROOT, "production", "quizzes", "02-ch2-moving-objects-have-energy-quiz.docx"),
                  "Unit 1 · Chapter 2 · Quiz · Moving Objects Have Energy · 12 points")

def build_practice():
    doc = g.new_doc()
    header(doc, "Practice — Moving Objects Have Energy", "Chapter 2 · Grade 4 · PRACTICE — NOT GRADED")
    g._para(doc, "A. What is moving?", size=16, bold=True, color=CH2_PRIMARY, after=2)
    g._caption(doc, "Write the object that is in motion.", after=2)
    grid(doc, ["Scenario", "Object in motion"], [4.4, 3.0], [
        ["A basketball falls through the hoop.", ""],
        ["A marble rolls toward another marble.", ""],
        ["Wind spins a turbine blade.", ""],
        ["A shopping cart bumps a box.", ""]], row_h=620)
    g._para(doc, "B. What changed?", size=16, bold=True, color=CH2_PRIMARY, after=2)
    grid(doc, ["Scenario", "Change caused"], [4.4, 3.0], [
        ["The basketball hits the net.", ""],
        ["The marble hits the target marble.", ""],
        ["Wind spins the blades.", ""],
        ["The cart bumps the box.", ""]], row_h=620)
    page_break(doc)
    g._para(doc, "C. Evidence", size=16, bold=True, color=CH2_PRIMARY, after=2)
    g._caption(doc, "Circle YES or NO. Is it evidence that the moving object had energy?", after=2)
    grid(doc, ["Observation", "Evidence?"], [5.2, 2.2], [
        ["The target marble changed position.", "YES    NO"],
        ["The marble is round.", "YES    NO"],
        ["The block slid forward after the ball hit it.", "YES    NO"],
        ["I think the ball is heavy.", "YES    NO"]], row_h=640)
    g._para(doc, "D. Before and after", size=16, bold=True, color=CH2_PRIMARY, after=2)
    g._caption(doc, "Draw one moving object causing a change.", after=2)
    draw_pair(doc, 2600)
    g._para(doc, "E. Explain", size=16, bold=True, color=CH2_PRIMARY, after=2)
    g._para(doc, "How can a moving object show evidence of energy?", size=BODY, bold=True, after=2)
    lines(doc, 3)
    return finish(doc, os.path.join(ROOT, "production", "worksheets", "02-ch2-moving-objects-have-energy-practice.docx"),
                  "Unit 1 · Chapter 2 · Practice · Moving Objects Have Energy · Not graded")

def build_marbles():
    doc = g.new_doc()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Inches(11)
    sec.page_height = Inches(8.5)
    sec.left_margin = Inches(0.55)
    sec.right_margin = Inches(0.55)
    header(doc, "Moving Marbles", "Chapter 2 · Week 1 · Optional After-Quiz Extension · NOT GRADED")
    g._para(doc, "Big Question: How does a moving marble show that it has energy?", size=BODY, bold=True, after=2)
    g._para(doc, "Predict. What will happen when the moving marble reaches the still marble?", size=BODY, after=2)
    lines(doc, 2)
    g._para(doc, "Three rolls. Keep the setup the same. Do not measure speed.", size=BODY, bold=True, after=2)
    grid(doc, ["Trial", "What moved before?", "What changed after?", "Evidence"],
         [1.1, 2.8, 3.0, 3.0], [["1", "", "", ""], ["2", "", "", ""], ["3", "", "", ""]], row_h=780)
    page_break(doc)
    g._para(doc, "Draw the marbles before and after one roll.", size=BODY, bold=True, after=2)
    # landscape usable width is wider than the portrait FULL_W used inside draw_pair.
    t = doc.add_table(rows=2, cols=2); g._no_borders(t)
    half = Inches(4.8)
    for i, lab in enumerate(("BEFORE", "AFTER")):
        c = t.cell(0, i); c.width = half
        g._margins(c, top=30, left=80, bottom=30, right=80); g._shade(c, PRIMARY_HEX)
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        g._run(p, lab, size=CAP, bold=True, color=WHITE)
    g._row_height(t.rows[0], 360, rule="exact")
    for i in range(2):
        c = t.cell(1, i); c.width = half
        g._margins(c, top=40, left=80, bottom=40, right=80)
        for edge in ("top", "left", "bottom", "right"):
            g._border(c, **{edge: dict(sz=8, color=PRIMARY_HEX)})
    g._row_height(t.rows[1], 3200, rule="exact")
    g._para(doc, "Explain. The moving marble had energy of motion. What did you observe?", size=BODY, bold=True, after=2)
    lines(doc, 2)
    return finish(doc, os.path.join(ROOT, "production", "worksheets", "02-ch2-moving-marbles-student.docx"),
                  "Unit 1 · Chapter 2 · Moving Marbles · Optional after the quiz · Not graded")

def build_recovery():
    doc = g.new_doc()
    header(doc, "Absent-Student Recovery — Moving Objects Have Energy",
           "Chapter 2 · Grade 4 · Use with the Reader · Not a classmate report")
    g._para(doc, "A moving object has energy of motion because it can cause a change. The change is the evidence.",
            size=BODY, bold=True, after=4)
    g._para(doc, "Read printed pages 7–12. Then answer.", size=BODY, after=2)
    g._para(doc, "1. What is motion?", size=BODY, bold=True, after=1)
    lines(doc, 2)
    g._para(doc, "2. What is energy of motion?", size=BODY, bold=True, after=1)
    lines(doc, 2)
    g._para(doc, "3. Marble demonstration. One marble rolls into a still marble.", size=BODY, bold=True, after=2)
    draw_pair(doc, 1800)
    g._para(doc, "What changed?", size=BODY, bold=True, after=1)
    lines(doc, 2)
    g._para(doc, "4. How do we know the moving marble had energy? Start with “We know because…”", size=BODY, bold=True, after=1)
    lines(doc, 2)
    g._para(doc, "5. Name the change caused by a swinging hammer, a swinging pendulum, or the wind.", size=BODY, bold=True, after=1)
    lines(doc, 2)
    g._para(doc, "6. A ball does not start its own motion. It is put into motion by outside motion. What change can that moving ball cause?",
            size=BODY, bold=True, after=1)
    lines(doc, 2)
    g._para(doc, "7. Pick the wind or a swinging pendulum. What useful change does it cause?",
            size=BODY, bold=True, after=1)
    lines(doc, 2)
    g._para(doc, "", after=8)
    box = doc.add_table(rows=1, cols=1)
    g._no_borders(box)
    cell = box.cell(0, 0)
    g._shade(cell, "F6F7F9")
    g._margins(cell, top=60, left=100, bottom=60, right=100)
    g._border(cell, left=dict(sz=16, color="0E8A97"))
    p = cell.paragraphs[0]
    g._run(p, "TEACHER FOLLOW-UP", size=CAP, bold=True, color=CH2_PRIMARY)
    for line in (
        "Student finishes: Guided Notes N1–N3 and CHECK",
        "Quiz makeup: within 2 school days",
        "Moving Marbles: optional",
    ):
        p2 = cell.add_paragraph()
        g._run(p2, line, size=BODY)
    return finish(doc, os.path.join(ROOT, "production", "worksheets", "02-ch2-absent-student-recovery.docx"),
                  "Unit 1 · Chapter 2 · Absent-student recovery · Moving Objects Have Energy")

def build_flex():
    """Student card sort. Answers are not on this page."""
    doc = g.new_doc()
    header(doc, "Friday Flex — Moving Objects Have Energy",
           "Optional after the Chapter 2 quiz · Not Week 1 Friday · Not graded")
    g._para(doc, "Recover: finish Guided Notes, or use the absent-student recovery page.", size=BODY, after=4)
    g._para(doc, "Mystery files. For each story, name the moving object, the change, and the evidence.",
            size=BODY, bold=True, after=2)
    stories = [
        "1. A rolling cart pushes a box. The box slides across the floor.",
        "2. A swinging door bumps a ball. The ball rolls away.",
        "3. A toy car knocks over a stack of blocks. The blocks fall.",
    ]
    for story in stories:
        g._para(doc, story, size=BODY, after=1)
        lines(doc, 2)
    page_break(doc)
    g._para(doc, "Cut out the 9 cards. Match each object with the change it causes and the evidence that shows it has energy.",
            size=BODY, bold=True, after=2)
    g._para(doc, "Cut on the dashed lines. Do not write the matches on this page until you have sorted the cards.",
            size=CAP, italic=True, color=INK_SOFT, after=4)
    cards = [
        ("OBJECT", "Hammer"),
        ("CHANGE", "The blades spin"),
        ("EVIDENCE", "I saw the gears move"),
        ("OBJECT", "Pendulum"),
        ("CHANGE", "The nail goes into the wood"),
        ("EVIDENCE", "I saw the blades start to turn"),
        ("OBJECT", "Wind"),
        ("CHANGE", "The clock gears turn"),
        ("EVIDENCE", "I saw the nail move into the wood"),
    ]
    t = doc.add_table(rows=3, cols=3)
    g._no_borders(t)
    dash = dict(sz=12, val="dashed", color="0F5A5C")
    for i, (kind, label) in enumerate(cards):
        cell = t.cell(i // 3, i % 3)
        g._margins(cell, top=60, left=80, bottom=60, right=80)
        for edge in ("top", "left", "bottom", "right"):
            g._border(cell, **{edge: dash})
        p = cell.paragraphs[0]
        g._run(p, kind, size=10, bold=True, color=CH2_ACCENT, caps=True)
        p2 = cell.add_paragraph()
        g._run(p2, label, size=BODY, bold=True, color=INK)
        g._row_height(t.rows[i // 3], 900, rule="atLeast")
    g._para(doc, "", after=6)
    g._para(doc, "Sort mat. Place the matching cards in these columns.", size=BODY, bold=True, after=2)
    mat = doc.add_table(rows=4, cols=3)
    g._no_borders(mat)
    for i, lab in enumerate(("Object", "Change", "Evidence")):
        c = mat.cell(0, i)
        g._shade(c, PRIMARY_HEX)
        g._margins(c, top=40, left=60, bottom=40, right=60)
        g._run(c.paragraphs[0], lab, size=CAP, bold=True, color=WHITE)
    for r in range(1, 4):
        for c in range(3):
            cell = mat.cell(r, c)
            g._margins(cell, top=40, left=60, bottom=40, right=60)
            for edge in ("top", "left", "bottom", "right"):
                g._border(cell, **{edge: dict(sz=8, color="5B6770")})
            g._run(cell.paragraphs[0], "", size=BODY)
        g._row_height(mat.rows[r], 700, rule="atLeast")
    return finish(doc, os.path.join(ROOT, "production", "worksheets", "02-ch2-friday-flex.docx"),
                  "Unit 1 · Chapter 2 · Friday Flex · Cut-and-sort · Not graded")

if __name__ == "__main__":
    print(build_quiz())
    print(build_practice())
    print(build_marbles())
    print(build_recovery())
    print(build_flex())
