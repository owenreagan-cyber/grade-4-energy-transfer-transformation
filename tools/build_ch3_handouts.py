"""Chapter 3 practice, quiz, Ramp It Up sheet, recovery, and teacher records.

Reads the chapter spec. Does not edit frozen Chapter 1 or Chapter 2 files.
Run: .venv/bin/python tools/build_ch3_handouts.py
"""
import json
import os

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
SPEC = json.load(open(os.path.join(ROOT, "docs/chapter-03/CH3_CHAPTER_SPEC.json"), encoding="utf-8"))
PRIMARY = RGBColor(0x1B, 0x4F, 0xC2)
INK = RGBColor(0x1A, 0x1D, 0x26)
FONT = "Calibri"


def doc():
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    style = document.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(14)
    style.font.color.rgb = INK
    return document


def add(document, text, size=14, bold=False, color=INK, space=6, italic=False):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(space)
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    return p


def save(document, path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    document.save(path)
    print(path)


def _table_borders(table, color="1A1D26", sz="6"):
    """Visible, student-usable borders on every cell of a table."""
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement("w:%s" % edge)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), sz)
        node.set(qn("w:color"), color)
        borders.append(node)
    tblPr.append(borders)


def practice():
    document = doc()
    add(document, "Chapter 3  ·  Energy and Speed Are Related", 18, True, PRIMARY)
    add(document, "Practice  ·  not graded  ·  OPTIONAL", 12, True)
    add(document, "Name ______________________________    Date ______________")
    add(document, "Use one ball in every part. Do not write a speed number.")
    add(document, "1. A student rolls the same ball down a low ramp and a steeper ramp. The cup moves farther from the steeper ramp. What is the evidence that the ball had more energy of motion?", bold=True)
    add(document, "________________________________________________________________")
    add(document, "________________________________________________________________")
    add(document, "2. This chapter compares one ball at two speeds. Why would a different ball change the test?", bold=True)
    add(document, "________________________________________________________________")
    add(document, "________________________________________________________________")
    add(document, "3. Name one thing you measured in centimeters. Then name the two setups you compared to see which gave the same ball more energy of motion.", bold=True)
    add(document, "Measured: ______________________________________________________")
    add(document, "Compared: ______________________________________________________")
    save(document, os.path.join(ROOT, "production/worksheets/03-ch3-energy-speed-practice.docx"))
    key = """# Chapter 3 Practice Key

Not graded. OPTIONAL.

1. The farther cup move is the evidence. Accept: a bigger change, or the cup moved farther. Do not accept a speed number.
2. A different ball could be heavier or lighter, so the test would compare different objects instead of one object at two speeds.
3. Measured: ramp height in centimeters, or how far the cup moved in centimeters. Compared: a lower setup and a steeper setup (for example, 1 book and 3 books) for the same ball.

Reader pp. 13–16. Teacher Guide Lessons 4 and 5.
"""
    path = os.path.join(ROOT, "production/answer-keys/03-ch3-energy-speed-practice-key.md")
    open(path, "w", encoding="utf-8").write(key)
    print(path)


def quiz():
    document = doc()
    add(document, "Chapter 3  ·  Energy and Speed Are Related", 18, True, PRIMARY)
    add(document, "Quiz  ·  8 items  ·  12 points  ·  FLEX", 12, True)
    add(document, "Name ______________________________    Date ______________")
    add(document, "Score: ______ / 12 = ______ %", 16, True)
    items = [
        "1. (1 point) Speed is a measurement of the ________ an object travels over an amount of time.",
        "2. (1 point) Which comparison fits this chapter?  A. one ball at two speeds   B. a heavy ball and a light ball",
        "3. (2 points) What shows that the faster ball had more energy of motion?",
        "4. (2 points) In Ramp It Up 2.0, what did we measure in centimeters? Name both things.",
        "5. (1 point) What was the one thing you changed?",
        "6. (1 point) A correct graph uses which labels?  A. ramp height (cm) and average cup movement (cm)   B. speed and energy",
        "7. (2 points) The same ball rolls from a steeper ramp. What greater change do you look for?",
        "8. (2 points) How are energy and speed related?",
    ]
    for item in items:
        add(document, item, bold=True, space=2)
        add(document, "________________________________________________________________", space=10)
    save(document, os.path.join(ROOT, "production/quizzes/03-ch3-energy-speed-quiz.docx"))
    key = """# Chapter 3 Quiz Key

8 items · 12 points. Score: ______ / 12 = ______ %.
FLEX. Friday 2026-10-23 or makeup within 2 school days. Not a first lesson.

1. (1 point) distance. Reader p. 13.
2. (1 point) A. one ball at two speeds.
5. (1 point) ramp height. The ball stayed the same.
6. (1 point) A. ramp height (cm) and average cup movement (cm).

Scoring for the four 2-point items (3, 4, 7, 8):

3. (2 points) What shows that the faster ball had more energy of motion?
   2 = names a greater change (a farther cup move, a higher bounce, or a louder sound) with no speed number.
   1 = names a change but not a greater one.
   0 = incorrect, or gives a speed number.

4. (2 points) In Ramp It Up 2.0, what did we measure in centimeters? Name both things.
   2 = names both: ramp height in centimeters AND cup movement in centimeters.
   1 = names only one of the two.
   0 = neither, or gives a speed number.

7. (2 points) The same ball rolls from a steeper ramp. What greater change do you look for?
   2 = names a greater change for the same ball (the cup or bucket moving farther).
   1 = names a change but not a greater one.
   0 = incorrect.

8. (2 points) How are energy and speed related?
   2 = same object, and greater speed relates to greater energy of motion, with a greater change as the evidence.
   1 = states the relationship for the same object but omits the evidence, or states evidence without the relationship.
   0 = a mass-based comparison, a speed number, or no same object.

Do not accept kinetic, potential, mechanical, or a speed in meters per second.
"""
    path = os.path.join(ROOT, "production/answer-keys/03-ch3-energy-speed-quiz-key.md")
    open(path, "w", encoding="utf-8").write(key)
    print(path)
    blueprint = """# Chapter 3 Quiz Blueprint

8 items · 12 points · raw score and percentage.
Status: FLEX. It assesses ideas taught Monday 2026-10-19 through Thursday 2026-10-22.

| Item | Points | Anchor | Demand |
|---|---|---|---|
| 1 | 1 | V1 | vocabulary in context |
| 2 | 1 | N1 | same object |
| 3 | 2 | N1 | evidence |
| 4 | 2 | N2 | measured in centimeters |
| 5 | 1 | N2 | fair test |
| 6 | 1 | GRAPH | axis labels |
| 7 | 2 | N3 | steeper ramp |
| 8 | 2 | BQ_NOW | Big Question |

No mass comparison as a correct answer. No required speed or energy calculation.
"""
    path = os.path.join(ROOT, "production/quizzes/03-ch3-energy-speed-quiz-blueprint.md")
    open(path, "w", encoding="utf-8").write(blueprint)
    print(path)


def lab():
    document = doc()
    add(document, "Ramp It Up 2.0", 18, True, PRIMARY)
    add(document, "Chapter 3  ·  REQUIRED  ·  Tuesday and Wednesday", 12, True)
    add(document, "Name ______________________________    Date ______________")
    add(document, "QUESTION: How does changing ramp height affect the change caused by the same ball?", bold=True)
    add(document, "PREDICT: Which ramp height do you predict will move the cup farthest? Why?", bold=True)
    add(document, "________________________________________________________________")
    add(document, "FAIR TEST — keep the same:", bold=True)
    add(document, "· same ball   · same ramp   · same release spot   · let go, do not push")
    add(document, "· same cup   · same cup starting spot   · same measuring method")
    add(document, "CHANGE ONE THING: ramp height.  MEASURE: ramp height (cm) and cup movement (cm).")
    add(document, "Roll once to calibrate. Do not record that roll. Then record three trials at each height. That is nine measured trials.")
    add(document, "Ramp height is not a speed. Do not write a speed number.")
    table = document.add_table(rows=4, cols=6)
    headers = ["Condition / books", "Actual ramp height (cm)", "Trial 1 cup movement (cm)", "Trial 2 cup movement (cm)", "Trial 3 cup movement (cm)", "Average cup movement (cm)"]
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for index, books in enumerate(("1 book", "2 books", "3 books"), start=1):
        table.rows[index].cells[0].text = books
    _table_borders(table)
    add(document, "")
    add(document, "GRAPH: Plot your three averages. Do not label an axis as speed.", bold=True)
    add(document, "y-axis — Average cup movement (cm)", size=12, italic=True)
    graph = document.add_table(rows=7, cols=6)
    _table_borders(graph)
    for row in graph.rows:
        tr = row._tr
        trPr = tr.get_or_add_trPr()
        trHeight = OxmlElement("w:trHeight")
        trHeight.set(qn("w:val"), str(480))
        trHeight.set(qn("w:hRule"), "atLeast")
        trPr.append(trHeight)
        for cell in row.cells:
            cell.paragraphs[0].text = ""
    add(document, "")
    add(document, "x-axis — Actual ramp height (cm)", size=12, italic=True)
    add(document, "Challenge: move the cup as close as you can to 30 cm without going past.")
    add(document, "My closest result: ________ cm. Did it go past 30 cm?  yes / no")
    add(document, "READ THE DATA: Which ramp height had the greatest average cup movement?", bold=True)
    add(document, "________________________________________________________________")
    add(document, "COMPARE: Compare your 1-book and 3-book averages. Which moved the cup farther?", bold=True)
    add(document, "________________________________________________________________")
    add(document, "NOTICE THE PATTERN: As ramp height increased, what happened to the average cup movement?", bold=True)
    add(document, "________________________________________________________________")
    add(document, "CONNECT TO SPEED: The Reader says objects pick up more speed on steeper ramps. Which setup gave the same ball more speed?", bold=True)
    add(document, "________________________________________________________________")
    add(document, "CONNECT TO ENERGY: The Reader says the faster the ball is moving, the greater its energy of motion. Which setup gave the same ball more energy of motion?", bold=True)
    add(document, "________________________________________________________________")
    add(document, "CLAIM: When the same object moves faster, it has __________ energy of motion.", bold=True)
    add(document, "EVIDENCE: Write TWO actual average cup-movement values from your table that support your claim.", bold=True)
    add(document, "____________________  and  ____________________")
    add(document, "REASONING: Explain how those two values support your claim.", bold=True)
    add(document, "________________________________________________________________")
    add(document, "________________________________________________________________")
    add(document, "EXTEND YOUR THINKING: If you tested a ramp height between two of your tested heights, what cup movement would you predict? Explain using your graph.", bold=True)
    add(document, "________________________________________________________________")
    save(document, os.path.join(ROOT, "production/worksheets/03-ch3-ramp-it-up-student.docx"))
    key = """# Ramp It Up 2.0 Key

REQUIRED lab. There is no single correct centimeter table. Students record their own trials.

Look for:
- one ball
- vertical ramp height in centimeters for 1, 2, and 3 books
- one unrecorded calibration roll
- three recorded trials at each height (nine recorded trials)
- an average column
- a graph of ramp height (cm) against average cup movement (cm) with a visible grid
- the CONNECT TO SPEED and CONNECT TO ENERGY reasoning steps (not collapsed into ramp height = speed)
- no speed number and no meters per second
- the 30 cm challenge judged by closeness without going past

Expected claim: When the same object moves faster, it has greater energy of motion.
Acceptable wording: "more energy of motion", "greater energy".
Evidence should contain TWO actual average cup-movement values from the student's own table.
Reasoning should connect a larger average to a greater change.
Misconception to catch: ramp height IS the speed, or that a speed number was measured.

A steeper ramp should usually move the cup farther. If it does not, help the student check the fair test before changing the science claim.
"""
    path = os.path.join(ROOT, "production/answer-keys/03-ch3-ramp-it-up-key.md")
    open(path, "w", encoding="utf-8").write(key)
    print(path)
    spec = """# Ramp It Up 2.0 — Teacher Spec

Status: REQUIRED. Provenance: ADAPTED from AP 4.1.
Dates: Tuesday 2026-10-20 and Wednesday 2026-10-21.
Materials: cardboard ramp, 3 books, 1 ball, 1 cup, centimeter ruler.

%s

Fair-test controls: same ball, same ramp, same release spot, let go (no push), same cup, same cup starting spot, same measuring method. One changed variable: ramp height.

Do not build a second ramp lab. Drop It is Friday FLEX, not this lab.
""" % SPEC["anchors"][4]["protocolStatement"]
    path = os.path.join(ROOT, "production/labs/03-ch3-ramp-it-up-teacher-spec.md")
    open(path, "w", encoding="utf-8").write(spec)
    print(path)


def recovery():
    document = doc()
    add(document, "Chapter 3  ·  Absent-Student Recovery", 18, True, PRIMARY)
    add(document, "FLEX  ·  not a first lesson", 12, True)
    add(document, "Name ______________________________    Date ______________")
    add(document, "BIG IDEA: For the same object, greater speed is related to greater energy of motion. A greater change is the evidence.", bold=True)
    add(document, "READ: Reader pages 13–16.", bold=True)
    add(document, "1. What is speed? (Reader p. 13)")
    add(document, "________________________________________________________________")
    add(document, "2. Compare the same ball at two speeds. What greater change shows more energy of motion?")
    add(document, "________________________________________________________________")
    add(document, "3. In Ramp It Up 2.0, what two things did students measure in centimeters?")
    add(document, "________________________________________________________________")
    add(document, "4. What did the Reader say happens to objects on steeper ramps? (p. 15)")
    add(document, "________________________________________________________________")
    add(document, "5. How are energy and speed related? Use the same object in your answer.")
    add(document, "________________________________________________________________")
    add(document, "________________________________________________________________")
    add(document, "Missed the lab? Explain the relationship from the Reader. Your teacher may show you the class graph. Do not copy fake personal data.", bold=True)
    add(document, "Drop It is not required makeup.")
    add(document, "If you missed the quiz, makeup is within 2 school days.")
    save(document, os.path.join(ROOT, "production/worksheets/03-ch3-absent-student-recovery.docx"))


def pacing_and_playbook():
    lines = ["# Chapter 3 — Authoritative Pacing Plan", "", "Generated from `docs/chapter-03/CH3_CHAPTER_SPEC.json`. Do not keep a second timing system.", "", "30-minute block. Required sessions stay at or under 24 minutes.", ""]
    lines.append("| Session | Date | Status | Reader | Required min | Slides |")
    lines.append("|---|---|---|---|---|---|")
    for session in SPEC["sessions"]:
        lines.append("| %s | %s | %s | %s | %s | %s |" % (
            session["id"], session["date"], session["status"], session["readerPages"] or "—",
            session["requiredMinutes"], ", ".join(session["slideIds"]),
        ))
    lines.append("")
    lines.append(SPEC["calendarResolution"]["statement"])
    lines.append("")
    lines.append(SPEC["anchors"][4]["protocolStatement"])
    path = os.path.join(ROOT, "docs/chapter-03/CH3_PACING_PLAN.md")
    open(path, "w", encoding="utf-8").write("\n".join(lines) + "\n")
    print(path)
    play = ["# Chapter 3 Playbook Region", "", "The frozen unit playbook is not edited. This file is the Chapter 3 teacher region.", ""]
    for anchor in SPEC["anchors"]:
        text = " ".join(anchor.get("lines") or []) or anchor.get("protocolStatement") or ""
        play.append("## %s" % anchor["code"])
        play.append("status: %s" % anchor["status"])
        play.append("text: %s" % text)
        play.append("")
    path = os.path.join(ROOT, "docs/chapter-03/CH3_PLAYBOOK_REGION.md")
    open(path, "w", encoding="utf-8").write("\n".join(play))
    print(path)


def main():
    practice()
    quiz()
    lab()
    recovery()
    pacing_and_playbook()


if __name__ == "__main__":
    main()
