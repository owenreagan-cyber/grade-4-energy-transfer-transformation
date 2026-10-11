#!/usr/bin/env python3
"""Rendered-layout visual QA v2 (Phase 5C) — reusable for any 16:9 python-pptx deck + its LibreOffice PDF.

Why this exists: the Phase 6.1 / 5B checks (qa_render_geometry.py, qa_visual_pdf.py, gate_overflow) only test
PDF words against the *page* and against *other words*, and PPTX shapes against the *slide*. They never relate a
rendered word to the shape that owns it or to the pictures/cards around it, so three visible defects passed:
  * slide 1  — a photo drawn ABOVE a text box hid "…causing change?" (text still extracted by pdftotext)
  * slide 4/8 — a caption wrapped to 2 lines inside a 0.5in box; LibreOffice drew it over the photo and
                across the card border (no word overlapped another word, nothing left the page).

This checker measures independently of the builder (it reads only the PPTX geometry/z-order and PDF word boxes):

  T-B  pdf_text_completeness   every PPTX text token appears in the rendered PDF page (multiset)
  T-C  text_containment        each rendered word's glyph core lies inside the shape that owns it
  T-D  occlusion               no picture drawn ABOVE a word's owner covers that word
  T-D2 text_over_photo         a word drawn over a photo needs an opaque backing shape (chip / scrim)
  T-D3 card_edge_straddle      a word may not straddle the edge of a filled card/panel
  T-C2 safe_margin             words stay >= 0.30in from the slide edge (slide number exempt) [WARN]
  T-E  paired_geometry         WRITE/CHECK and question/reveal pairs keep identical geometry for shared text
  T-F  font_fidelity           only the intended families are embedded (no silent substitution)
  T-G  min_rendered_size       rendered glyph height floor for student-facing text [WARN]
  T-G2 image_resolution        effective DPI of each placed photo (crop-aware) [WARN < 120, FAIL < 72]
  T-H  pixel_diff              % changed pixels vs a reference render (report only)

Usage:
  .venv/bin/python tools/qa_slide_visual_v2.py --pptx X.pptx --pdf X.pdf [--png-dir D --ref-png-dir R]
         [--pairs 11-12,15-16] [--json out.json] [--md out.md] [--label NAME]
Exit 1 if any FAIL.
"""
import argparse, json, os, re, subprocess, sys
from collections import Counter
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

EMU_PT = 12700.0
WORD_RE = re.compile(r'<word xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)" yMax="([\d.]+)">(.*?)</word>')
PAGE_RE = re.compile(r'<page width="([\d.]+)" height="([\d.]+)">(.*?)</page>', re.S)
ALLOWED_FONTS = ("Georgia", "AvenirNext", "Avenir Next")
# LucidaGrande = the macOS glyph fallback for "→" (U+2192), which neither Avenir Next nor Georgia contains.
FALLBACK_FONTS = ("LucidaGrande",)
CORE_TRIM = 0.20      # pdftotext word boxes span ascent+descent (~1.37x size); trim to the glyph core
CONTAIN_TOL = 1.5     # pt
MARGIN_PT = 0.30 * 72
MIN_GLYPH_PT = 15.0   # rendered word-box height floor (≈11pt type) for student-facing words


def _norm(t):
    t = t.replace("&amp;", "&").replace("&quot;", '"').replace("&lt;", "<").replace("&gt;", ">").replace("&apos;", "'")
    return re.sub(r"[\s“”\"'‘’.,;:!?()…]+", "", t).lower()


def pdf_pages(pdf):
    out = subprocess.run(["pdftotext", "-bbox", pdf, "-"], capture_output=True, text=True, check=True).stdout
    res = []
    for m in PAGE_RE.finditer(out):
        words = [(float(a), float(b), float(c), float(d), t) for a, b, c, d, t in WORD_RE.findall(m.group(3))]
        res.append((float(m.group(1)), float(m.group(2)), words))
    return res


def pdf_fonts(pdf):
    # NOTE: LibreOffice writes one shared font resource dict, so `pdffonts -f N -l N` lists the whole document.
    # Font fidelity is therefore a document-level test (slide 0), not a per-page one.
    out = subprocess.run(["pdffonts", pdf], capture_output=True, text=True).stdout
    names = []
    for ln in out.splitlines()[2:]:
        if ln.strip():
            names.append(ln.split()[0].split("+")[-1])
    return names


def _rect(sh):
    return (sh.left / EMU_PT, sh.top / EMU_PT, (sh.left + sh.width) / EMU_PT, (sh.top + sh.height) / EMU_PT)


def _solid(sh):
    try:
        return sh.fill.type == 1  # MSO_FILL.SOLID
    except Exception:
        return False


def slide_model(slide, W, H):
    """Flatten a slide into z-ordered records (spTree order == paint order)."""
    recs = []
    for z, sh in enumerate(slide.shapes):
        if sh.left is None or sh.width is None:
            continue
        r = _rect(sh)
        full = r[0] <= 1 and r[1] <= 1 and r[2] >= W - 1 and r[3] >= H - 1
        rec = {"z": z, "rect": r, "name": sh.name, "full": full, "text": "", "tokens": [], "pic": False,
               "filled": False, "min_size": None}
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            rec["pic"] = True
            try:
                im = sh.image
                from PIL import Image
                import io
                pw, ph = Image.open(io.BytesIO(im.blob)).size
                cw = pw * (1 - sh.crop_left - sh.crop_right)
                chh = ph * (1 - sh.crop_top - sh.crop_bottom)
                rec["dpi"] = min(cw / (sh.width / 914400.0), chh / (sh.height / 914400.0))
                rec["img_sha"] = im.sha1
            except Exception:
                rec["dpi"] = None
        else:
            rec["filled"] = _solid(sh) and not full
        if getattr(sh, "has_text_frame", False) and sh.has_text_frame and sh.text_frame.text.strip():
            rec["text"] = sh.text_frame.text
            rec["tokens"] = [_norm(t) for t in sh.text_frame.text.split() if _norm(t)]
            sizes = [r_.font.size.pt for p in sh.text_frame.paragraphs for r_ in p.runs if r_.font.size]
            rec["min_size"] = min(sizes) if sizes else None
        recs.append(rec)
    return recs


def _core(w):
    x0, y0, x1, y1 = w[:4]
    h = y1 - y0
    return (x0 + 0.5, y0 + CORE_TRIM * h, x1 - 0.5, y1 - CORE_TRIM * h)


def _inside(a, b, tol=CONTAIN_TOL):
    return a[0] >= b[0] - tol and a[1] >= b[1] - tol and a[2] <= b[2] + tol and a[3] <= b[3] + tol


def _inter(a, b):
    ix = min(a[2], b[2]) - max(a[0], b[0]); iy = min(a[3], b[3]) - max(a[1], b[1])
    return ix * iy if ix > 0 and iy > 0 else 0.0


def _dist(a, b):
    dx = max(b[0] - a[2], a[0] - b[2], 0); dy = max(b[1] - a[3], a[1] - b[3], 0)
    return (dx * dx + dy * dy) ** 0.5


def owner_of(word, recs):
    tok = _norm(word[4])
    if not tok:
        return None
    cands = [r for r in recs if r["tokens"] and any(tok == t or tok in t for t in r["tokens"])]
    if not cands:
        return None
    core = _core(word)
    inside = [r for r in cands if _inside(core, r["rect"])]
    if inside:
        return max(inside, key=lambda r: r["z"])
    return min(cands, key=lambda r: _dist(core, r["rect"]))


def check_slide(n, recs, words, W, H):
    out = []
    def add(test, status, problem, region=None, owner=None):
        out.append({"slide": n, "test": test, "status": status, "detected_problem": problem,
                    "expected_region": region, "owner": owner})

    # T-B completeness (multiset of tokens)
    want = Counter(t for r in recs for t in r["tokens"])
    got = Counter(_norm(w[4]) for w in words if _norm(w[4]))
    # pdftotext may split/join on punctuation differently; compare by concatenated characters per token set
    missing = []
    for t, c in want.items():
        if got.get(t, 0) < c:
            # tolerate tokens glued to neighbours in PDF ("change?" vs "change")
            if not any(t in g for g in got):
                missing.append(t)
    if missing:
        add("pdf_text_completeness", "FAIL", "tokens absent from rendered page: %s" % missing[:8])

    pics = [r for r in recs if r["pic"]]
    filled = [r for r in recs if r["filled"] and (r["rect"][3] - r["rect"][1]) > 0.08 * 72]
    for w in words:
        if not _norm(w[4]):
            continue
        own = owner_of(w, recs)
        core = _core(w)
        if own is None:
            add("text_containment", "FAIL", "rendered word %r has no owning text shape" % w[4])
            continue
        oname = "%s %r" % (own["name"], own["text"][:40])
        if not _inside(core, own["rect"]):
            add("text_containment", "FAIL",
                "word %r rendered outside its text frame (word core y %.0f-%.0f pt, frame y %.0f-%.0f pt)"
                % (w[4], core[1], core[3], own["rect"][1], own["rect"][3]), region=own["name"], owner=oname)
        area = max((core[2] - core[0]) * (core[3] - core[1]), 1e-6)
        for p in pics:
            cov = _inter(core, p["rect"]) / area
            if cov <= 0.10:
                continue
            if p["z"] > own["z"]:
                add("occlusion", "FAIL", "word %r is covered by picture %s drawn above it (%.0f%%)"
                    % (w[4], p["name"], 100 * cov), region=own["name"], owner=oname)
            else:
                backed = (own["filled"] and _inside(core, own["rect"])) or any(
                    f["z"] > p["z"] and f["z"] <= own["z"] and _inside(core, f["rect"]) for f in filled)
                if not backed:
                    add("text_over_photo", "FAIL", "word %r sits on photo %s with no opaque backing"
                        % (w[4], p["name"]), region=own["name"], owner=oname)
        for f in filled:
            if f is own:
                continue
            if _inter(core, f["rect"]) > 0 and not _inside(core, f["rect"], tol=0.5):
                add("card_edge_straddle", "FAIL", "word %r straddles the edge of %s" % (w[4], f["name"]),
                    region=own["name"], owner=oname)
        is_slide_no = own["text"].strip().isdigit() and own["rect"][1] > H - 0.75 * 72
        if not is_slide_no and (core[0] < MARGIN_PT or core[1] < MARGIN_PT or core[2] > W - MARGIN_PT or core[3] > H - MARGIN_PT):
            add("safe_margin", "WARN", "word %r inside the 0.30in edge band" % w[4], owner=oname)
        if (w[3] - w[1]) < MIN_GLYPH_PT and not is_slide_no and not (own["min_size"] and own["min_size"] <= 16):
            add("min_rendered_size", "WARN", "word %r renders %.1fpt tall" % (w[4], w[3] - w[1]), owner=oname)
    for p in pics:
        if p.get("dpi") is not None and p["dpi"] < 72:
            add("image_resolution", "FAIL", "%s effective %.0f dpi" % (p["name"], p["dpi"]))
        elif p.get("dpi") is not None and p["dpi"] < 120:
            add("image_resolution", "WARN", "%s effective %.0f dpi" % (p["name"], p["dpi"]))
    # de-duplicate identical messages
    seen, uniq = set(), []
    for f in out:
        k = (f["test"], f["detected_problem"])
        if k not in seen:
            seen.add(k); uniq.append(f)
    return uniq


PRIMARY_FONT_FILES = ["/System/Library/Fonts/Avenir Next.ttc", "/System/Library/Fonts/Supplemental/Georgia Bold.ttf"]
REQUIRED_EMBEDS = ("AvenirNext-Regular", "AvenirNext-Bold", "Georgia-Bold")


def uncovered_glyphs(prs):
    """Characters in the deck that the intended fonts cannot draw (these legitimately need a fallback font)."""
    from PIL import ImageFont
    chars = {c for s in prs.slides for sh in s.shapes if getattr(sh, "has_text_frame", False) and sh.has_text_frame
             for c in sh.text_frame.text if ord(c) > 126}
    miss = set()
    for path in PRIMARY_FONT_FILES:
        f = ImageFont.truetype(path, 40)
        nd = f.getmask("\uffff").getbbox()
        miss |= {c for c in chars if f.getmask(c).getbbox() == nd}
    return sorted(miss)


def font_fidelity(prs, fonts):
    out = []
    base = lambda f: f.split("+")[-1]
    names = sorted({base(f) for f in fonts})
    missing = [r for r in REQUIRED_EMBEDS if r not in names and any(r.split("-")[0] in n for n in names) is False]
    foreign = [n for n in names if not any(a.replace(" ", "") in n for a in ALLOWED_FONTS)]
    need = uncovered_glyphs(prs)
    if missing:
        out.append("required font(s) not embedded (substituted?): %s" % missing)
    if foreign and not need:
        out.append("unintended font(s) embedded: %s" % foreign)
    elif len(foreign) > 2:
        out.append("too many fallback fonts %s for uncovered glyphs %s" % (foreign, need))
    res = [{"slide": 0, "test": "font_fidelity", "status": "FAIL", "owner": None, "expected_region": None,
            "detected_problem": m + " · embedded=%s" % names} for m in out]
    if not res and foreign:
        res.append({"slide": 0, "test": "font_fidelity", "status": "WARN", "owner": None, "expected_region": None,
                    "detected_problem": "glyph fallback %s used for %s (not in Avenir Next/Georgia)" % (foreign, need)})
    return res


def paired_geometry(models, pairs):
    out = []
    for a, b in pairs:
        ta = {r["text"]: r["rect"] for r in models[a] if r["text"]}
        tb = {r["text"]: r["rect"] for r in models[b] if r["text"]}
        moved = [t for t in ta if t in tb and max(abs(x - y) for x, y in zip(ta[t], tb[t])) > 0.75]
        pa = sorted(tuple(round(v) for v in r["rect"]) for r in models[a] if r["pic"])
        pb = sorted(tuple(round(v) for v in r["rect"]) for r in models[b] if r["pic"])
        if moved:
            out.append({"slide": b, "test": "paired_geometry", "status": "FAIL", "owner": "pair %d/%d" % (a, b),
                        "expected_region": None, "detected_problem": "shared text moved between pair: %s" % [m[:30] for m in moved[:4]]})
        if pa != pb:
            out.append({"slide": b, "test": "paired_geometry", "status": "FAIL", "owner": "pair %d/%d" % (a, b),
                        "expected_region": None, "detected_problem": "picture geometry differs between pair"})
    return out


def pixel_diff(png_dir, ref_dir, n):
    from PIL import Image, ImageChops
    import numpy as np
    out = {}
    for i in range(1, n + 1):
        a = os.path.join(png_dir, "slide-%02d.png" % i); b = os.path.join(ref_dir, "slide-%02d.png" % i)
        if not (os.path.exists(a) and os.path.exists(b)):
            continue
        A = Image.open(a).convert("L"); B = Image.open(b).convert("L").resize(A.size)
        d = np.asarray(ImageChops.difference(A, B))
        out[i] = round(100.0 * float((d > 24).mean()), 1)
    return out


def run(pptx, pdf, pairs=(), png_dir=None, ref_dir=None, slides=None):
    prs = Presentation(pptx)
    pages = pdf_pages(pdf)
    W, H = prs.slide_width / EMU_PT, prs.slide_height / EMU_PT
    models, findings = {}, []
    for n, s in enumerate(prs.slides, 1):
        models[n] = slide_model(s, W, H)
    if len(pages) != len(models):
        findings.append({"slide": 0, "test": "page_count", "status": "FAIL", "owner": None, "expected_region": None,
                         "detected_problem": "PDF pages %d != slides %d" % (len(pages), len(models))})
    for n in models:
        if slides and n not in slides:
            continue
        if n > len(pages):
            break
        pw, ph, words = pages[n - 1]
        sx, sy = pw / W, ph / H     # PDF pt per PPTX pt (LibreOffice exports 13.333in as 960pt)
        words = [(x0 / sx, y0 / sy, x1 / sx, y1 / sy, t) for x0, y0, x1, y1, t in words]
        findings += check_slide(n, models[n], words, W, H)
    findings += font_fidelity(prs, pdf_fonts(pdf))
    findings += paired_geometry(models, [p for p in pairs if p[0] in models and p[1] in models])
    diff = pixel_diff(png_dir, ref_dir, len(models)) if png_dir and ref_dir else {}
    return models, findings, diff


def summarize(findings, n):
    per = {i: "PASS" for i in range(1, n + 1)}
    for f in findings:
        s = f["slide"]
        if s in per:
            if f["status"] == "FAIL":
                per[s] = "FAIL"
            elif f["status"] == "WARN" and per[s] == "PASS":
                per[s] = "WARN"
    return per


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", required=True); ap.add_argument("--pdf", required=True)
    ap.add_argument("--pairs", default=""); ap.add_argument("--png-dir"); ap.add_argument("--ref-png-dir")
    ap.add_argument("--json"); ap.add_argument("--md"); ap.add_argument("--label", default="")
    ap.add_argument("--slides", default="")
    a = ap.parse_args()
    pairs = [tuple(int(x) for x in p.split("-")) for p in a.pairs.split(",") if p]
    slides = {int(x) for x in a.slides.split(",") if x} or None
    models, findings, diff = run(a.pptx, a.pdf, pairs, a.png_dir, a.ref_png_dir, slides)
    per = summarize(findings, len(models))
    fails = [f for f in findings if f["status"] == "FAIL"]
    for f in findings:
        if f["status"] != "PASS":
            print("%-4s s%02d %-22s %s" % (f["status"], f["slide"], f["test"], f["detected_problem"]))
    cnt = Counter(per[i] for i in per if not slides or i in slides)
    print("slides PASS/WARN/FAIL = %d/%d/%d" % (cnt["PASS"], cnt["WARN"], cnt["FAIL"]))
    if a.json:
        json.dump({"label": a.label, "pptx": a.pptx, "pdf": a.pdf, "per_slide": per, "findings": findings,
                   "pixel_diff_pct": diff}, open(a.json, "w"), indent=1)
    if a.md:
        with open(a.md, "w") as f:
            f.write("# Rendered-layout visual QA v2 — %s\n\nTool: `tools/qa_slide_visual_v2.py` · overall **%s**\n\n"
                    % (a.label or os.path.basename(a.pptx), "FAIL" if fails else "PASS"))
            f.write("Slides PASS/WARN/FAIL = %d/%d/%d\n\n" % (cnt["PASS"], cnt["WARN"], cnt["FAIL"]))
            f.write("| Slide | Status | Test | Detected problem | Owner |\n|---|---|---|---|---|\n")
            for x in findings:
                f.write("| %d | %s | %s | %s | %s |\n" % (x["slide"], x["status"], x["test"],
                        x["detected_problem"].replace("|", "/"), (x["owner"] or "").replace("|", "/")))
            if diff:
                f.write("\n## Pixel diff vs reference (% of pixels changed)\n\n")
                f.write(" · ".join("%d: %s%%" % (k, v) for k, v in sorted(diff.items())) + "\n")
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
