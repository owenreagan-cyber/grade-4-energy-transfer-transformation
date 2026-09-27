"""Build Chapter 5 Crash Lab student worksheet (Part A speed + Part B mass/weight)."""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from ch4_common import *

OUT = os.path.join(os.path.dirname(__file__), '..', 'production', 'worksheets')


def _data_table(d, rows):
    """rows = list of (label, ) — build a 3-trial + average + sound table."""
    t = d.add_table(rows=len(rows) + 1, cols=6)
    t.style = 'Table Grid'
    headers = ["Condition", "Trial 1 (cm)", "Trial 2 (cm)", "Trial 3 (cm)", "Average (cm)", "Sound (loud / quiet)"]
    for i, h in enumerate(headers):
        set_cell_text(t.cell(0, i), h, size=10, bold=True, color=RGBColor(255, 255, 255))
    for r, label in enumerate(rows, start=1):
        set_cell_text(t.cell(r, 0), label, size=11, bold=True)
    style_table(t)
    return t


def build():
    d = Document()
    setup_landscape(d)
    title(d, "Crash Lab — Collisions Transfer Energy",
          "Chapter 5 · Collisions Transfer Energy · Week 6")
    name_date(d)

    section(d, "BIG QUESTION")
    prompt(d, "What happens when objects collide? How are collisions predictable?")

    section(d, "OUR MODEL")
    add_par(d, "COLLISION  →  MOTION CAN CHANGE  →  ENERGY CAN TRANSFER  →  ENERGY DOES NOT DISAPPEAR",
            size=12, bold=True, color=ACCENT, space_after=8)

    # ---------------- PART A ----------------
    section(d, "PART A — SPEED")
    prompt(d, "FAIR TEST: SAME car, SAME target, SAME surface. We change ONLY the ramp height (1 / 2 / 3 books). "
              "A higher ramp means a faster car.")
    add_par(d, "PREDICT:", size=11, bold=True, space_after=2)
    prompt(d, "At which ramp height do you think the target will move the farthest? Why?")
    answer_lines(d, 2)

    prompt(d, "TEST: roll the SAME car down each ramp height. Measure how far the target moves (cm). 3 trials each.")
    _data_table(d, ["1 book", "2 books", "3 books"])
    add_par(d, " ", size=11, space_after=4)

    prompt(d, "COMPARE: At which ramp height did the target move the farthest?")
    answer_lines(d, 1)
    prompt(d, "EXPLAIN: How did the SPEED of the car affect the change it caused?")
    answer_lines(d, 2)

    # ---------------- PART B ----------------
    section(d, "PART B — MASS / WEIGHT")
    prompt(d, "FAIR TEST: SAME ramp height, SAME target, SAME surface. We change ONLY how much the car weighs "
              "(1 car / 2 cars taped / 3 cars taped). The cars roll at about the same speed.")
    add_par(d, "PREDICT:", size=11, bold=True, space_after=2)
    prompt(d, "Which car do you think will move the target the farthest? Why?")
    answer_lines(d, 2)

    prompt(d, "TEST: roll each mass down the SAME ramp height. Measure how far the target moves (cm). 3 trials each.")
    _data_table(d, ["1 car", "2 cars", "3 cars"])
    add_par(d, " ", size=11, space_after=4)

    prompt(d, "COMPARE: Which mass moved the target the farthest?")
    answer_lines(d, 1)
    prompt(d, "EXPLAIN: How did the WEIGHT of the car affect the change it caused?")
    answer_lines(d, 2)

    # ---------------- FINAL ----------------
    section(d, "PUT IT TOGETHER")
    prompt(d, "Did the energy disappear when the car hit the target? Where did the energy go?")
    answer_lines(d, 2)

    d.save(os.path.join(OUT, "05-ch5-crash-lab-student.docx"))
    print("built crash lab")


if __name__ == "__main__":
    build()
