#!/usr/bin/env python3
"""Negative / regression fixtures for tools/qa_slide_visual_v2.py (Phase 5C). Plain-assert script, no pytest.

Part 1 — synthetic fixture deck (built here, rendered font-faithfully with tools/render_slides_isolated.py):
  slide 1  F01 class: photo drawn ABOVE a long text box          -> must FAIL occlusion
  slide 2  F02/F03 class: 2-line caption in a 0.5in box under a photo, inside a card
                                                              -> must FAIL text_containment + card_edge_straddle
  slide 3  text placed on a photo with no backing              -> must FAIL text_over_photo
  slide 4  control: same content laid out with reserved regions -> must PASS (no FAIL)
  slide 5  control: label on an opaque chip above a photo       -> must PASS (permitted overlay)
Part 2 — historical regression: the frozen Phase 5B candidate PDF must FAIL on slides 1, 4 and 8, and the
         old checker (qa_visual_pdf / qa_render_geometry) is recorded as passing them (the blind spot).

Usage: .venv/bin/python tools/qa_slide_visual_v2_negative_tests.py [WORKDIR]
"""
import os, sys, tempfile
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import qa_slide_visual_v2 as Q
import render_slides_isolated as R
import qa_render_geometry as G
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
V1 = os.path.join(ROOT, "production", "candidates", "chapter-01", "ch1-premium-hybrid-v1")


def _tb(s, x, y, w, h, txt, size=28, bold=False):
    b = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = b.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    r = tf.paragraphs[0].add_run(); r.text = txt
    r.font.size = Pt(size); r.font.name = "Avenir Next"; r.font.bold = bold
    return b


def _card(s, x, y, w, h, fill=RGBColor(255, 255, 255)):
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    c.fill.solid(); c.fill.fore_color.rgb = fill; c.line.color.rgb = RGBColor(0x1F, 0x2A, 0x5A)
    return c


def build_fixture(work):
    img = os.path.join(work, "photo.png")
    Image.new("RGB", (1200, 900), (70, 90, 120)).save(img)
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    L = prs.slide_layouts[6]
    # 1 — occlusion (F01)
    s = prs.slides.add_slide(L)
    _tb(s, 0.9, 3.25, 11.9, 0.6, "Where can we observe evidence of energy causing change?")
    s.shapes.add_picture(img, Inches(9.6), Inches(1.0), Inches(3.3), Inches(4.6))
    # 2 — caption overflow into photo + across card edge (F02/F03)
    s = prs.slides.add_slide(L)
    _card(s, 0.75, 2.5, 3.5, 2.6)
    s.shapes.add_picture(img, Inches(1.2), Inches(3.18), Inches(2.6), Inches(1.3))
    _tb(s, 0.95, 4.5, 3.1, 0.5, "light fills the room every time")
    # 3 — text on photo, no backing
    s = prs.slides.add_slide(L)
    s.shapes.add_picture(img, Inches(1), Inches(1), Inches(6), Inches(4.5))
    _tb(s, 1.5, 2.0, 4.0, 0.6, "motion")
    # 4 — control: reserved regions
    s = prs.slides.add_slide(L)
    _card(s, 0.75, 2.0, 5.0, 4.8)
    s.shapes.add_picture(img, Inches(1.0), Inches(2.25), Inches(4.5), Inches(3.0))
    _tb(s, 1.0, 5.45, 4.5, 1.1, "light fills the room every time")
    _tb(s, 6.4, 3.25, 6.0, 0.6, "Where can we observe evidence?")
    # 5 — control: label on an opaque chip over a photo (permitted overlay)
    s = prs.slides.add_slide(L)
    s.shapes.add_picture(img, Inches(1), Inches(1), Inches(6), Inches(4.5))
    chip = _card(s, 1.2, 4.6, 2.4, 0.7, fill=RGBColor(0x1F, 0x2A, 0x5A))
    tf = chip.text_frame; r = tf.paragraphs[0].add_run(); r.text = "motion"
    r.font.size = Pt(28); r.font.name = "Avenir Next"; r.font.bold = True
    out = os.path.join(work, "fixture-visual-qa.pptx")
    prs.save(out)
    return out


def main():
    work = sys.argv[1] if len(sys.argv) > 1 else tempfile.mkdtemp(prefix="vqa2_fx_")
    os.makedirs(work, exist_ok=True)
    results = []

    def expect(name, cond):
        results.append((name, cond)); print("%s  %s" % ("ok  " if cond else "FAIL", name))

    deck = build_fixture(work)
    pdf, _, _ = R.render(deck, work)
    _, f, _ = Q.run(deck, pdf)
    tests = lambda n: {x["test"] for x in f if x["slide"] == n and x["status"] == "FAIL"}
    expect("fixture s1 occlusion detected", "occlusion" in tests(1))
    expect("fixture s2 containment detected", "text_containment" in tests(2))
    expect("fixture s2 card-edge straddle detected", "card_edge_straddle" in tests(2))
    expect("fixture s3 unbacked text-over-photo detected", "text_over_photo" in tests(3))
    expect("fixture s4 control passes", not tests(4))
    expect("fixture s5 chip-over-photo control passes", not tests(5))
    old = G.check(pdf, edge_pt=5.0)
    expect("blind spot reproduced: old checker sees no overlap on fixture s1/s2",
           not old[1]["overlap"] and not old[2]["overlap"])

    v1pptx = os.path.join(V1, "01-ch1-energy-causes-change-premium-hybrid-v1.pptx")
    v1pdf = os.path.join(V1, "01-ch1-energy-causes-change-premium-hybrid-v1.pdf")
    _, f1, _ = Q.run(v1pptx, v1pdf, slides={1, 4, 8})
    fail1 = {x["slide"] for x in f1 if x["status"] == "FAIL"}
    expect("5B regression: slide 1 FAIL (occlusion)", 1 in fail1)
    expect("5B regression: slide 4 FAIL (caption)", 4 in fail1)
    expect("5B regression: slide 8 FAIL (caption)", 8 in fail1)
    old1 = G.check(v1pdf, edge_pt=5.0)
    expect("5B blind spot: old checker had no overlap on slides 1/4/8",
           not any(old1[n]["overlap"] for n in (1, 4, 8)))

    bad = [n for n, ok in results if not ok]
    print("NEGATIVE TESTS: %d/%d passed" % (len(results) - len(bad), len(results)))
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
