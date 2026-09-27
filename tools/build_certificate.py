"""Energy Detective certificate — revised wording (Phase 13B)."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from ch1_notes_common import NAVY, NAVY_HEX, CORAL, GOLD, INK, GRAY, HEAD_FONT, BODY_FONT, set_run

OUT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..",
    "production", "activities", "unit-02", "energy-detective-certificate.docx"))
os.makedirs(os.path.dirname(OUT), exist_ok=True)

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(11.0)
sec.page_height = Inches(8.5)
sec.left_margin = sec.right_margin = Inches(0.9)
sec.top_margin = sec.bottom_margin = Inches(0.7)

def p(text="", size=14, bold=False, color=INK, font=BODY_FONT, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=6, caps=False, italic=False):
    par = doc.add_paragraph()
    par.alignment = align
    par.paragraph_format.space_before = Pt(before)
    par.paragraph_format.space_after = Pt(after)
    if text:
        set_run(par.add_run(text), size=size, bold=bold, color=color, font=font, caps=caps)
    return par

p("OFFICIAL CERTIFICATE OF ACCOMPLISHMENT", size=12, bold=True, color=GRAY, after=2, caps=True)
p("Energy Detective", size=40, bold=True, color=NAVY, font=HEAD_FONT, after=4)
p("Mastery Badge", size=20, bold=True, color=CORAL, after=12)
p("This certificate is proudly awarded to", size=13, color=GRAY, after=2)
p("_______________________________________________", size=16, bold=True, color=INK, after=12)
p("For demonstrating science mastery in energy transfer, energy transformation, and the benefits and limitations of natural resources in Grade 4 Science.", size=13, italic=True, color=INK, after=14)

p("Core Scientific Competencies Mastered", size=13, bold=True, color=NAVY, after=6, caps=True)
for line in [
    "Identifying observable evidence of energy (energy → change → evidence)",
    "Using fair tests to investigate speed and energy of motion",
    "Tracing energy pathways through power systems",
    "Comparing benefits and limitations of renewable and nonrenewable resources",
]:
    p("•  " + line, size=13, color=INK, after=4)

p("", after=6)
p("Date: ______________________            Science Instructor: ______________________", size=13, color=INK, after=0)

doc.save(OUT)
print("saved", OUT)
