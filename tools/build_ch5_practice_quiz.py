"""Build Chapter 5 practice worksheet and quiz DOCX."""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from ch4_common import *

WS = os.path.join(os.path.dirname(__file__), '..', 'production', 'worksheets')
QZ = os.path.join(os.path.dirname(__file__), '..', 'production', 'quizzes')


def build_practice():
    d = Document()
    setup_portrait(d)
    title(d, "Collisions Transfer Energy — Practice",
          "Chapter 5 · Collisions Transfer Energy · Week 6")
    name_date(d)

    section(d, "A. WHAT CHANGED?")
    prompt(d, "When objects collide, what can change? Circle ALL that can change.")
    add_par(d, "MOTION (speed/direction)      SHAPE      SOUND", size=12, bold=True, space_before=2, space_after=4)
    prompt(d, "Give ONE example of each from the Reader.")
    answer_lines(d, 2)

    section(d, "B. SPEED (same car)")
    prompt(d, "The SAME car rolls slowly, then fast, into the SAME target.")
    add_par(d, "1. Which roll moves the target farther? Why?", size=11, space_before=2, space_after=2)
    answer_lines(d, 2)
    add_par(d, "2. We changed only the speed. What did we keep the same? "
              "(name the car, the target, and the surface)", size=11, space_before=2, space_after=2)
    answer_lines(d, 2)

    section(d, "C. MASS / WEIGHT (same speed)")
    prompt(d, "Two cars roll at about the SAME speed into the SAME target. One car weighs more than the other.")
    add_par(d, "1. Which car has more energy of motion? Why?", size=11, space_before=2, space_after=2)
    answer_lines(d, 2)
    add_par(d, "2. Which car moves the target farther?", size=11, space_before=2, space_after=2)
    answer_lines(d, 1)

    section(d, "D. FAIR TEST")
    prompt(d, "This plan is NOT a fair test. Fix it:")
    add_par(d, "\u201cI changed the ramp AND the car at the same time.\u201d", size=12, italic=True, space_after=2)
    prompt(d, "Your fix:")
    answer_lines(d, 2)

    section(d, "E. ENERGY DOES NOT DISAPPEAR")
    prompt(d, "A car hits a cup. The cup moves, and we hear a sound. Where did the car's energy of motion go?")
    answer_lines(d, 3)

    d.save(os.path.join(WS, "05-ch5-collisions-practice.docx"))
    print("built practice")


def build_quiz():
    d = Document()
    setup_portrait(d)
    title(d, "Chapter 5 Quiz — Collisions Transfer Energy",
          "Grade 4 · Energy Transfer and Transformation")
    name_date(d)
    add_par(d, "Circle or write the best answer.", size=11, italic=True, color=GRAY, space_after=8)

    def q(num, text):
        add_par(d, f"{num}. {text}", size=12, space_before=4, space_after=2)

    section(d, "PART 1 — Multiple Choice")
    q(1, "A collision happens when two objects ___.")
    add_par(d, "a) disappear   b) come together with impact   c) change color   d) float", size=11, space_after=6)
    q(2, "When a moving car hits a still car, what happens to some of the moving car's energy of motion?")
    add_par(d, "a) it disappears   b) it transfers to the still car   c) it turns into nothing   d) it stops existing", size=11, space_after=6)
    q(3, "The SAME car rolls faster into a target. What happens?")
    add_par(d, "a) the target moves LESS   b) the target moves MORE   c) nothing changes   d) the car disappears", size=11, space_after=6)
    q(4, "Two cars roll at about the SAME speed into the same target. One car weighs more. Which car moves the target farther?")
    add_par(d, "a) the heavier car   b) the lighter car   c) both the same   d) neither", size=11, space_after=6)
    q(5, "In a collision, some energy of motion becomes ___ energy that we hear.")
    add_par(d, "a) light   b) sound   c) stored   d) electric", size=11, space_after=6)
    q(6, "What NEVER happens to energy in a collision?")
    add_par(d, "a) it transfers   b) it changes form   c) it disappears   d) it moves to the target", size=11, space_after=6)

    section(d, "PART 2 — The Model")
    q(7, "Fill in the missing step of the collision model:")
    add_par(d, "COLLISION  →  MOTION CAN CHANGE  →  ENERGY CAN TRANSFER  →  ____________________", size=12, space_after=6)

    section(d, "PART 3 — Data")
    q(8, "A student rolls the SAME car down three ramp heights and measures target displacement.")
    t = d.add_table(rows=4, cols=2)
    t.style = 'Table Grid'
    set_cell_text(t.cell(0, 0), "Ramp height", size=11, bold=True, color=RGBColor(255, 255, 255))
    set_cell_text(t.cell(0, 1), "Average target displacement (cm)", size=11, bold=True, color=RGBColor(255, 255, 255))
    for r, (h, v) in enumerate([("1 book", "13"), ("2 books", "26"), ("3 books", "40")], start=1):
        set_cell_text(t.cell(r, 0), h, size=11)
        set_cell_text(t.cell(r, 1), v, size=11)
    style_table(t)
    prompt(d, "At which ramp height did the car move the target the farthest? What does this show about speed and energy?")
    answer_lines(d, 2)

    section(d, "PART 4 — Fair Test")
    q(9, "Which is a fair test of how SPEED affects a collision?")
    add_par(d, "a) change the car AND the ramp   b) change only the ramp height (same car)   "
              "c) change only the car (same ramp)   d) change everything", size=11, space_after=6)

    section(d, "PART 5 — Short Explanation")
    q(10, "A car hits a cup. The cup moves and we hear a sound. Write where the energy went. "
           "Use the model: \u201cI know the energy did not disappear because ___.\u201d")
    answer_lines(d, 3)

    d.save(os.path.join(QZ, "05-ch5-collisions-quiz.docx"))
    print("built quiz")


if __name__ == "__main__":
    build_practice()
    build_quiz()
    print("done")
