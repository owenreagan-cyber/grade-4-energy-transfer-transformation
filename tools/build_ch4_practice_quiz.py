"""Build Chapter 4 practice worksheet and quiz DOCX."""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from ch4_common import *

WS = os.path.join(os.path.dirname(__file__), '..', 'production', 'worksheets')
QZ = os.path.join(os.path.dirname(__file__), '..', 'production', 'quizzes')

def build_practice():
    d = Document()
    setup_portrait(d)
    title(d, "Energy Transfer — Practice",
          "Chapter 4 · Energy Transfer · Weeks 4–5")
    name_date(d)

    section(d, "A. FIND THE SOURCE AND TARGET")
    prompt(d, "For each example, write the SOURCE (where energy starts) and the TARGET (what changes).")
    t = d.add_table(rows=5, cols=3)
    t.style = 'Table Grid'
    for i, h in enumerate(["Example", "Source", "Target"]):
        set_cell_text(t.cell(0, i), h, size=11, bold=True, color=RGBColor(255,255,255))
    for r, ex in enumerate(["bell rings → paper moves", "flashlight → card lights up",
                            "warm water → spoon warms", "battery → bulb glows"], start=1):
        set_cell_text(t.cell(r, 0), ex, size=11)
    style_table(t)
    add_par(d, " ", size=11, space_after=8)

    section(d, "B. WHAT IS THE EVIDENCE?")
    prompt(d, "1. What is the evidence that energy transferred in ONE of the examples above?")
    answer_lines(d, 2)
    prompt(d, "2. Do we see energy itself, or do we see a change? Explain.")
    answer_lines(d, 2)

    section(d, "C. TRANSFER / TRANSFORMATION / BOTH")
    prompt(d, "TRANSFER = energy moves.  TRANSFORMATION = energy changes form.  BOTH = moves AND changes form.")
    add_par(d, "Sort each system. Circle one.", size=11, italic=True, space_after=4)
    t2 = d.add_table(rows=5, cols=2)
    t2.style = 'Table Grid'
    set_cell_text(t2.cell(0, 0), "System", size=11, bold=True, color=RGBColor(255,255,255))
    set_cell_text(t2.cell(0, 1), "TRANSFER / TRANSFORMATION / BOTH", size=11, bold=True, color=RGBColor(255,255,255))
    for r, ex in enumerate(["Flashlight", "Buzzer", "Warm water warms a spoon", "Motor"], start=1):
        set_cell_text(t2.cell(r, 0), ex, size=11)
        set_cell_text(t2.cell(r, 1), "TRANSFER    TRANSFORMATION    BOTH", size=11)
    style_table(t2)
    add_par(d, " ", size=11, space_after=8)

    section(d, "D. FIX THE EXPLANATION")
    prompt(d, "This statement is wrong. Fix it:")
    add_par(d, "\u201cEnergy disappeared when the buzzer made sound.\u201d", size=12, italic=True, space_after=2)
    prompt(d, "Your fix:")
    answer_lines(d, 2)

    section(d, "E. SHORT EXPLAIN")
    prompt(d, "Choose one example. How does it provide evidence that energy transferred? "
              "Name the source, the change, and why the change is evidence.")
    answer_lines(d, 3)

    d.save(os.path.join(WS, "04-ch4-energy-transfer-practice.docx"))
    print("built practice")

def build_quiz():
    d = Document()
    setup_portrait(d)
    title(d, "Chapter 4 Quiz — Energy Transfer",
          "Grade 4 · Energy Transfer and Transformation")
    name_date(d)
    add_par(d, "Circle or write the best answer.", size=11, italic=True, color=GRAY, space_after=8)

    def q(num, text):
        add_par(d, f"{num}. {text}", size=12, space_before=4, space_after=2)

    section(d, "PART 1 — Multiple Choice")
    q(1, "Energy transfer means energy ___.")
    add_par(d, "a) disappears   b) moves from place to place or object to object   c) changes form   d) stops", size=11, space_after=6)
    q(2, "What is the best evidence that energy transferred?")
    add_par(d, "a) a guess   b) an observable change in a target   c) the energy itself   d) a prediction", size=11, space_after=6)
    q(3, "A cup is tapped, and a paper target moves even though nothing touched it. This is evidence that ___ transferred energy.")
    add_par(d, "a) sound   b) heat   c) light   d) electricity", size=11, space_after=6)
    q(4, "A flashlight lights up a card across the room. This is evidence that ___ transferred energy.")
    add_par(d, "a) sound   b) light   c) heat   d) motion", size=11, space_after=6)
    q(5, "Warm water is placed next to a cool cup. The water cools and the cup warms. This is evidence that ___ transferred energy.")
    add_par(d, "a) sound   b) light   c) heat   d) electricity", size=11, space_after=6)
    q(6, "A battery is connected to a bulb by wires, and the bulb lights up. This is evidence that ___ transferred energy.")
    add_par(d, "a) electric current   b) sound   c) light   d) heat", size=11, space_after=6)

    section(d, "PART 2 — Diagram")
    q(7, "Label the energy transfer path for a flashlight shining on a card.")
    add_par(d, "Write SOURCE, PATH, and TARGET on the correct boxes.")
    t = d.add_table(rows=1, cols=3)
    t.style = 'Table Grid'
    set_cell_text(t.cell(0, 0), "[  flashlight  ]", size=11)
    set_cell_text(t.cell(0, 1), "[  -----> light ----->  ]", size=11)
    set_cell_text(t.cell(0, 2), "[  card  ]", size=11)
    style_table(t)
    add_par(d, " ", size=11, space_after=6)

    section(d, "PART 3 — Data")
    q(8, "A student places a cool metal cup in warm water and records the temperature.")
    t2 = d.add_table(rows=3, cols=3)
    t2.style = 'Table Grid'
    for i, h in enumerate(["Material", "Start (°C)", "Later (°C)"]):
        set_cell_text(t2.cell(0, i), h, size=11, bold=True, color=RGBColor(255,255,255))
    set_cell_text(t2.cell(1, 0), "Warm water", size=11)
    set_cell_text(t2.cell(1, 1), "38", size=11)
    set_cell_text(t2.cell(1, 2), "34", size=11)
    set_cell_text(t2.cell(2, 0), "Cool metal cup", size=11)
    set_cell_text(t2.cell(2, 1), "21", size=11)
    set_cell_text(t2.cell(2, 2), "27", size=11)
    style_table(t2)
    prompt(d, "Which material got warmer? Which got cooler? What does this show about energy?")
    answer_lines(d, 2)

    section(d, "PART 4 — Transfer / Transformation / BOTH")
    prompt(d, "TRANSFER = energy moves.  TRANSFORMATION = energy changes form.  BOTH = moves AND changes form.")
    q(9, "A flashlight: energy moves from the battery to the bulb AND changes form to light. This is ___.")
    add_par(d, "a) TRANSFER only   b) TRANSFORMATION only   c) BOTH", size=11, space_after=6)
    q(10, "A buzzer: energy moves through the wires AND changes form to sound. This is ___.")
    add_par(d, "a) TRANSFER only   b) TRANSFORMATION only   c) BOTH", size=11, space_after=6)

    section(d, "PART 5 — Short Explanation")
    q(11, "Choose one system (sound, light, heat, or electricity). Write a claim that energy transferred. "
           "Use the model: \u201cI know ___ because ___.\u201d Include the change you would observe as evidence.")
    answer_lines(d, 3)

    d.save(os.path.join(QZ, "04-ch4-energy-transfer-quiz.docx"))
    print("built quiz")

if __name__ == "__main__":
    build_practice()
    build_quiz()
    print("done")
