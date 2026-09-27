"""Build the premium Chapter 1 Interactive Notebook (INB) comparison prototype.

3 inserts, each a 7.0" x 9.0" dashed-trim container for composition-notebook insertion.
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from ch1_notes_common import *

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..",
                   "production", "comparisons", "chapter-01-notes", "inb",
                   "01-ch1-inb-comparison-master.docx")
OUT = os.path.abspath(OUT)

INSERT_W = 10080   # 7.0"
INSERT_H = 12960   # 9.0"
INNER = 10080 - 2 * 120   # usable inside dashed cell (cell margins 120 dxa each side)

doc = Document()
set_page(doc, left=0.75, right=0.75, top=0.9, bottom=0.9)
# trim default empty paragraph(s)
for _p in list(doc.paragraphs):
    if not _p.text.strip():
        _p._p.getparent().remove(_p._p)


def dashed_container():
    """Return (outer_table, cell) — a 7x9 dashed single-cell container."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = 1  # center
    set_table_borders(t, val='dashed', sz=10, color=DASH_HEX)
    cell = t.cell(0, 0)
    set_cell_width(cell, INSERT_W)
    set_cell_margins(cell, top=0, left=120, bottom=60, right=120)
    set_row_height(t.rows[0], INSERT_H, rule='atLeast')
    return t, cell


def glue_strip(cell):
    g = cell.add_table(rows=1, cols=1)
    no_table_borders(g)
    gc = g.cell(0, 0)
    shade_cell(gc, GLUE_HEX)
    set_cell_width(gc, INNER)
    set_cell_margins(gc, top=20, left=80, bottom=20, right=80)
    set_row_height(g.rows[0], 620, rule='exact')
    p = gc.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("GLUE — TOP EDGE   "), size=10, bold=True, color=GRAY, caps=True)
    set_run(p.add_run("(this strip tucks into the notebook spine)"), size=10, italic=True, color=GRAY)
    # dashed fold line under the glue strip
    d = cell.add_paragraph()
    dPr = d._p.get_or_add_pPr()
    sp = OxmlElement('w:spacing'); sp.set(qn('w:after'), '20'); dPr.append(sp)
    add_paragraph_bottom_border(d, sz=10, val='dashed', color=DASH_HEX)
    return g


def title_block(cell, title, subtitle):
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(title), size=17, bold=True, color=NAVY, font=HEAD_FONT, caps=True)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p2.add_run(subtitle), size=10, italic=True, color=GRAY)


def section(cell, text, size=13):
    p = cell.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    sp = OxmlElement('w:spacing'); sp.set(qn('w:before'), '40'); sp.set(qn('w:after'), '10'); pPr.append(sp)
    set_run(p.add_run(text), size=size, bold=True, color=NAVY, caps=True)
    return p


def blank_in(p, text_before, blank="______________", text_after="", size=12, bold_term=None):
    if bold_term:
        set_run(p.add_run(bold_term + "  "), size=size, bold=True, color=INK)
    set_run(p.add_run(text_before + " "), size=size, color=INK)
    r = p.add_run(" " + blank + " ")
    r.underline = True
    set_run(r, size=size, color=INK)
    if text_after:
        set_run(p.add_run(" " + text_after), size=size, color=INK)


# ============================================================
# INSERT 1 — Energy Causes Change
# ============================================================
t, c = dashed_container()
glue_strip(c)
title_block(c, "Energy Causes Change", "Chapter 1 · Grade 4 · Insert 1 of 3")
section(c, "Big Question")
q = c.add_paragraph()
set_run(q.add_run("Where can we observe evidence of energy causing change?"), size=12, italic=True, bold=True, color=INK)

section(c, "Vocabulary + Draw")
vocab = c.add_table(rows=3, cols=1)
no_table_borders(vocab)
vterms = [
    ("energy", "Energy is the ability to cause", "."),
    ("change", "Change means something becomes", "."),
    ("evidence", "Evidence is something we observe that helps us", "what happened."),
]
for i, (term, before, after) in enumerate(vterms):
    vc = vocab.cell(i, 0)
    set_cell_width(vc, INNER)
    set_cell_margins(vc, top=30, left=20, bottom=30, right=20)
    p = vc.paragraphs[0]
    blank_in(p, before, "________________", after, size=12, bold_term=term + " —")
    add_paragraph_bottom_border(p, sz=8, val='single', color=INK_HEX)

section(c, "The Model")
m = c.add_table(rows=1, cols=5)
no_table_borders(m)
box_w = 2720
arr_w = 720
words = ["ENERGY", "CHANGE", "EVIDENCE"]
for idx in range(5):
    mcell = m.cell(0, idx)
    w = box_w if idx % 2 == 0 else arr_w
    set_cell_width(mcell, w)
    set_cell_margins(mcell, top=40, left=10, bottom=40, right=10)
    vcenter(mcell)
    mp = mcell.paragraphs[0]
    mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if idx % 2 == 0:
        shade_cell(mcell, NAVY_HEX)
        set_run(mp.add_run(words[idx // 2]), size=13, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    else:
        set_run(mp.add_run("→"), size=15, bold=True, color=CORAL)

section(c, "One Example — Before & After")
ba = c.add_table(rows=1, cols=2)
no_table_borders(ba)
for idx, lab in enumerate(["BEFORE", "AFTER"]):
    bc = ba.cell(0, idx)
    set_cell_width(bc, INNER // 2 - 40)
    set_cell_margins(bc, top=30, left=60, bottom=10, right=60)
    bp = bc.paragraphs[0]
    bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(bp.add_run(lab), size=11, bold=True, color=CORAL, caps=True)
    set_row_height(ba.rows[0], 1550, rule='exact')
    # thin border around the two boxes
    set_cell_border(bc, top=dict(sz=8, val='single', color=NAVY_HEX),
                    left=dict(sz=8, val='single', color=NAVY_HEX),
                    bottom=dict(sz=8, val='single', color=NAVY_HEX),
                    right=dict(sz=8, val='single', color=NAVY_HEX))
wp = c.add_paragraph()
wpPr = wp._p.get_or_add_pPr()
sp = OxmlElement('w:spacing'); sp.set(qn('w:before'), '30'); sp.set(qn('w:after'), '10'); wpPr.append(sp)
set_run(wp.add_run("What changed?  "), size=11, bold=True, color=INK)
r = wp.add_run("________________________")
r.underline = True
set_run(r, size=11, color=INK)

section(c, "Forms of Energy — Memory Cues")
f = c.add_table(rows=1, cols=5)
no_table_borders(f)
cues = [("MOTION", "rolling"), ("LIGHT", "bright"), ("SOUND", "loud"), ("HEAT", "warm"), ("ELECTRICAL", "powers")]
for idx, (form, cue) in enumerate(cues):
    fc = f.cell(0, idx)
    set_cell_width(fc, INNER // 5)
    set_cell_margins(fc, top=30, left=20, bottom=30, right=20)
    shade_cell(fc, LIGHT_HEX)
    fp = fc.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(fp.add_run(form), size=11, bold=True, color=NAVY)
    fp2 = fc.add_paragraph()
    fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(fp2.add_run(cue), size=10, italic=True, color=GRAY)

# ============================================================
# INSERT 2 — Stored Energy — Ready for Change
# ============================================================
t2, c2 = dashed_container()
glue_strip(c2)
title_block(c2, "Stored Energy — Ready for Change", "Chapter 1 · Grade 4 · Insert 2 of 3")
section(c2, "Key Idea")
k = c2.add_paragraph()
set_run(k.add_run("Stored energy can cause a change later."), size=13, bold=True, color=INK)

section(c2, "Three Examples — What Change Could Happen Later?")
ex = c2.add_table(rows=3, cols=2)
no_table_borders(ex)
examples = [
    ("battery", "in a flashlight", "the light turns on"),
    ("stretched rubber band", "held still", "it snaps back and moves"),
    ("raised object", "held up high", "it can fall down"),
]
for i, (obj, state, later) in enumerate(examples):
    ec = ex.cell(i, 0)
    set_cell_width(ec, 5000)
    set_cell_margins(ec, top=60, left=80, bottom=60, right=80)
    ep = ec.paragraphs[0]
    set_run(ep.add_run(obj), size=13, bold=True, color=NAVY)
    ep2 = ec.add_paragraph()
    set_run(ep2.add_run(state), size=10, italic=True, color=GRAY)
    # lift-the-flap reveal cell
    rc = ex.cell(i, 1)
    set_cell_width(rc, INNER - 5000)
    set_cell_margins(rc, top=40, left=80, bottom=40, right=80)
    shade_cell(rc, LIGHT_HEX)
    rp = rc.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(rp.add_run("LIFT THE FLAP"), size=10, bold=True, color=CORAL, caps=True)
    # dashed fold line at top of flap
    add_paragraph_bottom_border(rp, sz=10, val='dashed', color=DASH_HEX)
    rp2 = rc.add_paragraph()
    rp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(rp2.add_run(later), size=11, bold=True, color=INK)
    set_row_height(ex.rows[i], 900, rule='atLeast')

note = c2.add_paragraph()
notePr = note._p.get_or_add_pPr()
sp = OxmlElement('w:spacing'); sp.set(qn('w:before'), '40'); sp.set(qn('w:after'), '10'); notePr.append(sp)
set_run(note.add_run("Flat option (no flap): "), size=11, bold=True, color=GRAY)
set_run(note.add_run("skip folding — write your prediction on the line under each example."), size=11, italic=True, color=GRAY)

section(c2, "Draw One Example")
d = c2.add_table(rows=1, cols=1)
no_table_borders(d)
dc = d.cell(0, 0)
set_cell_width(dc, INNER)
set_cell_margins(dc, top=20, left=60, bottom=20, right=60)
set_cell_border(dc, top=dict(sz=8, val='single', color=NAVY_HEX),
                left=dict(sz=8, val='single', color=NAVY_HEX),
                bottom=dict(sz=8, val='single', color=NAVY_HEX),
                right=dict(sz=8, val='single', color=NAVY_HEX))
set_row_height(d.rows[0], 1400, rule='exact')
dp = dc.paragraphs[0]
dp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run(dp.add_run("draw your stored-energy example here"), size=10, italic=True, color=GRAY)

# ============================================================
# INSERT 3 — Energy Detectives — Evidence
# ============================================================
t3, c3 = dashed_container()
glue_strip(c3)
title_block(c3, "Energy Detectives — Evidence", "Chapter 1 · Grade 4 · Insert 3 of 3")
section(c3, "Observation vs Explanation")
ov = c3.add_paragraph()
set_run(ov.add_run("Observation = "), size=11, bold=True, color=NAVY)
set_run(ov.add_run("what I see.    "), size=11, color=INK)
set_run(ov.add_run("Explanation = "), size=11, bold=True, color=NAVY)
set_run(ov.add_run("what it means."), size=11, color=INK)

section(c3, "Scavenger Hunt Record")
h = c3.add_table(rows=5, cols=1)
no_table_borders(h)
hrows = [
    ("WHAT I OBSERVED BEFORE", 1),
    ("WHAT HAPPENED", 1),
    ("WHAT CHANGED", 1),
    ("WHAT IS MY EVIDENCE?", 1),
    ("HOW DOES THIS SUPPORT THE IDEA THAT ENERGY WAS INVOLVED?", 2),
]
for i, (label, lines) in enumerate(hrows):
    hc = h.cell(i, 0)
    set_cell_width(hc, INNER)
    set_cell_margins(hc, top=40, left=20, bottom=20, right=20)
    hp = hc.paragraphs[0]
    set_run(hp.add_run(label), size=10, bold=True, color=NAVY, caps=True)
    for _ in range(lines):
        lp = hc.add_paragraph()
        lpPr = lp._p.get_or_add_pPr()
        sp = OxmlElement('w:spacing'); sp.set(qn('w:after'), '30'); sp.set(qn('w:line'), '360'); lpPr.append(sp)
        add_paragraph_bottom_border(lp, sz=8, val='single', color=INK_HEX)
    set_row_height(h.rows[i], 700 + lines * 360, rule='exact')

section(c3, "Remember")
rm = c3.add_table(rows=1, cols=3)
no_table_borders(rm)
for idx, w in enumerate(["ENERGY", "CHANGE", "EVIDENCE"]):
    rc = rm.cell(0, idx)
    set_cell_width(rc, INNER // 3)
    set_cell_margins(rc, top=30, left=10, bottom=30, right=10)
    shade_cell(rc, NAVY_HEX)
    rp = rc.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(rp.add_run(w), size=11, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

doc.save(OUT)
print("saved", OUT)
