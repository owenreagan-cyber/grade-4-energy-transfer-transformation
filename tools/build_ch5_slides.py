"""Build Chapter 5 Collisions Transfer Energy slide deck (PPTX, 16:9)."""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = os.path.join(os.path.dirname(__file__), '..', 'production', 'slides')

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x2E, 0x6E, 0x8E)
AMBER = RGBColor(0xD9, 0xA4, 0x41)
LIGHT = RGBColor(0xF5, 0xF3, 0xEC)
INK = RGBColor(0x1A, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x55, 0x55, 0x55)
PALE = RGBColor(0xD5, 0xDE, 0xEA)

CUE_COLORS = {
    "NOTEBOOK": RGBColor(0x6B, 0x4F, 0xA0),
    "WRITE": RGBColor(0x2E, 0x6E, 0x8E),
    "DRAW": RGBColor(0x3E, 0x8E, 0x5A),
    "LABEL": RGBColor(0xD9, 0xA4, 0x41),
    "REVEAL": RGBColor(0xC0, 0x4A, 0x4A),
}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def textbox(slide, x, y, w, h, text, size=24, color=INK, bold=False, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = color; r.font.name = "Arial"
    return tb


def cue_badge(slide, cue):
    if not cue:
        return
    c = CUE_COLORS.get(cue, GRAY)
    sp = rect(slide, 12.25, 0.25, 0.9, 0.5, c, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = sp.text_frame; tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = cue
    r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = "Arial"


def model_strip(slide, items, y=5.6, x=0.8, w=11.7, box_h=0.7, gap=0.3):
    n = len(items)
    box_w = (w - (n - 1) * gap) / n
    cx = x
    for i, it in enumerate(items):
        sp = rect(slide, cx, y, box_w, box_h, NAVY if i == 0 else (AMBER if i == n - 1 else TEAL),
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        tf = sp.text_frame; tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = it
        r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = "Arial"
        if i < n - 1:
            ar = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(cx + box_w + 0.02),
                                        Inches(y + box_h / 2 - 0.12), Inches(gap - 0.04), Inches(0.24))
            ar.fill.solid(); ar.fill.fore_color.rgb = GRAY; ar.line.fill.background(); ar.shadow.inherit = False
        cx += box_w + gap


def header(slide, title, subtitle, cue=None, num=None):
    rect(slide, 0, 0, 13.333, 1.0, NAVY)
    textbox(slide, 0.5, 0.12, 11.6, 0.75, title, size=30, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        textbox(slide, 0.5, 0.86, 11.6, 0.3, subtitle, size=13, color=PALE)
    if num:
        textbox(slide, 11.4, 0.32, 0.7, 0.4, str(num), size=14, color=PALE, align=PP_ALIGN.RIGHT)
    cue_badge(slide, cue)


def body(slide, lines, x=0.8, y=1.5, w=11.7, h=3.9, size=24, gap=8):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.color.rgb = INK; r.font.name = "Arial"
    return tb


def concept(title, subtitle, lines, cue=None, num=None, strip=None, size=24):
    s = add_slide()
    rect(s, 0, 0, 13.333, 7.5, LIGHT)
    header(s, title, subtitle, cue, num)
    body(s, lines, size=size)
    if strip:
        model_strip(s, strip)
    return s


MODEL = ["COLLISION", "MOTION CAN CHANGE", "ENERGY CAN TRANSFER", "ENERGY DOES NOT DISAPPEAR"]

# ---------------- Title + Big Question ----------------
s = add_slide(); rect(s, 0, 0, 13.333, 7.5, NAVY)
textbox(s, 0.8, 2.0, 11.7, 1.2, "Collisions Transfer Energy", size=46, color=WHITE, bold=True)
textbox(s, 0.8, 3.2, 11.7, 0.6, "Chapter 5 · Grade 4 · Week 6", size=20, color=PALE)
model_strip(s, MODEL, y=5.0)

s = add_slide(); rect(s, 0, 0, 13.333, 7.5, LIGHT)
rect(s, 0, 2.6, 13.333, 2.2, NAVY)
textbox(s, 0.8, 2.75, 11.7, 1.9, "How is energy involved in collisions?",
        size=34, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
textbox(s, 0.8, 2.0, 11.7, 0.5, "Big Question", size=16, color=GRAY, align=PP_ALIGN.CENTER)
textbox(s, 11.4, 0.32, 0.7, 0.4, "2", size=14, color=GRAY, align=PP_ALIGN.RIGHT)

# ---------------- Opening phenomenon ----------------
concept("Clap. Stomp. Tap.", "Opening phenomenon",
        ["Clap your hands. Stomp your foot. Tap a pencil.", "Every time, something changed."], num=3)
concept("I notice… I wonder…", "Opening phenomenon",
        ["I notice: a sound, a feeling, maybe a shape changed.",
         "I wonder: how is energy involved?"], num=4)

# ---------------- Concept ----------------
concept("What is a collision?", "Definition",
        ["A collision happens when two objects come together with impact.",
         "collide = to come together with impact",
         "collision = an instance of colliding"],
        cue="WRITE", num=5, strip=["TWO OBJECTS", "COME TOGETHER", "WITH IMPACT"])
concept("Notebook — collision vocab", "Insert 1",
        ["Fill in the blanks: collide, collision, sound.",
         "Draw one collision and label what changed."],
        cue="NOTEBOOK", num=6)

# ---------------- Reader ----------------
concept("Changes in motion", "Reader pp. 25–26",
        ["A bowling ball hits the pins.", "Some energy of motion transfers to the pins.",
         "The pins move — the ball keeps less."],
        num=7, strip=["BALL MOVES", "TRANSFERS", "PINS MOVE"])
concept("Changes in shape", "Reader p. 27",
        ["Clay flattens on a wall. An egg cracks.", "A collision can change an object's shape."], num=8)
concept("Changes we hear — sound", "Reader p. 28",
        ["Collisions often make a sound.", "Some energy of motion becomes sound energy.",
         "sound = a form of energy from a vibrating object"],
        num=9, strip=["MOTION", "→", "SOUND"])
concept("Speed affects collisions", "Reader p. 29",
        ["The SAME object moving faster has greater energy of motion.",
         "Faster → transfers more energy → greater change."],
        num=10, strip=["SAME OBJECT", "FASTER", "GREATER CHANGE"])
concept("Weight affects collisions", "Reader p. 30",
        ["At about the SAME speed, the heavier object has more energy of motion.",
         "Heavier → greater change."],
        num=11, strip=["SAME SPEED", "HEAVIER", "GREATER CHANGE"])

# ---------------- AP 10.1 ----------------
concept("Check — Lesson 10", "AP 10.1",
        ["Look at the picture (foot + ball).", "What objects collide?",
         "What energy is there before? After?"],
        cue="WRITE", num=12)

# ---------------- Fair test ----------------
concept("Fair test — change ONE thing", "Investigating collisions",
        ["Part A: change ONLY ramp height (same car, same target, same surface).",
         "Part B: change ONLY mass/weight (same ramp, same target, same surface).",
         "We measure how far the TARGET moves."],
        num=13, strip=["CHANGE ONE THING", "KEEP THE REST", "MEASURE THE CHANGE"])
concept("Notebook — model + fair test", "Insert 2",
        ["Copy the model:", "COLLISION → MOTION CAN CHANGE → ENERGY CAN TRANSFER → ENERGY DOES NOT DISAPPEAR"],
        cue="NOTEBOOK", num=14, strip=MODEL)

# ---------------- Crash Lab Part A (Speed) ----------------
concept("Crash Lab Part A — SPEED", "Predict",
        ["Same car, three ramp heights.", "Which height will move the target farthest?", "Why?"], num=15)
concept("Part A setup", "Fair test",
        ["SAME car, SAME target, SAME surface.",
         "Change ONLY ramp height: 1 / 2 / 3 books.", "Higher ramp → faster car (inferred)."],
        cue="LABEL", num=16, strip=["SAME CAR", "VARY RAMP", "MEASURE TARGET"])
concept("Part A test", "Test + record",
        ["Roll the car from each ramp height.", "Measure target displacement (cm).",
         "3 trials each. Listen for the sound (loud/quiet)."],
        cue="WRITE", num=17)
concept("Part A compare + explain", "Compare",
        ["Highest ramp → target moved farthest.", "Faster car → more energy of motion → greater change."],
        cue="REVEAL", num=18, strip=["HIGHER RAMP", "FASTER CAR", "MORE CHANGE"])

# ---------------- Crash Lab Part B (Mass/Weight) ----------------
concept("Crash Lab Part B — MASS / WEIGHT", "Predict",
        ["Same ramp, three masses.", "Which mass will move the target farthest?", "Why?"], num=19)
concept("Part B setup", "Fair test",
        ["SAME ramp height, SAME target, SAME surface.",
         "Change ONLY mass: 1 / 2 / 3 cars taped.", "Cars roll at about the same speed."],
        cue="LABEL", num=20, strip=["SAME RAMP", "VARY MASS", "MEASURE TARGET"])
concept("Part B test", "Test + record",
        ["Roll each mass down the SAME ramp height.", "Measure target displacement (cm).",
         "3 trials each. Listen for the sound (loud/quiet)."],
        cue="WRITE", num=21)
concept("Part B compare + explain", "Compare",
        ["Heaviest car → target moved farthest.", "Same speed, heavier car → more energy of motion."],
        cue="REVEAL", num=22, strip=["SAME SPEED", "HEAVIER", "MORE CHANGE"])

# ---------------- Synthesis ----------------
concept("Notebook — Part A + B evidence", "Insert 2",
        ["Write your evidence: the target moved farthest when ____.",
         "Circle the ONE thing you changed in each part."],
        cue="NOTEBOOK", num=23)
concept("Energy does NOT disappear", "Synthesis",
        ["A car hits a cup. The cup moves. We hear a sound.",
         "Energy moved to the cup + became sound (+ a tiny bit of heat).",
         "Energy did NOT disappear."],
        cue="REVEAL", num=24, strip=MODEL)
concept("Notebook — lift the flap", "Insert 3",
        ["Q: Did the energy disappear?", "Lift the flap → A: No — it transferred."],
        cue="NOTEBOOK", num=25)
concept("Misconception check", "Synthesis",
        ["\u201cDid the energy vanish when the car stopped?\u201d",
         "NO — it transferred to the target and became sound."],
        cue="REVEAL", num=26)
concept("Predictable collisions", "Synthesis",
        ["PREDICT → TEST → OBSERVE → COMPARE → EXPLAIN",
         "Same setup → same kind of outcome."],
        num=27, strip=["PREDICT", "TEST", "OBSERVE", "COMPARE", "EXPLAIN"])
concept("Review / exit", "Chapter 5",
        ["Name ONE change a collision causes and where the energy went."], num=28)

prs.save(os.path.join(OUT, "05-ch5-collisions-slides.pptx"))
print("built slides:", len(prs.slides._sldIdLst))
