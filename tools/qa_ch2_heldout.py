"""Held-out attacks. Text changes are caught by the golden baseline.

Status, date, export, and raster checks use their own tools.
Temp copies only. The working tree is not modified.
"""
import os, shutil, subprocess, sys, tempfile

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
PY = sys.executable
TOOLS = os.path.join(ROOT, "tools")
NOTES = "production/premium-notes/02-ch2-moving-objects-have-energy-guided-notes.docx"
SLIDES = "production/slides/02-ch2-moving-objects-have-energy-slides.pptx"
QUIZ = "production/quizzes/02-ch2-moving-objects-have-energy-quiz.docx"
PRACTICE = "production/worksheets/02-ch2-moving-objects-have-energy-practice.docx"
MARBLES = "production/worksheets/02-ch2-moving-marbles-student.docx"
FLEX = "production/worksheets/02-ch2-friday-flex.docx"
RECOVERY = "production/worksheets/02-ch2-absent-student-recovery.docx"
PLAYBOOK = "production/teacher-playbook/unit1-teacher-playbook.docx"
KEY = "production/answer-keys/02-ch2-moving-objects-have-energy-quiz-key.md"
CAL = "docs/playbook/UNIT1_PACING_CALENDAR.md"


def fresh():
    tmp = tempfile.mkdtemp(prefix="ch2hold_")
    for rel in (NOTES, SLIDES, QUIZ, PRACTICE, MARBLES, FLEX, RECOVERY, PLAYBOOK, KEY, CAL,
                "production/final-8.5-week/MASTER_CURRICULUM_MAP.md",
                "production/final-8.5-week/MASTER_RESOURCE_INDEX.md",
                "production/final-8.5-week/WEEK1_2_TEACHER_QUICKSTART.md",
                "production/final-package/ENERGY_UNIT_WEEK_BY_WEEK_TEACHER_MAP.md",
                "production/final-package/ENERGY_UNIT_MEDIA_PLAYBACK_GUIDE.md",
                "production/final-8.5-week/VIDEO_PRESENTATION_INTEGRATION_MATRIX.md",
                "qa/baselines/ch2_golden_text.json",
                "production/premium-notes/02-ch2-moving-objects-have-energy-guided-notes.pdf",
                "production/premium-notes/02-ch2-moving-objects-have-energy-guided-notes-key.pdf",
                "production/slides/02-ch2-moving-objects-have-energy-slides.pdf",
                "production/answer-keys/02-ch2-moving-objects-have-energy-practice-key.md",
                "production/answer-keys/02-ch2-moving-marbles-key.md",
                "production/answer-keys/02-ch2-friday-flex-key.md",
                "production/premium-notes/02-ch2-moving-objects-have-energy-guided-notes-key.docx"):
        dest = os.path.join(tmp, rel)
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        shutil.copy2(os.path.join(ROOT, rel), dest)
    return tmp


def run(cmd):
    return subprocess.run(cmd, cwd=ROOT, capture_output=True, text=True).returncode


def golden(tmp):
    return run([PY, os.path.join(TOOLS, "qa_ch2_golden.py"), "--root", tmp]) != 0


def add_para(path, sentence):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph
    d = Document(path)
    host = d.paragraphs[6]
    el = OxmlElement("w:p")
    host._p.addnext(el)
    Paragraph(el, host._parent).add_run(sentence)
    d.save(path)


def slide_box(path, number, sentence):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation(path)
    box = prs.slides[number - 1].shapes.add_textbox(Inches(0.5), Inches(6.3), Inches(8), Inches(0.35))
    box.text_frame.text = sentence
    prs.save(path)


def speaker(path, number, old, new):
    from pptx import Presentation
    prs = Presentation(path)
    box = prs.slides[number - 1].notes_slide.notes_text_frame
    if old not in box.text:
        raise SystemExit("missing speaker on %d: %s" % (number, old))
    box.text = box.text.replace(old, new, 1)
    prs.save(path)


def docx_sub(path, old, new):
    from docx import Document
    d = Document(path)
    hit = False
    blobs = list(d.paragraphs)
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                blobs.extend(cell.paragraphs)
    for p in blobs:
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, new, 1)
                hit = True
                break
        if hit:
            break
    if not hit:
        raise SystemExit("docx text not found: " + old)
    d.save(path)


def sub_file(path, old, new):
    text = open(path, encoding="utf-8").read()
    if old not in text:
        raise SystemExit("missing text %s" % old[:50])
    open(path, "w", encoding="utf-8").write(text.replace(old, new, 1))


def main():
    green = []
    red = []

    def bad(name, mutate, check):
        tmp = fresh()
        try:
            mutate(tmp)
            if check(tmp):
                print("[CAUGHT]", name)
            else:
                print("[GREEN] ", name)
                green.append(name)
        except Exception as exc:
            print("[ERROR] ", name, exc)
            green.append(name)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def doc(tmp, rel):
        return os.path.join(tmp, rel)

    bad("X1", lambda t: add_para(doc(t, NOTES), "The classroom goldfish caused the change."), golden)
    bad("X2", lambda t: slide_box(doc(t, SLIDES), 11, "Stored energy is the real answer here."), golden)
    bad("X4", lambda t: add_para(doc(t, PRACTICE), "Count the speed and call it momentum."), golden)
    bad("X5", lambda t: add_para(doc(t, FLEX), "This card is worth five bonus points."), golden)
    bad("W3", lambda t: add_para(doc(t, MARBLES), "Measure the distance in centimeters."), golden)
    bad("B1p", lambda t: slide_box(doc(t, SLIDES), 11, "Moving objects cause changes."), golden)
    bad("B1q", lambda t: slide_box(doc(t, SLIDES), 12, "Any moving object causes a change."), golden)
    bad("B3", lambda t: slide_box(doc(t, SLIDES), 17, "All moving objects have energy of motion."), golden)
    bad("B4", lambda t: slide_box(doc(t, SLIDES), 22, "A swinging hammer drives a nail into the wood."), golden)
    bad("S2", lambda t: speaker(doc(t, SLIDES), 18, "ASK Fill in the blank.", "ASK The blank is motion."), golden)
    bad("S3", lambda t: speaker(doc(t, SLIDES), 22, "ASK Write or draw the change each moving object causes.",
                                "SAY The answers are nail, gears, and blades."), golden)
    bad("L5", lambda t: speaker(doc(t, SLIDES), 17, "NEXT Record N2 on notes.", "NEXT Tell them the answer is motion."), golden)
    bad("L6", lambda t: speaker(doc(t, SLIDES), 21, "NEXT Record N3 table on notes.", "NEXT The nail is the change."), golden)
    bad("O1", lambda t: speaker(doc(t, SLIDES), 12, "REQUIRED", "as time permits"), golden)
    bad("O2", lambda t: speaker(doc(t, SLIDES), 12, "REQUIRED", "nice-to-have"), golden)
    bad("O3", lambda t: speaker(doc(t, SLIDES), 12, "REQUIRED", "drop if running behind"), golden)
    bad("O4", lambda t: speaker(doc(t, SLIDES), 12, "REQUIRED", "if we have time"), golden)
    bad("FR1", lambda t: open(doc(t, CAL), "a").write("\n### Friday\n- Chapter 2 Session 1 required slides.\n"),
        lambda t: run([PY, os.path.join(TOOLS, "qa_ch2_schedule.py"), "--root", t]) != 0)
    bad("FR2", lambda t: open(doc(t, CAL), "a").write("\nOct 16 Chapter 2 Session 1 required\n"),
        lambda t: run([PY, os.path.join(TOOLS, "qa_ch2_schedule.py"), "--root", t]) != 0)
    bad("FR3", lambda t: open(doc(t, CAL), "a").write("\n2026-10-16 Moving Objects Have Energy Session 2 required\n"),
        lambda t: run([PY, os.path.join(TOOLS, "qa_ch2_schedule.py"), "--root", t]) != 0)
    bad("FR4", lambda t: docx_sub(doc(t, PLAYBOOK), "Thu 10/15", "Fri 10/16"), golden)
    bad("FR6", lambda t: open(doc(t, CAL), "a").write("\n10/16 Ch2 Session 1 required\n"),
        lambda t: run([PY, os.path.join(TOOLS, "qa_ch2_schedule.py"), "--root", t]) != 0)
    bad("FR7", lambda t: open(doc(t, CAL), "a").write("\n| Friday | Chapter 2 Session 2 required |\n"),
        lambda t: run([PY, os.path.join(TOOLS, "qa_ch2_schedule.py"), "--root", t]) != 0)
    bad("GOLD-1", lambda t: add_para(doc(t, NOTES), "An unrelated sentence was inserted here."), golden)
    bad("GOLD-2", lambda t: slide_box(doc(t, SLIDES), 25, "Ignore the evidence and guess."), golden)
    bad("GOLD-3", lambda t: add_para(doc(t, QUIZ), "Question 9. What is kinetic energy?"), golden)
    bad("GOLD-4", lambda t: sub_file(doc(t, KEY), "change", "speed"), golden)
    bad("GOLD-5", lambda t: add_para(doc(t, RECOVERY), "Skip the notes and invent the data."), golden)
    bad("GOLD-6", lambda t: docx_sub(doc(t, PLAYBOOK), "Wed 10/14", "Fri 10/16"), golden)
    bad("GOLD-7", lambda t: speaker(doc(t, SLIDES), 3, "ASK", "ASK something else"), golden)
    bad("GOLD-8", lambda t: docx_sub(doc(t, NOTES), "My first thinking", ""), golden)

    # Clean control: negation is not a Friday assignment.
    import qa_ch2_schedule as sched
    if sched.hits("Chapter 2 never moves to Friday."):
        print("[RED]   C1")
        red.append("C1")
    else:
        print("[CLEAN] C1")
    if run([PY, os.path.join(TOOLS, "qa_ch2_schedule.py")]) != 0:
        print("[RED]   schedule-good")
        red.append("schedule-good")
    else:
        print("[CLEAN] schedule-good")

    # Post-freeze cannot be satisfied from the working tree.
    code = run([PY, os.path.join(TOOLS, "qa_ch2_golden.py"), "--mode", "post-freeze"])
    if code == 0:
        print("[GREEN] MA1 post-freeze armed itself")
        green.append("MA1")
    else:
        print("[CAUGHT] MA1 post-freeze refuses an uncommitted baseline")
    print("BAD GREEN", len(green), "CLEAN RED", len(red))
    if green or red:
        sys.exit(1)


if __name__ == "__main__":
    main()
