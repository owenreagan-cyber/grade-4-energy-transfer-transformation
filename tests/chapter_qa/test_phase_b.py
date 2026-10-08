"""Phase B auditor, render, visual, PDF, and OOXML evidence tests."""

from __future__ import annotations

import io
import shutil
import sys
import unittest
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "tools"))
sys.path.insert(0, str(ROOT / "tests" / "chapter_qa"))

from phase_b_support import ROOT as REPO  # noqa: E402
from phase_b_support import build_identity, cleanup_work, evidence_file, write_pdf  # noqa: E402

from chapter_qa.cross_artifact import emit_cross_evidence  # noqa: E402
from chapter_qa.emit import seal  # noqa: E402
from chapter_qa.render_completeness import assess_renders, emit_render_evidence  # noqa: E402
from chapter_qa.source_audit import derivation, emit_source_evidence  # noqa: E402
from chapter_qa.visual_qa import CATEGORIES, emit_visual_evidence, page_record  # noqa: E402
from chapter_qa.release_judge import judge  # noqa: E402
from chapter_qa.gates import gate_by_id, load_gates  # noqa: E402
from qa_pdf_quality import assess_pdf, classify_dpi, emit_pdf_evidence  # noqa: E402
from qa_ooxml import inspect_package  # noqa: E402
from qa_status_authority import emit_status_evidence, find_conflicts  # noqa: E402
from render_chapter_artifacts import guard_output_directory, mark_contact_sheet, planned_png_name  # noqa: E402

# Known Chapter 2 human-facing callout. The frozen playbook is not modified.
CH2_RETROSPECTIVE = (
    "Chapter 2 is a Phase 7.3 remediation candidate and is not frozen."
)


def _render(artifact="slides", page=1, artifact_hash="a" * 64, render_hash="e" * 64):
    relative = "qa/work/_phase_b/renders/%s-p%03d-%s.png" % (artifact, page, render_hash[:8])
    path = REPO / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    data = b"\x89PNG\r\n" + render_hash.encode()
    path.write_bytes(data)
    import hashlib
    digest = hashlib.sha256(data).hexdigest()
    return {
        "artifactId": artifact,
        "artifactHash": artifact_hash,
        "pageOrSlide": page,
        "renderPath": relative,
        "renderHash": digest,
        "navigation": False,
    }


def _expected(count=3, artifact_hash="a" * 64):
    return [_render(page=index, artifact_hash=artifact_hash, render_hash="e" * 64) for index in range(1, count + 1)]


def _box(**kwargs):
    base = {"kind": "text", "x": 0.7, "y": 0.7, "w": 2.0, "h": 0.4, "text": "moving", "fontPt": 18,
            "pageWidthIn": 8.5, "pageHeightIn": 11}
    base.update(kwargs)
    return base


class SourceTests(unittest.TestCase):
    def tearDown(self):
        cleanup_work()

    def test_source_derived_is_not_self_agreement(self):
        claim = {
            "claimId": "bq",
            "kind": "big_question",
            "candidateText": "What can moving water do?",
            "sourceText": "What can moving water do?",
            "sourceLocation": "Reader p. 1",
            "sourceId": "reader-1",
            "candidateId": "notes-1",
        }
        self.assertEqual(derivation(claim), "SOURCE_DERIVED")
        self.assertEqual(derivation({**claim, "sourceId": "same", "candidateId": "same"}), "UNRESOLVED")
        approved = {**claim, "sourceText": "", "approvedText": claim["candidateText"], "approvalId": "owner-decision-9"}
        self.assertEqual(derivation(approved), "APPROVED_DERIVED")

    def test_generated_agreement_cannot_pass_source_authority(self):
        path = evidence_file("source.json")
        stored, result = emit_source_evidence([{
            "claimId": "bq",
            "kind": "big_question",
            "candidateText": "same words",
            "sourceText": "same words",
            "sourceLocation": "Reader p. 1",
            "sourceId": "model-1",
            "candidateId": "model-1",
        }], build_identity(), "SOURCE_AUTHORITY", path)
        self.assertEqual(result.status, "FAIL")
        self.assertEqual(stored["status"], "FAIL")
        self.assertTrue(any(item["code"] == "generated_agreement" for item in stored["findings"]))

    def test_source_authority_can_pass_from_reader_text(self):
        path = evidence_file("source-pass.json")
        _stored, result = emit_source_evidence([{
            "claimId": "bq",
            "kind": "big_question",
            "candidateText": "What can moving water do?",
            "sourceText": "What can moving water do?",
            "sourceLocation": "Reader p. 1",
            "sourceId": "reader-1",
            "candidateId": "notes-1",
        }], build_identity(), "SOURCE_AUTHORITY", path, auditor_role="SOURCE_SCHOLAR")
        self.assertEqual(result.status, "PASS")

    def test_approved_text_does_not_satisfy_source_authority(self):
        path = evidence_file("approved.json")
        _stored, result = emit_source_evidence([{
            "claimId": "word",
            "kind": "vocabulary",
            "candidateText": "energy of motion",
            "sourceText": "",
            "approvedText": "energy of motion",
            "sourceLocation": "approved decision",
            "sourceId": "decision-1",
            "candidateId": "notes-1",
        }], build_identity(), "SOURCE_AUTHORITY", path)
        self.assertEqual(result.status, "FAIL")

    def test_curriculum_accepts_approved_derivation(self):
        path = evidence_file("curriculum.json")
        _stored, result = emit_source_evidence([{
            "claimId": "word",
            "kind": "vocabulary",
            "candidateText": "energy of motion",
            "sourceText": "",
            "approvedText": "energy of motion",
            "sourceLocation": "approved decision",
            "sourceId": "decision-1",
            "candidateId": "notes-1",
            "approvalId": "owner-decision-9",
        }], build_identity(), "CURRICULUM_CORRECTNESS", path, auditor_role="CURRICULUM_AUDITOR")
        self.assertEqual(result.status, "PASS")

    def test_empty_claims_cannot_pass(self):
        path = evidence_file("empty-source.json")
        _stored, result = emit_source_evidence([], build_identity(), "SOURCE_AUTHORITY", path)
        self.assertEqual(result.status, "FAIL")

    def test_caller_pass_is_discarded(self):
        path = evidence_file("forced.json")
        data = {
            "evidenceId": "forced",
            "gateId": "SOURCE_AUTHORITY",
            "status": "PASS",
            "performed": False,
            "expectedInputs": [],
            "observedInputs": [],
            "thingBeingCertified": "candidate-curriculum:%s" % ("a" * 64),
            "independentEvidenceSource": "authoritative-reader-tg:%s" % ("b" * 64),
            "findings": [],
            "evidencePaths": [path],
            "execution": {
                "succeeded": False, "detectorException": False, "requiredInputsPresent": False,
                "missingRender": False, "missingEvidence": True, "stale": False,
                "gateSpecificFailure": False, "unavailableEvidence": False,
            },
            "inventory": [{"name": "sourceClaims", "expected": 0, "observed": 0}],
        }
        stored, result = seal(data, build_identity())
        self.assertEqual(result.status, "FAIL")
        self.assertEqual(stored["status"], "FAIL")


class CrossArtifactTests(unittest.TestCase):
    def tearDown(self):
        cleanup_work()

    def _pair(self, left_text="moving", right_text="moving"):
        return (
            {"artifactClass": "guided_notes", "label": "notes", "contentHash": "n" * 64,
             "anchors": {"N1": {"instructional_text": left_text, "vocabulary": "motion", "anchor_code": "N1", "answers": "motion", "scoring": "1/1", "blanks": "motion"}}},
            {"artifactClass": "slides", "label": "slides", "contentHash": "s" * 64,
             "anchors": {"N1": {"instructional_text": right_text, "vocabulary": "motion", "anchor_code": "N1", "answers": "motion", "scoring": "1/1", "blanks": "motion"}}},
        )

    def test_agreement_is_not_source_authority(self):
        left, right = self._pair()
        with self.assertRaises(ValueError):
            emit_cross_evidence(left, right, build_identity(), "SOURCE_AUTHORITY", evidence_file("x.json"))
        with self.assertRaises(ValueError):
            emit_cross_evidence(left, right, build_identity(), "CURRICULUM_CORRECTNESS", evidence_file("y.json"))

    def test_matching_anchors_may_pass_sync(self):
        left, right = self._pair()
        _stored, result = emit_cross_evidence(left, right, build_identity(), "NOTES_SLIDES_SYNC", evidence_file("sync.json"), auditor_role="CROSS_ARTIFACT_AUDITOR")
        self.assertEqual(result.status, "PASS")

    def test_answer_mismatch_fails(self):
        left, right = self._pair()
        right["anchors"]["N1"]["answers"] = "speed"
        _stored, result = emit_cross_evidence(left, right, build_identity(), "ASSESSMENT_ALIGNMENT", evidence_file("answers.json"))
        self.assertEqual(result.status, "FAIL")

    def test_empty_artifacts_fail(self):
        left = {"artifactClass": "playbook", "label": "playbook", "contentHash": "p" * 64, "anchors": {}}
        right = {"artifactClass": "chapter_spec", "label": "spec", "contentHash": "q" * 64, "anchors": {}}
        _stored, result = emit_cross_evidence(left, right, build_identity(), "PLAYBOOK_ALIGNMENT", evidence_file("empty-cross.json"))
        self.assertEqual(result.status, "FAIL")


class RenderCompletenessTests(unittest.TestCase):
    def tearDown(self):
        cleanup_work()

    def test_expected_three_observed_three_may_proceed(self):
        expected = _expected(3)
        assessed = assess_renders(expected, expected)
        self.assertFalse(assessed["failsGate"])
        _stored, result = emit_render_evidence(expected, expected, build_identity(), evidence_file("renders.json"), auditor_role="VISUAL_QA_AUDITOR")
        self.assertEqual(result.status, "PASS")

    def test_expected_three_observed_two_fails(self):
        expected = _expected(3)
        _stored, result = emit_render_evidence(expected, expected[:2], build_identity(), evidence_file("two.json"))
        self.assertEqual(result.status, "FAIL")

    def test_expected_three_observed_zero_fails(self):
        _stored, result = emit_render_evidence(_expected(3), [], build_identity(), evidence_file("zero.json"), output_empty=True)
        self.assertEqual(result.status, "FAIL")

    def test_empty_directory_and_command_failure(self):
        _stored, empty = emit_render_evidence(_expected(3), [], build_identity(), evidence_file("dir.json"), output_empty=True)
        _stored, crashed = emit_render_evidence(_expected(3), [], build_identity(), evidence_file("crash.json"), command_failed=True)
        self.assertEqual(empty.status, "FAIL")
        self.assertEqual(crashed.status, "FAIL")

    def test_duplicate_extra_and_wrong_hash(self):
        expected = _expected(3)
        duplicate = [_render(page=1), _render(page=1, render_hash="f" * 64), _render(page=2)]
        extra = _expected(3) + [_render(page=4, render_hash="f" * 64)]
        wrong = [_render(page=index, artifact_hash="9" * 64) for index in range(1, 4)]
        self.assertEqual(emit_render_evidence(expected, duplicate, build_identity(), evidence_file("dup.json"))[1].status, "FAIL")
        self.assertEqual(emit_render_evidence(expected, extra, build_identity(), evidence_file("extra.json"))[1].status, "FAIL")
        self.assertEqual(emit_render_evidence(expected, wrong, build_identity(), evidence_file("wrong.json"))[1].status, "FAIL")

    def test_contact_sheet_is_not_exhaustive(self):
        sheet = mark_contact_sheet("qa/work/_phase_b/sheet.png")
        self.assertFalse(sheet["exhaustive"])
        observed = [{**item, "navigation": True} for item in _expected(3)]
        self.assertTrue(assess_renders(_expected(3), observed)["failsGate"])


class VisualTests(unittest.TestCase):
    def tearDown(self):
        cleanup_work()

    def _page(self, boxes, method="rendered-page", number=1, crops=None, image=None, dpi=300):
        return page_record({
            "artifactId": "notes",
            "artifactPath": "candidate/notes.docx",
            "artifactHash": "a" * 64,
            "renderPath": "renders/notes-p001-300dpi.png",
            "renderHash": "e" * 64,
            "pageOrSlide": number,
            "dpi": dpi,
            "width": 2550,
            "height": 3300,
            "inspectionMethod": method,
            "cropEvidence": crops or [],
            "reviewerRole": "VISUAL_QA_AUDITOR",
        }, boxes=boxes, image=image)

    def test_categories_are_classified(self):
        self.assertEqual(len(CATEGORIES), 20)
        self.assertIn(CATEGORIES["MID_WORD_SPLIT"], ("DETERMINISTIC", "MODEL_VISUAL_REVIEW", "BOTH", "DEFERRED"))
        self.assertEqual(CATEGORIES["DIAGRAM_UNREADABLE"], "MODEL_VISUAL_REVIEW")
        self.assertEqual(CATEGORIES["MISSING_GLYPH"], "DEFERRED")

    def test_clean_page_may_pass(self):
        page = self._page([_box()])
        _stored, result = emit_visual_evidence([page], build_identity(), evidence_file("visual.json"), expected_count=1, auditor_role="VISUAL_QA_AUDITOR")
        self.assertEqual(result.status, "PASS")

    def test_safe_area_overlap_font_split_and_writing_space(self):
        safe = self._page([_box(x=0.1)])
        self.assertTrue(any(item["code"] == "TEXT_OUTSIDE_SAFE_AREA" for item in safe["findings"]))
        overlap = self._page([_box(text="one"), _box(text="two")])
        self.assertTrue(any(item["code"] == "TEXT_OVERLAP" for item in overlap["findings"]))
        tiny = self._page([_box(fontPt=8, text="tiny")])
        self.assertTrue(any(item["code"] == "FONT_TOO_SMALL" for item in tiny["findings"]))
        split = self._page([_box(text="mov", w=0.8), _box(text="ing", x=1.6, w=0.8)])
        self.assertTrue(any(item["code"] == "MID_WORD_SPLIT" for item in split["findings"]))
        writing = self._page([_box(kind="writing", text="", h=0.2)])
        self.assertTrue(any(item["code"] == "WRITING_SPACE_TOO_SMALL" for item in writing["findings"]))
        for page in (safe, overlap, tiny, split, writing):
            status = emit_visual_evidence([page], build_identity(), evidence_file("bad-%s.json" % page["findings"][0]["code"]), expected_count=1)[1].status
            self.assertEqual(status, "FAIL")

    def test_no_render_input_fails(self):
        page = self._page(None)
        self.assertIn("no render input", page["findings"][0]["message"])
        _stored, result = emit_visual_evidence([page], build_identity(), evidence_file("no-input.json"), expected_count=1)
        self.assertEqual(result.status, "FAIL")

    def test_xml_and_contact_sheet_cannot_pass(self):
        for method in ("source-xml", "text-extraction", "contact-sheet"):
            page = self._page([_box()], method=method)
            status = emit_visual_evidence([page], build_identity(), evidence_file(method + ".json"), expected_count=1)[1].status
            self.assertEqual(status, "FAIL", method)

    def test_margin_ink_and_missing_crop(self):
        from PIL import Image
        image = Image.new("RGB", (200, 200), "white")
        image.putpixel((0, 0), (0, 0, 0))
        page = self._page([_box()], image=image, dpi=72)
        self.assertTrue(any(item["code"] == "TEXT_OUTSIDE_SAFE_AREA" for item in page["findings"]))
        self.assertTrue(any("600 DPI" in item["message"] for item in page["findings"]))


class PdfTests(unittest.TestCase):
    def tearDown(self):
        cleanup_work()

    def test_valid_corrupt_count_and_size(self):
        folder = REPO / "qa" / "work" / "_phase_b"
        folder.mkdir(parents=True, exist_ok=True)
        valid = folder / "valid.pdf"
        write_pdf(valid, [(612, 792, b"BT /F1 24 Tf 72 100 Td (Hello) Tj ET")])
        corrupt = folder / "corrupt.pdf"
        corrupt.write_bytes(b"%PDF-1.4\nthis is not a pdf")
        short = folder / "short.pdf"
        write_pdf(short, [(612, 792, b"BT /F1 24 Tf 72 100 Td (Hi) Tj ET")])
        small = folder / "small.pdf"
        write_pdf(small, [(200, 200, b"BT /F1 12 Tf 10 10 Td (X) Tj ET")])
        valid_assess = assess_pdf(str(valid), {"pageCount": 1, "pageSizePts": (612, 792), "label": "valid"})
        self.assertFalse(valid_assess["failsGate"])
        self.assertEqual(assess_pdf(str(corrupt), {"pageCount": 1, "label": "bad"})["findings"][0]["code"], "corrupt_pdf")
        self.assertTrue(assess_pdf(str(short), {"pageCount": 2, "pageSizePts": (612, 792), "label": "short"})["failsGate"])
        self.assertTrue(assess_pdf(str(small), {"pageCount": 1, "pageSizePts": (612, 792), "label": "small"})["failsGate"])
        _stored, result = emit_pdf_evidence(str(valid), build_identity(), evidence_file("pdf.json"), {"pageCount": 1, "pageSizePts": (612, 792), "label": "valid"}, auditor_role="PDF_OOXML_FORENSICS")
        self.assertEqual(result.status, "PASS")

    def test_dpi_bands_and_placements(self):
        low = classify_dpi(10, 10, "print-instructional", essential=True)
        high = classify_dpi(300, 300, "print-instructional", essential=True)
        self.assertEqual(low["band"], "fail")
        self.assertEqual(low["finding"]["code"], "low_effective_dpi_essential")
        self.assertIsNone(high["finding"])
        folder = REPO / "qa" / "work" / "_phase_b"
        folder.mkdir(parents=True, exist_ok=True)
        pdf = folder / "dpi.pdf"
        write_pdf(pdf, [(612, 792, b"BT /F1 24 Tf 72 100 Td (Hello) Tj ET")])
        _stored, result = emit_pdf_evidence(
            str(pdf),
            build_identity(),
            evidence_file("dpi.json"),
            {"pageCount": 1, "pageSizePts": (612, 792), "label": "placement", "placements": [{
                "pixelWidth": 50, "pixelHeight": 50, "placedWidthIn": 5, "placedHeightIn": 5, "essential": True,
            }]},
            auditor_role="PDF_OOXML_FORENSICS",
        )
        self.assertEqual(result.status, "FAIL")


class OoxmlTests(unittest.TestCase):
    def tearDown(self):
        cleanup_work()

    def test_valid_docx_and_corrupt_package(self):
        from docx import Document
        folder = REPO / "qa" / "work" / "_phase_b"
        folder.mkdir(parents=True, exist_ok=True)
        docx_path = folder / "valid.docx"
        document = Document()
        document.add_paragraph("Hello")
        document.save(docx_path)
        assessed = inspect_package(str(docx_path))
        self.assertFalse(assessed["failsGate"], assessed["findings"])
        corrupt = folder / "corrupt.docx"
        corrupt.write_bytes(b"not a zip")
        self.assertTrue(inspect_package(str(corrupt))["commandFailed"])

    def test_broken_relationship_and_corrupt_xml(self):
        from docx import Document
        folder = REPO / "qa" / "work" / "_phase_b"
        folder.mkdir(parents=True, exist_ok=True)
        original = folder / "base.docx"
        document = Document()
        document.add_paragraph("Hello")
        document.save(original)
        broken = folder / "broken.docx"
        _rewrite(original, broken, "word/_rels/document.xml.rels", lambda text: text.replace("styles.xml", "missing-styles.xml"))
        self.assertTrue(any(item["code"] == "broken_relationship" for item in inspect_package(str(broken))["findings"]))
        corrupt_xml = folder / "bad-xml.docx"
        _rewrite(original, corrupt_xml, "word/document.xml", lambda _text: "<not xml")
        self.assertTrue(any(item["code"] == "corrupt_xml" for item in inspect_package(str(corrupt_xml))["findings"]))

    def test_hidden_slide_and_note_mismatch(self):
        from pptx import Presentation
        folder = REPO / "qa" / "work" / "_phase_b"
        folder.mkdir(parents=True, exist_ok=True)
        hidden_path = folder / "hidden.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide._element.set("show", "0")
        presentation.save(hidden_path)
        hidden = inspect_package(str(hidden_path))
        self.assertTrue(hidden["summary"]["hiddenSlides"], hidden["summary"])
        mismatch_path = folder / "notes.pptx"
        deck = Presentation()
        first = deck.slides.add_slide(deck.slide_layouts[6])
        deck.slides.add_slide(deck.slide_layouts[6])
        first.notes_slide.notes_text_frame.text = "Only one note."
        deck.save(mismatch_path)
        mismatch = inspect_package(str(mismatch_path), notes_required=True)
        self.assertTrue(any(item["code"] == "speaker_note_mismatch" for item in mismatch["findings"]))
        matched = folder / "matched.pptx"
        one = Presentation()
        only = one.slides.add_slide(one.slide_layouts[6])
        only.notes_slide.notes_text_frame.text = "Say this."
        one.save(matched)
        self.assertFalse(inspect_package(str(matched), notes_required=True)["failsGate"])


class StatusTests(unittest.TestCase):
    def tearDown(self):
        cleanup_work()

    def test_historical_narrative_does_not_conflict(self):
        findings = find_conflicts("FROZEN", [{
            "kind": "HISTORICAL_NARRATIVE",
            "surfaceId": "changelog",
            "text": "Historical note: this chapter was a Phase 7.3 remediation candidate before it was frozen.",
        }])
        self.assertEqual(findings, [])

    def test_chapter_2_retrospective_human_status_conflicts(self):
        findings = find_conflicts("FROZEN", [{
            "kind": "HUMAN_FACING_STATUS",
            "surfaceId": "playbook-callout",
            "text": CH2_RETROSPECTIVE,
        }])
        self.assertEqual(findings[0]["severity"], "MEDIUM")
        self.assertEqual(findings[0]["code"], "stale_human_status")
        _stored, result = emit_status_evidence("FROZEN", [{
            "kind": "HUMAN_FACING_STATUS",
            "surfaceId": "playbook-callout",
            "text": CH2_RETROSPECTIVE,
        }], build_identity(), evidence_file("status.json"))
        self.assertEqual(result.status, "FAIL")

    def test_clean_frozen_status_may_pass(self):
        _stored, result = emit_status_evidence("FROZEN", [
            {"kind": "AUTHORITATIVE_STATUS", "surfaceId": "manifest", "text": "FROZEN"},
            {"kind": "BUILD_LABEL", "surfaceId": "build", "text": "phase-b-test"},
            {"kind": "HISTORICAL_NARRATIVE", "surfaceId": "history", "text": CH2_RETROSPECTIVE},
        ], build_identity(), evidence_file("status-pass.json"), auditor_role="CROSS_ARTIFACT_AUDITOR")
        self.assertEqual(result.status, "PASS")


class AdversarySmokeTests(unittest.TestCase):
    def tearDown(self):
        cleanup_work()

    def test_eight_bad_cases_are_not_pass_and_clean_controls_pass(self):
        identity = build_identity()
        expected = _expected(3)
        deleted = emit_render_evidence(expected, expected[:2], identity, evidence_file("smoke-delete.json"))[1]
        wrong = emit_render_evidence(expected, [_render(page=i, artifact_hash="9" * 64) for i in range(1, 4)], identity, evidence_file("smoke-hash.json"))[1]
        folder = REPO / "qa" / "work" / "_phase_b"
        folder.mkdir(parents=True, exist_ok=True)
        corrupt = folder / "smoke-corrupt.pdf"
        corrupt.write_bytes(b"%PDF-1.4\nbad")
        corrupt_result = emit_pdf_evidence(str(corrupt), identity, evidence_file("smoke-pdf.json"), {"pageCount": 1, "label": "corrupt"})[1]
        from pptx import Presentation
        hidden_path = folder / "smoke-hidden.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide._element.set("show", "0")
        presentation.save(hidden_path)
        hidden = inspect_package(str(hidden_path))
        from docx import Document
        base = folder / "smoke-base.docx"
        document = Document()
        document.add_paragraph("Hello")
        document.save(base)
        broken = folder / "smoke-broken.docx"
        _rewrite(base, broken, "word/_rels/document.xml.rels", lambda text: text.replace("styles.xml", "missing-styles.xml"))
        broken_result = inspect_package(str(broken))
        stale = emit_status_evidence("FROZEN", [{
            "kind": "HUMAN_FACING_STATUS", "surfaceId": "callout", "text": CH2_RETROSPECTIVE,
        }], identity, evidence_file("smoke-status.json"))[1]
        empty = emit_render_evidence(expected, [], identity, evidence_file("smoke-empty.json"), output_empty=True)[1]
        low_pdf = folder / "smoke-low.pdf"
        write_pdf(low_pdf, [(612, 792, b"BT /F1 24 Tf 72 100 Td (Hello) Tj ET")])
        low = emit_pdf_evidence(
            str(low_pdf), identity, evidence_file("smoke-dpi.json"),
            {"pageCount": 1, "pageSizePts": (612, 792), "label": "low", "placements": [{
                "pixelWidth": 40, "pixelHeight": 40, "placedWidthIn": 4, "placedHeightIn": 4, "essential": True,
            }]},
            auditor_role="PDF_OOXML_FORENSICS",
        )[1]
        for result in (deleted, wrong, corrupt_result, stale, empty, low):
            self.assertNotEqual(result.status, "PASS")
        self.assertTrue(hidden["failsGate"])
        self.assertTrue(broken_result["failsGate"])
        clean_render = emit_render_evidence(expected, expected, identity, evidence_file("clean-render.json"), auditor_role="VISUAL_QA_AUDITOR")[1]
        valid_pdf = folder / "clean.pdf"
        write_pdf(valid_pdf, [(612, 792, b"BT /F1 24 Tf 72 100 Td (Hello) Tj ET")])
        clean_pdf = emit_pdf_evidence(str(valid_pdf), identity, evidence_file("clean-pdf.json"), {"pageCount": 1, "pageSizePts": (612, 792), "label": "clean"}, auditor_role="PDF_OOXML_FORENSICS")[1]
        clean_docx = inspect_package(str(base))
        clean_status = emit_status_evidence("FROZEN", [
            {"kind": "AUTHORITATIVE_STATUS", "surfaceId": "manifest", "text": "FROZEN"},
        ], identity, evidence_file("clean-status.json"), auditor_role="CROSS_ARTIFACT_AUDITOR")[1]
        for result in (clean_render, clean_pdf, clean_status):
            self.assertEqual(result.status, "PASS")
        self.assertFalse(clean_docx["failsGate"])
        self.assertFalse(clean_docx["satisfiesPdfQuality"])


class JudgeIntegrationTests(unittest.TestCase):
    def tearDown(self):
        cleanup_work()

    def _passing_set(self):
        identity = build_identity()
        expected = _expected(1)
        render = emit_render_evidence(expected, expected, identity, evidence_file("j-render.json"), auditor_role="VISUAL_QA_AUDITOR")[0]
        page = page_record({
            "artifactId": "slides", "artifactPath": "candidate/slides.pptx", "artifactHash": "a" * 64,
            "renderPath": "renders/slides-p001-300dpi.png", "renderHash": "e" * 64,
            "pageOrSlide": 1, "dpi": 300, "width": 100, "height": 100,
            "inspectionMethod": "rendered-page", "reviewerRole": "VISUAL_QA_AUDITOR",
        }, boxes=[_box()])
        visual = emit_visual_evidence([page], identity, evidence_file("j-visual.json"), expected_count=1, auditor_role="VISUAL_QA_AUDITOR")[0]
        folder = REPO / "qa" / "work" / "_phase_b"
        folder.mkdir(parents=True, exist_ok=True)
        pdf = folder / "judge.pdf"
        write_pdf(pdf, [(612, 792, b"BT /F1 24 Tf 72 100 Td (Hello) Tj ET")])
        pdf_record = emit_pdf_evidence(str(pdf), identity, evidence_file("j-pdf.json"), {"pageCount": 1, "pageSizePts": (612, 792), "label": "judge"}, auditor_role="PDF_OOXML_FORENSICS")[0]
        status = emit_status_evidence("FROZEN", [
            {"kind": "AUTHORITATIVE_STATUS", "surfaceId": "manifest", "text": "FROZEN"},
        ], identity, evidence_file("j-status.json"), auditor_role="CROSS_ARTIFACT_AUDITOR")[0]
        from chapter_qa.evidence import evidence_from_dict
        gates = [gate_by_id(load_gates(), item["gateId"]) for item in (render, visual, pdf_record, status)]
        records = [evidence_from_dict(item) for item in (render, visual, pdf_record, status)]
        return identity, gates, records

    def test_complete_set_may_pass(self):
        identity, gates, records = self._passing_set()
        result = judge(gates, records, identity)
        self.assertEqual(result.status, "PASS")

    def test_missing_render_completeness_fails(self):
        identity, gates, records = self._passing_set()
        kept = [(gate, record) for gate, record in zip(gates, records) if gate.gateId != "RENDER_COMPLETENESS"]
        result = judge([gate for gate, _record in kept], [record for _gate, record in kept], identity)
        # The required set still includes the gate with no evidence.
        result = judge(gates, [record for record in records if record.gateId != "RENDER_COMPLETENESS"], identity)
        self.assertEqual(result.status, "FAIL")
        self.assertIn("MISSING_GATE: RENDER_COMPLETENESS", result.blockingReasons)

    def test_builder_only_visual_evidence_fails(self):
        identity, gates, records = self._passing_set()
        from chapter_qa.evidence import evidence_from_dict
        from dataclasses import replace
        visual = next(record for record in records if record.gateId == "RENDERED_VISUAL_QA")
        visual = replace(visual, auditorRole="BUILDER")
        others = [record for record in records if record.gateId != "RENDERED_VISUAL_QA"]
        result = judge(gates, others + [visual], identity)
        self.assertEqual(result.status, "FAIL")

    def test_stale_pdf_and_blocking_pdf_fail(self):
        identity, gates, records = self._passing_set()
        from dataclasses import replace
        pdf = next(record for record in records if record.gateId == "PDF_QUALITY")
        stale = replace(pdf, invalidationReasons=("CANDIDATE_CHANGED",))
        others = [record for record in records if record.gateId != "PDF_QUALITY"]
        self.assertEqual(judge(gates, others + [stale], identity).status, "FAIL")
        blocking = replace(pdf, findings=({"severity": "BLOCKING", "code": "corrupt_pdf", "message": "corrupt"},))
        # A BLOCKING finding with an otherwise passing execution still fails.
        from chapter_qa.evidence import Finding
        blocking = replace(pdf, findings=(Finding("BLOCKING", "corrupt_pdf", "corrupt"),))
        self.assertEqual(judge(gates, others + [blocking], identity).status, "FAIL")


class RenderPipelineTests(unittest.TestCase):
    def test_names_are_deterministic_and_production_is_protected(self):
        self.assertEqual(planned_png_name("My Notes.docx", 1, 300), "my-notes-docx-p001-300dpi.png")
        self.assertEqual(planned_png_name("My Notes.docx", 1, 600), "my-notes-docx-p001-600dpi.png")
        # A DOCX and a PDF with the same basename must not collide.
        self.assertNotEqual(planned_png_name("My Notes.docx", 1, 300), planned_png_name("My Notes.pdf", 1, 300))
        with self.assertRaises(ValueError):
            guard_output_directory(str(REPO / "production" / "releases"), str(REPO))
        sheet = mark_contact_sheet("qa/work/_phase_b/contact.png")
        self.assertEqual(sheet["role"], "navigation")
        self.assertFalse(sheet["exhaustive"])


def _rewrite(source: Path, dest: Path, member: str, transform) -> None:
    with zipfile.ZipFile(source) as incoming, zipfile.ZipFile(dest, "w") as outgoing:
        for info in incoming.infolist():
            data = incoming.read(info.filename)
            if info.filename == member:
                data = transform(data.decode("utf-8")).encode("utf-8")
            outgoing.writestr(info, data)


if __name__ == "__main__":
    unittest.main()
